"""Round 7B presentation of prepared Round 7A values; no Survey engine."""
from copy import deepcopy
from io import BytesIO
import math

import numpy as np

from well_profile import validate_well_inputs

CHART_ANCHOR = "Temperature profile chart: Not provided."
SCHEMATIC_ANCHOR = "Well schematic: Not provided."
SURVEY_TITLE = "Directional Survey Data"
SURVEY_ANCHOR = "Directional survey data: Not provided."
RED, PANEL_BG, GRID_GRAY, EDGE_GRAY, FRAME = "#D90000", "#F0F0F0", "#C9C9C9", "#D9D9D9", "#333333"


def canonical_report_source(state):
    geometry, geothermal = state.get("well_geometry", {}), state.get("geothermal_config", {})
    kind = geometry.get("type")
    source = {"type": kind}
    if kind == "Vertical":
        source["td_m"] = geometry.get("td_m")
    elif kind == "Directional":
        source["survey"] = geometry["survey"].to_dict("records")
    return {"geometry": source, "geothermal": {key: geothermal.get(key)
            for key in ("source", "surface_temp", "value")}}


def prepared_well_report(state):
    """Project prepared results in canonical station order, never integrate."""
    geometry, geothermal = state.get("well_geometry", {}), state.get("geothermal_config", {})
    validate_well_inputs(geometry, geothermal)
    kind = geometry.get("type")
    if kind is None or state.get("_well_profile_issues"):
        # Round 7A owns readiness; prepared-context callers may inspect drafts.
        return None
    if kind == "Vertical":
        td = state.get("geo_tvd")
        rows, departure = [[0., 0., 0., 0., 0.], [td, td, 0., 0., 0.]], [0., 0.]
    else:
        stations, resolved = geometry["survey"], state.get("_survey_result")
        if resolved is None or len(resolved) != len(stations) or not np.array_equal(stations["MD"], resolved["MD"]):
            raise ValueError("Round 7A Survey must be rebuilt before Word output.")
        rows = [[float(station["MD"]), float(result["TVD"]), float(station["Deviation"]),
                 float(station["Azimuth"]), float(result["DLS (deg/100 m)"])]
                for (_, station), (_, result) in zip(stations.iterrows(), resolved.iterrows())]
        departure = resolved["Horizontal Departure"].tolist()
    payload = {"type": kind, "rows": rows, "departure": departure,
               "surface": geothermal.get("surface_temp"), "bhst": state.get("bhst"),
               "gradient": state.get("geo_gradient")}
    if (len(rows) < 2 or any(value is None or not math.isfinite(value)
            for value in [payload["surface"], payload["bhst"], payload["gradient"], *departure,
                          *(value for row in rows for value in row)]) or rows[-1][1] <= 0):
        raise ValueError("Complete Round 7A Geometry and Geothermal before Word output.")
    return payload


def select_survey_rows(rows):
    """Owner-locked nearest-unused actual stations; ties choose lower MD."""
    if len(rows) <= 20:
        return rows
    available = set(range(1, len(rows)-1))
    selected = [0, len(rows)-1]
    first, last = rows[0][0], rows[-1][0]
    for k in range(1, 19):
        target = first+k*(last-first)/19
        index = min(available, key=lambda i: (abs(rows[i][0]-target), rows[i][0]))
        selected.append(index)
        available.remove(index)
    return [rows[i] for i in sorted(selected, key=lambda i: rows[i][0])]


# v2.2 FINAL presentation helpers only. Engineering remains in well_profile.py.
def _clean(value):
    return float(f"{value:.10g}")


def nice_step_ge(raw):
    if raw <= 0:
        return 1.
    exponent = math.floor(math.log10(raw))
    for multiplier in (1, 2, 2.5, 5, 10):
        value = multiplier*10.**exponent
        if value >= raw*(1-1e-12):
            return _clean(value)
    return _clean(10*10.**exponent)


def next_nice_step(step):
    exponent = math.floor(math.log10(step))
    for multiplier in (1, 2, 2.5, 5, 10):
        value = multiplier*10.**exponent
        if value > step*1.0000001:
            return _clean(value)
    return _clean(10*10.**exponent)


def fmt_ticks(values, scale=1.):
    values = [_clean(value/scale) for value in values]
    texts = [f"{value:.6f}".rstrip("0").rstrip(".") for value in values]
    decimals = max(len(text.split(".")[1]) if "." in text else 0 for text in texts)
    return ["0" if abs(value) < 1e-12 else f"{value:.{decimals}f}" for value in values]


def resolve_depth_axis(raw_max):
    def pick(maximum):
        step = nice_step_ge(maximum/8.)
        while maximum/step > 12:
            step = next_nice_step(step)
        return step
    step = pick(raw_max)
    maximum = _clean(math.ceil(raw_max/step-1e-9)*step)
    for _ in range(5):
        new_step = pick(maximum)
        if new_step == step:
            break
        step = new_step
        maximum = _clean(math.ceil(raw_max/step-1e-9)*step)
    for multiplier in (1, 2, 5, 10, 20, 50, 100):
        label_step = _clean(step*multiplier)
        if maximum/label_step <= 5+1e-9:
            break
    ticks = [_clean(k*label_step) for k in range(int(math.floor(maximum/label_step+1e-9))+1)]
    factor = 1. if maximum <= 1000 else 10.**(len(str(int(maximum)))-2)
    return dict(depth_max=maximum, depth_step=step,
                depth_grid=[_clean(k*step) for k in range(1, int(round(maximum/step)))],
                depth_label_step=label_step, depth_ticks=ticks, depth_factor=factor,
                depth_tick_labels=fmt_ticks(ticks, factor),
                depth_label="TVD (m)" if factor == 1 else f"(x {factor:g})   TVD (m)",
                outer_ticks=[0., _clean(maximum/2), maximum],
                outer_tick_labels=fmt_ticks([0., _clean(maximum/2), maximum]))


def pick_temp_axis(low, high):
    step = nice_step_ge(max(high-low, 1e-6)/4.)
    for _ in range(60):
        minimum = _clean(math.floor(low/step+1e-9)*step)
        if minimum+4*step >= high-1e-9:
            return step, minimum, _clean(minimum+4*step)
        step = next_nice_step(step)
    return step, minimum, _clean(minimum+4*step)


def build_chart_spec(payload):
    """v2.2 axes and plotting data from resolved values, not a new engine."""
    tvd = np.array([row[1] for row in payload["rows"]])
    spec = resolve_depth_axis(float(np.max(tvd)))
    z = np.linspace(float(np.min(tvd)), float(np.max(tvd)), 250)
    # The sole profile is the prepared Surface/BHST line. v2.2 holds an
    # endpoint outside its defined TVD interval; no extra temperature source.
    temperature = np.interp(z, [0., tvd[-1]], [payload["surface"], payload["bhst"]])
    step, low, high = pick_temp_axis(float(np.min(temperature)), float(np.max(temperature)))
    spec.update(temp_step=step, temp_xmin=low, temp_xmax=high, temp_data=(temperature, z),
                temp_legend="Geoth.Profile", dep_legend="Hori. Depart.",
                dep_data=(np.array(payload["departure"]), tvd))
    dep_step = .25 if max(payload["departure"]) < 1e-6 else nice_step_ge(max(payload["departure"])/4.)
    spec.update(dep_step=dep_step, dep_xmin=0., dep_xmax=1. if max(payload["departure"]) < 1e-6 else _clean(4*dep_step))
    for track in ("dep", "temp"):
        minimum, step = spec[f"{track}_xmin"], spec[f"{track}_step"]
        spec[f"{track}_grid"] = [_clean(minimum+k*step) for k in (1, 2, 3)]
        spec[f"{track}_ticks"] = [_clean(minimum+k*step) for k in (0, 2, 4)]
        spec[f"{track}_tick_labels"] = fmt_ticks(spec[f"{track}_ticks"])
    return spec


def render_chart_png(payload):
    """Fresh OO/Agg figure and BytesIO per build; no pyplot/global figure."""
    try:
        from matplotlib.figure import Figure
        from matplotlib.backends.backend_agg import FigureCanvasAgg
    except ImportError as exc:
        raise ValueError("Word chart rendering requires matplotlib==3.10.8.") from exc
    spec = build_chart_spec(payload)
    figure = Figure(figsize=(10., 8.), constrained_layout=True)
    FigureCanvasAgg(figure)
    grid = figure.add_gridspec(1, 2, width_ratios=[1, 1.22])
    left = figure.add_subplot(grid[0, 0])
    right = figure.add_subplot(grid[0, 1], sharey=left)
    try:
        for ax, track in ((left, "dep"), (right, "temp")):
            ax.set_facecolor(PANEL_BG)
            for side, color, width in (("left", FRAME, 1.3), ("top", FRAME, 1.3),
                                       ("right", EDGE_GRAY, 1.), ("bottom", EDGE_GRAY, 1.)):
                ax.spines[side].set_color(color)
                ax.spines[side].set_linewidth(width)
            ax.xaxis.tick_top()
            ax.xaxis.set_label_position("top")
            ax.set_xlabel("m" if track == "dep" else "degF", fontsize=11, labelpad=8)
            ax.tick_params(labelsize=10, direction="out", color=FRAME)
            ax.set_ylim(spec["depth_max"], 0)
            ax.set_yticks(spec["depth_ticks"])
            if track == "dep":
                ax.set_yticklabels(spec["depth_tick_labels"])
                ax.set_ylabel(spec["depth_label"], fontsize=11)
            else:
                for label in ax.get_yticklabels():
                    label.set_visible(False)
            for y in spec["depth_grid"]:
                ax.axhline(y, color=GRID_GRAY, lw=1., zorder=1)
            ax.set_xlim(spec[f"{track}_xmin"], spec[f"{track}_xmax"])
            ax.set_xticks(spec[f"{track}_ticks"])
            ax.set_xticklabels(spec[f"{track}_tick_labels"])
            for x in spec[f"{track}_grid"]:
                ax.axvline(x, color=GRID_GRAY, lw=1., zorder=1)
            ax.plot(*spec[f"{track}_data"], color=RED, lw=6., label=spec[f"{track}_legend"], zorder=3, clip_on=False)
            ax.legend(loc="upper left", fontsize=10, frameon=True, framealpha=1,
                      facecolor="white", edgecolor="black")
        outer = left.secondary_yaxis(-.30)
        outer.set_yticks(spec["outer_ticks"])
        outer.set_yticklabels(spec["outer_tick_labels"], fontsize=10)
        outer.set_ylabel("m", fontsize=11, rotation=0)
        outer.yaxis.set_label_coords(-.30, 1.01)
        outer.tick_params(direction="out", length=4, color=FRAME, labelsize=10)
        outer.spines["left"].set_color(FRAME)
        output = BytesIO()
        figure.savefig(output, format="png", dpi=200, facecolor="white")
        return output.getvalue()
    finally:
        figure.clear()


def _postprocess_centralizer(doc, prepared):
    """Guarded Section IX replacement in the rendered copy; preserve its page transition."""
    import re
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn
    from centralizer_report import NOTE_BODY, WARNING, NOT_APPLICABLE, render_chart_png
    headings = {}
    for title in ("IX. Centralizer Placement Recommendation", "X. Cementing Procedure"):
        found = [p for p in doc.paragraphs if p.text == title]
        if len(found) != 1:
            raise ValueError(f"Centralizer report requires exactly one heading: {title}")
        headings[title] = found[0]._p
    elements = list(doc._element.body)
    start, end = [elements.index(headings[title]) for title in headings]
    if start >= end:
        raise ValueError("Centralizer Section IX must precede Section X.")
    region = elements[start+1:end]
    tables = [table for table in doc.tables if table._tbl in region]
    if (len(tables) != 2 or len(tables[0].rows) != 4 or len(tables[1].rows) != 4 or
            tables[0].cell(0, 0).text != "Centralizer Data" or
            tables[1].cell(0, 0).text.strip() != "Centralizer Placement"):
        raise ValueError("Section IX requires the two approved Centralizer tables and sample-row structure.")
    data, placement = tables
    labels = [[cell.text.replace("\xa0", " ") for cell in row.cells] for row in data.rows[1:]]
    if [row[::2] for row in labels] != [["Top Cent. MD", "NB of Cent. Used"],
            ["Bottom Cent. MD", "NB of Rigid Cent. Used"], ["Casing Shoe", "NB of Spring Cent. Used"]]:
        raise ValueError("Centralizer Data labels differ from the approved template.")
    if [cell.text.strip() for cell in placement.rows[1].cells] != ["Bottom MD (m)", "Number",
            "Cent. / Joint", "Centralizer Type", "Min. Standoff (%)", "Depth (m)"]:
        raise ValueError("Centralizer Placement must retain its exact six columns.")
    paragraphs = [Paragraph(e, doc._body) for e in region if e.tag == qn("w:p")]
    notes = [p for p in paragraphs if re.fullmatch(r"NOTE \d+:\s+" + re.escape(NOTE_BODY), p.text)]
    charts = [p for p in paragraphs if p._p.xpath('.//w:drawing')]
    if len(notes) != 1 or len(charts) != 1 or len(charts[0]._p.xpath('.//w:drawing')) != 1:
        raise ValueError("Section IX requires exactly one API-10D note and one existing chart drawing.")
    note, chart = notes[0], charts[0]
    if (chart._p is not elements[end-1] or len(chart._p.xpath('.//w:br[@w:type="page"]')) != 1 or
            not start < elements.index(data._tbl) < elements.index(placement._tbl) <
            elements.index(note._p) < elements.index(chart._p) < end):
        raise ValueError("Section IX note/chart order or existing Section-X page transition differs.")
    drawing = chart._p.xpath('.//w:drawing')[0]
    inline = drawing.find(qn("wp:inline"))
    if inline is None or len(inline.xpath('.//a:blip')) != 1:
        raise ValueError("Section IX requires the approved inline Pipe Standoff drawing.")
    extent = inline.find(qn("wp:extent"))
    transform = inline.xpath('.//a:xfrm/a:ext')
    if extent is None or len(transform) != 1:
        raise ValueError("Cannot resolve the existing Pipe Standoff layout box.")
    width, height = int(extent.get("cx")), int(extent.get("cy"))
    if width <= 0 or height <= 0:
        raise ValueError("Pipe Standoff layout box must be positive.")
    blip = inline.xpath('.//a:blip')[0]
    old_rid = blip.get(qn("r:embed"))

    def replace_text(paragraph, text):
        runs = paragraph.runs
        (runs[0] if runs else paragraph.add_run()).text = text
        for run in runs[1:]:
            run.text = ""

    if not prepared["applicable"]:
        for table in tables:
            table._tbl.getparent().remove(table._tbl)
        replace_text(note, NOT_APPLICABLE)
        drawing.getparent().remove(drawing)
    else:
        view = prepared["view"]
        if not prepared.get("note"):
            raise ValueError("Section IX requires its sequential Centralizer NoteCounter note.")
        png = render_chart_png(prepared)
        rid, image = doc.part.get_or_add_image(BytesIO(png))
        scale = min(width/image.px_width, height/image.px_height)
        cx, cy = int(image.px_width*scale), int(image.px_height*scale)
        blip.set(qn("r:embed"), rid)
        for node in (extent, transform[0]):
            node.set("cx", str(cx)); node.set("cy", str(cy))
        inline.find(qn("wp:docPr")).set("descr", "Pipe Standoff")
        values = ((f"{view['top_cent_md']:.1f} m", str(view["n_cent"])),
                  (f"{view['bottom_cent_md']:.1f} m", str(view["n_rigid"])),
                  (f"{view['shoe']:.1f} m", str(view["n_spring"])))
        for row, pair in zip(data.rows[1:], values):
            for cell, value in zip((row.cells[1], row.cells[3]), pair):
                replace_text(cell.paragraphs[0], value)
        row_template = deepcopy(placement.rows[2]._tr)
        for row in list(placement.rows)[2:]:
            placement._tbl.remove(row._tr)
        for values in view["zone_rows"]:
            placement._tbl.append(deepcopy(row_template))
            texts = (f"{values['bottom_md']:.1f}", str(values["number"]), values["cent_per_joint"],
                     values["kind_display"], f"{values['min_stdf_pct']:.1f}", f"{values['min_depth']:.1f}")
            for cell, text in zip(placement.rows[-1].cells, texts):
                replace_text(cell.paragraphs[0], text)
        replace_text(note, prepared["note"])
        if prepared["warning"]:
            warning = deepcopy(note._p)
            note._p.addprevious(warning)
            replace_text(Paragraph(warning, doc._body), WARNING)
    if old_rid not in doc._element.xpath('.//@r:embed'):
        doc.part.drop_rel(old_rid)


PUMPING_CHART_ANCHOR = "Pumping simulation charts: Not provided."
PUMPING_NOT_APPLICABLE = "Pumping simulation charts: Not applicable."


def _content_width(doc, paragraph):
    from docx.section import Section
    from docx.oxml.ns import qn
    elements = list(doc._element.body)
    for element in elements[elements.index(paragraph._p):]:
        section_xml = element if element.tag == qn("w:sectPr") else element.find(".//"+qn("w:sectPr"))
        if section_xml is not None:
            section = Section(section_xml, doc.part)
            return section.page_width-section.left_margin-section.right_margin
    raise ValueError("Cannot resolve the chart's existing page geometry.")


def render_schematic_png(payload):
    """The owner renderer owns drawing; its temporary output never enters state."""
    from pathlib import Path
    from tempfile import TemporaryDirectory
    from well_schematic_engine import render_well_schematic
    with TemporaryDirectory(prefix="cement-well-schematic-") as directory:
        path = render_well_schematic(payload, str(Path(directory) / "schematic.png"), strict=True)
        return Path(path).read_bytes()


def postprocess_well_report(rendered_bytes, payload, vertical_note, cement_report=None, centralizer_report=None, *, pumping_report=None, schematic_report=None):
    """Modify only exact anchors in a rendered, trusted template copy."""
    from docx import Document
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn
    doc = Document(BytesIO(rendered_bytes))
    paragraphs = [Paragraph(element, doc._body) for element in doc._element.body.xpath(".//w:p")]
    schematic_anchor = None
    if schematic_report is not None:
        matches = [p for p in paragraphs if p.text == SCHEMATIC_ANCHOR]
        if len(matches) != 1 or matches[0]._p.getparent() is not doc._element.body:
            raise ValueError(f"Word report requires exactly one body anchor: {SCHEMATIC_ANCHOR}")
        schematic_anchor = matches[0]
        elements = list(doc._element.body)
        ii = [p for p in doc.paragraphs if p.text.strip() == "II. Well Data Summary"]
        iii = [p for p in doc.paragraphs if p.text.strip() == "III. Geothermal Temperature Profile"]
        if (len(ii) != 1 or len(iii) != 1 or not
                elements.index(ii[0]._p) < elements.index(schematic_anchor._p) < elements.index(iii[0]._p)):
            raise ValueError("The Well schematic anchor must be inside the existing Section II.")
    cement_anchor = None
    if cement_report is not None:
        from cement_chart_report import CHART_ANCHOR as cement_text
        matches = [p for p in paragraphs if p.text == cement_text]
        if len(matches) != 1 or matches[0]._p.getparent() is not doc._element.body:
            raise ValueError(f"Word report requires exactly one anchor: {cement_text}")
        cement_anchor = matches[0]
        elements = list(doc._element.body)
        iv = [p for p in doc.paragraphs if p.text.strip() == "IV. Fluids Sequence"]
        v = [p for p in doc.paragraphs if p.text.strip() == "V. Cementing Program"]
        if (len(iv) != 1 or len(v) != 1 or not
                elements.index(iv[0]._p) < elements.index(cement_anchor._p) < elements.index(v[0]._p)):
            raise ValueError("The diagnostic chart anchor must be inside the existing Section IV.")
        if cement_report["kind"] is not None and cement_report["model"] is None:
            raise ValueError("A prepared Round 8A chart model is required.")
    pumping_anchor = None
    if pumping_report is not None:
        matches = [p for p in paragraphs if p.text == PUMPING_CHART_ANCHOR]
        if len(matches) != 1 or matches[0]._p.getparent() is not doc._element.body:
            raise ValueError(f"Word report requires exactly one anchor: {PUMPING_CHART_ANCHOR}")
        pumping_anchor = matches[0]
        if cement_anchor is None or not elements.index(cement_anchor._p) < elements.index(pumping_anchor._p) < elements.index(v[0]._p):
            raise ValueError("The Pumping chart anchor must follow the diagnostic chart inside the existing Section IV.")
        if pumping_report["applicable"] and pumping_report.get("figure") is None:
            raise ValueError("A fresh Pumping Simulation figure is required.")
    anchors = {}
    for text in (CHART_ANCHOR, SURVEY_ANCHOR):
        matches = [paragraph for paragraph in paragraphs if paragraph.text == text]
        if len(matches) != 1 or matches[0]._p.getparent() is not doc._element.body:
            raise ValueError(f"Word report requires exactly one anchor: {text}")
        anchors[text] = matches[0]
    tables = [table for table in doc.tables if any(cell.text == SURVEY_TITLE for row in table.rows for cell in row.cells)]
    if len(tables) != 1:
        raise ValueError("Word report requires exactly one Directional Survey Data table.")
    table = tables[0]
    if (len(table.rows) < 4 or [cell.text.strip() for cell in table.rows[1].cells] !=
            ["MD", "TVD", "Deviation", "Azimuth", "Dogleg Sev."] or
            [cell.text for cell in table.rows[2].cells] != ["(m)", "(m)", "(deg)", "(deg)", "(deg/m)"]):
        raise ValueError("Directional Survey Data table structure differs from the approved template.")
    if payload is None:
        raise ValueError("Complete Round 7A inputs before building the chart and Survey report.")
    if payload["type"] == "Vertical" and not vertical_note:
        raise ValueError("Vertical report requires its sequential NoteCounter note.")
    chart = anchors[CHART_ANCHOR]
    elements = list(doc._element.body)
    headings = {}
    for title in ("VIII. Directional Well Survey Data", "IX. Centralizer Placement Recommendation"):
        matches = [p for p in doc.paragraphs if p.text.strip() == title]
        if len(matches) != 1:
            raise ValueError(f"Word report requires exactly one section anchor: {title}")
        headings[title] = matches[0]._p
    survey_index = elements.index(anchors[SURVEY_ANCHOR]._p)
    ix = headings["IX. Centralizer Placement Recommendation"]
    page_break = ix.getprevious()
    if (page_break is None or page_break.tag != qn("w:p") or
            page_break.xpath('.//w:br[@w:type="page"]') == [] or Paragraph(page_break, doc._body).text.strip()):
        raise ValueError("Section VIII requires the existing manual page break immediately before Section IX.")
    break_index = elements.index(page_break)
    if not (elements.index(headings["VIII. Directional Well Survey Data"]) <
            elements.index(table._tbl) < survey_index < break_index):
        raise ValueError("Section VIII Survey anchors are out of the approved document order.")
    spacers = []
    for element in elements[survey_index+1:break_index]:
        if (element.tag != qn("w:p") or Paragraph(element, doc._body).text.strip() or
                element.xpath(".//w:drawing | .//w:br | .//w:fldChar | .//w:instrText")):
            raise ValueError("Section VIII spacer region differs from the approved structure.")
        if element.xpath('./w:pPr/w:pStyle[@w:val="Heading1"]'):
            spacers.append(element)
    if len(spacers) != 6:
        raise ValueError("Section VIII requires exactly six empty Heading1 spacer paragraphs.")
    png = render_chart_png(payload)
    chart.clear().add_run().add_picture(BytesIO(png), width=_content_width(doc, chart))
    row_template = deepcopy(table.rows[3]._tr)
    for row in list(table.rows)[3:]:
        table._tbl.remove(row._tr)
    for values in select_survey_rows(payload["rows"]):
        table._tbl.append(deepcopy(row_template))
        row = table.rows[-1]
        texts = [f"{values[0]:.1f}", f"{values[1]:.1f}", f"{values[2]:.2f}",
                 f"{values[3]:.2f}", f"{values[4]/100:.3f}"]
        for cell, text in zip(row.cells, texts):
            paragraph = cell.paragraphs[0]
            runs = paragraph.runs
            (runs[0] if runs else paragraph.add_run()).text = text
            for run in runs[1:]:
                paragraph._p.remove(run._r)
    survey_anchor = anchors[SURVEY_ANCHOR]
    if payload["type"] == "Vertical":
        survey_anchor.clear().add_run(vertical_note)
    else:
        survey_anchor._p.getparent().remove(survey_anchor._p)
    for spacer in spacers:
        doc._element.body.remove(spacer)
    if cement_anchor is not None:
        from cement_chart_report import render_chart_png as render_cement_png, NOT_APPLICABLE
        if cement_report["kind"] is None:
            # Preserve paragraph/run properties while replacing the exact text.
            runs = cement_anchor.runs
            (runs[0] if runs else cement_anchor.add_run()).text = NOT_APPLICABLE
            for run in runs[1:]:
                cement_anchor._p.remove(run._r)
        else:
            png = render_cement_png(cement_report["model"])
            run = cement_anchor.clear().add_run()
            picture = run.add_picture(BytesIO(png), width=_content_width(doc, cement_anchor))
            picture._inline.docPr.set("descr", "Static Well Security and Fluid Sequence")
    if pumping_anchor is not None:
        from docx.text.run import Run
        texts = pumping_anchor._p.xpath('.//w:t')
        for text in texts:
            text.text = ""
        if not pumping_report["applicable"]:
            texts[0].text = PUMPING_NOT_APPLICABLE
        else:
            output = BytesIO()
            pumping_report["figure"].savefig(output, format="png", dpi=200, facecolor="white")
            run = Run(texts[0].getparent(), pumping_anchor)
            picture = run.add_picture(BytesIO(output.getvalue()), width=_content_width(doc, pumping_anchor))
            picture._inline.docPr.set("descr", "Pumping Simulation Charts")
    if centralizer_report is not None:
        _postprocess_centralizer(doc, centralizer_report)
    if schematic_anchor is not None:
        from docx.text.run import Run
        texts = schematic_anchor._p.xpath('.//w:t')
        for text in texts:
            text.text = ""
        run = Run(texts[0].getparent(), schematic_anchor)
        picture = run.add_picture(BytesIO(render_schematic_png(schematic_report)),
                                  width=_content_width(doc, schematic_anchor))
        picture._inline.docPr.set("descr", "Well Schematic")
    result = BytesIO()
    doc.save(result)
    return result.getvalue()
