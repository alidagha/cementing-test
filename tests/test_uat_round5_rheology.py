"""Round 5-1: CEMCADE Bingham and three independent Lab datasets."""
from copy import deepcopy
import test_audit_regressions as audit
from engineering_tools import compute_phase_status


def test_three_independent_selectors_exist():
    case = audit.AuditRegressions()
    app = case.configured_app('CSG 9 5/8"')
    case.phase(app, 'phase7')
    labels = [w.label for w in app.checkbox]
    for label in ('Surface — Ramp-Down', 'BHCT — Ramp-Up', 'BHCT — Ramp-Down'):
        assert f'{label} - Main' in labels


def test_old_project_without_rheology_is_not_export_ready():
    case = audit.AuditRegressions()
    app = case.configured_app('CSG 9 5/8"')
    project = audit.round_trip(app.session_state.to_dict())
    project['lab_qc_params']['Main'].pop('rheology', None)
    restored = case.app(project)
    assert compute_phase_status(restored.session_state)['phase7']['level'] != 'ok'
    case.assert_export_blocked(restored)

from io import BytesIO
from itertools import combinations
from pathlib import Path
from unittest.mock import patch
import hashlib
import re
import subprocess
import zipfile
import pytest
from docx import Document
import phase_10_procedure as report
from project_io import replace_project_state
from project_state import prepare_calculations, purge_inactive_slurries, lab_source_signature, refresh_fluids
from engineering_tools import lab_review_signature
from rheology import (calculate_cemcade_bingham, build_rheology, validate_rheology_inputs,
                      RPM_ORDER, RHEOLOGY_DATASETS)

CORPUS = [
    ('A', [168,116,64,41,24,8,6], '159.218','9.47','0.960'),
    ('B', [175,124,68,45,29,12,9], '163.158','13.14','0.960'),
    ('C', [150,108,70,52,38,18,14], '123.126','26.90','0.903'),
    ('D', [156,106,56,36,21,8], '150.637','5.61',None),
    ('E', [170,120,70,50,15], '165.302','8.96','0.920'),
    *[(f'6={n}', [168,116,64,41,24,n], pv, ty, '0.979' if n==9 else None)
      for n,pv,ty in [(8,'159.619','9.13'),(9,'161.410','8.07'),(10,'161.124','8.24'),
                     (12,'159.884','9.15'),(15,'158.530','10.04'),(17,'157.417','10.71'),(20,'155.933','11.77')]],
    ('Blind1', [192,139,82,55,36,15,11], '173.353','21.01','0.927'),
    ('Blind2', [220,151,82,52,29,8,5], '214.435','7.43','0.966'),
    ('Blind3', [120,95,70,55,43,24,18], '84.369','37.75','0.852'),
    ('negative', [150,100,50,30,15,2,1], '150.088','0.00','0.992'),
]

@pytest.mark.parametrize('name,dials,pv,ty,iod',CORPUS,ids=[c[0] for c in CORPUS])
def test_authority_corpus(name,dials,pv,ty,iod):
    r=calculate_cemcade_bingham(zip(RPM_ORDER,dials))
    if name=='Blind1':
        assert abs(float(f"{r['pv_cp']:.3f}")-float(pv)) <= .002000000001
    else:
        assert f"{r['pv_cp']:.3f}" == pv
    assert f"{r['ty_lbf_100ft2']:.2f}" == ty
    if iod is not None:assert f"{r['iod']:.3f}" == iod
    assert r['fallback_to_newtonian'] == (name=='negative')
    assert r['is_valid_bingham'] == (name!='negative')


def dataset(dials=None,sec='8',minutes='12'):
    return {'selected':True, 'readings':dict(zip(map(str,RPM_ORDER),map(str,dials or CORPUS[0][1]))),
            'gel_10_sec':sec, 'gel_10_min':minutes}


def test_filtering_order_minimum_and_failure_safety():
    rows=list(zip(RPM_ORDER,CORPUS[4][1]))
    r=calculate_cemcade_bingham(rows)
    assert calculate_cemcade_bingham(rows+[(6,''),(3,0),(-1,2),(1,-2)]) == r
    # Fresh, reproducible fitting and preserved accumulation order.
    assert calculate_cemcade_bingham(rows) == r
    with pytest.raises(ValueError):calculate_cemcade_bingham([(300,10),(200,0)])
    assert calculate_cemcade_bingham([(300,168),(200,116)])['valid_point_count']==2
    with pytest.raises((ValueError,OverflowError)):calculate_cemcade_bingham([(300,float('inf')),(200,10)])
    result,issues=build_rheology({'rheology':{'surface_down':dataset([1e300,1e300])}})
    assert issues and not result['surface_down']['is_valid_bingham']

KEYS=[key for key,_ in RHEOLOGY_DATASETS]
SELECTIONS=[c for n in (1,2,3) for c in combinations(KEYS,n)]

@pytest.mark.parametrize('selected',SELECTIONS)
def test_independent_selection_and_real_docx_direct_restore(selected):
    case=audit.AuditRegressions();app=case.configured_app('CSG 9 5/8"')
    project=audit.round_trip(app.session_state.to_dict())
    qc=project['lab_qc_params']['Main']
    qc['rheology']={key:dict(dataset(CORPUS[i][1],str(8+i),'-'), selected=key in selected)
                    for i,key in enumerate(KEYS)}
    qc['review_signature']=lab_review_signature(qc,project['lab_grid_dfs']['Main'])
    result,issues=build_rheology(qc);assert not issues
    # Saved derived payloads are never the fitting authority.
    project['lab_payload_Main']={'rheology': {'stale_saved_fit':9999}}
    assert prepare_calculations(project)==[]
    assert project['lab_payload_Main']['rheology']==result
    restored=case.app(audit.round_trip(project));case.export(restored)
    docbytes=restored.session_state['_compiled_doc_bytes']
    Path('/tmp/r5-'+ '-'.join(selected)+'.docx').write_bytes(docbytes)
    table=next(t for t in Document(BytesIO(docbytes)).tables if t.rows[0].cells[0].text=='Rheology Test')
    assert table.rows[1].cells[2].text == table.rows[1].cells[3].text == 'Tstart: 150 degF'
    for i,key in enumerate(KEYS,1):
        assert table.rows[3].cells[i].text == (str(CORPUS[i-1][1][0]) if key in selected else '-')
        assert table.rows[10].cells[i].text == (f"{result[key]['pv_cp']:.3f}" if key in selected else '-')
        assert table.rows[11].cells[i].text == (f"{result[key]['ty_lbf_100ft2']:.2f}" if key in selected else '-')
        assert table.rows[12].cells[i].text == (str(7+i) if key in selected else '-')
        assert table.rows[13].cells[i].text == '-'
    assert restored.session_state['lab_qc_params']['Main']['rheology']==qc['rheology']


def test_first_edit_navigation_drafts_and_review_invalidation():
    case=audit.AuditRegressions();app=case.configured_app('CSG 9 5/8"');case.phase(app,'phase7')
    # Existing reviewed Surface reading, edited on first attempt then immediate navigation.
    label='300 RPM - Surface — Ramp-Down - Main'
    next(w for w in app.text_input if w.label==label).set_value('169')
    case.phase(app,'phase1');case.phase(app,'phase7')
    assert next(w for w in app.text_input if w.label==label).value=='169'
    assert compute_phase_status(app.session_state)['phase7']['level']!='ok'
    box=next(w for w in app.checkbox if w.label=='Surface — Ramp-Down - Main')
    box.uncheck().run();assert next(b for b in app.button if b.label=='Confirm measured lab results').disabled
    next(w for w in app.checkbox if w.label=='Surface — Ramp-Down - Main').check().run()
    assert next(w for w in app.text_input if w.label==label).value=='169'
    for key,name in RHEOLOGY_DATASETS[1:]:
        next(w for w in app.checkbox if w.label==name+' - Main').check().run()
    assert next(b for b in app.button if b.label=='Confirm measured lab results').disabled
    assert len(app.session_state['lab_qc_params']['Main']['rheology'])==3
    saved=audit.round_trip(app.session_state.to_dict());restored=case.app(saved);case.phase(restored,'phase7')
    assert next(w for w in restored.text_input if w.label==label).value=='169'
    # Both BHCT dataset conditions consume the same canonical QC temperature.
    next(w for w in restored.number_input if w.label=='BHCT (°F) - Main').set_value(155).run()
    assert sum('155 °F (BHCT)' in c.value for c in restored.caption)==2
    for field in ('gel_10_sec','gel_10_min'):
        changed=deepcopy(saved['lab_qc_params']['Main']);changed['rheology']['surface_down'][field]='9'
        assert lab_review_signature(changed,saved['lab_grid_dfs']['Main']) != lab_review_signature(saved['lab_qc_params']['Main'],saved['lab_grid_dfs']['Main'])


@pytest.mark.parametrize('defect',['none','partial','gel','fallback'])
def test_invalid_rheology_blocks_confirm_status_and_export(defect):
    case=audit.AuditRegressions();app=case.configured_app('CSG 9 5/8"')
    p=audit.round_trip(app.session_state.to_dict());qc=p['lab_qc_params']['Main']
    row=dataset()
    if defect=='partial':row['readings']={'300':'168'}
    if defect=='gel':row['gel_10_sec']=''
    if defect=='fallback':row=dataset(CORPUS[-1][1])
    qc['rheology']={} if defect=='none' else {'surface_down':row}
    qc['review_signature']=lab_review_signature(qc,p['lab_grid_dfs']['Main'])
    app=case.app(p);case.phase(app,'phase7')
    assert next(b for b in app.button if b.label=='Confirm measured lab results').disabled
    assert compute_phase_status(app.session_state)['phase7']['level']!='ok'
    assert prepare_calculations(deepcopy(app.session_state.to_dict()))
    case.assert_export_blocked(app)
    if defect=='fallback':
        case.phase(app,'phase7')
        assert any('Ty must be greater than 0' in w.value for w in app.warning)
        assert any(m.label=='Ty (lbf/100ft²)' and m.value=='0.00' for m in app.metric)


@pytest.mark.parametrize('bad',[
    [], {'unknown':{}}, {'surface_down':{'mode':'both'}},
    {'surface_down':{'selected':'yes'}}, {'surface_down':{'selected':1}},
    {'surface_down':{'readings':{'600':10}}}, {'surface_down':{'readings':{'300':True}}},
    {'surface_down':{'readings':{'300':'NaN'}}}, {'surface_down':{'readings':{'300':'Infinity'}}},
    {'surface_down':{'gel_10_sec':'1+2'}}, {'surface_down':{'gel_10_min':False}},
])
def test_active_and_inactive_trust_boundary_rejects_malformed(bad):
    for inactive in (False,True):
        project={'job_type':'CSG 9 5/8"'}
        if inactive:project['inactive_slurry_drafts']={'Tail':{'lab_qc_params':{'rheology':bad}}}
        else:project['lab_qc_params']={'Main':{'rheology':bad}}
        with pytest.raises(ValueError):audit.round_trip(project)


def test_multi_slurry_archive_reactivation_and_word():
    case=audit.AuditRegressions();app=case.configured_app('CSG 9 5/8"')
    p=audit.round_trip(app.session_state.to_dict())
    p['fluids_config']['active']=['Lead','Tail','Displacement Fluid']
    for s,top,i in [('Lead',0,0),('Tail',2850,1)]:
        p['fluids_config']['params'][s]=deepcopy(p['fluids_config']['params']['Main'])
        for key in ('cement_params','cement_additives_dfs','lab_qc_params','lab_grid_dfs'):
            p[key][s]=deepcopy(p[key]['Main'])
        p['cement_params'][s].update(top_mode='Surface' if s=='Lead' else 'Depth (m MD)',top_depth=top)
        p['lab_qc_params'][s]['rheology']={KEYS[i]:dataset(CORPUS[i][1])}
        p['lab_qc_params'][s]['review_signature']=lab_review_signature(p['lab_qc_params'][s],p['lab_grid_dfs'][s])
        p['lab_source_signatures'][s]=lab_source_signature(p,s)
    refresh_fluids(p)
    for slurry in ('Lead','Tail'):
        p['lab_source_signatures'][slurry]=lab_source_signature(p,slurry)
    p['_rheo_tail_surface_down_selected_default_project']=False
    purge_inactive_slurries(p,['Lead','Displacement Fluid'])
    assert not any(k.startswith('_rheo_tail_') for k in p)
    restored=audit.round_trip(p);purge_inactive_slurries(restored,['Lead','Tail','Displacement Fluid'])
    assert restored['lab_qc_params']['Tail']['rheology']==p['inactive_slurry_drafts']['Tail']['lab_qc_params']['rheology']
    assert 'Tail' not in restored['lab_source_signatures']
    restored['fluids_config']['active']=['Lead','Tail','Displacement Fluid']
    app=case.app(restored);case.phase(app,'phase7')
    # Existing reactivation requires explicit keep/sync + confirm.
    next(b for b in app.button if b.label=='Keep reviewed lab entries').click().run()
    for _ in range(2):
        next(b for b in app.button if b.label=='Confirm measured lab results' and not b.disabled).click().run()
    case.export(app)
    doc=Document(BytesIO(app.session_state['_compiled_doc_bytes']))
    tables=[t for t in doc.tables if t.rows[0].cells[0].text=='Rheology Test']
    assert len(tables)==2
    assert tables[0].rows[10].cells[1].text=='159.218'
    assert tables[1].rows[10].cells[2].text=='163.158'
    Path('/tmp/r5-multi.docx').write_bytes(app.session_state['_compiled_doc_bytes'])


def test_template_only_authorized_rheology_text_changes():
    original=BytesIO(subprocess.check_output(['git','show','9205d84:master_template.docx']))
    completed=BytesIO(subprocess.check_output(['git','show','ecb0a6e:master_template.docx']))
    # Round 5-2 authorizes other Lab cells; retain the complete 5-1 scope proof.
    with zipfile.ZipFile(original) as old,zipfile.ZipFile(completed) as new:
        assert old.namelist()==new.namelist()
        for name in old.namelist():
            if name!='word/document.xml':assert old.read(name)==new.read(name)
        a,b=old.read('word/document.xml'),new.read('word/document.xml')
        tb=lambda x:next(m for m in re.finditer(rb'<w:tbl>.*?</w:tbl>',x,re.S) if b'Rheology Test' in m.group())
        before,after=tb(a),tb(b)
        assert a[:before.start()]==b[:after.start()] and a[before.end():]==b[after.end():]
        clean=after.group().replace(b'<w:t>{{ s.lab.bhct }} degF</w:t>',b'<w:t>degF</w:t>',1)
        clean=re.sub(rb'{{ s.lab.rheology\.get.*? }}',b'-',clean)
        assert clean==before.group()
        with zipfile.ZipFile('master_template.docx') as current:
            assert tb(current.read('word/document.xml')).group()==after.group()


def test_hidden_invalid_draft_excluded_from_runtime_validity_but_not_json_trust():
    qc={'rheology': {'surface_down':dataset(), 'bhct_up':{
        'selected':False,'readings':{'300':'unfinished'},'gel_10_sec':'unfinished'}}}
    result,issues=build_rheology(qc)
    assert not issues and result['surface_down']['is_valid_bingham']
    assert not result['bhct_up']['selected']
    with pytest.raises(ValueError):validate_rheology_inputs(qc['rheology'])


def test_filter_preserves_original_order_and_fresh_fit():
    rows=list(zip(RPM_ORDER,CORPUS[0][1]))
    observed=[]
    import rheology
    original=rheology.f32
    def record(value):
        observed.append(value)
        return original(value)
    with patch.object(rheology,'f32',record):
        a=calculate_cemcade_bingham(rows)
    # The input float32 stores are made in the incoming grid order.
    assert [observed[6*i] for i in range(7)] == list(RPM_ORDER)
    assert calculate_cemcade_bingham(rows) == a


def test_review_fingerprints_each_real_rheology_source_field():
    qc={'rheology':{'surface_down':dataset()}}
    result,_=build_rheology(qc)
    assert f"{result['surface_down']['pv_cp']:.3f}"=='159.218'
    # Review fingerprint covers each real source field, not saved payload results.
    import pandas as pd
    grid=pd.DataFrame()
    baseline=lab_review_signature(qc,grid)
    for path,value in [('selected',False),('readings',{'300':'169'}),('gel_10_sec','9'),('gel_10_min','13')]:
        changed=deepcopy(qc);changed['rheology']['surface_down'][path]=value
        assert lab_review_signature(changed,grid)!=baseline


def test_malformed_live_gel_is_controlled_and_blocks_confirmation():
    case=audit.AuditRegressions();app=case.configured_app('CSG 9 5/8"')
    next(w for w in app.text_input if w.label=='10 Sec Gel - Surface — Ramp-Down - Main').set_value('abc').run()
    assert not app.exception
    assert any('finite numeric' in w.value for w in app.warning)
    assert next(b for b in app.button if b.label=='Confirm measured lab results').disabled


@pytest.mark.parametrize('selected',SELECTIONS)
def test_checkbox_combinations_through_real_callbacks(selected):
    case=audit.AuditRegressions();app=case.configured_app('CSG 9 5/8"')
    p=audit.round_trip(app.session_state.to_dict())
    p['lab_qc_params']['Main']['rheology']={key:dict(dataset(CORPUS[i][1]),selected=False)
                                          for i,key in enumerate(KEYS)}
    app=case.app(p);case.phase(app,'phase7')
    assert next(b for b in app.button if b.label=='Confirm measured lab results').disabled
    for key,label in RHEOLOGY_DATASETS:
        if key in selected:
            next(w for w in app.checkbox if w.label==label+' - Main').check().run()
    assert {key for key,row in app.session_state['lab_qc_params']['Main']['rheology'].items()
            if row['selected']} == set(selected)
    assert len([m for m in app.metric if m.label=='PV (cP)'])==len(selected)
    for i,key in enumerate(KEYS):
        if key in selected:assert any(m.value==CORPUS[i][2] for m in app.metric if m.label=='PV (cP)')
    assert not next(b for b in app.button if b.label=='Confirm measured lab results').disabled
    next(b for b in app.button if b.label=='Confirm measured lab results').click().run()
    assert compute_phase_status(app.session_state)['phase7']['level']=='ok'
