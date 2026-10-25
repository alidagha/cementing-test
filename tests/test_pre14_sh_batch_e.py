"""PRE14-SH E: preserve numeric sources and fail closed at their consumers."""
import ast
from copy import deepcopy
from io import BytesIO
import subprocess
from unittest.mock import patch

from docx import Document
import pandas as pd
import pytest

import engineering_tools as eng
import phase_10_procedure as report
import project_state as ps
import test_audit_regressions as audit

CASE = audit.AuditRegressions()
DOC_KEYS = ('_compiled_doc_bytes', '_compiled_doc_filename', '_compiled_doc_signature')
NUMERIC_HARDWARE = ('ID (in)', 'Joint (m)', 'Weight (ppf)', 'Collapse (psi)', 'Burst (psi)')


@pytest.mark.parametrize('value', ['1/2', '۱/۲', '1 / 2', '80/84'])
def test_l12_slash_is_not_a_scalar_decimal(value):
    with pytest.raises(ValueError):
        eng.clean_number(value)


@pytest.mark.parametrize('raw,expected', [('1.2', 1.2), ('1,2', 1.2), ('1,234.5', 1234.5),
                                        ('۱۲٫۵', 12.5), ('١٢.٥', 12.5), (0, 0.), (None, 0.)])
def test_l12_existing_scalar_formats_and_separate_slash_parsers(raw, expected):
    assert eng.clean_number(raw) == expected
    assert eng.require_positive_density('80/84') == 82.
    assert eng.parse_fractional_size('9 5/8') == 9.625
    assert eng.parse_fractional_size('1/2') == .5


@pytest.mark.parametrize('kwargs', [
    {'powders': [{'percent': 100., 'density_pcf': 80.}]},
    {'cmt_sg': 1., 'slurry_weight_pcf': 150.},
    # Direct engine input isolates nonpositive yield with positive water;
    # this adds no dosage rule to the established UI validators.
    {'powders': [{'percent': -200., 'density_pcf': 128.}]},
])
def test_m9_invalid_water_or_yield_cannot_return_field_or_lab_quantities(kwargs):
    inputs = {'slurry_weight_pcf': 118., 'slurry_volume_bbl': 50., **kwargs}
    with pytest.raises(ValueError, match='mix water|yield'):
        eng.calculate_slurry_from_components(**inputs)


@pytest.mark.parametrize('density,volume,salt', [(104., 500., 0.), (118., 50., 18.),
                                               (152., 731., 0.), (180., 0., 0.)])
def test_m9_valid_complete_results_are_exactly_equal_to_approved_baseline(density, volume, salt):
    tree = ast.parse(subprocess.check_output(['git', 'show',
        'f7d9009705ec8983b6ff90f0fff9fc261e03d0b4:engineering_tools.py']))
    original = next(n for n in tree.body if isinstance(n, ast.FunctionDef)
                    and n.name == 'calculate_slurry_from_components')
    namespace = vars(eng).copy()
    exec(compile(ast.Module(body=[original], type_ignores=[]), '<approved mass balance>', 'exec'), namespace)
    inputs = dict(slurry_weight_pcf=density, slurry_volume_bbl=volume, salt_pct=salt)
    assert eng.calculate_slurry_from_components(**inputs) == namespace['calculate_slurry_from_components'](**inputs)


def test_m9_valid_zero_water_remains_zero_with_positive_exact_yield():
    result = eng.calculate_slurry_from_components(124.8, 50., cmt_sg=2.)
    assert result['water_vol_per_sack'] == result['field_water_bbl'] == result['lab_water_gr'] == 0.
    assert result['yield_ft3_per_sk'] == 110. / 124.8
    assert result['field_sacks'] > 0 and result['lab_cmt_gr'] > 0


@pytest.mark.parametrize('field', NUMERIC_HARDWARE)
@pytest.mark.parametrize('source', ['hardware_table', 'hardware_editor_draft'])
def test_m3_malformed_nonblank_hardware_is_retained_for_repair_and_not_geometry(field, source):
    project = audit.valid_export_chart_fixture('CSG 9 5/8"')
    project['job_type'] = 'CSG 9 5/8"'
    project['fluids_config'] = {'active': ['Main'], 'params': {
        'Main': {'volume': 50., 'density': '118', 'pump_rate': '4'}}}
    original = project['hardware_table'].copy(deep=True)
    row = original.index[original['Description'].eq('Casing')][0]
    project[source] = original.astype({field: object})
    project[source].at[row, field] = 'abc'
    app = CASE.app(audit.round_trip(project))
    CASE.phase(app, 'phase2_3')
    assert any(f'Row {row+1}' in w.value and field in w.value for w in app.warning)
    assert app.session_state['hardware_editor_draft'].at[row, field] == 'abc'
    assert ps.hardware_draft_pending(app.session_state)
    assert not app.session_state['hardware_table']['Description'].eq('Casing').any()
    state = audit.round_trip(app.session_state.to_dict())
    assert state['hardware_editor_draft'].at[row, field] == 'abc'
    assert any('pending hardware' in issue for issue in ps.prepare_calculations(state))
    CASE.phase(app, 'phase1')
    restored = CASE.app(audit.round_trip(app.session_state.to_dict()))
    CASE.phase(restored, 'phase2_3')
    assert restored.session_state['hardware_editor_draft'].at[row, field] == 'abc'


def test_m3_editor_repair_commits_before_navigation_then_restores_exact_geometry():
    project = audit.valid_export_chart_fixture('CSG 9 5/8"')
    project['job_type'] = 'CSG 9 5/8"'
    table = project['hardware_table'].astype({'ID (in)': object})
    project['hardware_table'] = table
    row = table.index[table['Description'].eq('Casing')][0]
    table.at[row, 'ID (in)'] = 'bad ID'
    app = CASE.app(project)
    CASE.phase(app, 'phase2_3')
    assert ps.hardware_draft_pending(app.session_state)
    CASE.hardware_event(app, {'edited_rows': {str(row): {'ID (in)': '8.535'}}}, navigate='phase1')
    assert app.session_state['hardware_editor_draft'].at[row, 'ID (in)'] == '8.535'
    CASE.phase(app, 'phase2_3')
    assert not ps.hardware_draft_pending(app.session_state)
    restored = audit.round_trip(app.session_state.to_dict())
    casing = restored['hardware_table'].loc[lambda df: df['Description'].eq('Casing')].iloc[0]
    assert casing['ID (in)'] == 8.535


@pytest.mark.parametrize('raw', [None, '', 0., '0', '8.535'])
def test_m3_blank_zero_and_numeric_text_keep_existing_meanings(raw):
    project = audit.valid_export_chart_fixture('CSG 9 5/8"')
    project['job_type'] = 'CSG 9 5/8"'
    table = project['hardware_table'].astype({'ID (in)': object})
    project['hardware_table'] = table
    row = table.index[table['Description'].eq('Casing')][0]
    table.at[row, 'ID (in)'] = raw
    app = CASE.app(project)
    CASE.phase(app, 'phase2_3')
    casing = app.session_state['hardware_table'].loc[lambda df: df['Description'].eq('Casing')].iloc[0]
    assert casing['ID (in)'] == (0. if raw in (None, '', 0., '0') else 8.535)
    assert not ps.hardware_draft_pending(app.session_state)


@pytest.mark.parametrize('raw', ['abc', '1e999', float('inf'), True, False])
def test_m3_nonblank_invalid_numbers_have_no_zero_sentinel(raw):
    from phase_2_3_well_data import hardware_number
    assert hardware_number(raw) is None


def test_m3_nullable_hardware_cells_keep_blank_defaults():
    from phase_2_3_well_data import hardware_number
    assert hardware_number(pd.NA) == hardware_number(float('nan')) == 0.


@pytest.fixture(scope='module')
def reviewed_main():
    app = CASE.configured_app('CMT PLUG')
    state = deepcopy(app.session_state.to_dict())
    assert ps.prepare_calculations(state) == []
    with patch.object(report.st, 'session_state', deepcopy(state)):
        assert report.build_master_context(calculations_prepared=True)['slurries']
    return state


@pytest.mark.parametrize('defect', ['record', 'key', None, 'None', 'abc'])
def test_l8_prepared_missing_or_unreadable_volume_is_f3_not_later_error(reviewed_main, defect):
    state = deepcopy(reviewed_main)
    if defect == 'record':
        state['fluid_data'].pop('Main')
    elif defect == 'key':
        state['fluid_data']['Main'].pop('volume')
    else:
        state['fluid_data']['Main']['volume'] = defect
    state.update(zip(DOC_KEYS, (b'old document', 'old.docx', 'old signature')))
    with patch.object(report.st, 'session_state', state):
        with pytest.raises(ValueError, match='Unreadable numeric volume.*Main.*volume'):
            report.build_master_context(calculations_prepared=True)
    assert not any(k in state for k in DOC_KEYS)


@pytest.mark.parametrize('prepared', [False, True])
@pytest.mark.parametrize('defect', ['record', 'key', None, 'None', 'abc'])
def test_l8_invalid_authoritative_source_rejects_without_stale_derived_fallback(reviewed_main, prepared, defect):
    state = deepcopy(reviewed_main)
    if defect == 'record':
        state['fluids_config']['params'].pop('Main')
        state['fluid_data'].pop('Main')
    elif defect == 'key':
        state['fluids_config']['params']['Main'].pop('volume')
    else:
        state['fluids_config']['params']['Main']['volume'] = defect
    original = deepcopy(state['fluids_config'])
    state.update(zip(DOC_KEYS, (b'old document', 'old.docx', 'old signature')))
    with patch.object(report.st, 'session_state', state):
        with pytest.raises(ValueError, match='(?:Phase IV.*Main.*volume|Unreadable numeric volume.*Main.*volume)'):
            report.build_master_context(calculations_prepared=prepared)
    assert state['fluids_config'] == original
    assert not any(k in state for k in DOC_KEYS)


@pytest.mark.parametrize('prepared', [False, True])
def test_l8_valid_numeric_text_and_legacy_fallback_recover(reviewed_main, prepared):
    state = deepcopy(reviewed_main)
    state['fluids_config']['params'].pop('Main')
    state['fluid_data']['Main']['volume'] = '50.0'
    with patch.object(report.st, 'session_state', state):
        context = report.build_master_context(calculations_prepared=prepared)
    assert context['slurries'][0]['volume_bbl'] == 50.
    assert float(context['fluids_train'][0]['volume']) == 50.
    if prepared:
        assert state['fluid_data']['Main']['volume'] == '50.0'


def test_l8_invalid_then_repaired_volume_can_build_a_real_current_docx(reviewed_main):
    app = CASE.app(deepcopy(reviewed_main))
    CASE.phase(app, 'phase4')
    app.number_input(key='vol_Main').set_value(0.).run()
    CASE.phase(app, 'phase10')
    assert next(b for b in app.button if b.label == 'Build Word Document').disabled
    assert not any(k in app.session_state for k in DOC_KEYS)
    CASE.phase(app, 'phase4')
    app.number_input(key='vol_Main').set_value(50.).run()
    CASE.export(app)
    doc = Document(BytesIO(app.session_state['_compiled_doc_bytes']))
    assert doc.tables and any('Main' in cell.text for t in doc.tables for r in t.rows for cell in r.cells)


@pytest.mark.parametrize('fluid', ['Pre Flush', 'Spacer Ahead', 'Spacer Behind', 'Diesel', 'Displacement Fluid'])
def test_l8_each_consumed_operational_fluid_names_missing_volume_before_later_gates(reviewed_main, fluid):
    state = deepcopy(reviewed_main)
    if fluid not in state['fluids_config']['active']:
        state['fluids_config']['active'].append(fluid)
    record = {'volume': None, 'density': '80', 'pump_rate': '4'}
    state['fluids_config']['params'][fluid] = deepcopy(record)
    state['fluid_data'][fluid] = deepcopy(record)
    state.update(zip(DOC_KEYS, (b'old', 'old.docx', 'old signature')))
    with patch.object(report.st, 'session_state', state):
        with pytest.raises(ValueError, match=rf'Unreadable numeric volume.*{fluid}.*volume'):
            report.build_master_context(calculations_prepared=True)
    assert not any(k in state for k in DOC_KEYS)


def test_l8_normal_preparation_rebuilds_missing_derived_volume_from_valid_canonical_source(reviewed_main):
    state = deepcopy(reviewed_main)
    canonical = deepcopy(state['fluids_config']['params'])
    state['fluid_data']['Main'].pop('volume')
    state['fluid_data']['Dormant fluid'] = {'volume': None}
    with patch.object(report.st, 'session_state', state):
        context = report.build_master_context(calculations_prepared=False)
    assert context['slurries'][0]['volume_bbl'] == 50.
    assert state['fluids_config']['params'] == canonical


def test_m3_open_hole_normalization_and_other_valid_rows_survive_restore():
    project = audit.valid_export_chart_fixture('CSG 9 5/8"')
    project['job_type'] = 'CSG 9 5/8"'
    table = project['hardware_table']
    hole = table.index[table['Description'].eq('Open Hole Size')][0]
    for field, value in {'Joint (m)': 12.2, 'Weight (ppf)': 50., 'Grade': 'L-80',
                         'Collapse (psi)': 100., 'Burst (psi)': 200.}.items():
        table.at[hole, field] = value
    original = table.loc[~table['Description'].eq('Open Hole Size')].copy(deep=True)
    app = CASE.app(audit.round_trip(project))
    CASE.phase(app, 'phase2_3')
    restored = audit.round_trip(app.session_state.to_dict())['hardware_table']
    result = restored.loc[restored['Description'].eq('Open Hole Size')].iloc[0]
    assert all(result[field] == 0. for field in NUMERIC_HARDWARE if field != 'ID (in)')
    assert result['Grade'] == '-'
    assert result['ID (in)'] == table.at[hole, 'ID (in)']
    pd.testing.assert_frame_equal(restored.loc[~restored['Description'].eq('Open Hole Size')],
                                  original, check_dtype=False)
