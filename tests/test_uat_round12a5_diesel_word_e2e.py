"""A-5: Diesel uses existing Procedure, Word, persistence and review authorities."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import hashlib
import re
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn
from docxtpl import DocxTemplate
import pytest

import materials_db
import phase_4_fluids as phase4
import phase_10_procedure as report
import project_state as ps
import pumping_simulation as pumping
import phase_9_centralizer as phase9
import test_audit_regressions as audit
from centralizer_fixtures import execute_centralizer
from standalone_spacer import PENDING_INTEGRATION


JOBS = ('CSG 9 5/8"', 'LNR 7"', 'TIE BACK LNR 7"', 'CMT PLUG', 'CMT SQUEEZE')
TEMPLATE_SHA = '0c3d911e471aae22c0b18fecc53ed99caa7e7479d94b22640bc19439da92e331'
DIESEL = {'volume': 20., 'density': '52.5', 'pump_rate': '4.0', 'material_name': 'Diesel'}
CASE = audit.AuditRegressions()


@pytest.fixture(scope='module')
def projects():
    cache = {}
    def get(job=JOBS[0], active=True, mud='OBM', include_standalone_spacer=False):
        key = (job, include_standalone_spacer)
        expected_issues = [PENDING_INTEGRATION] if include_standalone_spacer else []
        if key not in cache:
            app = CASE.configured_app(job, optional_fluids=True,
                                      include_standalone_spacer=include_standalone_spacer)
            state = audit.round_trip(app.session_state.to_dict())
            state['mud_type'] = 'OBM'
            state['fluids_config']['active'].insert(0, 'Diesel')
            state['fluids_config']['params']['Diesel'] = deepcopy(DIESEL)
            state['lab_qc_params']['Diesel'] = {'rheology': audit.valid_export_rheology_fixture()}
            assert ps.prepare_calculations(state) == expected_issues
            execute_centralizer(state, physics='dynamic')
            cache[key] = state
        state = deepcopy(cache[key])
        state['mud_type'] = mud
        if not active:
            state['fluids_config']['active'].remove('Diesel')
        ps.reconcile_fluid_applicability(state)
        assert ps.prepare_calculations(state) == expected_issues
        if not phase9.result_is_current(state):
            execute_centralizer(state, physics='dynamic')
        return state
    return get


def sync(state):
    with patch.object(report.st, 'session_state', state):
        return report.synchronize_report_texts()


def context(state):
    with patch.object(report.st, 'session_state', state):
        return report.build_master_context()


def word(ctx, postprocess=False):
    template = DocxTemplate(report.TEMPLATE_PATH)
    template.render(report._word_quantity_context(ctx))
    stream = BytesIO()
    template.save(stream)
    data = stream.getvalue()
    if postprocess:
        data = report.postprocess_well_report(data, ctx['well_profile_report'], ctx['note_vertical'],
                                             ctx['cement_chart_report'], ctx['centralizer_report'])
    return Document(BytesIO(data))


def sequence_table(doc):
    body = list(doc.element.body)
    start = next(i for i, node in enumerate(body) if node.tag == qn('w:p')
                 and ''.join(node.itertext()) and ''.join(t.text or '' for t in node.iter(qn('w:t'))).strip() == 'IV. Fluids Sequence')
    end = next(i for i, node in enumerate(body[start + 1:], start + 1)
               if node.tag == qn('w:p') and ''.join(t.text or '' for t in node.iter(qn('w:t'))).strip() == 'V. Cementing Program')
    tables = [table for table in doc.tables if table._tbl in body[start + 1:end]]
    return next(table for table in tables if any('Main' == cell.text.strip()
                for row in table.rows for cell in row.cells))


@pytest.mark.parametrize('job', JOBS)
def test_exact_diesel_procedure_pumping_step(projects, job):
    state = projects(job)
    text = sync(state)['procedure_text']['generated']
    assert text.count('Pump 20.0 bbl Diesel as pre flush.(52.5 pcf)') == 1


@pytest.mark.parametrize('field,value', [('volume', 24.), ('density', '53.0'), ('pump_rate', '2.0')])
def test_diesel_changes_do_not_drift_manual_executive_summary(projects, field, value):
    state = projects()
    before = sync(state)['exec_summary_text']
    state['exec_summary_text'] += '\nOwner-reviewed manual summary.'
    manual = state['exec_summary_text']
    state['fluids_config']['params']['Diesel'][field] = value
    ps.refresh_fluids(state)
    after = sync(state)['exec_summary_text']
    assert after['generated'] == before['generated']
    assert 'Diesel' not in after['generated']
    assert after['signature'] == before['signature']
    assert not after['pending'] and state['exec_summary_text'] == manual


def test_protected_template_generic_sequence_contract_already_accepts_diesel(projects):
    assert hashlib.sha256(Path(report.TEMPLATE_PATH).read_bytes()).hexdigest() == TEMPLATE_SHA
    state = projects()
    ctx = context(state)
    rows = [row for row in ctx['fluids_train'] if row['name'] == 'Diesel']
    assert len(rows) == 1
    row = rows[0]
    assert row['material_name'] == 'Diesel' and row['volume'] == 20. and row['density'] == '52.5'
    assert row['pump_rate'] == '4.0' and row['duration_min'] == 5. and row['cumul_time_min'] == 5.
    doc = word(ctx)
    table = sequence_table(doc)
    diesel = [r for r in table.rows if any(c.text.strip() == 'Diesel' for c in r.cells)]
    assert len(diesel) == 1
    print('Diesel Word row:', [cell.text for cell in diesel[0].cells])
    assert hashlib.sha256(Path(report.TEMPLATE_PATH).read_bytes()).hexdigest() == TEMPLATE_SHA


@pytest.mark.parametrize('source', ['canonical', 'legacy'])
def test_malformed_active_diesel_volume_is_already_controlled_before_word(projects, source):
    state = projects()
    state.update(_compiled_doc_bytes=b'old', _compiled_doc_filename='old.docx', _compiled_doc_signature='old')
    if source == 'canonical':
        state['fluids_config']['params']['Diesel']['volume'] = 'malformed saved Diesel volume'
    else:
        state['fluids_config']['params'].pop('Diesel')
        state['fluid_data']['Diesel']['volume'] = 'malformed saved Diesel volume'
    with pytest.raises(ValueError, match='Phase IV: Diesel: unreadable fluid volume'):
        context(state)
    assert not any(key in state for key in ('_compiled_doc_bytes', '_compiled_doc_filename', '_compiled_doc_signature'))


def numbered_steps(text):
    return [(int(number), body) for number, body in re.findall(r'(?m)^(\d+)\. (.+)$', text)]


def assert_diesel_word_scope(doc, active):
    table = sequence_table(doc)
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    diesel_rows = [row for row in rows if 'Diesel' in row]
    assert len(diesel_rows) == int(active)
    if active:
        assert diesel_rows == [['Diesel', 'Diesel', '52.5', '20.0', '4.0', '00:05', '00:05']]
    # Inspect every body table/paragraph. Only the existing sequence table and
    # exactly one numbered Procedure paragraph may carry the report concept.
    procedure = []
    for node in doc.element.body:
        text = ''.join(t.text or '' for t in node.iter(qn('w:t')))
        if 'Diesel' not in text:
            continue
        if node.tag == qn('w:tbl'):
            assert node == table._tbl
        else:
            assert re.fullmatch(r'\d+\. Pump 20\.0 bbl Diesel as pre flush\.\(52\.5 pcf\)', text.strip())
            procedure.append(text.strip())
    assert len(procedure) == int(active)
    if active:
        names = [row[0] for row in rows if row and row[0] in materials_db.HYDRAULIC_EXECUTION_ORDER]
        assert names == sorted(names, key=materials_db.HYDRAULIC_EXECUTION_ORDER.index)
        assert names.count('Diesel') == 1


@pytest.mark.parametrize('job', JOBS)
@pytest.mark.parametrize('preflush', [True, False])
def test_existing_step_counter_order_and_job_archetype_are_preserved(projects, job, preflush):
    state = projects(job)
    if not preflush:
        state['fluids_config']['active'].remove('Pre Flush')
        assert ps.prepare_calculations(state) == []
    active = sync(state)['procedure_text']['generated']
    without = deepcopy(state)
    without['fluids_config']['active'].remove('Diesel')
    ps.refresh_fluids(without)
    old = sync(without)['procedure_text']['generated']
    steps, old_steps = numbered_steps(active), numbered_steps(old)
    assert [number for number, _ in steps] == list(range(1, len(steps) + 1))
    assert len(steps) == len(old_steps) + 1
    diesel = next((number, body) for number, body in steps if 'Diesel' in body)
    assert diesel[1] == 'Pump 20.0 bbl Diesel as pre flush.(52.5 pcf)'
    if preflush:
        assert 'as preflush.' in steps[diesel[0]][1] or 'as pre flush.' in steps[diesel[0]][1]
        assert 'Diesel' not in steps[diesel[0]][1]
    else:
        assert all('as preflush.' not in body and 'as pre flush.' not in body
                   for _, body in steps if 'Diesel' not in body)
        assert 'Pre Flush' not in state['fluids_config']['active']
    old_pumps = [body for _, body in old_steps if body.startswith(('Pump ', 'Mix and pump '))]
    pumps = [body for _, body in steps if body.startswith(('Pump ', 'Mix and pump ')) and 'Diesel' not in body]
    assert pumps == old_pumps
    old_number = {body: number for number, body in old_steps}
    for number, body in steps:
        if body in old_number:
            assert number == old_number[body] + int(old_number[body] >= diesel[0])
    heading = ('IX.' if job == 'CMT PLUG' else 'VIII.' if job == 'CMT SQUEEZE' else 'X.') + ' Cementing Procedure'
    assert active.splitlines()[0] == old.splitlines()[0] == heading


@pytest.mark.parametrize('job', JOBS)
def test_notes_remain_unique_sequential_and_have_no_diesel_note(projects, job):
    state = projects(job)
    with_diesel = context(state)
    without_diesel = context(projects(job, active=False))
    notes = with_diesel['all_notes']
    numbers = [int(re.match(r'NOTE (\d+):', note).group(1)) for note in notes]
    assert numbers == list(range(1, len(notes) + 1))
    assert len(set(numbers)) == len(numbers)
    assert all('Diesel' not in note for note in notes)
    assert len(notes) == len(without_diesel['all_notes'])
    assert [note for note in notes if note != with_diesel['note_maxpump']] == [
        note for note in without_diesel['all_notes'] if note != without_diesel['note_maxpump']]
    assert with_diesel['note_maxpump'] != without_diesel['note_maxpump']
    for text in (with_diesel['procedure_text'], without_diesel['procedure_text']):
        assert not any('Diesel' in line for line in text.splitlines() if 'NOTE' in line.upper())
    rendered = word(with_diesel, postprocess=True)
    actual = [int(re.match(r'NOTE (\d+):', p.text.strip()).group(1)) for p in rendered.paragraphs
              if re.match(r'NOTE \d+:', p.text.strip())]
    assert actual == numbers


@pytest.mark.parametrize('job', JOBS)
def test_actual_docx_contains_diesel_only_in_existing_sequence_and_procedure(projects, job):
    state = projects(job)
    ctx = context(state)
    assert_diesel_word_scope(word(ctx), True)
    assert [row['name'] for row in ctx['fluids_train']] == state['fluids_config']['active']
    assert 'Diesel' not in ctx['exec_summary']
    assert 'Diesel' not in [row['name'] for row in ctx['slurries']]
    assert ctx['has_preflush'] and ctx['has_spacer']


@pytest.mark.parametrize('field,value', [('volume', 24.), ('density', '53.0'), ('pump_rate', '2.0')])
@pytest.mark.parametrize('job', [JOBS[0], 'CMT PLUG'])
def test_active_procedure_edits_preserve_manual_text_until_existing_review(projects, job, field, value):
    state = projects(job)
    before = sync(state)
    state['procedure_text'] += '\nOwner-reviewed operational instruction.'
    manual = state['procedure_text']
    state['fluids_config']['params']['Diesel'][field] = value
    ps.refresh_fluids(state)
    after = sync(state)
    assert after['procedure_text']['signature'] != before['procedure_text']['signature']
    assert after['procedure_text']['pending']
    assert state['procedure_text'] == manual
    if field in ('volume', 'density') or job == 'CMT PLUG':
        assert after['procedure_text']['generated'] != before['procedure_text']['generated']
    assert after['exec_summary_text']['signature'] == before['exec_summary_text']['signature']
    with pytest.raises(ValueError, match='Review the preserved manual report text'):
        context(state)
    spec = after['procedure_text']
    ps.accept_report_text(state, 'procedure_text', spec['generated'], spec['signature'])
    assert not sync(state)['procedure_text']['pending']
    assert state['procedure_text'] == manual
    ps.accept_report_text(state, 'procedure_text', spec['generated'], spec['signature'], replace=True)
    assert state['procedure_text'] == spec['generated']
    assert state['report_text_state']['procedure_text']['history'][-1] == manual


@pytest.mark.parametrize('field,value', [('volume', 24.), ('density', '53.0'), ('pump_rate', '2.0')])
def test_canonical_diesel_callback_invalidates_compiled_word_and_stales_consumers(projects, field, value):
    state = projects()
    before = context(state)
    pumping_source = pumping.source_signature(state)
    state.update(_compiled_doc_bytes=b'previous real-build cache marker', _compiled_doc_filename='old.docx',
                 _compiled_doc_signature=report.document_signature(before, str(report.TEMPLATE_PATH)))
    state['_a5_edit'] = value
    with patch.object(phase4.st, 'session_state', state):
        phase4._commit_fluid_widget('Diesel', field, '_a5_edit')
    assert not any(key in state for key in ('_compiled_doc_bytes', '_compiled_doc_filename', '_compiled_doc_signature'))
    assert state['fluids_config']['params']['Diesel'][field] == value
    assert pumping.source_signature(state) != pumping_source
    # Centralizer consumes Diesel volume/density, not its own pumping rate.
    assert phase9.result_is_current(state) == (field == 'pump_rate')


def test_inactive_wbm_drafts_leave_report_text_context_and_word_content_unchanged(projects):
    state = projects(active=False, mud='WBM')
    before_specs = sync(state)
    state['exec_summary_text'] += '\nReviewed summary.'
    state['procedure_text'] += '\nReviewed procedure.'
    before = context(state)
    before_word = word(before).element.body.xml
    state['fluids_config']['params']['Diesel'].update(volume='unreadable dormant draft', density='51-54', pump_rate='bad')
    state['lab_qc_params']['Diesel'] = {'rheology': {'bhct_down': {'selected': False, 'readings': {}}}}
    assert ps.prepare_calculations(state) == []
    after_specs = sync(state)
    for key in ('procedure_text', 'exec_summary_text'):
        assert after_specs[key]['signature'] == before_specs[key]['signature']
        assert after_specs[key]['generated'] == before_specs[key]['generated']
        assert not after_specs[key]['pending']
    after = context(state)
    after.pop('generated_timestamp'); before.pop('generated_timestamp')
    assert ps.fingerprint(after) == ps.fingerprint(before)
    assert word(context(state)).element.body.xml == before_word
    assert_diesel_word_scope(word(context(state)), False)


@pytest.mark.parametrize('job', JOBS[:3])
def test_real_build_and_active_obm_save_restore_keep_diesel_and_current_consumers(projects, job):
    from types import SimpleNamespace
    from test_uat_round10c_integration import save_project, restore_project, build
    state = projects(job)
    raw = save_project(SimpleNamespace(session_state=state))
    restored = restore_project(raw)
    saved = restored.session_state.to_dict()
    assert 'Diesel' in saved['fluids_config']['active']
    assert saved['fluids_config']['params']['Diesel'] == state['fluids_config']['params']['Diesel']
    assert saved['lab_qc_params']['Diesel'] == state['lab_qc_params']['Diesel']
    assert '_centralizer_result' not in saved
    execute_centralizer(restored.session_state, physics='dynamic')
    data = build(restored)
    current = restored.session_state.to_dict()
    assert phase9.result_is_current(current)
    assert current['_pumping_simulation_source'] == pumping.source_signature(current)
    assert current['fluid_data']['Diesel']['material_name'] == 'Diesel'
    assert_diesel_word_scope(Document(BytesIO(data)), True)
    assert hashlib.sha256(Path(report.TEMPLATE_PATH).read_bytes()).hexdigest() == TEMPLATE_SHA
    # Retire the actual compiled output on the first active operational edit.
    CASE.phase(restored, 'phase4')
    restored.number_input(key='vol_Diesel').set_value(24.).run()
    assert '_compiled_doc_bytes' not in restored.session_state
    assert restored.session_state['fluids_config']['params']['Diesel']['volume'] == 24.


def test_real_inactive_wbm_restore_obm_selection_and_export_do_not_reactivate_diesel(projects):
    from types import SimpleNamespace
    from test_uat_round10c_integration import save_project, restore_project, build
    state = projects(active=False, mud='WBM')
    raw = save_project(SimpleNamespace(session_state=state))
    app = restore_project(raw)
    assert 'Diesel' not in app.session_state['fluids_config']['active']
    for key in ('fluids_config', 'lab_qc_params'):
        source = app.session_state[key]['params'] if key == 'fluids_config' else app.session_state[key]
        expected = state[key]['params'] if key == 'fluids_config' else state[key]
        assert source['Diesel'] == expected['Diesel']
    CASE.phase(app, 'phase2_3')
    next(w for w in app.selectbox if w.label == 'Mud Type').set_value('OBM').run()
    CASE.phase(app, 'phase4')
    assert not app.checkbox(key='chk_fluid_Diesel').value
    assert app.session_state['fluids_config']['params']['Diesel'] == state['fluids_config']['params']['Diesel']
    assert app.session_state['lab_qc_params']['Diesel'] == state['lab_qc_params']['Diesel']
    execute_centralizer(app.session_state, physics='dynamic')
    data = build(app)
    assert_diesel_word_scope(Document(BytesIO(data)), False)


# Captured by loading phase_10_procedure.py from exact A-4-2 and comparing
# generated text, sources, complete context and actual rendered body XML.
NO_DIESEL_A42 = {
    'CSG 9 5/8"': ('a5134954af6038660d25feb27bd7b19ecce4844658d9ec1d5c268353ec7e79ba',
                   '2040622d6a65ed778dd147ae7552d014b016711a095604c0a643764ed3d243ce'),
    'LNR 7"': ('8db5ea1d997b98f9b151226485dcea9bed2cf2849bce58557db81d5a00e3ec66',
               'bc60c6ee2adf46ba38333934ccdb25d4107ed52fd24c01329955a5ec8288baff'),
    'TIE BACK LNR 7"': ('7b28fa5f211df4f078be57bc8ec9e6bcf3a3c709b70c87c7dd0132c7365ff289',
                        '1fa15bb27a81295be7729a963b087e4e47defdc25e545e7d0d61fbf50833b1da'),
    'CMT PLUG': ('2d3b3558e5bdb118ea7e4d0dc97d8e09f37c122509d47506ea174c6718c7a805',
                 '3b17a454c16d9fdd07c1f1b4fc4a35b94900730b9cb4b65e528cf53d06175d34'),
    'CMT SQUEEZE': ('1b728a841796aed381a0b3f488924dedc6e8f4f65d7fe51b287711ea489ddaf4',
                    '7148fb2eb84ba260748c3eeb83948e939deac7b18858aadf2eeafdef62ab6556'),
}


@pytest.mark.parametrize('job', JOBS)
@pytest.mark.parametrize('mud', ['WBM', 'OBM'])
def test_no_diesel_report_signatures_match_exact_a42(projects, job, mud):
    # Preserve the original four-stage source program and exact A-4-2 text
    # snapshots. This is a source-projection test, not a Standalone Word build.
    state = projects(job, active=False, mud=mud, include_standalone_spacer=True)
    specs = sync(state)
    assert (specs['exec_summary_text']['signature'], specs['procedure_text']['signature']) == NO_DIESEL_A42[job]
    assert 'Diesel' not in specs['exec_summary_text']['generated']
    assert 'Diesel' not in specs['procedure_text']['generated']
    with pytest.raises(ValueError, match='Standalone Spacer.*pending'):
        context(state)


def test_excluded_operational_fluids_cannot_leak_cumulative_timing_into_summary(projects):
    state = projects()
    state['fluids_config']['active'].append('Scavenger')
    state['fluids_config']['params']['Scavenger'] = {'volume': 12., 'density': '118', 'pump_rate': '4'}
    ps.refresh_fluids(state)
    before = sync(state)['exec_summary_text']
    for name in ('Diesel', 'Scavenger'):
        state['fluids_config']['params'][name]['volume'] += 4.
    ps.refresh_fluids(state)
    after = sync(state)['exec_summary_text']
    assert after == before
    without = deepcopy(state)
    without['fluids_config']['active'] = [name for name in without['fluids_config']['active']
                                          if name not in ('Diesel', 'Scavenger')]
    ps.refresh_fluids(without)
    assert sync(without)['exec_summary_text'] == after
