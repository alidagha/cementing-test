"""Round 3 Batch 2: shared Pre-flush names and ordinary Spacer choices."""
from copy import deepcopy
from io import BytesIO
import unittest
import json

import pandas as pd
from docx import Document
import materials_db
import phase_6_spacer as spacer
import test_audit_regressions as audit
from engineering_tools import compute_phase_status
from project_state import prepare_calculations, refresh_fluids


NAME_OPTIONS = ["Fresh Water", "Salt Saturated Water", "Fresh Water Mix with Chemical Wash",
                "Salt Saturated Water Mix with Chemical Wash", "Custom"]
CHEMICALS = ["NaCl", "Spacer", "Surfactant", "Anti Foam", "Weighting Agent", "Mud", "Magneset Thinner"]
FORMULATIONS = ["Fresh Water Only", "Water + Chemical Wash", "Brine (Water + NaCl)",
                "Combined (Water + NaCl + Wash)"]


class Round3FluidsUAT(unittest.TestCase):
    app = audit.AuditRegressions.app
    phase = audit.AuditRegressions.phase
    healthy = audit.AuditRegressions.healthy

    def project(self, name="Salt Saturated Water", extra_spacers=False):
        active = ["Pre Flush"] + (["Spacer", "Spacer Ahead", "Spacer Behind"] if extra_spacers else []) + ["Displacement Fluid"]
        return {"mud_density": "80-82", "effective_mud_density": 81.,
                "fluids_config": {"active": active, "params": {s: {
                    "volume": 25., "density": "85-87", "pump_rate": "3-5",
                    "material_name": name if s == "Pre Flush" else materials_db.DEFAULT_MATERIAL_NAMES[s]}
                    for s in active}},
                "preflush_config": {"type": FORMULATIONS[-1], "nacl_multiplier": 126., "wash_multiplier": 3.}}

    def selector(self, app):
        return next(w for w in app.selectbox if w.label == "Material Name - Pre Flush")

    def custom(self, app):
        return next(w for w in app.text_input if w.label == "Custom Material Name - Pre Flush")

    def name(self, app):
        return app.session_state["fluids_config"]["params"]["Pre Flush"]["material_name"]

    def test_existing_formulation_characterization(self):
        # Fixed baseline engineering outputs; naming must not alter them.
        info = {"volume": 25., "density": "85-87", "effective_density": 86.}
        app = self.app(self.project()); self.phase(app, "phase4"); self.phase(app, "phase6")
        for kind, salt, wash, salt_ratio, wash_ratio in (
            (FORMULATIONS[0], 0., 0., 0., 0.), (FORMULATIONS[1], 0., 75., 0., 3.),
            (FORMULATIONS[2], 3150., 0., 126., 0.), (FORMULATIONS[3], 3150., 75., 126., 3.)):
            with self.subTest(formulation=kind):
                expected = {"type": kind, "volume_bbl": 25., "density_pcf": "85-87", "effective_density": 86.,
                            "water_bbl": 25., "nacl_lbs": salt, "wash_gal": wash,
                            "nacl_multiplier": salt_ratio, "wash_multiplier": wash_ratio}
                cfg = {"type": kind, "nacl_multiplier": 126., "wash_multiplier": 3.}
                self.assertEqual(spacer.compute_preflush_payload(cfg, info), expected)
                next(w for w in app.selectbox if w.label == "Pre-flush Fluid Type").set_value(kind).run()
                self.assertEqual(app.session_state["preflush_calc"], expected)
                self.assertEqual(self.name(app), "Salt Saturated Water")

    def test_shared_presets_first_edit_both_navigation_directions(self):
        app = self.app(self.project()); self.phase(app, "phase4")
        self.assertEqual(self.selector(app).options, NAME_OPTIONS)
        self.assertEqual(self.selector(app).value, "Salt Saturated Water")
        for preset in NAME_OPTIONS[:-1]:
            with self.subTest(preset=preset):
                self.selector(app).set_value(preset)
                self.phase(app, "phase6")
                self.assertEqual(self.name(app), preset)
                self.assertEqual(self.selector(app).value, preset)
                self.phase(app, "phase4")
                self.assertEqual(self.selector(app).value, preset)
        self.phase(app, "phase6")
        for preset in NAME_OPTIONS[:-1]:
            with self.subTest(reverse_preset=preset):
                self.selector(app).set_value(preset)
                self.phase(app, "phase4")
                self.assertEqual(self.name(app), preset)
                self.assertEqual(self.selector(app).value, preset)
                self.phase(app, "phase6")

    def test_custom_first_entry_arbitrary_current_name_and_independent_formulation(self):
        app = self.app(self.project()); self.phase(app, "phase4")
        self.selector(app).set_value("Custom").run()
        self.assertEqual(self.name(app), "")
        self.custom(app).set_value("Field Wash Alpha")
        self.phase(app, "phase6")
        self.assertEqual(self.selector(app).value, "Custom")
        self.assertEqual(self.custom(app).value, "Field Wash Alpha")
        before = deepcopy(app.session_state["preflush_calc"])
        self.custom(app).set_value("Field Wash Beta")
        self.phase(app, "phase4")
        self.assertEqual(self.custom(app).value, "Field Wash Beta")
        self.assertEqual(self.name(app), "Field Wash Beta")
        self.selector(app).set_value(NAME_OPTIONS[2])
        self.phase(app, "phase6")
        self.assertEqual(app.session_state["preflush_calc"], before)
        self.assertEqual(app.session_state["preflush_config"], self.project()["preflush_config"])
        next(w for w in app.selectbox if w.label == "Pre-flush Fluid Type").set_value(FORMULATIONS[0]).run()
        self.assertEqual(self.name(app), NAME_OPTIONS[2])
        for name in ("Field Wash Alpha", "  Field Wash ALPHA  "):
            with self.subTest(name=name):
                restored = self.app(audit.round_trip(self.project(name)))
                self.phase(restored, "phase4"); self.phase(restored, "phase6")
                self.assertEqual(self.selector(restored).value, "Custom")
                self.assertEqual(self.custom(restored).value, name)
                self.assertEqual(self.name(restored), name)

    def test_current_json_preserves_preset_custom_blank_and_single_canonical_source(self):
        for name in (NAME_OPTIONS[3], "Field Wash Alpha", ""):
            with self.subTest(name=name):
                app = self.app(self.project(name)); self.phase(app, "phase4"); self.phase(app, "phase6")
                restored_project = audit.round_trip(app.session_state.to_dict())
                self.assertEqual(restored_project["fluids_config"]["params"]["Pre Flush"]["material_name"], name)
                self.assertEqual(restored_project["preflush_config"], self.project()["preflush_config"])
                for forbidden in ("preflush_name", "preflush_material_name", "preflush_name_mode"):
                    self.assertNotIn(forbidden, restored_project)
                    self.assertNotIn(forbidden, restored_project["preflush_config"])
                restored = self.app(restored_project)
                for phase in ("phase6", "phase4"):
                    self.phase(restored, phase)
                    self.assertEqual(self.selector(restored).value, name if name in NAME_OPTIONS[:-1] else "Custom")
                    if name not in NAME_OPTIONS[:-1]: self.assertEqual(self.custom(restored).value, name)

    def test_blank_custom_blocks_status_build_download_and_invalidates_compiled_word(self):
        case = audit.AuditRegressions()
        app = case.configured_app(audit.BATCH1_JOBS[0], optional_fluids=True)
        case.export(app)
        self.phase(app, "phase6")
        self.selector(app).set_value("Custom").run()
        self.assertEqual(self.name(app), "")
        self.assertNotEqual(compute_phase_status(app.session_state)["phase4"]["level"], "ok")
        self.assertNotIn("_compiled_doc_bytes", app.session_state)
        case.assert_export_blocked(app)
        self.phase(app, "phase6")
        self.custom(app).set_value("Field Wash Alpha").run()
        case.export(app)

    def test_blank_name_is_preflush_specific_and_normal_prepare_refreshes_name(self):
        state = self.project("Field Wash Alpha")
        refresh_fluids(state)
        state["fluid_data"]["Pre Flush"]["material_name"] = "STALE"
        # No cement slurry is needed to prove fluid reconstruction itself.
        prepare_calculations(state)
        self.assertEqual(state["fluid_data"]["Pre Flush"]["material_name"], "Field Wash Alpha")
        for bad in ("", "   ", None, "Custom"):
            with self.subTest(bad=bad):
                invalid = deepcopy(state)
                invalid["fluids_config"]["params"]["Pre Flush"]["material_name"] = bad
                invalid["_compiled_doc_bytes"] = b"stale"
                issues = prepare_calculations(invalid)
                self.assertTrue(any("Pre Flush" in issue and "name" in issue.lower() for issue in issues))
                self.assertNotIn("_compiled_doc_bytes", invalid)
        other = self.project(extra_spacers=True)
        other["fluids_config"]["params"]["Spacer"]["material_name"] = ""
        refresh_fluids(other)
        self.assertEqual(other["fluid_data"]["Spacer"]["material_name"], "")

    def test_spacer_choice_catalog_and_all_three_editors(self):
        self.assertEqual(materials_db.SPACER_CHEMICALS, CHEMICALS)
        self.assertEqual(spacer.REQUIRED_SPACER_COLS, ["Chemical", "User Input (% or gal)", "Weighting Agent Type"])
        app = self.app(self.project(extra_spacers=True)); self.phase(app, "phase4"); self.phase(app, "phase6")
        editors = [e for e in app.dataframe if "_editor_spacer_tbl_" in e.proto.id]
        self.assertEqual(len(editors), 3)
        for editor in editors:
            self.assertEqual(json.loads(editor.proto.columns)["Chemical"]["type_config"]["options"], CHEMICALS)

    def test_spacer_first_rows_navigation_and_current_restore_preserve_exact_text(self):
        app = self.app(self.project(extra_spacers=True)); self.phase(app, "phase4"); self.phase(app, "phase6")
        for s in ("Spacer", "Spacer Ahead", "Spacer Behind"):
            for chemical in CHEMICALS[-2:]:
                with self.subTest(spacer=s, chemical=chemical):
                    editor = next(e for e in app.dataframe if
                                  "_editor_spacer_tbl_"+s.lower().replace(" ", "_")+"_" in e.proto.id)
                    # Submit the actual editor event together with navigation,
                    # so its on_change callback must commit before unmounting.
                    app.radio(key="_app_mode_key").set_value("phase4")
                    widgets = app._tree.get_widget_states()
                    event = widgets.widgets.add()
                    event.id = editor.proto.id
                    event.string_value = json.dumps({"edited_rows": {}, "added_rows": [{"Chemical": chemical,
                        "User Input (% or gal)": .125, "Weighting Agent Type": "-"}], "deleted_rows": []})
                    app._run(widgets); self.healthy(app)
                    self.phase(app, "phase6")
                    row = app.session_state["spacer_dfs"][s].iloc[-1]
                    self.assertEqual(row["Chemical"], chemical)
                    self.assertEqual(row["User Input (% or gal)"], .125)
        restored = self.app(audit.round_trip(app.session_state.to_dict())); self.phase(restored, "phase6")
        for s in ("Spacer", "Spacer Ahead", "Spacer Behind"):
            self.assertEqual(list(restored.session_state["spacer_dfs"][s]["Chemical"]), CHEMICALS[-2:])
            self.assertEqual(list(restored.session_state["spacer_dfs"][s]["User Input (% or gal)"]), [.125, .125])
        self.assertNotIn("Magneser Thinner", materials_db.SPACER_CHEMICALS)

    def test_actual_word_name_and_detailed_rows_independent_for_five_jobs(self):
        case = audit.AuditRegressions()
        for job in audit.BATCH1_JOBS:
            with self.subTest(job=job):
                app = case.configured_app(job, optional_fluids=True)
                self.phase(app, "phase6")
                self.selector(app).set_value("Custom").run()
                self.custom(app).set_value("Field Wash Alpha").run()
                baseline = deepcopy(app.session_state["preflush_calc"])
                app.session_state["spacer_dfs"]["Spacer"] = pd.DataFrame([
                    {"Chemical": chemical, "User Input (% or gal)": 1.25, "Weighting Agent Type": "-"}
                    for chemical in CHEMICALS[-2:]], columns=spacer.REQUIRED_SPACER_COLS)
                for name in ("Field Wash Alpha", NAME_OPTIONS[3]):
                    if name != "Field Wash Alpha":
                        self.phase(app, "phase6"); self.selector(app).set_value(name).run()
                    case.export(app)
                    self.assertEqual(app.session_state["fluid_data"]["Pre Flush"]["material_name"], name)
                    self.assertEqual(app.session_state["preflush_calc"], baseline)
                    doc = Document(BytesIO(app.session_state["_compiled_doc_bytes"]))
                    fluid_table = next(t for t in doc.tables if any(
                        any(c.text == "Type" for c in r.cells) and any(c.text == "Name" for c in r.cells)
                        for r in t.rows))
                    row = next(r for r in fluid_table.rows if any(c.text == "Pre Flush" for c in r.cells))
                    self.assertIn(name, [c.text for c in row.cells])
                    preflush = next(t for t in doc.tables if "Pre Flush Data" in t.rows[0].cells[0].text)
                    text = "\n".join(c.text for r in preflush.rows for c in r.cells)
                    for detailed in ("Fresh Water", "Salt", "Chemical Wash"): self.assertIn(detailed, text)
                    self.assertNotIn(name, text)
                    all_text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
                    for chemical in CHEMICALS[-2:]: self.assertIn(chemical, all_text)
                    self.assertNotIn("Magneser Thinner", all_text)


if __name__ == "__main__": unittest.main()
