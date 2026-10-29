"""B4 consumes canonical CT placement, stages and current measured lab sources."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
import re
import pytest
from docx import Document
from docxtpl import DocxTemplate
import materials_db as db
import project_state as ps
import phase_10_procedure as report
import well_report
import well_schematic_adapter as schematic
from ct_stages import ct_stage_projection, stage_note
import test_audit_regressions as audit
from test_round13b_b3_ct_schematic import state, PLACEHOLDER

CASE=audit.AuditRegressions()

@pytest.fixture(scope='module')
def valid_ct():
    s=state();s['ct_stage_config']['count']=3;s['bhct']=150.
    app=CASE.app(s)
    app.text_input(key='_w_client').set_value('B4 CT & Client <review>').run()
    CASE.phase(app,'phase5')
    app.session_state['lab_qc_params']={'Main':{'rheology':audit.rheology_fixture(),**audit.thickening_fixture(),**audit.compressive_fixture()}}
    CASE.phase(app,'phase7')
    for label,value in [('Free Water Collected (45° angle) (ml) - Main',0.),('Surface Sample Hours - Main',8.)]:
        next(w for w in app.number_input if w.label==label).set_value(value).run()
    next(w for w in app.text_input if w.label=='70 Bc (HH:MM) - Main').set_value('03:30').run()
    button=next(w for w in app.button if w.label=='Confirm measured lab results')
    assert not button.disabled, [w.value for w in app.warning]+[w.value for w in app.error]+[w.value for w in app.caption]
    button.click().run()
    return app.session_state.to_dict()


def context(s,prepared=False):
    with patch.object(report.st,'session_state',s):
        return report.build_master_context(calculations_prepared=prepared)


def test_current_ct_is_admitted_and_dedicated_procedure_uses_sources(valid_ct):
    s=deepcopy(valid_ct)
    assert ps.prepare_calculations(s)==[]
    ctx=context(s)
    text=ctx['procedure'] if 'procedure' in ctx else s['procedure_text']
    assert ctx['job_type']==db.CT_JOB
    assert '3527.0 m MD' in text and '3470.0 m' in ctx['executive_summary']
    for term in ['Coiled Tubing','3000 psi','DAQ','30 m above','Displacement Volume: Calculate at rig site.','0.7-1.2','WOC']:
        assert term in text,term
    assert stage_note(ct_stage_projection(s)) in text
    assert stage_note(ct_stage_projection(s)) in ctx['slurries'][0]['note_mix']
    assert '7 + 7 + 6' in text and '4.9 + 4.9 + 4.2' in text
    assert not re.search(r'drill.?pipe|swage|stands|bump|plug release',text,re.I)
    steps=[int(x) for x in re.findall(r'(?m)^(\d+)\. ',text)]
    assert steps==list(range(1,len(steps)+1))


@pytest.mark.parametrize('kind',['complete','partial'])
def test_actual_ct_docx_uses_existing_section_ii_and_canonical_context(valid_ct,kind,tmp_path):
    s=deepcopy(valid_ct);s['hardware_table']=state(kind)['hardware_table']
    ctx=context(s)
    doc=DocxTemplate('master_template.docx');doc.render(ctx);buf=BytesIO();doc.save(buf)
    result=well_report.postprocess_well_report(buf.getvalue(),ctx['well_profile_report'],ctx['note_vertical'],ctx['cement_chart_report'],ctx['centralizer_report'],
        pumping_report=__import__('pumping_simulation').prepare_pumping_simulation(s),
        schematic_report=schematic.prepare_schematic(s))
    out=Document(BytesIO(result));xml=ZipFile(BytesIO(result)).read('word/document.xml').decode()
    from test_uat_round9g_centralizer_word import ix_region
    from docx.oxml.ns import qn
    region=ix_region(out)
    section_text='\n'.join(''.join(t.text or '' for t in e.iter(qn('w:t'))) for e in region)
    assert section_text.count('Not applicable (Plug/Squeeze).')==1
    assert not any(t._tbl in region for t in out.tables)
    assert not any(e.xpath('.//w:drawing') for e in region)
    assert len(region[-1].xpath('.//w:br[@w:type="page"]'))==1
    assert db.CT_JOB in xml and 'B4 CT &amp; Client &lt;review&gt;' in xml
    assert 'Spacer Formulation Data' not in xml and 'Spacer Additives Data' not in xml
    assert '{{' not in xml and '{%' not in xml
    assert '3527.0 m MD' in xml and '3470.0 m' in xml
    assert 'Displacement Volume: Calculate at rig site.' in xml
    if kind=='partial':assert PLACEHOLDER in '\n'.join(p.text for p in out.paragraphs)
    else:assert out.inline_shapes
    (tmp_path/f'ct-{kind}.docx').write_bytes(result)


@pytest.mark.parametrize('prepared',[False,True])
@pytest.mark.parametrize('fault',['depth','top','stage','rate','lab','geometry'])
def test_all_ct_report_entry_points_fail_closed_and_invalidate(valid_ct,prepared,fault):
    s=deepcopy(valid_ct);s['_compiled_doc_bytes']=b'stale'
    if fault=='depth':s['hardware_table'].loc[s['hardware_table']['Description']=='CT','MD (m)']='3500-3527'
    elif fault=='top':s['cement_params']['Main']['top_depth']=3527.
    elif fault=='stage':s['ct_stage_config']['count']=2;s['fluids_config']['params']['Main']['volume']=30.
    elif fault=='rate':s['fluids_config']['params']['Main']['pump_rate']='0'
    elif fault=='lab':s['lab_qc_params']['Main']['reviewed']=False
    elif fault=='geometry':s['hardware_table'].loc[0,'ID (in)']=20.
    with pytest.raises(ValueError):context(s,prepared)
    assert '_compiled_doc_bytes' not in s


def test_stage_depth_and_operational_edits_preserve_manual_text_for_review(valid_ct):
    s=deepcopy(valid_ct);context(s)
    s['procedure_text']+='\nEngineer reviewed CT instruction.'
    with patch.object(report.st,'session_state',s):
        spec=report.synchronize_report_texts()['procedure_text']
        ps.accept_report_text(s,'procedure_text',spec['generated'],spec['signature'])
        s['ct_stage_config']['count']=2
        specs=report.synchronize_report_texts()
    assert specs['procedure_text']['pending']
    assert s['procedure_text'].endswith('Engineer reviewed CT instruction.')
    with pytest.raises(ValueError,match='preserved manual'):context(s)


@pytest.mark.parametrize('job',[db.CT_JOB,'CMT PLUG','CMT SQUEEZE'])
def test_explicit_nonapplicable_centralizer_contract(job):
    import centralizer_report as bridge
    expected={'applicable':False,'analysis':None,'view':None,'source':None,'warning':False}
    assert bridge.prepare_report({'job_type':job})==expected


@pytest.mark.parametrize('job',['CSG 9 5/8"','LNR 7"','TIE BACK LNR 7"'])
def test_applicable_jobs_still_require_executed_centralizer(job):
    import centralizer_report as bridge
    with pytest.raises(ValueError,match='Execute a current successful Centralizer'):
        bridge.prepare_report({'job_type':job})

COMBINATIONS=[(),('Spacer Ahead',),('Spacer Behind',),('Spacer Ahead','Spacer Behind'),('Pre Flush',),('Pre Flush','Spacer Ahead')]

@pytest.mark.parametrize('job',db.JOB_TYPES)
@pytest.mark.parametrize('selected',COMBINATIONS)
def test_all_job_operational_procedure_contract(valid_ct,job,selected):
    from test_round13b_b4_operational_spacers import operational
    s=deepcopy(valid_ct) if job==db.CT_JOB else operational(job)
    cfg=s['fluids_config'];cfg['active']=[*selected,'Main','Displacement Fluid']
    for name in selected:
        cfg['params'][name]={'volume':7.5,'density':'80','pump_rate':'0.7-1.2','material_name':name+' field & material'}
    s['preflush_config']={'type':'Fresh Water Only','nacl_multiplier':126.,'wash_multiplier':3.}
    ps.refresh_fluids(s);ps.refresh_preflush(s)
    before=deepcopy(s.get('cement_params'))
    with patch.object(report.st,'session_state',s):
        text=report.generate_official_procedure(job,s['fluid_data'],cfg['active'],s['cement_params'],s.get('cement_additives_dfs',{}),s.get('preflush_calc'),s.get('geo_md'),s.get('total_pump_time_min'),s.get('placement_config'))
    assert s['cement_params']==before
    positions=[]
    for name in cfg['active']:
        if name in ('Spacer Ahead','Spacer Behind'):
            record=s['fluid_data'][name]
            line=next(l for l in text.splitlines() if 'as '+name in l)
            assert record['material_name'] in line and '0.7-1.2' in line
            if job == db.CT_JOB:
                assert f"calculated duration {record['duration_min']:.2f} min." in line and 'HH:MM' not in line
            else:
                assert record['duration_str'] in line
            positions.append((cfg['active'].index(name),text.index(line)))
    pump_main=next(l for l in text.splitlines() if ('as Main' in l or 'Main cement slurry' in l))
    for index,position in positions:
        assert (position<text.index(pump_main))==(index<cfg['active'].index('Main'))
    if 'Pre Flush' in selected:assert 'pre flush' in text.lower()
    assert 'Hematite' not in text


def test_ct_ui_build_restore_and_operational_edits_retire_download(valid_ct,tmp_path):
    from test_uat_round10c_integration import save_project,restore_project
    app=CASE.app(deepcopy(valid_ct));CASE.phase(app,'phase4')
    for name in ('Pre Flush','Spacer Ahead','Spacer Behind'):
        app.checkbox(key='chk_fluid_'+name).check().run()
        app.number_input(key='vol_'+name).set_value(7.5).run()
        app.text_input(key='den_'+name).set_value('80').run()
        app.text_input(key='rate_'+name).set_value('0.7-1.2').run()
    assert all(app.checkbox(key='chk_fluid_'+name).value for name in ('Pre Flush','Spacer Ahead'))
    CASE.phase(app,'phase6')
    assert not any('_editor_spacer_' in e.proto.id for e in app.dataframe)
    assert any(w.label=='Pre-flush Fluid Type' for w in app.selectbox)
    saved=save_project(app);app=restore_project(saved)
    CASE.phase(app,'phase7')
    # The established Main review survives compatible operational-only edits.
    assert app.session_state['lab_qc_params']['Main']['reviewed']
    CASE.phase(app,'phase10')
    assert not app.error,[e.value for e in app.error]
    button=next(b for b in app.button if b.label=='Build Word Document')
    assert not button.disabled
    button.click().run(timeout=120);CASE.healthy(app)
    assert not app.error,[e.value for e in app.error]
    blob=app.session_state['_compiled_doc_bytes'];doc=Document(BytesIO(blob))
    text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert all(name in text for name in ('Spacer Ahead','Spacer Behind','Pre Flush','CMT PLUG WITH CT'))
    assert 'Spacer Additives Data' not in text and 'Spacer Formulation Data' not in text
    (tmp_path/'ct-ui-actual-build.docx').write_bytes(blob)
    CASE.phase(app,'phase4');app.number_input(key='vol_Spacer Ahead').set_value(8.).run()
    CASE.phase(app,'phase10')
    assert '_compiled_doc_bytes' not in app.session_state


def test_unrelated_guardrails_and_all_frozen_authorities_are_unchanged():
    import ast,subprocess
    before=ast.parse(subprocess.check_output(['git','show','f185c0f88bdcb9fc5342dcbf5eb306dd29c0cee6:test_regression_guardrails.py']))
    after=ast.parse(Path('test_regression_guardrails.py').read_text())
    approved=ast.parse(subprocess.check_output(['git','show','e4db7795135fc3039fe9ce281586ec80f42b75c5:test_regression_guardrails.py']))
    fixture_names={'test_master_context_contains_placement_keys','test_blank_volume_basis_exports_as_na',
                   'test_preflush_payload_recomputed_at_export_time','test_spacer_densities_sort_numerically'}
    fixture_nodes={node.name:node for node in approved.body
                   if isinstance(node,ast.FunctionDef) and node.name in fixture_names}
    assert set(fixture_nodes)==fixture_names
    before.body=[fixture_nodes.get(node.name,node) if isinstance(node,ast.FunctionDef) else node
                 for node in before.body]
    for tree in (before,after):
        tree.body=[node for node in tree.body if not(isinstance(node,ast.FunctionDef) and node.name=='test_render_master_document_smoke')]
    assert ast.dump(before)==ast.dump(after)
    for name in ('well_schematic_engine.py','cement_charts_engine.py','centralizer_optimizer_engine.py','standalone_spacer.py','ct_stages.py','placement.py','phase_5_cement.py','phase_7_lab.py','rheology.py','project_io.py','requirements.txt'):
        if name == 'well_schematic_engine.py':
            from test_round13b_b3_ct_schematic import assert_authorized_radial_renderer
            assert_authorized_radial_renderer(Path(name).read_bytes())
            continue
        if name == 'project_io.py':
            import hashlib
            assert hashlib.sha256(Path(name).read_bytes()).hexdigest() == '20ec22a012d42de360cd84d54ffcbec923896b5a228e48b540d1e8f91a2518d2'
            continue
        if name == 'phase_7_lab.py':
            import hashlib
            assert hashlib.sha256(Path(name).read_bytes()).hexdigest() == 'f95e337cc8445a5b6cf8c582bb09c982a36e9d638608582d41fb50922f56903f'
            continue
        if name == 'phase_5_cement.py':
            import hashlib
            assert hashlib.sha256(Path(name).read_bytes()).hexdigest() == '3a30e2a0d42316b34d0c15b88ef7f2c539e8c1b38b48bd8a0ed614e6704469de'
            continue
        reference='f185c0f88bdcb9fc5342dcbf5eb306dd29c0cee6'
        assert Path(name).read_bytes()==subprocess.check_output(['git','show',f'{reference}:{name}'])
    from test_round13b_post_b4_word_quality import test_template_changes_only_approved_sample_text_payloads
    test_template_changes_only_approved_sample_text_payloads()


@pytest.mark.parametrize('tank',[None,'','Unknown tank'])
@pytest.mark.parametrize('prepared',[True,False])
def test_ct_tank_admission_preserves_draft_and_lab_but_retires_word(valid_ct,tank,prepared):
    s=deepcopy(valid_ct);s['cement_params']['Main']['tank_name']=tank
    s['_compiled_doc_bytes']=b'stale';review=s['lab_qc_params']['Main']['review_signature']
    with pytest.raises(ValueError,match='Phase V CT Main: select a Mixing Tank'):
        context(s,prepared)
    assert '_compiled_doc_bytes' not in s
    assert s['cement_params']['Main']['tank_name']==tank
    assert s['lab_qc_params']['Main']['reviewed']
    assert s['lab_qc_params']['Main']['review_signature']==review


@pytest.mark.parametrize('tank',db.get_available_tanks(db.CT_JOB))
def test_ct_report_uses_selected_catalog_tank_without_changing_formulation(valid_ct,tank):
    s=deepcopy(valid_ct);s['cement_params']['Main']['tank_name']=tank
    before=deepcopy(s['cement_params']);review=s['lab_qc_params']['Main']['review_signature']
    ctx=context(s)
    assert ctx['slurries'][0]['mixing_tank']==tank
    assert tank in ctx['slurries'][0]['note_mix'] and tank in s['procedure_text']
    assert s['cement_params']==before
    assert s['lab_qc_params']['Main']['reviewed']
    assert s['lab_qc_params']['Main']['review_signature']==review
