"""R3-06–09: real project custom identities use the existing Density/SG path."""
from copy import deepcopy
from io import BytesIO
import json
import unittest

import pandas as pd
from docx import Document
import test_audit_regressions as audit
import test_uat_round2_phase5_density as density_tests
from engineering_tools import (calculate_slurry_from_components, compute_phase_status)
from phase_5_cement import (build_components, _normalize_additive_rows, build_cement_tables,
                            calculate_base_results, REQUIRED_ADDITIVE_COLUMNS)
from phase_7_lab import build_lab_df_from_phase5
from phase_10_procedure import separate_lab_tables
from project_state import lab_source_signature


def custom(name='EXT-X', kind='Local Extender', state='Powder', mix='In Mix Water', sg=2.2):
    return {'Material Type': kind, 'Name': name, 'Physical State': state,
            'Mix Method': mix, 'User Input': 0.5 if state == 'Liquid' else 2., 'Density': sg}


class TrueCustomUAT(unittest.TestCase):
    app = audit.AuditRegressions.app
    phase = audit.AuditRegressions.phase
    healthy = audit.AuditRegressions.healthy
    configured_app = audit.AuditRegressions.configured_app
    export = audit.AuditRegressions.export
    assert_export_blocked = audit.AuditRegressions.assert_export_blocked
    editor = density_tests.Round2DensityUAT.editor
    edit_rows = density_tests.Round2DensityUAT.edit_rows

    def test_sg_engine_characterization_and_sensitivity(self):
        for state, values, factor in [('Powder', (2.1, 2.5), 62.4), ('Liquid', (1.1, 1.2), (62.4 / 7.48051945))]:
            for mix in ('In Mix Water', 'Dry Blend'):
                with self.subTest(state=state, mix=mix):
                    results = []; labs = []
                    for sg in values:
                        frame = pd.DataFrame([custom(state=state, mix=mix, sg=sg)])
                        powders, liquids, salt = build_components(frame)
                        component = (powders or liquids)[0]
                        self.assertAlmostEqual(component['density_pcf' if powders else 'density_ppg'], sg * factor)
                        if liquids:
                            self.assertAlmostEqual(component['lab_factor'], sg * (62.4 / 7.48051945) / 110.0)
                        results.append(calculate_slurry_from_components(118, 50, powders=powders, liquids=liquids, salt_pct=salt))
                        labs.append(build_lab_df_from_phase5(frame, recalculate_mass=True))
                    self.assertNotEqual(results[0]['field_water_bbl'], results[1]['field_water_bbl'])
                    self.assertNotEqual(labs[0]['Mass'].tolist(), labs[1]['Mass'].tolist())

    def test_incomplete_identity_and_sg_block_status_calculation_export(self):
        app = self.configured_app(audit.BATCH1_JOBS[0])
        invalid = [custom(kind='Other (Custom)'), custom(name='Other (Custom)'),
                   custom(kind=''), custom(name='')]
        invalid += [custom(sg=value) for value in (None, '', 0, -1, float('nan'), float('inf'), 'bad')]
        for row in invalid:
            with self.subTest(row=row):
                frame = pd.DataFrame([row])
                with self.assertRaises(ValueError): build_components(frame)
                with self.assertRaises(ValueError): build_lab_df_from_phase5(frame)
                app.session_state['cement_additives_dfs']['Main'] = frame
                self.assertNotEqual(compute_phase_status(app.session_state)['phase5']['level'], 'ok')
        self.phase(app, 'phase5')
        self.assert_export_blocked(app)

    def test_catalog_locks_and_custom_lookalikes_remain_exact(self):
        frame = pd.DataFrame([custom(name='Micro Silica', kind='Local Extender', state='Liquid', sg=None)])
        frame = _normalize_additive_rows(frame)
        self.assertEqual(frame.iloc[0][['Material Type','Name','Physical State','Density']].tolist(),
                         ['Extender','Micro Silica','Powder',137.3415375 / 62.4])
        frame.at[0,'Density'] = 2.4
        self.assertEqual(_normalize_additive_rows(frame).at[0,'Density'], 2.4)
        for name in ('Micro Silica Local', 'Hidense-X', 'NaCl-X'):
            powders, _, salt = build_components(pd.DataFrame([custom(name=name)]))
            self.assertTrue(powders[0]['in_solution']); self.assertEqual(salt, 0)
        self.assertEqual(build_components(pd.DataFrame([custom(kind='Extender')]))[0][0]['name'], 'EXT-X')

    def test_true_custom_text_first_commit_navigation_and_restore(self):
        app = self.configured_app(audit.BATCH1_JOBS[0])
        app.session_state['cement_additives_dfs']['Main'] = pd.DataFrame([
            custom(name='Other (Custom)', kind='Other (Custom)', sg=None)])
        self.phase(app, 'phase5')
        for field, value in [('Material Type','Local Extender'), ('Name','EXT-X')]:
            widget = next(w for w in app.text_input if w.label == f'Custom {field} - Main row 1')
            widget.set_value(value).run(); self.healthy(app)
            self.assertEqual(app.session_state['cement_additives_dfs']['Main'].at[0,field], value)
        self.edit_rows(app, {'0': {'Density': 2.2}})
        self.phase(app, 'phase7'); self.phase(app, 'phase5')
        saved = audit.round_trip(app.session_state.to_dict())
        restored = self.app(saved); self.phase(restored,'phase5')
        frame = restored.session_state['cement_additives_dfs']['Main']
        self.assertEqual(frame.iloc[0].to_dict(), custom())
        config = json.loads(self.editor(restored).proto.columns)
        self.assertIn('EXT-X', config['Name']['type_config']['options'])
        self.assertIn('Local Extender', config['Material Type']['type_config']['options'])
        self.assertEqual(list(self.editor(restored).proto.column_order), REQUIRED_ADDITIVE_COLUMNS + ['Density'])
        next(w for w in restored.text_input if w.label == 'Custom Name - Main row 1').set_value('RET-Y')
        restored.radio(key='_app_mode_key').set_value('phase2_3').run()
        self.phase(restored,'phase5')
        self.assertEqual(restored.session_state['cement_additives_dfs']['Main'].at[0,'Name'], 'RET-Y')
        self.edit_rows(restored, {'0': {'Name': 'Micro Silica'}})
        frame = restored.session_state['cement_additives_dfs']['Main']
        self.assertEqual(frame.iloc[0][['Material Type','Name','Physical State','Density']].tolist(),
                         ['Extender','Micro Silica','Powder',137.3415375 / 62.4])

    def test_custom_controls_follow_row_deletion_and_sentinel_reselection(self):
        app = self.configured_app(audit.BATCH1_JOBS[0])
        app.session_state['cement_additives_dfs']['Main'] = pd.DataFrame([
            custom(name='FIRST-X',kind='First Type'), custom(name='SECOND-X',kind='Second Type')])
        self.phase(app,'phase5')
        widgets = app._tree.get_widget_states()
        event = widgets.widgets.add(); event.id = self.editor(app).proto.id
        event.string_value = json.dumps({'edited_rows':{},'added_rows':[],'deleted_rows':[0]})
        app._run(widgets); self.healthy(app)
        self.assertEqual(next(w for w in app.text_input if w.label=='Custom Name - Main row 1').value,'SECOND-X')
        self.assertEqual(next(w for w in app.text_input if w.label=='Custom Material Type - Main row 1').value,'Second Type')
        self.phase(app,'phase2_3'); self.phase(app,'phase5')
        self.edit_rows(app, {'0':{'Name':'Other (Custom)'}})
        self.assertEqual(next(w for w in app.text_input if w.label=='Custom Name - Main row 1').value,'')
        self.assertEqual(app.session_state['cement_additives_dfs']['Main'].at[0,'Name'],'Other (Custom)')
        self.assertNotEqual(compute_phase_status(app.session_state)['phase5']['level'],'ok')

    def test_occurrence_aware_custom_lab_classification_and_values(self):
        frame = pd.DataFrame([custom(name='DUP-X',kind='Type A',mix='Dry Blend'),
                              custom(name='DUP-X',kind='Type B'),
                              custom(name='Local cement liquid',kind='Type C',state='Liquid')])
        lab = build_lab_df_from_phase5(frame, recalculate_mass=True)
        lab['Lot No'] = ['C', 'A', 'B', 'L', 'W']
        conv, adds = separate_lab_tables(lab, frame)
        self.assertEqual([(r['Name'],r['Material Type']) for r in conv if r['Name']],
                         [('CEMENT G DELIJAN','Cement'), ('DUP-X','Type A')])
        self.assertEqual([(r['Name'],r['Material Type']) for r in adds if r['Name']],
                         [('DUP-X','Type B'),('Local cement liquid','Type C'),('Base Water','Base Water')])
        for expected, actual in zip(lab.iloc[1:4].to_dict('records'), [conv[1],adds[0],adds[1]]):
            self.assertEqual(actual['Mass'],expected['Mass'])
            self.assertEqual(actual['Concentration'],expected['Concentration'])
            self.assertEqual(actual['Design Unit'],expected['Unit'])
            self.assertEqual(actual['Lot Number'],expected['Lot No'])

    def test_generated_tables_use_actual_identity_and_existing_water_basis(self):
        frame = pd.DataFrame([custom(name='DRY-X',kind='Local Blend',mix='Dry Blend'), custom(state='Liquid')])
        p = {'cmt_sg':3.2,'dead_vol':10.,'tank_name':'Tank A','base_cement':'CEMENT G DELIJAN','auto_calc':True,'mix_water':30.,'total_sacks':100.}
        calc = calculate_base_results(p,50,118,*build_components(frame))
        p.update(mix_water=calc['field_water_bbl'], total_sacks=calc['field_sacks'])
        blend, adds, note = build_cement_tables(p,frame,calc)
        self.assertEqual(blend.iloc[1]['Name'],'DRY-X'); self.assertEqual(blend.iloc[1]['Material Type'],'Local Blend')
        self.assertEqual(adds.iloc[0]['Name'],'EXT-X'); self.assertEqual(adds.iloc[0]['Material Type'],'Local Extender')
        self.assertIn(f"{p['mix_water'] + p['dead_vol']:.1f}",note)
        self.assertNotIn('Other (Custom)',blend.to_string()+adds.to_string()+note)

    def test_custom_round_trip_drafts_review_drift_and_actual_word_all_workflows(self):
        from project_state import purge_inactive_slurries
        for job in audit.BATCH1_JOBS:
            with self.subTest(job=job):
                app = self.configured_app(job)
                frame = pd.DataFrame([custom(name='DRY-X',kind='Local Blend',mix='Dry Blend'),
                                      custom(name='LQ-X',kind='Local Liquid',state='Liquid',sg=1.15)])
                app.session_state['cement_additives_dfs']['Main'] = frame.copy()
                self.phase(app,'phase5'); self.phase(app,'phase7')
                next(b for b in app.button if b.label=='Sync with Phase V').click().run()
                next(b for b in app.button if b.label=='Confirm measured lab results').click().run()
                self.assertEqual(compute_phase_status(app.session_state)['phase7']['level'],'ok')
                before = lab_source_signature(app.session_state,'Main')
                for field,value in [('Material Type','Another Type'),('Name','Another Name'),('Physical State','Liquid'),
                                    ('Mix Method','In Mix Water'),('User Input',3.),('Density',2.4)]:
                    saved = frame.at[0,field]; app.session_state['cement_additives_dfs']['Main'].at[0,field] = value
                    self.assertNotEqual(lab_source_signature(app.session_state,'Main'),before)
                    self.assertNotEqual(compute_phase_status(app.session_state)['phase7']['level'],'ok')
                    app.session_state['cement_additives_dfs']['Main'].at[0,field] = saved
                project = audit.round_trip(app.session_state.to_dict())
                restored = self.app(project); self.phase(restored,'phase5'); self.phase(restored,'phase7')
                pd.testing.assert_frame_equal(restored.session_state['cement_additives_dfs']['Main'],frame)
                self.assertEqual(restored.session_state['lab_grid_dfs']['Main']['Mass'].tolist(),app.session_state['lab_grid_dfs']['Main']['Mass'].tolist())
                self.export(restored)
                doc = Document(BytesIO(restored.session_state['_compiled_doc_bytes']))
                text = '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
                self.assertIn('Local Blend',text); self.assertIn('DRY-X',text)
                self.assertIn('Local Liquid',text); self.assertIn('LQ-X',text); self.assertNotIn('Other (Custom)',text)
                conventional = next(t for t in doc.tables if t.rows[0].cells[0].text == 'Conventional')
                split = next(i for i,r in enumerate(conventional.rows) if r.cells[0].text == 'Additives')
                conventional_text = '\n'.join(c.text for r in conventional.rows[:split] for c in r.cells)
                additives_text = '\n'.join(c.text for r in conventional.rows[split:] for c in r.cells)
                self.assertIn('DRY-X',conventional_text); self.assertNotIn('LQ-X',conventional_text)
                self.assertIn('LQ-X',additives_text); self.assertNotIn('DRY-X',additives_text)
                for name in ('DRY-X','LQ-X'):
                    measured = restored.session_state['lab_grid_dfs']['Main'].query('Material == @name').iloc[0]
                    rendered = next(r for r in conventional.rows if r.cells[1].text == name)
                    self.assertEqual([c.text for c in rendered.cells[3:]],
                                     [measured['Concentration'],measured['Unit'],measured['Mass']])
                state = audit.round_trip(project)
                state['fluids_config']['active'].remove('Main'); purge_inactive_slurries(state,state['fluids_config']['active'])
                pd.testing.assert_frame_equal(state['inactive_slurry_drafts']['Main']['cement_additives_dfs'],frame)
                state['fluids_config']['active'].append('Main'); purge_inactive_slurries(state,state['fluids_config']['active'])
                pd.testing.assert_frame_equal(state['cement_additives_dfs']['Main'],frame)
                self.assertNotIn('Main',state.get('lab_source_signatures',{}))

if __name__ == '__main__': unittest.main()
