"""Owner-approved 14C presentation and retained measured-source boundaries."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import hashlib
import subprocess
from zipfile import ZipFile

import pandas as pd
import pytest
from docx import Document

import engineering_tools as eng
import project_state as ps
import test_audit_regressions as audit
from test_round13a_a2_spacer_lab import state as spacer_state
from test_uat_round5_thickening import FIELDS

CASE = audit.AuditRegressions()
BASE = 'b21d1fc94bcc5e1853800b802c48f77026fefd6c'


@pytest.fixture(scope='module')
def cement():
    return audit.round_trip(CASE.configured_app('CSG 9 5/8"').session_state.to_dict())


def test_original_cement_qc_controls_and_visual_order(cement):
    app = CASE.app(deepcopy(cement))
    CASE.phase(app, 'phase7')
    titles = [x.value for x in app.markdown if x.value.startswith('#### ')]
    expected = ['#### Rheology Test', '#### Fluid Loss and Free Water Test',
                '#### Thickening Time Test', '#### Compressive Strength Test']
    assert [x for x in titles if x in expected] == expected
    assert not any('Enter Time (HH:MM)' in w.label for w in app.text_input)
    hours = [w for w in app.number_input if w.label == 'Surface Sample Hours - Main']
    assert len(hours) == 1 and hours[0].key.startswith('_qc_surface_hours_in_main_')
    sequence = list(app.tabs[0])
    thickening = next(i for i, w in enumerate(sequence) if w.type == 'markdown' and w.value == '#### Thickening Time Test')
    compressive = next(i for i, w in enumerate(sequence) if w.type == 'markdown' and w.value == '#### Compressive Strength Test')
    hours_position = next(i for i, w in enumerate(sequence) if getattr(w, 'key', None) == hours[0].key)
    assert thickening < hours_position < compressive
    assert all(w.label != 'IOD' for w in app.metric)
    for label in ('Surface — Ramp-Down', 'BHCT — Ramp-Up', 'BHCT — Ramp-Down'):
        selected = next(w for w in app.checkbox if w.label == label + ' - Main')
        if not selected.value:
            selected.check().run()
        readings = [w for w in app.text_input if f'RPM - {label} - Main' in w.label]
        assert [w.label.split(' RPM')[0] for w in readings] == ['300', '200', '100', '60', '30', '6', '3']
    # The actual AppTest tree places RPM fields directly in their shared tab,
    # rather than inside the old seven horizontal columns.
    for tab in app.tabs:
        direct = [child for child in tab.children.values() if getattr(child, 'label', '').endswith('RPM - BHCT — Ramp-Down - Main')]
        assert len(direct) == 7
    result = eng.validate_lab_rheology(cement['lab_qc_params']['Main'], 'Main')
    assert result['bhct_down']['iod'] is not None


def test_phase6_neutral_original_heading():
    app = CASE.app(spacer_state())
    CASE.phase(app, 'phase6')
    assert '### Spacer Additives Data' in [w.value for w in app.markdown]
    assert not any('Spacer Additives Data</div>' in w.value for w in app.markdown)


def test_new_water_lot_only_and_explicit_clear_roundtrip():
    state = spacer_state()
    grid = ps.refresh_spacer_lab(state)
    assert grid.loc[grid.Material == 'Fresh Water', 'Lot No'].item() == 'Lab Sample'
    for lot in ('operator lot & <1>', ''):
        grid.loc[grid.Material == 'Fresh Water', 'Lot No'] = lot
        state['lab_grid_dfs']['Spacer'] = grid
        restored = audit.round_trip(state)
        rebuilt = ps.refresh_spacer_lab(restored)
        assert rebuilt.loc[rebuilt.Material == 'Fresh Water', 'Lot No'].item() == lot
        assert (rebuilt.loc[rebuilt.Material != 'Fresh Water', 'Lot No'] == '').all()
        grid = rebuilt


def test_water_lot_actual_editor_clear_and_review_retirement():
    from test_round13a_a2_spacer_lab import measured, edit_lot_with_navigation
    state = measured(spacer_state())
    app = CASE.app(state)
    CASE.phase(app, 'phase7')
    editor = next(w for w in app.dataframe if '_editor_lab_tbl_spacer_' in w.proto.id)
    assert editor.value.iloc[0]['Lot No'] == 'Lab Sample'
    app.session_state['_compiled_doc_bytes'] = b'stale'
    edit_lot_with_navigation(app, 0, '')
    assert not app.session_state['lab_qc_params']['Spacer']['reviewed']
    assert '_compiled_doc_bytes' not in app.session_state
    restored = CASE.app(audit.round_trip(app.session_state.to_dict()))
    CASE.phase(restored, 'phase7')
    editor = next(w for w in restored.dataframe if '_editor_lab_tbl_spacer_' in w.proto.id)
    assert editor.value.iloc[0]['Lot No'] == ''
    assert eng.compute_phase_status(restored.session_state)['phase7']['level'] != 'ok'


def test_final_template_sha_rejects_one_byte_and_unrelated_parts():
    import ast
    tree = ast.parse(Path('tests/test_uat_round8b_report.py').read_text())
    identity = next(n for n in tree.body if isinstance(n, ast.FunctionDef)
                    and n.name == 'test_template_dependency_and_protected_sources_identical')
    # The expected pin is a fixed literal in the actual existing protected test.
    expected = next(n.value for n in ast.walk(identity) if isinstance(n, ast.Constant)
                    and isinstance(n.value, str) and len(n.value) == 64)
    actual = Path('master_template.docx').read_bytes()
    assert hashlib.sha256(actual).hexdigest() == expected
    assert hashlib.sha256(bytes([actual[0] ^ 1]) + actual[1:]).hexdigest() != expected


def test_surface_hours_original_callback_and_independent_admission(cement):
    app = CASE.app(deepcopy(cement))
    CASE.phase(app, 'phase7')
    legacy = app.session_state['lab_qc_params']['Main']['thickening_enter_time']
    hours = next(w for w in app.number_input if w.label == 'Surface Sample Hours - Main')
    key = hours.key
    hours.set_value(None).run()
    assert next(w for w in app.button if w.label == 'Confirm measured lab results').disabled
    assert eng.compute_phase_status(app.session_state)['phase7']['level'] != 'ok'
    CASE.assert_export_blocked(app)
    CASE.phase(app, 'phase7')
    next(w for w in app.number_input if w.key == key).set_value(9.25)
    CASE.phase(app, 'phase1')  # Pending blur commits the original collection field.
    assert app.session_state['lab_qc_params']['Main']['surface_hardened_hours'] == 9.25
    restored = CASE.app(audit.round_trip(app.session_state.to_dict()))
    CASE.phase(restored, 'phase7')
    assert next(w for w in restored.number_input if w.label == 'Surface Sample Hours - Main').value == 9.25
    assert restored.session_state['lab_qc_params']['Main']['thickening_enter_time'] == legacy
    next(w for w in restored.button if w.label == 'Confirm measured lab results').click().run()
    assert eng.compute_phase_status(restored.session_state)['phase7']['level'] == 'ok'


@pytest.mark.parametrize('legacy', ['absent', None, '', '01:25', 'unfinished legacy'])
def test_hidden_legacy_not_completion_authority(legacy):
    qc = {**FIELDS, 'free_water': 0., 'free_water_45': 0., 'surface_hardened_hours': 8.}
    if legacy == 'absent':
        qc.pop('thickening_enter_time')
    else:
        qc['thickening_enter_time'] = legacy
    before = deepcopy(qc)
    eng.validate_thickening_test(qc)
    eng.validate_lab_collection_results(qc)
    assert qc == before
    for key in ('thickening_test_start', 'thickening_cell', 'thickening_30_bc', 'thickening_50_bc', 'thickening_time'):
        bad = {**qc, key: None}
        with pytest.raises(ValueError):
            eng.validate_thickening_test(bad)
    with pytest.raises(ValueError):
        eng.validate_lab_collection_results({**qc, 'surface_hardened_hours': None})


def test_template_only_three_text_cells_and_real_render(cement, tmp_path):
    old = subprocess.check_output(['git', 'show', BASE + ':master_template.docx'])
    from test_round13b_post_b4_word_quality import approved_template_xml
    with ZipFile(BytesIO(old)) as original, ZipFile('master_template.docx') as current:
        assert original.namelist() == current.namelist()
        for name in original.namelist():
            if name != 'word/document.xml':
                assert original.read(name) == current.read(name), name
        approved_template_xml(current.read('word/document.xml'))
    state = deepcopy(cement)
    state['lab_qc_params']['Main'].pop('thickening_enter_time', None)
    app = CASE.app(audit.round_trip(state))
    CASE.phase(app, 'phase7')
    # Confirm actual current measurements after changing the reviewed QC record.
    next(w for w in app.button if w.label == 'Confirm measured lab results').click().run()
    assert eng.compute_phase_status(app.session_state)['phase7']['level'] == 'ok'
    CASE.export(app)
    data = app.session_state['_compiled_doc_bytes']
    (tmp_path / 'uat14c-csg.docx').write_bytes(data)
    table = next(t for t in Document(BytesIO(data)).tables if t.cell(0, 0).text == 'Thickening Time Test')
    assert ' '.join(table.rows[2].cells[2].text.split()) == 'Initial Temp (degF)'
    assert table.rows[3].cells[2].text == '70'
    assert table.rows[3].cells[3].text == 'Surface'
    assert table.rows[2].cells[3].text == 'Initial Pressure'
    assert 'Enter Time' in [c.text for c in table.rows[1].cells]
    qc = app.session_state['lab_qc_params']['Main']
    assert table.rows[3].cells[4].text == qc['thickening_test_start']
    assert [table.rows[3].cells[i].text for i in (6, 7, 8)] == [qc[k] for k in ('thickening_30_bc', 'thickening_50_bc', 'thickening_time')]
