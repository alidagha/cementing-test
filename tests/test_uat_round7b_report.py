"""Round 7B report projections; Round 7A remains the engineering authority."""
from copy import deepcopy
from io import BytesIO
from unittest.mock import patch
from concurrent.futures import ThreadPoolExecutor
from zipfile import ZipFile
import builtins
import json
import hashlib

import numpy as np
import pandas as pd
import pytest
from docx import Document
from docx.oxml.ns import qn
from docxtpl import DocxTemplate
from PIL import Image

import phase_10_procedure as report
import test_audit_regressions as audit
import well_report as output
from centralizer_fixtures import ensure_current_centralizer
from project_state import refresh_well_derived


@pytest.fixture(scope="module")
def project():
    case = audit.AuditRegressions()
    app = case.configured_app('CSG 9 5/8"')
    restored = audit.round_trip(app.session_state.to_dict())
    ensure_current_centralizer(restored)
    return restored


def directional(project, count=31):
    state = deepcopy(project)
    state["well_geometry"] = {"version": 1, "type": "Directional", "survey": pd.DataFrame(
        [[i*3000./(count-1), 20., i*90./(count-1)] for i in range(count)],
        columns=["MD", "Deviation", "Azimuth"])}
    state["well_auto_fields"]["bhsp"] = False
    return state


def context(state):
    state = deepcopy(state)
    ensure_current_centralizer(state)
    with patch.object(report.st, "session_state", deepcopy(state)):
        return report.build_master_context()


def source_context(state):
    # Pure source/signature checks may inspect a newly changed, unreviewed
    # condition; actual export tests retain the normal prepare/review gate.
    prepared = deepcopy(state)
    refresh_well_derived(prepared)
    ensure_current_centralizer(prepared)
    with patch.object(report.st, "session_state", prepared):
        return report.build_master_context(calculations_prepared=True)


def actual_word(state):
    case = audit.AuditRegressions()
    app = case.app(audit.round_trip(state))
    case.export(app)
    return Document(BytesIO(app.session_state["_compiled_doc_bytes"]))


def rendered(ctx):
    template = DocxTemplate(report.TEMPLATE_PATH)
    template.render(report._word_quantity_context(ctx))
    buffer = BytesIO()
    template.save(buffer)
    return buffer.getvalue()


def compiled(ctx):
    return output.postprocess_well_report(rendered(ctx), ctx["well_profile_report"], ctx["note_vertical"])


@pytest.mark.parametrize("source", ["BHST", "Temperature Gradient"])
def test_incomplete_temperature_context_preserves_missing_projections(project, source):
    state = deepcopy(project)
    state["geothermal_config"].update(source=source, value=None)
    refresh_well_derived(state)
    assert state["_well_profile_issues"]
    with patch.object(output, "render_chart_png", side_effect=AssertionError("Incomplete chart rendered")) as renderer:
        with patch.object(report.st, "session_state", state):
            ctx = report.build_master_context(calculations_prepared=True)
        assert ctx["well_data"]["geo_gradient"] is None
        assert ctx["well_data"]["bhst"] is None
        assert ctx["well_profile_report"] is None
        assert state["geothermal_config"]["value"] is None
        with pytest.raises(ValueError, match="Complete Round 7A inputs"):
            compiled(ctx)
        renderer.assert_not_called()


@pytest.mark.parametrize("geometry", [
    {"version": 1, "type": None},
    {"version": 1, "type": "Vertical", "td_m": None},
    {"version": 1, "type": "Directional", "survey": pd.DataFrame(
        [[0., 0., 0.]], columns=["MD", "Deviation", "Azimuth"])},
])
def test_incomplete_geometry_context_does_not_fabricate_report_data(project, geometry):
    state = deepcopy(project)
    state["well_geometry"] = deepcopy(geometry)
    refresh_well_derived(state)
    with patch.object(output, "render_chart_png", side_effect=AssertionError("Incomplete chart rendered")) as renderer:
        assert output.prepared_well_report(state) is None
        assert state["well_data"]["geo_md"] is None
        assert state["well_data"]["geo_tvd"] is None
        assert state["well_data"]["geo_gradient"] is None
        assert state["well_data"]["bhst"] == state["geothermal_config"]["value"]
        with patch.object(report.st, "session_state", state), pytest.raises(ValueError, match="Phase IX"):
            report.build_master_context(calculations_prepared=True)
        renderer.assert_not_called()


@pytest.mark.parametrize("source", ["BHST", "Temperature Gradient"])
def test_fresh_incomplete_restore_direct_export_remains_blocked(project, source):
    state = deepcopy(project)
    state["geothermal_config"].update(source=source, value=None)
    case = audit.AuditRegressions()
    with patch.object(output, "render_chart_png", side_effect=AssertionError("Incomplete chart rendered")) as renderer:
        app = case.app(audit.round_trip(state))
        case.assert_export_blocked(app)
        assert not app.exception
        assert app.session_state["geo_gradient"] is None
        assert "_compiled_doc_bytes" not in app.session_state
        renderer.assert_not_called()


def chart_paragraph(doc):
    # Existing unrelated images remain; the chart occupies the Section III anchor.
    in_section = False
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("III."):
            in_section = True
        elif paragraph.text.startswith("IV."):
            break
        if in_section and paragraph._p.xpath(".//pic:pic"):
            return paragraph
    raise AssertionError("Section III combined chart missing")


def survey_table(doc):
    return next(t for t in doc.tables if t.rows[0].cells[0].text == "Directional Survey Data")


def test_vertical_word_has_combined_chart_and_real_survey_rows(project):
    doc = actual_word(project)
    assert "Temperature profile chart: Not provided." not in [p.text for p in doc.paragraphs]
    assert len(chart_paragraph(doc)._p.xpath(".//pic:pic")) == 1
    assert len(doc.inline_shapes) == len(Document(report.TEMPLATE_PATH).inline_shapes)+4
    assert len(doc._element.body.xpath('.//wp:docPr[@descr="Well Schematic"]')) == 1
    rows = [[c.text for c in row.cells] for row in survey_table(doc).rows[3:]]
    assert rows == [["0.0", "0.0", "0.00", "0.00", "0.000"],
                    ["3000.0", "3000.0", "0.00", "0.00", "0.000"]]


def test_directional_word_prints_twenty_actual_stations(project):
    doc = actual_word(directional(project))
    rows = survey_table(doc).rows[3:]
    assert len(rows) == 20
    assert rows[0].cells[0].text == "0.0"
    assert rows[-1].cells[0].text == "3000.0"


def test_full_survey_source_changes_document_signature(project):
    state = directional(project)
    before = report.document_signature(context(state), report.TEMPLATE_PATH)
    state["well_geometry"]["survey"].loc[1, "Azimuth"] += .01
    assert report.document_signature(context(state), report.TEMPLATE_PATH) != before


def test_full_geothermal_source_changes_document_signature(project):
    state = deepcopy(project)
    before = report.document_signature(context(state), report.TEMPLATE_PATH)
    state["geothermal_config"]["surface_temp"] += 1e-8
    assert report.document_signature(context(state), report.TEMPLATE_PATH) != before


def test_vertical_note_uses_global_sequence_after_lab_notes(project):
    ctx = context(project)
    assert ctx["all_notes"][-2].endswith("The Well is Considered Vertical")
    assert ctx["all_notes"][-1].endswith("Recommended Centralizer Performance Data is Based on API 10D Specifications.")
    numbers = [int(note.split(":", 1)[0].split()[1]) for note in ctx["all_notes"]]
    assert numbers == list(range(1, numbers[-1]+1))


@pytest.mark.parametrize("count", [2, 19, 20, 21, 31, 101])
def test_rows_are_real_ordered_unique_and_only_word_is_thinned(project, count):
    state = directional(project, count)
    ctx = context(state)
    payload = ctx["well_profile_report"]
    rows = payload["rows"]
    selected = output.select_survey_rows(rows)
    assert len(selected) == min(count, 20)
    assert selected[0] == rows[0] and selected[-1] == rows[-1]
    assert len({row[0] for row in selected}) == len(selected)
    assert all(row in rows for row in selected)
    assert [row[0] for row in selected] == sorted(row[0] for row in selected)
    spec = output.build_chart_spec(payload)
    assert len(spec["dep_data"][0]) == count
    np.testing.assert_array_equal(spec["dep_data"][1], [row[1] for row in rows])
    doc = Document(BytesIO(compiled(ctx)))
    printed = [[c.text for c in row.cells] for row in survey_table(doc).rows[3:]]
    assert printed == [[f"{md:.1f}", f"{tvd:.1f}", f"{inc:.2f}", f"{azi:.2f}", f"{dls/100:.3f}"]
                       for md, tvd, inc, azi, dls in selected]
    assert len(payload["rows"]) == count
    assert not ctx["note_vertical"]
    assert all("Considered Vertical" not in p.text for p in doc.paragraphs)
    pd.testing.assert_frame_equal(state["well_geometry"]["survey"], directional(project, count)["well_geometry"]["survey"])


def test_exact_target_algorithm_and_lower_md_tie():
    # Targets 10..180; first target has a tie at 9/11. 11 is still available
    # for subsequent targets, proving selection removes only the winning row.
    rows = [[md, md, 0., 0., 0.] for md in [0, 9, 11, *range(20, 181, 10), 185, 190]]
    assert [r[0] for r in output.select_survey_rows(rows)] == [0, 9, *range(20, 181, 10), 190]
    rows = [[md, md, 0., 0., 0.] for md in range(0, 3001, 100)]
    assert [r[0] for r in output.select_survey_rows(rows)] == [0, 200, 300, 500, 600, 800, 900,
              1100, 1300, 1400, 1600, 1700, 1900, 2100, 2200, 2400, 2500, 2700, 2800, 3000]


@pytest.mark.parametrize("raw,maximum,step,label,factor,labels", [
    (3927,4000,500,1000,100,["0","10","20","30","40"]),
    (5000,5000,1000,1000,100,["0","10","20","30","40","50"]),
    (6000,6000,1000,2000,100,["0","20","40","60"]),
    (8000,8000,1000,2000,100,["0","20","40","60","80"]),
    (2000,2000,250,500,100,["0","5","10","15","20"]),
    (1000,1000,200,200,1,["0","200","400","600","800","1000"]),
    (500,500,100,100,1,["0","100","200","300","400","500"]),
])
def test_v22_depth_acceptance(raw, maximum, step, label, factor, labels):
    spec = output.resolve_depth_axis(raw)
    assert (spec["depth_max"], spec["depth_step"], spec["depth_label_step"], spec["depth_factor"],
            spec["depth_tick_labels"]) == (maximum, step, label, factor, labels)
    assert all(0 < grid < maximum and grid/step == round(grid/step) for grid in spec["depth_grid"])


@pytest.mark.parametrize("low,high,minimum,maximum", [(80,240,50,250), (80,400,0,400),
    (80,600,0,800), (80,800,0,800), (60,300,0,400), (70,150,50,150)])
def test_v22_temperature_acceptance(low, high, minimum, maximum):
    step, left, right = output.pick_temp_axis(low, high)
    assert (left, right) == (minimum, maximum)
    assert (right-left)/step == 4


def test_v22_reference_labels_png_dimensions_and_colors():
    payload = {"rows": [[0.,0.,0.,0.,0.], [3927.,3927.,0.,0.,0.]], "departure": [0.,0.],
               "surface": 80.24, "bhst": 80.24+1.24*3927*3.28084/100}
    spec = output.build_chart_spec(payload)
    assert spec["depth_grid"] == list(range(500, 4000, 500))
    assert spec["temp_ticks"] == [50.,150.,250.]
    assert spec["temp_grid"] == [100.,150.,200.]
    assert spec["dep_grid"] == [.25,.5,.75]
    assert spec["dep_tick_labels"] == ["0","0.5","1.0"]
    assert spec["depth_label"] == "(x 100)   TVD (m)"
    assert (spec["dep_legend"], spec["temp_legend"]) == ("Hori. Depart.", "Geoth.Profile")
    assert output.fmt_ticks([0,250,500],100) == ["0","2.5","5.0"]
    image = Image.open(BytesIO(output.render_chart_png(payload)))
    assert image.size == (2000,1600)
    assert image.info["dpi"] == pytest.approx((200,200), abs=.01)
    pixels = np.array(image.convert("RGB"))
    for color, minimum in ((output.GRID_GRAY,500), (output.PANEL_BG,5000), (output.RED,500)):
        rgb = [int(color[i:i+2],16) for i in (1,3,5)]
        assert np.all(np.abs(pixels.astype(int)-rgb) <= 1,axis=-1).sum() > minimum


@pytest.mark.parametrize("kind", ["Vertical", "Directional"])
@pytest.mark.parametrize("source", ["BHST", "Temperature Gradient"])
def test_chart_uses_prepared_source_and_no_png_persists(project, kind, source):
    state = deepcopy(project) if kind == "Vertical" else directional(project)
    if source == "Temperature Gradient":
        state["geothermal_config"].update(source=source, value=1.3)
    ctx = source_context(state)
    payload = ctx["well_profile_report"]
    before = json.dumps({k:audit.namespace["serialize_item"](v) for k,v in state.items()}, sort_keys=True)
    spec = output.build_chart_spec(payload)
    assert spec["temp_data"][0][0] == state["geothermal_config"]["surface_temp"]
    assert spec["temp_data"][0][-1] == payload["bhst"]
    assert compiled(ctx)
    assert json.dumps({k:audit.namespace["serialize_item"](v) for k,v in state.items()}, sort_keys=True) == before
    saved = audit.round_trip(state)
    assert not any("png" in k.lower() or k == "well_profile_report" for k in saved)


def test_chart_retains_actual_nonmonotonic_tvd_and_endpoint_convention():
    payload = {"rows": [[0,0,0,0,0], [1000,900,80,0,0], [2000,800,110,0,0]],
               "departure": [0,500,1300], "surface": 80., "bhst": 200.}
    spec = output.build_chart_spec(payload)
    np.testing.assert_array_equal(spec["dep_data"][1], [0,900,800])
    assert spec["depth_max"] >= 900
    assert spec["temp_data"][0][-1] == 200.  # v2.2 endpoint hold beyond bottom TVD.


def test_concurrent_report_builds_are_isolated(project):
    contexts = [context(project), context(directional(project))]
    sequential = [compiled(ctx) for ctx in contexts]
    with ThreadPoolExecutor(max_workers=2) as pool:
        parallel = list(pool.map(compiled, contexts))
    def png(raw):
        doc = Document(BytesIO(raw))
        node = chart_paragraph(doc)._p.xpath(".//a:blip")[0]
        return doc.part.related_parts[node.get(qn("r:embed"))].blob
    assert [hashlib.sha256(png(raw)).hexdigest() for raw in parallel] == [
        hashlib.sha256(png(raw)).hexdigest() for raw in sequential]
    assert png(parallel[0]) != png(parallel[1])


@pytest.mark.parametrize("anchor", [output.CHART_ANCHOR, output.SURVEY_ANCHOR, output.SURVEY_TITLE])
@pytest.mark.parametrize("fault", ["missing", "duplicate"])
def test_exact_anchor_cardinality_fails_before_renderer(project, anchor, fault):
    ctx = context(project)
    doc = Document(BytesIO(rendered(ctx)))
    if anchor == output.SURVEY_TITLE:
        element = survey_table(doc)._tbl
    else:
        element = next(p._p for p in doc.paragraphs if p.text == anchor)
    if fault == "missing":
        element.getparent().remove(element)
    else:
        element.addnext(deepcopy(element))
    buffer = BytesIO()
    doc.save(buffer)
    with patch.object(output, "render_chart_png", side_effect=AssertionError("must validate anchors first")):
        with pytest.raises(ValueError, match="exactly one"):
            output.postprocess_well_report(buffer.getvalue(), ctx["well_profile_report"], ctx["note_vertical"])


@pytest.mark.parametrize("index", [0, 1, 3, 30])
def test_printed_and_unprinted_source_changes_signature_without_render(project, index):
    state = directional(project)
    with patch.object(output, "render_chart_png", side_effect=AssertionError("no render during status")):
        before = report.document_signature(context(state), report.TEMPLATE_PATH)
        state["well_geometry"]["survey"].loc[index,"Azimuth"] += .001
        assert report.document_signature(context(state), report.TEMPLATE_PATH) != before
    if index == 1:
        assert 100. not in [r[0] for r in output.select_survey_rows(context(state)["well_profile_report"]["rows"])]


@pytest.mark.parametrize("change", [{"source":"Temperature Gradient","value":1.3},
                                      {"surface_temp":80.00000001}, {"value":201.}])
def test_each_canonical_geothermal_source_stales_document(project, change):
    before = report.document_signature(context(project), report.TEMPLATE_PATH)
    state = deepcopy(project)
    state["geothermal_config"].update(change)
    assert report.document_signature(source_context(state), report.TEMPLATE_PATH) != before
    if "value" in change:
        with pytest.raises(ValueError, match="review Phase VII"):
            context(state)


def test_missing_renderer_is_controlled_and_import_is_lazy(project):
    original = builtins.__import__
    def blocked(name, *args, **kwargs):
        if name.startswith("matplotlib"):
            raise ImportError("renderer unavailable")
        return original(name, *args, **kwargs)
    with patch("builtins.__import__", side_effect=blocked):
        ctx = context(project)  # Context/status does not import Matplotlib.
        with pytest.raises(ValueError, match="matplotlib==3.10.8"):
            output.render_chart_png(ctx["well_profile_report"])


def test_unrelated_xml_package_and_survey_cell_properties_are_preserved(project):
    ctx = context(project)
    before = rendered(ctx)
    after = output.postprocess_well_report(before, ctx["well_profile_report"], ctx["note_vertical"])
    old, new = Document(BytesIO(before)), Document(BytesIO(after))
    def xml(element):
        from lxml import etree
        return None if element is None else etree.tostring(element, method="c14n")
    for a,b in zip(old.tables, new.tables):
        if a.rows[0].cells[0].text != output.SURVEY_TITLE:
            assert xml(a._tbl) == xml(b._tbl)
    a,b = survey_table(old), survey_table(new)
    assert xml(a._tbl.tblPr) == xml(b._tbl.tblPr)
    assert xml(a._tbl.tblGrid) == xml(b._tbl.tblGrid)
    for i in range(3):
        assert xml(a.rows[i]._tr) == xml(b.rows[i]._tr)
    for row in b.rows[3:]:
        for original, clone in zip(a.rows[3].cells, row.cells):
            assert xml(original._tc.tcPr) == xml(clone._tc.tcPr)
            assert xml(original.paragraphs[0]._p.pPr) == xml(clone.paragraphs[0]._p.pPr)
            assert xml(original.paragraphs[0].runs[0]._r.rPr) == xml(clone.paragraphs[0].runs[0]._r.rPr)
    assert [xml(s._sectPr) for s in old.sections] == [xml(s._sectPr) for s in new.sections]
    with ZipFile(BytesIO(before)) as original, ZipFile(BytesIO(after)) as changed:
        for name in original.namelist():
            if name not in ("word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"):
                if name.endswith(".xml"):
                    from lxml import etree
                    assert xml(etree.fromstring(original.read(name))) == xml(etree.fromstring(changed.read(name))), name
                else:
                    assert original.read(name) == changed.read(name), name
    assert output.CHART_ANCHOR not in "\n".join(p.text for p in new.paragraphs)
    assert output.SURVEY_ANCHOR not in "\n".join(p.text for p in new.paragraphs)


def test_vertical_note_participates_in_collision_validation(project):
    ctx = context(project)
    state = deepcopy(project)
    # Existing manual numbering validation must see the new final note.
    state["procedure_text"] = ctx["note_vertical"]
    with pytest.raises(ValueError, match="NOTE.*conflicts"):
        report.validate_manual_numbering(state, {"procedure_text":{"generated":""},
             "exec_summary_text":{"generated":state.get("exec_summary_text","")}}, ctx["all_notes"])


@pytest.mark.parametrize("kind", ["Vertical", "Directional"])
def test_only_six_authorized_spacers_removed_break_and_ix_unchanged(project, kind):
    from lxml import etree
    ctx = context(project if kind == "Vertical" else directional(project))
    before = rendered(ctx)
    old = Document(BytesIO(before))
    body = list(old._element.body)
    anchor = next(p._p for p in old.paragraphs if p.text == output.SURVEY_ANCHOR)
    ix = next(p._p for p in old.paragraphs if p.text.strip() == "IX. Centralizer Placement Recommendation")
    page_break = ix.getprevious()
    region = body[body.index(anchor)+1:body.index(page_break)]
    spacers = [p for p in region if p.xpath('./w:pPr/w:pStyle[@w:val="Heading1"]')]
    assert len(spacers) == 6
    def xml(element):
        return etree.tostring(element, method="c14n")
    after = Document(BytesIO(output.postprocess_well_report(before, ctx["well_profile_report"], ctx["note_vertical"])))
    new_ix = next(p._p for p in after.paragraphs if p.text.strip() == "IX. Centralizer Placement Recommendation")
    assert xml(new_ix) == xml(ix)
    assert xml(new_ix.getprevious()) == xml(page_break)
    permitted_old = {anchor, survey_table(old)._tbl, next(p._p for p in old.paragraphs if p.text == output.CHART_ANCHOR), *spacers}
    permitted_new = {survey_table(after)._tbl, chart_paragraph(after)._p}
    if kind == "Vertical":
        permitted_new.add(next(p._p for p in after.paragraphs if p.text == ctx["note_vertical"]))
    assert [xml(e) for e in body if e not in permitted_old] == [
        xml(e) for e in after._element.body if e not in permitted_new]


@pytest.mark.parametrize("count", [5, 7])
def test_spacing_cardinality_mismatch_stops_before_render(project, count):
    ctx = context(project)
    doc = Document(BytesIO(rendered(ctx)))
    body = list(doc._element.body)
    anchor = next(p._p for p in doc.paragraphs if p.text == output.SURVEY_ANCHOR)
    ix = next(p._p for p in doc.paragraphs if p.text.strip() == "IX. Centralizer Placement Recommendation")
    spacers = [p for p in body[body.index(anchor)+1:body.index(ix.getprevious())]
               if p.xpath('./w:pPr/w:pStyle[@w:val="Heading1"]')]
    assert len(spacers) == 6
    if count == 5:
        spacers[0].getparent().remove(spacers[0])
    else:
        spacers[0].addnext(deepcopy(spacers[0]))
    buffer = BytesIO(); doc.save(buffer)
    with patch.object(output, "render_chart_png", side_effect=AssertionError("must validate spacing first")):
        with pytest.raises(ValueError, match="exactly six"):
            output.postprocess_well_report(buffer.getvalue(), ctx["well_profile_report"], ctx["note_vertical"])
