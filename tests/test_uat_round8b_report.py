"""Round 8B consumes prepared engineering; only diagnostic presentation changes."""
from copy import deepcopy
from io import BytesIO
from concurrent.futures import ThreadPoolExecutor
from unittest.mock import patch
from zipfile import ZipFile
import importlib
import importlib.util
import hashlib
import json
from pathlib import Path

import numpy as np
import pytest
from docx import Document
from PIL import Image

import cement_chart_model as model
import phase_10_procedure as report
import well_report
import test_audit_regressions as audit
from test_uat_round8a_chart_model import chart_state, hardware
from test_uat_round7b_report import rendered
from project_state import prepare_calculations

ANCHOR = 'Pressure and fluid sequence charts: Not provided.'
PUMPING = 'Pumping simulation charts: Not provided.'
NA = 'Pressure and fluid sequence charts: Not applicable for this job type.'


def renderer():
    assert importlib.util.find_spec('cement_chart_report'), 'Round 8B presentation renderer missing'
    return importlib.import_module('cement_chart_report')


def prepared(job='CSG 9 5/8"'):
    state = chart_state()
    state["job_type"] = job
    if job != 'CSG 9 5/8"':
        state.update(audit.chart_fixture(job, 1000.))
    if job.startswith('LNR'):
        state['hardware_table'] = hardware(('Previous Casing','500','13.375',12.25),
            ('Liner','1000','7',6.276),('Open Hole Size','1000','12.25',12.25),
            ('Liner Hanger','300','9.625',8.535),('Drill Pipes','300','5',4.276))
    if job.startswith('TIE BACK'):
        from placement import hardware_choices, HOST_DESCRIPTIONS
        state['placement_config']['host_row'] = next(iter(hardware_choices(state['hardware_table'], HOST_DESCRIPTIONS)))
    state['_cement_chart_model'] = model.build_chart_model(state)
    return state


@pytest.fixture(scope='module')
def project():
    return audit.round_trip(audit.AuditRegressions().configured_app('CSG 9 5/8"').session_state.to_dict())


def context(state):
    with patch.object(report.st, 'session_state', deepcopy(state)):
        return report.build_master_context()


def final(ctx):
    return well_report.postprocess_well_report(rendered(ctx), ctx['well_profile_report'],
        ctx['note_vertical'], ctx['cement_chart_report'])


@pytest.mark.parametrize('job', ['CSG 9 5/8"','LNR 7"','TIE BACK LNR 7"'])
def test_render_jobs_consume_prepared_arrays_and_real_geometry(job):
    state = prepared(job)
    original = deepcopy(state['_cement_chart_model'])
    output = renderer()
    with patch.object(model, 'resolve_chart_geometry', side_effect=AssertionError('renderer recalculated geometry')):
        figure = output.build_chart_figure(original)
    try:
        left, right = figure.axes
        assert left.get_xlabel() == '(x 1000) psi'
        assert left.get_ylabel() == 'TVD (m)'
        lines = {line.get_label(): line for line in left.lines}
        np.testing.assert_array_equal(lines['Hydrostatic'].get_xdata(), np.array(original['hydrostatic']['pressure_psi'])/1000)
        np.testing.assert_array_equal(lines['Hydrostatic'].get_ydata(), original['hydrostatic']['tvd_m'])
        for label,key in [('Pore','pore_psi'),('Frac','frac_psi')]:
            np.testing.assert_array_equal(lines[label].get_xdata(), np.array(original['pore_frac'][key])/1000)
            np.testing.assert_array_equal(lines[label].get_ydata(), original['pore_frac']['tvd_m'])
        assert right.get_ylabel() == 'MD (m)'
        assert not any('NOTE' in t.get_text() for t in [*figure.texts,*left.texts,*right.texts])
        assert original == state['_cement_chart_model']
    finally:
        figure.clear()
    png = output.render_chart_png(original)
    assert Image.open(BytesIO(png)).size == (2300,1680)


def test_resolved_outer_geometry_is_projection_not_new_math():
    state = prepared()
    actual = state['_cement_chart_model']['enclosing_string']
    assert [(s['top_md'],s['bottom_md'],s['od_in'],s['id_in']) for s in actual] == [
        (0.,500.,13.375,12.25), (500.,1000.,12.25,12.25)]
    assert model.resolve_chart_geometry(state)['enclosing_string'] == actual


def test_all_final_fluids_shoe_track_and_excluded_stages():
    state = prepared()
    names = ['Pre Flush','Spacer','Spacer Ahead','Scavenger','Lead #1','Tail','Spacer Behind','Displacement Fluid']
    state['fluids_config']['active'] = names
    for name in names:state['fluids_config']['params'][name]={'volume':10.,'density':'118','pump_rate':'4'}
    m = model.build_chart_model(state)
    figure = renderer().build_chart_figure(m)
    try:
        labels=[t.get_text() for t in figure.axes[1].texts]
        for name in ['Pre Flush','Spacer','Spacer Ahead','Scavenger','Lead #1','Tail','Mud']:
            assert name in labels
        assert 'Spacer Behind' not in labels and 'Displacement Fluid' not in labels
        assert not any(text.startswith('Shoe Track:') for text in labels)
        assert any(p.get_gid() == 'shoe-track' for p in figure.axes[1].patches)
    finally:figure.clear()


def test_surface_returns_are_not_drawn_inside_well():
    state=prepared();state['fluids_config']['params']['Main']['volume']=10000.
    m=model.build_chart_model(state)
    assert m['returned_bbl']['Main']>0
    figure=renderer().build_chart_figure(m)
    try:
        bands=[p for p in figure.axes[1].patches if p.get_gid() and p.get_gid().startswith('annulus:')]
        assert all(p.get_y()>=0 and p.get_y()+p.get_height()<=m['shoe_md'] for p in bands)
        assert not any('returned' in t.get_text().lower() for t in figure.axes[1].texts)
    finally:figure.clear()


def test_overrides_do_not_change_model_or_png(project):
    state=deepcopy(project);assert not prepare_calculations(state)
    a=renderer().render_chart_png(state['_cement_chart_model'])
    state['cement_params']['Main'].update(top_mode='Depth (m MD)',top_depth=1234.,report_bottom_depth=2345.)
    assert not prepare_calculations(state)
    assert renderer().render_chart_png(state['_cement_chart_model'])==a


def test_fresh_restore_direct_context_and_actual_word(project):
    state=audit.round_trip(project);ctx=context(state)
    rebuilt=deepcopy(state)
    assert not prepare_calculations(rebuilt)
    assert ctx['cement_chart_report']['model']==rebuilt['_cement_chart_model']
    doc=Document(BytesIO(final(ctx)))
    assert len(doc.inline_shapes)==4
    paragraphs=[p.text for p in doc.paragraphs]
    assert ANCHOR not in paragraphs and paragraphs.count(PUMPING)==1
    anchor = next(p for p in doc.paragraphs if p._p.xpath('.//wp:docPr[@descr="Static Well Security and Fluid Sequence"]'))
    assert len(anchor._p.xpath('.//w:drawing'))==1
    assert not any('{{' in p.text or '{%' in p.text for p in doc.paragraphs)


@pytest.mark.parametrize('job',['CMT PLUG','CMT SQUEEZE'])
def test_special_jobs_not_applicable_actual_word(job):
    state=audit.round_trip(audit.AuditRegressions().configured_app(job).session_state.to_dict())
    doc=Document(BytesIO(final(context(state))))
    text=[p.text for p in doc.paragraphs]
    assert text.count(NA)==1 and ANCHOR not in text and text.count(PUMPING)==1
    assert len(doc.inline_shapes)==3


@pytest.mark.parametrize('fault',['missing','duplicate'])
def test_exact_anchor_guard_before_any_render(project,fault):
    ctx=context(project);doc=Document(BytesIO(rendered(ctx)))
    p=next(p for p in doc.paragraphs if p.text==ANCHOR)
    if fault=='missing':p.text='wrong anchor'
    else:doc.add_paragraph(ANCHOR)
    b=BytesIO();doc.save(b)
    with patch.object(well_report,'render_chart_png',side_effect=AssertionError('anchor guard too late')):
        with pytest.raises(ValueError,match='exactly one.*Pressure'):
            well_report.postprocess_well_report(b.getvalue(),ctx['well_profile_report'],ctx['note_vertical'],ctx['cement_chart_report'])


def test_decreasing_tvd_still_blocks_export_and_no_partial_word(project):
    import pandas as pd
    state=deepcopy(project)
    state['well_geometry']={'version':1,'type':'Directional','survey':pd.DataFrame(
        [[0.,0.,0.],[1500.,100.,0.],[3000.,80.,0.]],columns=['MD','Deviation','Azimuth'])}
    state['well_auto_fields']['bhsp']=False
    issues=prepare_calculations(state)
    assert any('TVD decreases' in issue for issue in issues)
    assert state['_cement_chart_model']['fluid_sequence_ready']
    assert not state['_cement_chart_model']['static_security_ready']
    with pytest.raises(ValueError,match='unsupported'):
        renderer().render_chart_png(state['_cement_chart_model'])
    assert '_compiled_doc_bytes' not in state


def test_json_has_no_images_models_and_source_change_invalidates(project):
    state=audit.round_trip(project);assert not prepare_calculations(state)
    ctx=context(state)
    state['_compiled_doc_bytes']=final(ctx)
    old=report.document_signature(ctx,report.TEMPLATE_PATH)
    saved=audit.round_trip(state)
    assert not any('chart' in k and k.startswith('_') for k in saved)
    assert 'cement_chart_report' not in saved
    state['pore_gradient_emw_pcf']=75.
    assert not prepare_calculations(state) and '_compiled_doc_bytes' not in state
    assert report.document_signature(context(state),report.TEMPLATE_PATH)!=old


def test_no_global_pyplot_or_reference_engine_import_and_concurrency():
    out=renderer();source=Path(out.__file__).read_text()
    assert 'import matplotlib.pyplot' not in source and 'cement_charts_engine' not in source
    models=[prepared()['_cement_chart_model'],prepared('TIE BACK LNR 7"')['_cement_chart_model']]
    sequential=[out.render_chart_png(m) for m in models]
    with ThreadPoolExecutor(max_workers=2) as pool:parallel=list(pool.map(out.render_chart_png,models))
    assert parallel==sequential and sequential[0]!=sequential[1]


def test_incomplete_model_is_controlled_and_no_image_rendered(project):
    state=deepcopy(project);state['pore_gradient_emw_pcf']=None
    with patch.object(report.st,'session_state',state),pytest.raises(ValueError,match='Gradient'):
        report.build_master_context()
    assert '_compiled_doc_bytes' not in state
    with pytest.raises(ValueError,match='prepared'):
        renderer().render_chart_png(None)


def test_template_dependency_and_protected_sources_identical():
    import subprocess
    for name in ['master_template.docx','requirements.txt','well_profile.py','placement.py','rheology.py','bhct_helper.py','test_regression_guardrails.py']:
        assert Path(name).read_bytes()==subprocess.check_output(['git','show','a61d61e:'+name])


def test_complete_word_preserves_unrelated_sections_xml_and_notes(project):
    from lxml import etree
    ctx=context(project)
    intermediate=well_report.postprocess_well_report(rendered(ctx),ctx['well_profile_report'],ctx['note_vertical'])
    before=Document(BytesIO(intermediate));after=Document(BytesIO(final(ctx)))
    def xml(node):return etree.tostring(node,method='c14n')
    old=list(before._element.body);new=list(after._element.body)
    assert len(old)==len(new)
    changed=[]
    for index,(a,b) in enumerate(zip(old,new)):
        if xml(a)!=xml(b):changed.append(index)
    assert len(changed)==1
    assert ANCHOR in ''.join(old[changed[0]].itertext())
    assert len(new[changed[0]].xpath('.//w:drawing'))==1
    assert [xml(s._sectPr) for s in before.sections]==[xml(s._sectPr) for s in after.sections]
    assert ctx['all_notes']==context(project)['all_notes']
    with ZipFile(BytesIO(intermediate)) as a,ZipFile(BytesIO(final(ctx))) as b:
        for name in a.namelist():
            if name not in ('word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'):
                assert a.read(name)==b.read(name),name


def test_png_is_only_in_memory_and_renderer_import_is_lazy(tmp_path,monkeypatch):
    import builtins
    output=renderer();m=prepared()['_cement_chart_model']
    monkeypatch.chdir(tmp_path)
    assert output.render_chart_png(m).startswith(b'\x89PNG')
    assert list(tmp_path.iterdir())==[]
    original=builtins.__import__
    def unavailable(name,*args,**kwargs):
        if name.startswith('matplotlib'):raise ImportError('test missing renderer')
        return original(name,*args,**kwargs)
    with patch('builtins.__import__',side_effect=unavailable):
        assert output.prepared_cement_report(prepared())['model']==m
        with pytest.raises(ValueError,match='matplotlib'):
            output.render_chart_png(m)


def test_ui_build_and_direct_rebuilt_chart_image_are_identical(project):
    case=audit.AuditRegressions();app=case.app(audit.round_trip(project));case.export(app)
    direct=final(context(audit.round_trip(project)))
    def image(data):
        doc=Document(BytesIO(data))
        p=next(p for p in doc.paragraphs if p._p.xpath('.//wp:docPr[@descr="Static Well Security and Fluid Sequence"]'))
        rid=p._p.xpath('.//a:blip')[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
        return doc.part.related_parts[rid].blob
    assert image(direct)==image(app.session_state['_compiled_doc_bytes'])


@pytest.mark.parametrize('fluid,color', [('Mud','#8B1010'),('Scavenger','#C6D2BD')])
def test_round8c_fluid_colors_are_shared_by_all_bands_and_key(fluid,color):
    from matplotlib.colors import to_hex
    state=prepared()
    state['fluids_config']['active'].insert(0,'Scavenger')
    state['fluids_config']['params']['Scavenger']={'volume':10.,'density':'118','pump_rate':'4'}
    figure=renderer().build_chart_figure(model.build_chart_model(state))
    try:
        axis=figure.axes[1]
        bands=[p for p in axis.patches if p.get_gid()=='annulus:'+fluid]
        assert bands and all(to_hex(p.get_facecolor()).upper()==color for p in bands)
        labels=[t.get_text() for t in axis.texts]
        keys=[line for line in axis.lines if line.get_transform()==axis.transAxes]
        assert to_hex(keys[labels.index(fluid)].get_color()).upper()==color
        assert renderer().COLORS[fluid]==color
    finally:figure.clear()


def test_round8c_panel_titles_share_rendered_baseline():
    figure=renderer().build_chart_figure(prepared()['_cement_chart_model'])
    try:
        figure.canvas.draw()
        left,right=figure.axes
        assert left.title.get_text()=='Static Well Security'
        assert right.title.get_text()=='Fluid Sequence'
        assert left.title.get_position()[1]==right.title.get_position()[1]
        bounds=[axis.title.get_window_extent(figure.canvas.get_renderer()) for axis in (left,right)]
        assert bounds[0].y0==pytest.approx(bounds[1].y0,abs=0.01)
        assert bounds[0].y1==pytest.approx(bounds[1].y1,abs=0.01)
    finally:figure.clear()
