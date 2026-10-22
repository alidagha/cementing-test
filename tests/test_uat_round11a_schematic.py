"""Round 11A: canonical well construction and fresh Word schematic integration."""
from copy import deepcopy
import hashlib
import importlib.util
from pathlib import Path
import pickle

import pandas as pd
import pytest

import materials_db
import pumping_simulation as pumping
from placement import hardware_choices
from test_uat_round10a_pumping_simulation import program_state

TEMPLATE_SHA = "d0f2c324cef018b87b486a3cc80ac5c2ab2274b84eb4b2514e56fb374e766487"


def remaining_case(job, change):
    import test_audit_regressions as audit
    state = {"job_type": job, **audit.valid_export_chart_fixture(job)}
    if change == "add_hole":
        row = {**state["hardware_table"].iloc[0].to_dict(), "Description": "Open Hole Size",
               "MD (m)": "3500", "Size (in)": "12.25", "ID (in)": 12.25}
        state["hardware_table"] = pd.concat([state["hardware_table"], pd.DataFrame([row])], ignore_index=True)
    elif change == "shorten":
        state["hardware_table"].loc[0, "MD (m)"] = "2950"
    state["fluids_config"] = {"active": ["Main", "Displacement Fluid"]}
    state["cement_params"] = {"Main": {"top_mode": "Surface", "top_depth": 0., "top_job_type": job}}
    return state


def test_plug_without_main_preserves_every_canonical_independent_interval():
    import well_schematic_adapter as schematic
    from placement import slurry_intervals
    state = remaining_case("CMT PLUG", "unchanged")
    state["fluids_config"]["active"] = ["Lead", "Tail", "Displacement Fluid"]
    state["cement_params"] = {name: {"top_mode": "Depth (m MD)", "top_depth": depth,
                                    "top_job_type": "CMT PLUG"} for name, depth in (("Lead", 1000.), ("Tail", 2500.))}
    before = pickle.dumps(state)
    intervals = slurry_intervals(state["hardware_table"], state["job_type"], state["placement_config"],
                                 ["Lead", "Tail"], state["cement_params"])
    payload = schematic.prepare_schematic(state)
    assert payload["cement_plugs"] == [{"top_md": row["top_depth"], "bottom_md": row["bottom_depth"]}
                                        for row in intervals.values()]
    assert pickle.dumps(state) == before


@pytest.mark.parametrize("job", ['CSG 9 5/8"', 'LNR 7"', 'TIE BACK LNR 7"'])
def test_compatible_current_open_hole_segments_and_unique_clearance_host(job):
    import well_schematic_adapter as schematic
    import well_schematic_engine as engine
    state = remaining_case(job, "add_hole")
    before = pickle.dumps(state)
    payload = schematic.prepare_schematic(state)
    assert len(payload["open_holes"]) == 1
    hole = payload["open_holes"][0]
    assert (hole["top_md"], hole["bottom_md"], hole["hole_size_in"]) == (
        3300. if "TIE BACK" in job else 1500., 3500., 12.25)
    assert payload["casings"][hole["host_index"]]["size_in"] == 13.375
    assert payload["final_relevant_depth_m"] == 3500.
    assert engine.check_well_design(payload) == []
    assert pickle.dumps(state) == before


@pytest.mark.parametrize("target", ["previous-casing:20", "open-hole"])
def test_boundary_squeeze_requires_explicit_presentation_identity(target, tmp_path):
    import well_schematic_adapter as schematic
    import well_schematic_engine as engine
    state = remaining_case("CMT SQUEEZE", "add_hole")
    with pytest.raises(ValueError, match="Squeeze Target for Well Schematic"):
        schematic.prepare_schematic(state)
    inputs = schematic.additional_drawing_inputs(state)
    choices = inputs["squeeze_targets"]
    assert {row["target_id"] for row in choices.values()} == {"previous-casing:20", "open-hole"}
    selected = next(token for token, row in choices.items() if row["target_id"] == target)
    state["schematic_config"] = {"version": 1, "squeeze_target_id": selected}
    before = pickle.dumps(state)
    payload = schematic.prepare_schematic(state)
    squeeze = payload["squeezes"][0]
    assert squeeze["depth_m"] == 3000.
    if target == "open-hole":
        assert squeeze == {"depth_m": 3000., "open_hole_index": 0, "target_kind": "point"}
    else:
        assert squeeze == {"depth_m": 3000., "target_index": 0, "target_kind": "shoe"}
    engine.render_well_schematic(payload, str(tmp_path / "squeeze.png"), dpi=90)
    assert pickle.dumps(state) == before


def test_unique_squeeze_and_known_plug_bore_require_no_duplicate_input():
    import well_schematic_adapter as schematic
    for job in ("CMT SQUEEZE", "CMT PLUG"):
        assert schematic.additional_drawing_inputs(remaining_case(job, "unchanged")) == {}


@pytest.mark.parametrize("job", ["CMT PLUG", "CMT SQUEEZE"])
def test_missing_operation_wall_consumes_only_scoped_diameter(job, tmp_path):
    import well_schematic_adapter as schematic
    import well_schematic_engine as engine
    import test_audit_regressions as audit
    state = remaining_case(job, "shorten")
    original_hardware = state["hardware_table"].copy(deep=True)
    original_placement = deepcopy(state["placement_config"])
    inputs = schematic.additional_drawing_inputs(state)
    if job == "CMT PLUG":
        assert inputs == {"plug_bores": [{"top_md": 2950., "bottom_md": 3000.}]}
        config = {"version": 1, "plug_bore_sections": [{"top_md": 2950., "bottom_md": 3000., "diameter_in": 8.5}]}
    else:
        assert inputs == {"squeeze_wall_depth": 3000.}
        config = {"version": 1, "squeeze_wall": {"depth_m": 3000., "diameter_in": 8.5}}
    state["schematic_config"] = config
    restored = audit.round_trip(state)
    assert restored["schematic_config"] == config
    payload = schematic.prepare_schematic(restored)
    assert len(payload["casings"]) == 1 and payload["casings"][0]["shoe_md"] == 2950.
    assert payload["open_holes"] == []  # No invented construction or wall classification.
    if job == "CMT PLUG":
        assert payload["plug_bores"] == config["plug_bore_sections"]
        assert payload["cement_plugs"] == [{"top_md": 0., "bottom_md": 3000.}]
    else:
        assert payload["squeezes"] == [{"depth_m": 3000., "wall_diameter_in": 8.5, "target_kind": "point"}]
    engine.render_well_schematic(payload, str(tmp_path / "operation.png"), dpi=90)
    pd.testing.assert_frame_equal(state["hardware_table"], original_hardware)
    assert state["placement_config"] == original_placement


@pytest.mark.parametrize("bad", [[], {"version": True}, {"version": 2}, {"version": 1, "payload": {}},
    {"version": 1, "squeeze_target_id": 3},
    {"version": 1, "squeeze_wall": {"depth_m": True, "diameter_in": 8.5}},
    {"version": 1, "squeeze_wall": {"depth_m": 3000., "diameter_in": 0}},
    {"version": 1, "plug_bore_sections": [{"top_md": 3000., "bottom_md": 2950., "diameter_in": 8.5}]},
    {"version": 1, "plug_bore_sections": [{"top_md": 2950., "bottom_md": 3000., "diameter_in": "8.5"}]}])
def test_schematic_config_trust_boundary_rejects_malformed_authority(bad):
    import test_audit_regressions as audit
    with pytest.raises(ValueError, match="schematic_config"):
        audit.round_trip({"job_type": "CMT PLUG", "schematic_config": bad})


def test_explicit_missing_bore_cannot_override_canonical_bore():
    import well_schematic_adapter as schematic
    state = remaining_case("CMT PLUG", "unchanged")
    state["schematic_config"] = {"version": 1, "plug_bore_sections": [
        {"top_md": 2950., "bottom_md": 3000., "diameter_in": 8.5}]}
    with pytest.raises(ValueError, match="contradicts.*canonical bore"):
        schematic.prepare_schematic(state)


def test_explicit_schematic_target_cannot_replace_an_existing_app_target():
    import well_schematic_adapter as schematic
    state = remaining_case("CMT SQUEEZE", "add_hole")
    choices = schematic.additional_drawing_inputs(state)["squeeze_targets"]
    token = next(token for token, row in choices.items() if row["target_id"] == "open-hole")
    state["schematic_config"] = {"version": 1, "squeeze_target_id": token}
    state["placement_config"]["target_string_id"] = "previous-casing:20"
    with pytest.raises(ValueError, match="contradicts.*canonical App target"):
        schematic.prepare_schematic(state)


def test_obsolete_schematic_selection_requires_explicit_clear_without_silent_redirect():
    import test_audit_regressions as audit
    import well_schematic_adapter as schematic
    state = remaining_case("CMT SQUEEZE", "add_hole")
    choices = schematic.additional_drawing_inputs(state)["squeeze_targets"]
    token = next(token for token, row in choices.items() if row["target_id"] == "open-hole")
    state["schematic_config"] = {"version": 1, "squeeze_target_id": token}
    state["hardware_table"].loc[0, "MD (m)"] = "3100"
    case = audit.AuditRegressions()
    app = case.app(state)
    case.phase(app, "phase2_3")
    with pytest.raises(ValueError, match="Squeeze Target for Well Schematic"):
        schematic.prepare_schematic(app.session_state.to_dict())
    assert not any(w.label == "Squeeze Target for Well Schematic" for w in app.selectbox)
    next(b for b in app.button if b.label == "Clear obsolete schematic target").click().run()
    assert "squeeze_target_id" not in app.session_state["schematic_config"]
    assert schematic.prepare_schematic(app.session_state.to_dict())["squeezes"] == [
        {"depth_m": 3000., "target_index": 0, "target_kind": "point"}]


def test_conditional_target_ui_commits_only_presentation_state_and_survives_restore():
    import test_audit_regressions as audit
    import well_schematic_adapter as schematic
    case = audit.AuditRegressions()
    app = case.app(remaining_case("CMT SQUEEZE", "unchanged"))
    case.phase(app, "phase2_3")
    assert not any(w.label == "Squeeze Target for Well Schematic" for w in app.selectbox)
    assert not any(w.label == "Squeeze Depth (m MD)" for w in app.number_input)
    app = case.app(remaining_case("CMT SQUEEZE", "add_hole"))
    case.phase(app, "phase2_3")
    placement = deepcopy(app.session_state["placement_config"])
    hardware = app.session_state["hardware_table"].copy(deep=True)
    choices = schematic.additional_drawing_inputs(app.session_state)["squeeze_targets"]
    token = next(token for token, row in choices.items() if row["target_id"] == "open-hole")
    selector = next(w for w in app.selectbox if w.label == "Squeeze Target for Well Schematic")
    assert selector.value == ""
    app.session_state["_compiled_doc_bytes"] = b"obsolete-download"
    selector.set_value(token).run()
    assert "_compiled_doc_bytes" not in app.session_state
    assert app.session_state["schematic_config"]["squeeze_target_id"] == token
    assert app.session_state["placement_config"] == placement
    pd.testing.assert_frame_equal(app.session_state["hardware_table"], hardware)
    restored = audit.round_trip(app.session_state.to_dict())
    assert restored["schematic_config"] == app.session_state["schematic_config"]
    case.phase(app, "phase1"); case.phase(app, "phase2_3")
    assert next(w for w in app.selectbox if w.label == "Squeeze Target for Well Schematic").value == token


@pytest.mark.parametrize("job", ["CMT PLUG", "CMT SQUEEZE"])
def test_only_missing_operation_dimension_appears_below_hardware(job):
    import test_audit_regressions as audit
    case = audit.AuditRegressions()
    app = case.app(remaining_case(job, "shorten"))
    case.phase(app, "phase2_3")
    labels = [widget.label for widget in app.number_input]
    expected = "Plug bore diameter (in) — 2950–3000 m" if job == "CMT PLUG" else "Squeeze wall diameter (in) — 3000 m"
    assert expected in labels
    next(widget for widget in app.number_input if widget.label == expected).set_value(8.5).run()
    assert "schematic_config" in app.session_state
    assert "Squeeze Depth (m MD)" not in [widget.label for widget in app.number_input]


def test_schematic_target_scope_rejects_obsolete_selection_and_canonical_target_wins():
    import well_schematic_adapter as schematic
    state = remaining_case("CMT SQUEEZE", "add_hole")
    token = next(token for token, row in schematic.additional_drawing_inputs(state)["squeeze_targets"].items()
                 if row["target_id"] == "open-hole")
    state["schematic_config"] = {"version": 1, "squeeze_target_id": token}
    state["hardware_table"].loc[0, "MD (m)"] = "3100"
    with pytest.raises(ValueError, match="Squeeze Target for Well Schematic"):
        schematic.prepare_schematic(state)
    assert schematic.additional_drawing_inputs(state)["obsolete_squeeze_target"]
    state["placement_config"]["target_string_id"] = "previous-casing:20"
    assert schematic.prepare_schematic(state)["squeezes"][0]["target_index"] == 0


def test_still_valid_physical_squeeze_target_survives_shoe_edit_without_duplicate_input():
    import well_schematic_adapter as schematic
    state = remaining_case("CMT SQUEEZE", "add_hole")
    choices = schematic.additional_drawing_inputs(state)["squeeze_targets"]
    token = next(token for token, row in choices.items() if row["target_id"] == "previous-casing:20")
    state["schematic_config"] = {"version": 1, "squeeze_target_id": token}
    signature = schematic.source_signature(state)
    assert schematic.prepare_schematic(state)["squeezes"][0]["target_kind"] == "shoe"
    state["hardware_table"].loc[0, "MD (m)"] = "3100"
    assert schematic.source_signature(state) != signature
    assert schematic.additional_drawing_inputs(state) == {}
    assert schematic.prepare_schematic(state)["squeezes"] == [
        {"depth_m": 3000., "target_index": 0, "target_kind": "point"}]
    assert state["schematic_config"]["squeeze_target_id"] == token


def test_schematic_only_input_never_changes_frozen_preparation_or_calculation_sources():
    import pumping_simulation as pumping
    import cement_chart_model as chart
    import centralizer_adapter as centralizer
    import phase_9_centralizer as phase9
    from centralizer_fixtures import execute_centralizer
    state = program_state(fluids={"Tail": 5., "Displacement Fluid": 2.})
    state.update({chart.PORE_FIELD: 70., chart.FRAC_FIELD: 160.})
    execute_centralizer(state)
    before = (pumping.source_signature(state), repr(pumping.prepare_pumping_simulation(state)),
              chart.build_chart_model(state), centralizer.to_engine_geometry(state),
              phase9.source_signature(state), pickle.dumps(state['_centralizer_result']))
    state["schematic_config"] = {"version": 1, "squeeze_target_id": "0" * 64, "squeeze_wall": {"depth_m": 3000., "diameter_in": 8.5}}
    after = (pumping.source_signature(state), repr(pumping.prepare_pumping_simulation(state)),
             chart.build_chart_model(state), centralizer.to_engine_geometry(state),
             phase9.source_signature(state), pickle.dumps(state['_centralizer_result']))
    assert after == before
    assert phase9.result_is_current(state)


@pytest.mark.parametrize("case_name", ["plug_bore", "squeeze_wall", "casing_target", "hole_target"])
def test_real_word_build_consumes_explicit_drawing_authority_and_rebuilds_fresh(case_name, monkeypatch, tmp_path):
    from io import BytesIO
    from docx import Document
    import test_audit_regressions as audit
    import well_schematic_adapter as schematic
    import well_report
    from test_uat_round10c_integration import save_project, restore_project
    job = "CMT PLUG" if case_name == "plug_bore" else "CMT SQUEEZE"
    case = audit.AuditRegressions()
    app = case.configured_app(job)
    app.session_state["hardware_table"] = remaining_case(job, "shorten" if case_name in ("plug_bore", "squeeze_wall") else "add_hole")["hardware_table"]
    case.phase(app, "phase2_3")
    canonical = deepcopy(app.session_state["placement_config"])
    hardware = app.session_state["hardware_table"].copy(deep=True)
    if case_name in ("plug_bore", "squeeze_wall"):
        label = "Plug bore diameter (in) — 2950–3000 m" if job == "CMT PLUG" else "Squeeze wall diameter (in) — 3000 m"
        next(w for w in app.number_input if w.label == label).set_value(8.5).run()
    else:
        target = "previous-casing:20" if case_name == "casing_target" else "open-hole"
        choices = schematic.additional_drawing_inputs(app.session_state)["squeeze_targets"]
        token = next(token for token, row in choices.items() if row["target_id"] == target)
        next(w for w in app.selectbox if w.label == "Squeeze Target for Well Schematic").set_value(token).run()
    assert app.session_state["placement_config"] == canonical
    pd.testing.assert_frame_equal(app.session_state["hardware_table"], hardware)
    images = []
    real_render = well_report.render_schematic_png
    def observe(payload):
        blob = real_render(payload)
        images.append(blob)
        return blob
    monkeypatch.setattr(well_report, "render_schematic_png", observe)
    case.export(app)
    assert len(images) == 1
    document = Document(BytesIO(app.session_state["_compiled_doc_bytes"]))
    assert len(document._element.body.xpath('.//wp:docPr[@descr="Well Schematic"]')) == 1
    assert "Well schematic: Not provided." not in [p.text for p in document.paragraphs]
    (tmp_path / (case_name + ".docx")).write_bytes(app.session_state["_compiled_doc_bytes"])
    (tmp_path / (case_name + ".png")).write_bytes(images[0])
    from project_state import is_project_key
    saved = {key: audit.namespace["serialize_item"](value) for key, value in app.session_state.to_dict().items()
             if is_project_key(key)}
    assert not any(key.startswith("_") for key in saved)
    restored = audit.round_trip(app.session_state.to_dict())
    assert restored["schematic_config"] == app.session_state["schematic_config"]
    # Restore regenerates the existing Well Profile runtime projection, never schematic results.
    assert {key for key in restored if key.startswith("_")} <= {
        "_well_profile_source", "_survey_result", "_well_profile_issues"}
    raw = save_project(app)
    assert b"_well_schematic" not in raw and b"PNG" not in raw
    app = restore_project(raw)
    case.export(app)
    assert len(images) == 2 and images[0] == images[1]
    if case_name in ("plug_bore", "squeeze_wall"):
        case.phase(app, "phase2_3")
        next(w for w in app.number_input if w.label == label).set_value(9.).run()
        assert "_compiled_doc_bytes" not in app.session_state
        case.export(app)
        assert len(images) == 3 and images[-1] != images[-2]


def test_rejected_schematic_load_is_transactional_and_has_no_derived_resurrection():
    import json
    import test_audit_regressions as audit
    from project_io import decode_project, replace_project_state
    state = {"job_type": "CMT PLUG", "well_name": "unchanged", "schematic_config": {"version": 1},
             "_compiled_doc_bytes": b"current"}
    previous = pickle.dumps(state)
    raw = json.dumps({"job_type": "CMT PLUG", "schematic_config": {"version": 1,
        "squeeze_wall": {"depth_m": 3000., "diameter_in": False}}}).encode()
    with pytest.raises(ValueError, match="schematic_config"):
        replace_project_state(state, decode_project(raw, audit.namespace['deserialize_item']), 'bad-load')
    assert pickle.dumps(state) == previous


def test_round11_sources_and_protected_template():
    assert "Previous Liner Hanger" in materials_db.HARDWARE_DESCRIPTIONS
    assert importlib.util.find_spec("well_schematic_adapter") is not None
    assert importlib.util.find_spec("well_schematic_engine") is not None
    assert hashlib.sha256(Path("master_template.docx").read_bytes()).hexdigest() == TEMPLATE_SHA


def test_schematic_source_signature_reads_inputs_without_render_or_calculation(monkeypatch):
    import well_schematic_adapter as schematic
    signature = getattr(schematic, 'source_signature', None)
    assert callable(signature)
    state = program_state()
    before = pickle.dumps(state)
    def forbidden(*args, **kwargs):
        raise AssertionError('Source fingerprint performed report preparation')
    monkeypatch.setattr(schematic, 'prepare_schematic', forbidden)
    first = signature(state)
    state['_schematic_widget'] = 22
    state['unrelated_report_text'] = 'presentation only'
    assert signature(state) == first
    state.pop('_schematic_widget'); state.pop('unrelated_report_text')
    assert pickle.dumps(state) == before
    state['hardware_table'].loc[2, 'MD (m)'] = '1100'
    assert signature(state) != first


@pytest.mark.parametrize("depth", [2500., 3000., 3300.])
def test_renderer_accepts_receiving_liner_lap_interior_and_shoe(depth, tmp_path):
    import well_schematic_engine as engine
    from PIL import Image
    payload = connection_payload(depth)
    before = deepcopy(payload)
    path = tmp_path / "connection.png"
    engine.render_well_schematic(payload, str(path), dpi=110, strict=True)
    with Image.open(path) as image:
        assert image.format == "PNG" and image.width > 300 and image.height > 300
    assert engine.compute_total_depth(payload) == 3300.
    assert engine.resolve_depth_axis(payload)[0] == 3500.
    assert payload == before


@pytest.mark.parametrize("depth", [2499., 3301.])
def test_renderer_rejects_connection_outside_receiving_liner(depth):
    import well_schematic_engine as engine
    with pytest.raises(ValueError, match="receiving liner interval"):
        engine.validate_schematic_data(connection_payload(depth))


def test_same_size_lap_examples_and_distinct_geometric_hosts_remain_valid(tmp_path):
    import well_schematic_engine as engine
    for name in ("tieback_7in", "tieback_5in"):
        payload = engine.get_preset_schematic_data(name)
        engine.render_well_schematic(payload, str(tmp_path / (name + ".png")), dpi=100)
    payload = {"casings": [
        {"type": "casing", "size_in": 20., "id_in": 18., "top_md": 0., "shoe_md": 2000.},
        {"type": "liner", "size_in": 13.375, "id_in": 12.25, "top_md": 1500., "shoe_md": 4000., "host_index": 0},
        {"type": "liner", "size_in": 9.625, "id_in": 8.535, "top_md": 2500., "shoe_md": 4500., "host_index": 1},
        {"type": "tieback", "size_in": 7., "top_md": 0., "shoe_md": 3000., "host_index": 0,
         "landing_liner_index": 2}], "current_job": {"type": "tieback", "target_index": 3},
        "open_holes": [], "final_relevant_depth_m": 4500.}
    engine.render_well_schematic(payload, str(tmp_path / "separate-host.png"), dpi=100)
    assert engine.compute_total_depth(payload) == 4500.
    assert engine.resolve_depth_axis(payload)[0] == 4700.


@pytest.mark.parametrize("fault", ["missing", "outside", "ambiguous"])
def test_receiving_liner_missing_outside_or_unresolved_is_blocked(fault):
    import well_schematic_adapter as schematic
    state = program_state("tieback")
    table = state["hardware_table"]
    if fault == "missing":
        state["hardware_table"] = table.drop(index=2)
    elif fault == "outside":
        table.loc[2, "MD (m)"] = "500-1000"
    else:
        table.loc[1, "Size (in)"] = "13 3/8"
        table.loc[1, "ID (in)"] = 12.25
        table.loc[3, "Size (in)"] = "12.25"
        table.loc[3, "ID (in)"] = 12.25
        other = {**table.iloc[2].to_dict(), "Size (in)": "9 5/8", "ID (in)": 8.535}
        state["hardware_table"] = pd.concat([table, pd.DataFrame([other])], ignore_index=True)
        state["placement_config"]["host_row"] = next(iter(hardware_choices(state["hardware_table"], {"Previous Casing"})))
    with pytest.raises(ValueError, match="receiving|ambiguous"):
        schematic.prepare_schematic(state)


def test_multiple_receiving_candidates_use_current_certified_continuation():
    import well_schematic_adapter as schematic
    state = program_state("tieback")
    table = state["hardware_table"]
    table.loc[1, "Size (in)"], table.loc[1, "ID (in)"] = "13 3/8", 12.25
    table.loc[3, "Size (in)"], table.loc[3, "ID (in)"] = "12.25", 12.25
    other = {**table.iloc[2].to_dict(), "MD (m)": "200-900", "Size (in)": "9 5/8", "ID (in)": 8.535}
    state["hardware_table"] = pd.concat([table, pd.DataFrame([other])], ignore_index=True)
    state["placement_config"]["host_row"] = next(iter(hardware_choices(state["hardware_table"], {"Previous Casing"})))
    geometry = pumping.prepare_pumping_simulation(state)["inputs"]["geometry"]
    payload = schematic.prepare_schematic(state)
    tb = payload["casings"][payload["current_job"]["target_index"]]
    receiver = payload["casings"][tb["landing_liner_index"]]
    assert receiver["size_in"] == geometry.pipe_sections[-1].od_in == 7.
    assert receiver["top_md"] == 400.
    assert receiver["host_index"] != tb["host_index"]


def test_certified_interval_and_tieback_top_need_no_new_schematic_input():
    import well_schematic_adapter as schematic
    state = program_state("tieback")
    before = pickle.dumps(state)
    certified = pumping.prepare_pumping_simulation(state)["inputs"]["geometry"]
    payload = schematic.prepare_schematic(state)
    liner = next(c for c in payload["casings"] if c["type"] == "liner")
    tieback = next(c for c in payload["casings"] if c["type"] == "tieback")
    assert (liner["top_md"], liner["shoe_md"]) == (400., 1000.)
    assert (tieback["top_md"], tieback["shoe_md"]) == (0., 400.)
    assert certified.pipe_sections[0].z_top == tieback["top_md"]
    assert certified.pipe_sections[1].z_bot == liner["shoe_md"]
    assert pickle.dumps(state) == before


def certified_interior_state():
    from test_audit_regressions import valid_export_chart_fixture
    state = program_state("tieback")
    state.update(valid_export_chart_fixture('TIE BACK LNR 7"'))
    state['well_geometry']['td_m'] = 3000.
    return state


def connection_payload(depth=3000.):
    # Exact owner-certified fixture; the receiving string differs from the
    # surface geometric host. No solver result or signature is fabricated.
    return {"casings": [
        {"type": "casing", "size_in": 13.375, "id_in": 12.25, "top_md": 0., "shoe_md": 3300.},
        {"type": "liner", "size_in": 9.625, "id_in": 8.535, "top_md": 2500., "shoe_md": 3300., "host_index": 0},
        {"type": "tieback", "size_in": 7., "id_in": 6.276, "top_md": 0., "shoe_md": depth,
         "host_index": 0, "landing_liner_index": 1}],
        "open_holes": [], "current_job": {"type": "tieback", "target_index": 2},
        "final_relevant_depth_m": 3300.}


def test_protected_renderer_accepts_certified_interior_connection():
    import well_schematic_engine as engine
    state = certified_interior_state()
    geometry = pumping.prepare_pumping_simulation(state)["inputs"]["geometry"]
    assert geometry.track_depth_m == 3000. and geometry.td_m == 3300.
    payload = connection_payload()
    engine.validate_schematic_data(payload)
    assert engine.check_well_design(payload) == []


def test_adapter_preserves_certified_interior_connection_and_whole_history():
    import well_schematic_adapter as schematic
    state = certified_interior_state()
    before = pickle.dumps(state)
    payload = schematic.prepare_schematic(state)
    tb = payload["casings"][payload["current_job"]["target_index"]]
    receiver = payload["casings"][tb["landing_liner_index"]]
    assert (tb["size_in"], tb["top_md"], tb["shoe_md"]) == (7., 0., 3000.)
    assert (receiver["size_in"], receiver["top_md"], receiver["shoe_md"]) == (9.625, 2500., 3300.)
    assert tb["host_index"] != tb["landing_liner_index"]
    assert payload["final_relevant_depth_m"] == 3300.
    assert pickle.dumps(state) == before


def tieback_state(interval=False):
    state = program_state("tieback")
    state["well_geometry"]["td_m"] = 4000.
    table = state["hardware_table"]
    table.loc[0, "MD (m)"] = "2500"
    table.loc[1, "MD (m)"] = "3000"
    table.loc[2, "MD (m)"] = "2500-4000" if interval else "4000"
    table.loc[2, "ID (in)"] = 6.276
    table.loc[3, "MD (m)"] = "4000"
    hanger = {**table.iloc[2].to_dict(), "Description": "Previous Liner Hanger", "MD (m)": "2500"}
    state["hardware_table"] = pd.concat([table, pd.DataFrame([hanger])], ignore_index=True)
    state["placement_config"].update(
        target_row=next(iter(hardware_choices(state["hardware_table"], {"Tie Back"}))),
        host_row=next(iter(hardware_choices(state["hardware_table"], {"Previous Casing"}))))
    return state


def test_scalar_historical_liner_is_schematic_only_and_does_not_change_pumping():
    import well_schematic_adapter as schematic
    state = tieback_state()
    before = pickle.dumps(state)
    payload = schematic.prepare_schematic(state)
    liner = next(c for c in payload["casings"] if c["type"] == "liner")
    assert (liner["top_md"], liner["shoe_md"]) == (2500., 4000.)
    with pytest.raises(ValueError, match="explicit old-liner interval"):
        pumping.prepare_pumping_simulation(state)
    assert pickle.dumps(state) == before


def test_equivalent_interval_and_endpoint_schematic_presentations():
    import well_schematic_adapter as schematic
    scalar, interval = tieback_state(), tieback_state(True)
    assert schematic.prepare_schematic(scalar) == schematic.prepare_schematic(interval)
    assert pumping.prepare_pumping_simulation(interval)["inputs"]["geometry"].td_m == 4000.


@pytest.mark.parametrize("fault", ["mismatch", "missing_top", "duplicate_hanger", "wrong_lap",
                                    "duplicate_section", "gap", "reversed"])
def test_invalid_historical_sources_never_infer_or_repair(fault):
    import well_schematic_adapter as schematic
    state = tieback_state(True)
    table = state["hardware_table"]
    if fault == "mismatch":
        table.loc[2, "MD (m)"] = "2600-4000"
    elif fault == "missing_top":
        table.loc[2, "MD (m)"] = "4000"
        state["hardware_table"] = table.iloc[:-1].copy()
    elif fault == "duplicate_hanger":
        state["hardware_table"] = pd.concat([table, table.iloc[[-1]]], ignore_index=True)
    elif fault == "wrong_lap":
        table.loc[4, "MD (m)"] = "2400"
    elif fault == "duplicate_section":
        state["hardware_table"] = pd.concat([table, table.iloc[[2]]], ignore_index=True)
    elif fault == "gap":
        table.loc[2, "MD (m)"] = "2500-3300"
        state["hardware_table"] = pd.concat([table, pd.DataFrame([
            {**table.iloc[2].to_dict(), "MD (m)": "3400-4000"}])], ignore_index=True)
    else:
        table.loc[2, "MD (m)"] = "4000-2500"
    before = pickle.dumps(state)
    with pytest.raises(ValueError):
        schematic.prepare_schematic(state)
    assert pickle.dumps(state) == before


def test_existing_app_depth_and_unique_target_do_not_expose_redundant_controls():
    import test_audit_regressions as audit
    case = audit.AuditRegressions()
    app = case.app({"job_type": "CMT SQUEEZE", **audit.chart_fixture("CMT SQUEEZE")})
    case.phase(app, "phase2_3")
    assert "Target shoe / treatment depth source" in [widget.label for widget in app.selectbox]
    assert "Squeeze Target for Well Schematic" not in [widget.label for widget in app.selectbox]
    assert "Squeeze Depth (m MD)" not in [widget.label for widget in app.number_input]


@pytest.mark.parametrize("field,value", [("tie_back_liner_lap_m", True), ("squeeze_depth_m", "500"),
                                        ("squeeze_depth_m", -1), ("target_string_id", 12)])
def test_saved_schematic_authority_rejects_malformed_values(field, value):
    import test_audit_regressions as audit
    with pytest.raises(ValueError, match=field):
        audit.round_trip({"job_type": "CMT SQUEEZE", "placement_config": {
            "job_type": "CMT SQUEEZE", field: value}})


def test_actual_build_replaces_schematic_placeholder():
    from io import BytesIO
    from docx import Document
    import test_audit_regressions as audit
    case = audit.AuditRegressions()
    app = case.configured_app('CSG 9 5/8"')
    case.export(app)
    doc = Document(BytesIO(app.session_state["_compiled_doc_bytes"]))
    assert "Well schematic: Not provided." not in [p.text for p in doc.paragraphs]
    assert len(doc._element.body.xpath('.//wp:docPr[@descr="Well Schematic"]')) == 1


def test_certified_squeeze_uses_existing_depth_and_unique_physical_target():
    import test_audit_regressions as audit
    import well_schematic_adapter as schematic
    state = {"job_type": "CMT SQUEEZE", **audit.chart_fixture("CMT SQUEEZE")}
    before = pickle.dumps(state)
    payload = schematic.prepare_schematic(state)
    assert payload["squeezes"] == [{"depth_m": 3000., "target_index": 0, "target_kind": "shoe"}]
    assert pickle.dumps(state) == before


def local_well(case):
    """Owner handoff construction, with unknown IDs left unknown for drawing."""
    from phase_2_3_well_data import HARDWARE_COLUMNS
    previous = [("Previous Casing", "800", "20"), ("Previous Casing", "1600", "13 3/8"),
                ("Previous Casing", "3000", "9 5/8")]
    historical = [("Previous Liner Hanger", "2500", "7"), ("Previous Liner", "4000", "7")]
    liner5 = [("Liner Hanger", "3600", "5"), ("Liner", "4300", "5"),
              ("Liner", "4800", "5"), ("Open Hole Size", "5200", "6")]
    if case == "casing":
        job = 'CSG 9 5/8"'
        rows = previous[:2] + [("Casing", "2200", "9 5/8"), ("Casing", "3000", "9 5/8"),
                               ("Open Hole Size", "3500", "12 1/4")]
        target = 3
    elif case == "liner7":
        job = 'LNR 7"'
        rows = previous + [("Drill Pipes", "2500", "5"), ("Liner Hanger", "2500", "7"),
                           ("Liner", "4000", "7"), ("Open Hole Size", "4300", "8 1/2")]
        target = 5
    elif case == "liner5":
        job = 'LNR 5"'
        rows = previous + historical + [("Drill Pipes", "3600", "3 1/2")] + liner5
        target = 8
    else:
        job = "CMT PLUG" if case == "plug" else "CMT SQUEEZE"
        rows = previous + historical + liner5
        target = None
    data = []
    for desc, depth, size in rows:
        row = dict.fromkeys(HARDWARE_COLUMNS, 0.)
        row.update({"Description": desc, "MD (m)": depth, "Size (in)": size, "Grade": "-"})
        data.append(row)
    table = pd.DataFrame(data, columns=HARDWARE_COLUMNS)
    for description, weights in (("Casing", [47., 53.5]), ("Liner", [18., 21.])):
        indices = table.index[table['Description'] == description]
        if len(indices) == 2:
            table.loc[indices, 'Weight (ppf)'] = weights
            table.loc[indices, 'Grade'] = ['N-80', 'L-80']
    placement = {"job_type": job}
    if target is not None:
        placement["target_row"] = next(iter(hardware_choices(table.iloc[[target]], {rows[target][0]})))
    state = {"job_type": job, "well_name": "Owner handoff UAT", "hardware_table": table,
             "placement_config": placement}
    if case == "plug":
        placement.update(target_row="__manual__", manual_depth_m=5000.)
        state["cement_params"] = {"Main": {"top_mode": "Depth (m MD)", "top_depth": 3900., "top_job_type": job}}
    elif case == "squeeze_liner":
        placement.update(squeeze_depth_m=4000., target_string_id="previous-liner:7")
    elif case == "squeeze_oh":
        placement.update(squeeze_depth_m=5000., target_string_id="open-hole")
    return state


@pytest.mark.parametrize("case", ["casing", "liner7", "liner5", "plug", "squeeze_liner", "squeeze_oh"])
def test_owner_handoff_construction_groups_sections_and_renders_deterministically(case, tmp_path):
    import well_schematic_adapter as schematic
    import well_schematic_engine as engine
    state = local_well(case)
    before = pickle.dumps(state)
    payload = schematic.prepare_schematic(state)
    first = tmp_path / "first.png"
    second = tmp_path / "second.png"
    engine.render_well_schematic(payload, str(first), dpi=100)
    engine.render_well_schematic(payload, str(second), dpi=100)
    assert first.read_bytes() == second.read_bytes()
    assert pickle.dumps(state) == before
    assert engine.resolve_depth_axis(payload)[0] == payload["final_relevant_depth_m"] + 200.
    if case == "casing":
        assert len(payload["casings"]) == 3
        assert payload["casings"][-1]["shoe_md"] == 3000.
        assert payload["open_holes"][0]["top_md"] == 1600.
    elif case == "liner7":
        assert payload["casings"][-1]["top_md"] == 2500.
    elif case == "liner5":
        assert len(payload["casings"]) == 5
        assert payload["casings"][-1]["top_md"] == 3600.
        assert payload["casings"][-1]["shoe_md"] == 4800.
        assert payload["open_holes"][0]["top_md"] == 4000.
    elif case == "plug":
        assert payload["cement_plugs"] == [{"top_md": 3900., "bottom_md": 5000.}]
    elif case == "squeeze_liner":
        sq = payload["squeezes"][0]
        assert sq["target_kind"] == "shoe" and sq["depth_m"] == 4000.
        assert payload["casings"][sq["target_index"]]["size_in"] == 7.
    else:
        assert payload["squeezes"] == [{"depth_m": 5000., "open_hole_index": 0, "target_kind": "point"}]


@pytest.mark.parametrize("fault,match", [("two_holes", "reversed/duplicate"), ("missing_lap", "Lap is missing"),
    ("hanger_drill", "must match exactly"), ("squeeze_outside", "inside.*Open Hole"),
    ("ambiguous_squeeze", "ambiguous"), ("bad_target", "does not resolve"), ("bad_id", "smaller than OD")])
def test_genuine_missing_or_contradictory_schematic_authority_blocks(fault, match):
    import well_schematic_adapter as schematic
    state = local_well("squeeze_oh" if fault == "squeeze_outside" else "liner5")
    table = state["hardware_table"]
    if fault == "two_holes":
        state["hardware_table"] = pd.concat([table, table.iloc[[-1]]], ignore_index=True)
    elif fault == "missing_lap":
        state["hardware_table"] = table[table["Description"] != "Previous Liner Hanger"].copy()
    elif fault == "hanger_drill":
        table.loc[table["Description"] == "Drill Pipes", "MD (m)"] = "3590"
    elif fault == "squeeze_outside":
        state["placement_config"]["squeeze_depth_m"] = 5300.
    elif fault == "ambiguous_squeeze":
        state["job_type"] = "CMT SQUEEZE"
        state["placement_config"] = {"job_type": "CMT SQUEEZE", "target_row": "__manual__", "manual_depth_m": 2600.}
    elif fault == "bad_target":
        state["placement_config"]["target_row"] = "deleted-row"
    else:
        table.loc[table["Description"] == "Liner", "ID (in)"] = 5.
    with pytest.raises(ValueError, match=match):
        schematic.prepare_schematic(state)


@pytest.mark.parametrize("job", materials_db.JOB_TYPES)
def test_current_supported_job_fixtures_translate_without_new_inputs(job):
    import test_audit_regressions as audit
    import well_schematic_adapter as schematic
    if job == materials_db.CT_JOB:
        from unittest.mock import patch
        import phase_10_procedure as report
        from engineering_tools import compute_phase_status
        state = {"job_type": job, "_compiled_doc_bytes": b"old"}
        with patch.object(report.st, "session_state", state):
            with pytest.raises(ValueError, match="Phase IV"):
                report.build_master_context(calculations_prepared=True)
        assert "_compiled_doc_bytes" not in state
        assert compute_phase_status(state)["phase2_3"]["level"] != "ok"
        return
    state = {"job_type": job, **audit.valid_export_chart_fixture(job)}
    if job == "CMT PLUG":
        state["cement_params"] = {"Main": {"top_mode": "Depth (m MD)", "top_depth": 2000., "top_job_type": job}}
    payload = schematic.prepare_schematic(state)
    assert payload["casings"] and payload["current_job"]["type"]
    assert all(key not in state["placement_config"] for key in ("squeeze_depth_m", "target_string_id", "tie_back_liner_lap_m"))


def test_row_target_resolves_to_whole_multisection_string_without_changing_calculations():
    import well_schematic_adapter as schematic
    from placement import target_depth
    state = local_well("squeeze_liner")
    table = state["hardware_table"]
    first = table.loc[table["Description"] == "Liner"].iloc[[0]]
    token = next(iter(hardware_choices(first, {"Liner"})))
    state["placement_config"] = {"job_type": "CMT SQUEEZE", "target_row": token, "squeeze_depth_m": 4500.}
    payload = schematic.prepare_schematic(state)
    sq = payload["squeezes"][0]
    assert payload["casings"][sq["target_index"]]["shoe_md"] == 4800.
    assert sq["depth_m"] == 4500. and sq["target_kind"] == "point"
    # The calculation authority remains row-based and unchanged.
    assert target_depth(table, state["job_type"], state["placement_config"]) == 4300.


@pytest.fixture(scope="module", params=['CSG 9 5/8"', 'LNR 7"', 'TIE BACK LNR 7"', 'CMT PLUG', 'CMT SQUEEZE'])
def word_project(request):
    import test_audit_regressions as audit
    from centralizer_fixtures import ensure_current_centralizer
    from project_state import prepare_calculations
    app = audit.AuditRegressions().configured_app(request.param)
    ensure_current_centralizer(app.session_state)
    state = app.session_state.to_dict()
    assert prepare_calculations(state) == []
    return state


def schematic_picture(doc):
    matches = [p for p in doc.paragraphs if p._p.xpath('.//wp:docPr[@descr="Well Schematic"]')]
    assert len(matches) == 1
    paragraph = matches[0]
    relationships = paragraph._p.xpath('.//a:blip/@r:embed')
    assert len(relationships) == 1
    return paragraph, doc.part.related_parts[relationships[0]].blob


def assert_word_structure(data, applicable):
    from io import BytesIO
    from docx import Document
    from test_uat_round10b_pumping_word import drawing_paragraph
    doc = Document(BytesIO(data))
    schematic, png = schematic_picture(doc)
    body = list(doc._element.body)
    headings = {p.text.strip(): p._p for p in doc.paragraphs}
    assert body.index(headings['II. Well Data Summary']) < body.index(schematic._p) < body.index(headings['III. Geothermal Temperature Profile'])
    assert 'Well schematic: Not provided.' not in [p.text for p in doc.paragraphs]
    assert png.startswith(b'\x89PNG\r\n\x1a\n')
    diagnostic = drawing_paragraph(doc, 'Static Well Security and Fluid Sequence') if applicable else None
    pumping_images = doc._element.body.xpath('.//wp:docPr[@descr="Pumping Simulation Charts"]')
    centralizer_images = doc._element.body.xpath('.//wp:docPr[@descr="Pipe Standoff"]')
    if applicable:
        pumping_paragraph = drawing_paragraph(doc, 'Pumping Simulation Charts')
        assert len(pumping_images) == len(centralizer_images) == 1
        assert body.index(headings['IV. Fluids Sequence']) < body.index(diagnostic._p) < body.index(pumping_paragraph._p) < body.index(headings['V. Cementing Program'])
    else:
        assert not pumping_images and not centralizer_images
        assert sum(p.text == 'Pumping simulation charts: Not applicable.' for p in doc.paragraphs) == 1
    assert body.index(headings['IX. Centralizer Placement Recommendation']) < body.index(headings['X. Cementing Procedure'])
    return doc, png


def test_real_word_save_restore_fresh_build_and_resource_cleanup(word_project, monkeypatch, tmp_path):
    import test_audit_regressions as audit
    import well_schematic_engine as engine
    from test_uat_round10c_integration import save_project, restore_project
    from centralizer_fixtures import ensure_current_centralizer
    from project_state import is_project_key
    from io import BytesIO
    from PIL import Image
    calls, images, solver_calls = [], [], []
    original_run = pumping.run_pumping_simulation
    original_render = engine.render_well_schematic
    original_solver = pumping.fcm.build_three_charts
    def solver(**kwargs):
        solver_calls.append(1)
        return original_solver(**kwargs)
    def run(state):
        result = original_run(state)
        calls.append(result)
        return result
    def render(payload, path, **kwargs):
        result = original_render(payload, path, **kwargs)
        images.append((deepcopy(payload), Path(result), Path(result).read_bytes()))
        return result
    monkeypatch.setattr(pumping, 'run_pumping_simulation', run)
    monkeypatch.setattr(engine, 'render_well_schematic', render)
    monkeypatch.setattr(pumping.fcm, 'build_three_charts', solver)
    case = audit.AuditRegressions()
    baseline = engine.plt.get_fignums()
    app = case.app(deepcopy(word_project))
    assert not calls and not images
    case.export(app)
    first = app.session_state['_compiled_doc_bytes']
    applicable = word_project['job_type'] not in ('CMT PLUG', 'CMT SQUEEZE')
    doc, png = assert_word_structure(first, applicable)
    assert len(calls) == len(images) == 1
    assert png == images[0][2]
    assert not images[0][1].parent.exists()
    assert engine.plt.get_fignums() == baseline
    assert calls[0]['figure'] is None or not engine.plt.fignum_exists(calls[0]['figure'].number)
    before_placement = deepcopy(app.session_state['placement_config'])
    raw = save_project(app)
    assert all(token not in raw for token in (b'_well_schematic', b'Well Schematic', b'.png', b'Figure('))
    assert not any(is_project_key(key) for key in app.session_state.to_dict() if key.startswith('_well_schematic'))
    restored = restore_project(raw)
    assert restored.session_state['placement_config'] == before_placement
    assert not images[0][1].exists()
    ensure_current_centralizer(restored.session_state)
    case.export(restored)
    _, restored_png = assert_word_structure(restored.session_state['_compiled_doc_bytes'], applicable)
    assert restored_png == png
    assert len(calls) == len(images) == 2
    assert len(solver_calls) == (2 if applicable else 0)
    assert all(not path.parent.exists() for _, path, _ in images)
    assert engine.plt.get_fignums() == baseline
    with Image.open(BytesIO(png)) as image:
        inline = schematic_picture(doc)[0]._p.xpath('.//wp:extent')[0]
        assert int(inline.get('cx')) / int(inline.get('cy')) == pytest.approx(image.width / image.height, rel=1e-5)
    # Real artifacts remain outside the repository for review.
    (tmp_path / 'round11.docx').write_bytes(first)
    (tmp_path / 'schematic.png').write_bytes(png)


@pytest.mark.parametrize('fault', ['missing', 'duplicate', 'moved'])
def test_exact_schematic_anchor_failure_blocks_before_render(fault, monkeypatch):
    from io import BytesIO
    from docx import Document
    import well_report
    doc = Document('master_template.docx')
    paragraph = next(p for p in doc.paragraphs if p.text == well_report.SCHEMATIC_ANCHOR)
    if fault == 'missing':
        paragraph.text = 'Missing anchor'
    elif fault == 'duplicate':
        paragraph._p.addnext(deepcopy(paragraph._p))
    else:
        heading = next(p for p in doc.paragraphs if p.text == 'III. Geothermal Temperature Profile')
        heading._p.addnext(paragraph._p)
    raw = BytesIO(); doc.save(raw)
    def forbidden(*args, **kwargs):
        raise AssertionError('Invalid Word anchor reached schematic rendering')
    monkeypatch.setattr(well_report, 'render_schematic_png', forbidden)
    with pytest.raises(ValueError, match='Well schematic'):
        well_report.postprocess_well_report(raw.getvalue(), None, '', schematic_report=connection_payload())


def test_schematic_changes_only_its_anchor_in_rendered_document_and_keeps_styles(word_project):
    from io import BytesIO
    from zipfile import ZipFile
    from docx import Document
    from docx.oxml.ns import qn
    from lxml import etree
    from unittest.mock import patch
    from test_uat_round7b_report import rendered
    import phase_10_procedure as phase10
    import well_report
    import well_schematic_adapter as schematic
    from project_state import prepare_calculations
    state = deepcopy(word_project)
    assert prepare_calculations(state) == []
    with patch.object(phase10.st, 'session_state', state):
        context = phase10.build_master_context(calculations_prepared=True)
    raw = rendered(context)
    args = (raw, context['well_profile_report'], context['note_vertical'], context['cement_chart_report'])
    baseline = well_report.postprocess_well_report(*args)
    final = well_report.postprocess_well_report(*args, schematic_report=schematic.prepare_schematic(state))
    before, after = Document(BytesIO(baseline)), Document(BytesIO(final))
    old, new = list(before._element.body), list(after._element.body)
    assert len(old) == len(new)
    expected = next(i for i, element in enumerate(old) if
                    ''.join(t.text or '' for t in element.iter(qn('w:t'))) == well_report.SCHEMATIC_ANCHOR)
    assert [i for i, (a, b) in enumerate(zip(old, new)) if etree.tostring(a) != etree.tostring(b)] == [expected]
    anchor = next(p for p in before.paragraphs if p.text == well_report.SCHEMATIC_ANCHOR)
    picture = schematic_picture(after)[0]
    assert anchor._p.pPr.xml == picture._p.pPr.xml
    assert anchor.runs[0]._r.rPr.xml == picture.runs[0]._r.rPr.xml
    assert len(anchor._p.xpath('.//w:br')) == len(picture._p.xpath('.//w:br'))
    with ZipFile(BytesIO(baseline)) as a, ZipFile(BytesIO(final)) as b:
        added = set(b.namelist()) - set(a.namelist())
        assert len(added) == 1 and next(iter(added)).startswith('word/media/')
        allowed = {'word/document.xml', 'word/_rels/document.xml.rels', '[Content_Types].xml'}
        assert all(a.read(path) == b.read(path) for path in set(a.namelist()) & set(b.namelist()) - allowed)


@pytest.mark.parametrize('failure', ['renderer', 'png_write'])
def test_schematic_failure_closes_owned_figure_cleans_temp_and_invalidates_download(failure, monkeypatch):
    import test_audit_regressions as audit
    import well_schematic_engine as engine
    import well_report
    from streamlit.runtime.scriptrunner import script_runner
    app = audit.AuditRegressions().configured_app('CSG 9 5/8"')
    case = audit.AuditRegressions()
    case.export(app)
    assert app.session_state['_compiled_doc_bytes']
    directories = set(Path('/tmp').glob('cement-well-schematic-*'))
    unrelated = engine.plt.figure()
    baseline = engine.plt.get_fignums()
    observed = []
    native_cleanup = script_runner._clean_problem_modules
    def before_native_cleanup():
        observed.append(engine.plt.get_fignums())
        native_cleanup()
    monkeypatch.setattr(script_runner, '_clean_problem_modules', before_native_cleanup)
    def fail(*args, **kwargs):
        raise RuntimeError('schematic output failure')
    if failure == 'renderer':
        monkeypatch.setattr(engine, 'render_well_schematic', fail)
    else:
        from unittest.mock import patch
        original = engine.render_well_schematic
        def fail_during_owned_render(*args, **kwargs):
            with patch.object(engine.plt.Figure, 'savefig', side_effect=fail):
                return original(*args, **kwargs)
        monkeypatch.setattr(engine, 'render_well_schematic', fail_during_owned_render)
    try:
        button = next(b for b in app.button if b.label == 'Build Word Document')
        button.click().run(timeout=120)
        case.healthy(app)
        assert any('schematic output failure' in error.value for error in app.error)
        assert '_compiled_doc_bytes' not in app.session_state
        assert observed == [baseline], 'The Build must preserve unrelated figures until native Streamlit cleanup'
        assert not engine.plt.get_fignums()
        assert set(Path('/tmp').glob('cement-well-schematic-*')) == directories
    finally:
        engine.plt.close(unrelated)


def test_current_source_edit_invalidates_word_and_build_renders_fresh_schematic(monkeypatch):
    from io import BytesIO
    from docx import Document
    import test_audit_regressions as audit
    import well_schematic_engine as engine
    case = audit.AuditRegressions()
    app = case.configured_app('CSG 9 5/8"')
    calls = []
    original = engine.render_well_schematic
    def render(payload, *args, **kwargs):
        calls.append(deepcopy(payload))
        return original(payload, *args, **kwargs)
    monkeypatch.setattr(engine, 'render_well_schematic', render)
    case.export(app)
    signature = app.session_state['_compiled_doc_signature']
    first_png = schematic_picture(Document(BytesIO(app.session_state['_compiled_doc_bytes'])))[1]
    case.phase(app, 'phase10')
    assert len(calls) == 1
    assert app.session_state['_compiled_doc_signature'] == signature
    table = app.session_state['hardware_table'].copy(deep=True)
    table.loc[table['Description'] == 'Open Hole Size', 'MD (m)'] = '3400'
    app.session_state['hardware_table'] = table
    case.phase(app, 'phase10')
    assert '_compiled_doc_bytes' not in app.session_state
    assert len(calls) == 1
    case.export(app)
    second_png = schematic_picture(Document(BytesIO(app.session_state['_compiled_doc_bytes'])))[1]
    assert len(calls) == 2
    assert calls[1]['final_relevant_depth_m'] == 3400.
    assert app.session_state['_compiled_doc_signature'] != signature
    assert first_png != second_png


def test_squeeze_callback_navigation_and_save_restore_preserve_separate_authorities():
    import test_audit_regressions as audit
    import well_schematic_adapter as schematic
    from placement import target_depth
    from test_uat_round10c_integration import save_project, restore_project
    case = audit.AuditRegressions()
    app = case.configured_app('CMT SQUEEZE')
    case.export(app)
    original = deepcopy(app.session_state['placement_config'])
    app.session_state['hardware_table'] = remaining_case('CMT SQUEEZE', 'add_hole')['hardware_table']
    case.phase(app, 'phase2_3')
    choices = schematic.additional_drawing_inputs(app.session_state)['squeeze_targets']
    token = next(token for token, row in choices.items() if row['target_id'] == 'open-hole')
    selector = next(w for w in app.selectbox if w.label == 'Squeeze Target for Well Schematic')
    selector.set_value(token).run()
    assert '_compiled_doc_bytes' not in app.session_state
    case.phase(app, 'phase1'); case.phase(app, 'phase2_3')
    assert next(w for w in app.selectbox if w.label == 'Squeeze Target for Well Schematic').value == token
    config = deepcopy(app.session_state['placement_config'])
    assert config == original
    assert target_depth(app.session_state['hardware_table'], 'CMT SQUEEZE', config) == 3000.
    restored = restore_project(save_project(app))
    assert restored.session_state['placement_config'] == config
    assert restored.session_state['schematic_config'] == app.session_state['schematic_config']
    case.export(restored)


@pytest.mark.parametrize('field,bad', [('tie_back_liner_lap_m', []), ('squeeze_depth_m', False),
                                     ('target_string_id', {}), ('squeeze_depth_m', float('inf'))])
def test_malformed_schematic_saved_inputs_leave_session_transactionally_intact(field, bad):
    import json
    import test_audit_regressions as audit
    from project_io import decode_project, replace_project_state
    state = {'job_type': 'CMT SQUEEZE', 'well_name': 'current project', '_compiled_doc_bytes': b'current'}
    previous = pickle.dumps(state)
    raw = json.dumps({'job_type': 'CMT SQUEEZE', 'placement_config': {field: bad}}).encode()
    with pytest.raises(ValueError):
        replace_project_state(state, decode_project(raw, audit.namespace['deserialize_item']), 'bad-load')
    assert pickle.dumps(state) == previous


def test_previous_hanger_and_schematic_metadata_round_trip_without_derived_data():
    import test_audit_regressions as audit
    state = local_well('squeeze_liner')
    state['placement_config']['tie_back_liner_lap_m'] = 1200.
    state['_well_schematic_source'] = 'transient'
    state['_schematic_widget_draft'] = {'depth': 7.}
    restored = audit.round_trip(state)
    assert restored['placement_config'] == state['placement_config']
    pd.testing.assert_frame_equal(restored['hardware_table'], state['hardware_table'])
    assert not any(key.startswith('_') for key in restored)


def test_historical_tieback_interval_remains_visible_without_new_top_input(tmp_path):
    import well_schematic_adapter as schematic
    import well_schematic_engine as engine
    state = local_well('plug')
    row = {**state['hardware_table'].iloc[0].to_dict(), 'Description': 'Tie Back',
           'MD (m)': '1200-2500', 'Size (in)': '7'}
    state['hardware_table'] = pd.concat([state['hardware_table'], pd.DataFrame([row])], ignore_index=True)
    before = pickle.dumps(state)
    payload = schematic.prepare_schematic(state)
    tb = next(c for c in payload['casings'] if c['type'] == 'tieback')
    assert (tb['top_md'], tb['shoe_md']) == (1200., 2500.)
    assert len(payload['casings']) == 6
    assert payload['final_relevant_depth_m'] == 5200.
    engine.render_well_schematic(payload, str(tmp_path / 'whole-history.png'), dpi=110)
    assert pickle.dumps(state) == before


def test_certified_selected_host_and_unselected_casing_history_are_both_preserved(tmp_path):
    import well_schematic_adapter as schematic
    import well_schematic_engine as engine
    state = program_state('tieback')
    # The existing Round-10A explicit-host regression includes this unselected
    # casing. Do not alter it or substitute it for the selected calculation host.
    extra = {**state['hardware_table'].iloc[1].to_dict(), 'Description': 'Casing',
             'MD (m)': '0-1000', 'Size (in)': '13 3/8', 'ID (in)': 12.415}
    state['hardware_table'] = pd.concat([state['hardware_table'], pd.DataFrame([extra])], ignore_index=True)
    before = pickle.dumps(state)
    certified = pumping.prepare_pumping_simulation(state)['inputs']['geometry']
    assert certified.ann_sections[0].do_in == 8.535
    payload = schematic.prepare_schematic(state)
    casings = [c for c in payload['casings'] if c['type'] == 'casing']
    assert {(c['size_in'], c['shoe_md']) for c in casings} == {(9.625, 800.), (13.375, 1000.)}
    engine.render_well_schematic(payload, str(tmp_path / 'selected-host-history.png'), dpi=110)
    assert pickle.dumps(state) == before


def test_real_casing_id_clearance_remains_blocking():
    import well_schematic_adapter as schematic
    state = program_state()
    state['hardware_table'].loc[1, 'ID (in)'] = 9.5
    with pytest.raises(ValueError, match='diameter must exceed'):
        pumping.prepare_pumping_simulation(state)
    with pytest.raises(ValueError, match='cannot pass host'):
        schematic.prepare_schematic(state)


def test_invalid_schematic_depth_blocks_actual_word_and_retires_prior_document(monkeypatch):
    import test_audit_regressions as audit
    import well_schematic_engine as engine
    case = audit.AuditRegressions()
    app = case.configured_app('CMT SQUEEZE')
    case.export(app)
    assert app.session_state['_compiled_doc_bytes']
    rendered = []
    original = engine.render_well_schematic
    def observe(*args, **kwargs):
        rendered.append(1)
        return original(*args, **kwargs)
    monkeypatch.setattr(engine, 'render_well_schematic', observe)
    case.phase(app, 'phase2_3')
    next(w for w in app.number_input if w.label == 'Measured target depth (m MD)').set_value(3500.).run()
    case.phase(app, 'phase10')
    assert '_compiled_doc_bytes' not in app.session_state
    assert next(b for b in app.button if b.label == 'Build Word Document').disabled
    assert any('Well schematic' in error.value and 'operation-wall diameter' in error.value for error in app.error)
    assert not rendered
