"""Owner-reference coverage for the report-only NIDC operational summary."""
from copy import deepcopy
from io import BytesIO
import re
import unittest
from unittest.mock import patch

import pandas as pd
from docx import Document
from docxtpl import DocxTemplate
import test_audit_regressions as audit
import materials_db
import phase_10_procedure as report
from placement import (hardware_choices, hardware_context_label, target_row_metadata,
                       slurry_intervals, resolve_report_excess,
                       summary_config_for_job)
from project_state import accept_report_text, restore_previous_text
from engineering_tools import compute_phase_status

JOB = 'CSG 9 5/8"'
P1 = ('Enclosed are our recommendations for NIDC Cement Engineering and Planning Department intervention '
      'on the referenced well. The proposal includes well data, materials and resources requirements.')
P2 = ('NIDC has established a safety policy to which all NIDC personnel must adhere. A pre-job safety meeting '
      'will be held with client representatives and other on location personnel to familiarize everyone with existing '
      'hazards and safety procedures. We would appreciate close cooperation between the client representative and '
      'the NIDC representative to ensure a safe operation.')
DRY = pd.DataFrame([
    {'Material Type': 'Weighting Agent', 'Name': 'Hidense', 'Physical State': 'Powder', 'Mix Method': 'Dry Blend'},
    {'Material Type': 'Anti Foam', 'Name': 'TA-47', 'Physical State': 'Liquid', 'Mix Method': 'In Mix Water'}])


def hardware():
    return pd.DataFrame([{'Description': desc, 'MD (m)': str(depth), 'Size (in)': size}
                         for desc, depth, size in [('Casing', 3030, '9 5/8'), ('Liner', 3000, '7'),
                                                  ('Open Hole Size', 3200, '6 1/8'), ('Tie Back', 2838, '7')]])


def tokens():
    return {row['description']: token for token, row in hardware_choices(
        hardware(), {'Casing', 'Liner', 'Open Hole Size', 'Tie Back'}).items()}


def config(job=JOB, records=None, **fields):
    return {'version': 1, 'job_type': job, 'per_slurry': records or {}, **fields}


def generate(job=JOB, active=('Lead', 'Tail'), metadata=None, adds=None, excess=True, table=None):
    fluids = {name: {'volume': 79. if name == 'Tail' else 731., 'density': '123' if name == 'Tail' else '152'}
              for name in active}
    params = {name: {'top_mode': 'Surface'} if i == 0 else {'top_mode': 'Depth (m MD)', 'top_depth': 2000. + i * 400}
              for i, name in enumerate(active)}
    if active == ('Lead', 'Tail'):
        params['Tail']['top_depth'] = 2830.
    placement = {'job_type': job, 'target_row': '__manual__', 'manual_depth_m': 3030.,
                 'excess_csg_oh_pct': 100. if excess else None, 'excess_csg_csg_pct': 20. if excess else None}
    return report.generate_executive_summary(job, fluids, list(active), adds or {},
        hardware() if table is None else table, cement_params=params, placement_config=placement,
        executive_summary_config=metadata)


class ExecutiveSummaryReferences(unittest.TestCase):
    def test_locked_paragraphs_and_sga043(self):
        text = generate(adds={'Lead': DRY})
        self.assertEqual(text.split('\n\n')[:2], [P1, P2])
        self.assertEqual(text.split('\n\n')[2],
            '9 5/8" Casing will be set and cemented at 3030.0 m MD. '
            'Dry blend and neat cement design will be used to cement this casing; '
            '123 pcf tail cement slurry will be cemented from the shoe at 3030.0 m to 2830.0 m '
            'and 152 pcf lead cement slurry will be cemented from 2830.0 m to surface. '
            'The calculated cement volume is based on 100% excess for OH-CSG and 20% excess for CSG-CSG, '
            'giving 731.0 bbl for lead slurry and 79.0 bbl for tail slurry.')

    def test_all_current_jobs_safe_and_taxonomy_unchanged(self):
        self.assertEqual(len(materials_db.JOB_TYPES), 13)
        self.assertNotIn('TIE BACK LNR 9 5/8"', materials_db.JOB_TYPES)
        for job in materials_db.JOB_TYPES:
            with self.subTest(job=job):
                text = generate(job, ('Main',), excess=False)
                self.assertEqual(text.split('\n\n')[:2], [P1, P2])
                self.assertNotIn('None', text)
                self.assertNotIn('client request', text)
                self.assertIn('731.0 bbl', text)

    def test_design_classification_is_per_slurry(self):
        for adds, expected in (({}, 'Neat'), ({'Lead': DRY}, 'Dry blend and neat'),
                               ({'Lead': DRY, 'Tail': DRY}, 'Dry blend')):
            with self.subTest(expected=expected):
                self.assertIn(expected + ' cement design will be used', generate(adds=adds))
                if expected == 'Dry blend':
                    self.assertNotIn('Dry blend and neat', generate(adds=adds))

    def test_csg_main_shoe_surface_and_oh_only(self):
        cfg = config(records={'Main': {'volume_basis': 'excess', 'excess_csg_oh_pct': 25.}})
        text = generate(active=('Main',), metadata=cfg, excess=False)
        self.assertIn('from the shoe at 3030.0 m to surface', text)
        self.assertIn('25% excess for OH-CSG', text)
        self.assertNotIn('excess for CSG-CSG', text)

    def test_three_slurry_chain_and_report_volume_order(self):
        names = ('Lead #1', 'Lead #2', 'Tail')
        cfg = config(records={name: {'volume_basis': 'excess', 'excess_csg_oh_pct': value}
                              for name, value in zip(names, (100., 50., 25.))})
        text = generate(active=names, metadata=cfg)
        chain, volumes = text.split('The calculated cement volume', 1)
        self.assertEqual(re.findall(r'(lead #\d|tail) cement slurry will', chain), ['tail', 'lead #2', 'lead #1'])
        self.assertEqual(re.findall(r'for (lead #\d|tail) slurry', volumes), list(n.lower() for n in names))
        for percent in (100, 50, 25):
            self.assertIn(f'{percent}% excess for OH-CSG', volumes)

    def test_liner_lap_and_contexts(self):
        job = 'LNR 7"'
        cfg = config(job, {'Main': {'volume_basis': 'excess', 'liner_lap_m': 70.,
                                    'liner_lap_host_row': tokens()['Casing']}},
                     placement_context_rows=[tokens()['Liner'], tokens()['Open Hole Size']])
        text = generate(job, ('Main',), cfg)
        self.assertIn('inside 7" Liner and 6 1/8" Open Hole', text)
        self.assertIn('and 70.0 m liner lap inside 9 5/8" Casing', text)

    def test_liner_client_request_wet_stands_with_and_without_host(self):
        for host in ('', tokens()['Liner']):
            with self.subTest(host=host):
                cfg = config('LNR 5"', {'Main': {'volume_basis': 'client_request', 'wet_stands': 10., 'wet_inside_host_row': host}})
                text = generate('LNR 5"', ('Main',), cfg)
                self.assertIn('based on client request with 10 stands wet', text)
                self.assertEqual('with 10 stands wet inside 7" Liner' in text, bool(host))

    def test_plug_one_and_multiple_contexts(self):
        for contexts, expected in (([tokens()['Open Hole Size']], 'inside 6 1/8" Open Hole'),
                                   ([tokens()['Liner'], tokens()['Casing']], 'inside 7" Liner and 9 5/8" Casing')):
            with self.subTest(contexts=contexts):
                text = generate('CMT PLUG', ('Main',), config('CMT PLUG', placement_context_rows=contexts))
                self.assertIn(expected, text)
                self.assertIn('from 3030.0 m to surface', text)

    def test_squeeze_reference_scopes_and_volume_modes(self):
        for scope, token, basis, opening in (
            ('Liner', tokens()['Liner'], 'direct', '7" Liner will be squeeze cemented'),
            ('Liner Lap', tokens()['Liner'], 'client_request', '7" Liner Lap Squeeze will be cemented'),
            ('Casing Region', tokens()['Casing'], 'client_request', '9 5/8" Casing Region will be squeeze cemented')):
            with self.subTest(scope=scope):
                cfg = config('CMT SQUEEZE', {'Main': {'volume_basis': basis}}, squeeze_scope=scope, squeeze_context_row=token)
                text = generate('CMT SQUEEZE', ('Main',), cfg)
                self.assertIn(opening + ' at 3030.0 m MD.', text)
                self.assertNotIn('from 3030', text)
                self.assertIn('The calculated cement volume is 731.0 bbl.' if basis == 'direct' else 'based on client request', text)

    def test_tie_back_origin_size_design_and_optional_host(self):
        for job in ('TIE BACK LNR 5"', 'TIE BACK LNR 7"'):
            for origin in ('Target depth', 'Tie Back Sleeve', 'Shoe'):
                with self.subTest(job=job, origin=origin):
                    cfg = config(job, {'Main': {'volume_basis': 'client_request'}}, tie_back_origin=origin)
                    text = generate(job, ('Main',), cfg)
                    self.assertIn(job.removeprefix('TIE BACK LNR ') + ' Tie Back cementing will be ' +
                                  ('set' if origin == 'Target depth' else 'done'), text)
                    self.assertNotIn('inside', text)
                    self.assertIn('from ' + ('the Tie Back Sleeve at ' if origin == 'Tie Back Sleeve' else
                                            'the shoe at ' if origin == 'Shoe' else '') + '3030.0 m', text)
        cfg = config('TIE BACK LNR 7"', {'Main': {'volume_basis': 'client_request', 'wet_stands': 7.,
                                                'wet_inside_host_row': tokens()['Casing']}})
        text = generate('TIE BACK LNR 7"', ('Main',), cfg, {'Main': DRY})
        self.assertIn('7" Tie Back', text)
        self.assertIn('Dry blend cement design', text)
        self.assertIn('with 7 stands wet inside 9 5/8" Casing', text)

    def test_stale_tokens_omit_optional_labels(self):
        cfg = config('CMT PLUG', placement_context_rows=[tokens()['Liner']])
        changed = hardware(); changed.loc[1, 'Size (in)'] = '5'
        text = generate('CMT PLUG', ('Main',), cfg, table=changed)
        self.assertNotIn('7" Liner', text)
        self.assertIn('inside casing / open hole', text)
        self.assertIsNone(hardware_context_label(changed, tokens()['Liner']))

    def test_selected_target_metadata_uses_shared_tokens(self):
        row = target_row_metadata(hardware(), 'LNR 7"', {'job_type': 'LNR 7"', 'target_row': tokens()['Liner']})
        self.assertEqual((row['size'], row['depth']), ('7', 3000.))

    def test_job_provenance_and_inactive_drafts_do_not_change_generated_text(self):
        cfg = config('CMT SQUEEZE', {'Main': {'volume_basis': 'client_request'}}, squeeze_scope='Liner')
        self.assertEqual(generate(metadata=cfg), generate())
        cfg = config(records={'Main': {'volume_basis': 'client_request', 'wet_stands': 99.}})
        self.assertEqual(generate(metadata=cfg), generate(metadata=config()))
        self.assertEqual(summary_config_for_job(cfg, JOB, ['Lead', 'Tail']), {})

    def test_excess_shared_override_global_blank_zero(self):
        placement = {'job_type': JOB, 'excess_csg_oh_pct': 100., 'excess_csg_csg_pct': None}
        cfg = config(records={'Lead': {'excess_csg_oh_pct': 0., 'excess_csg_csg_pct': 25.}, 'Tail': {'excess_csg_oh_pct': None}})
        for name, expected in [('Lead', {'excess_oh': 0., 'excess_csg': 25.}), ('Tail', {'excess_oh': 100., 'excess_csg': None})]:
            self.assertEqual(resolve_report_excess(JOB, placement, cfg, name), expected)
            self.assertEqual(report._placement_excess_values(placement, JOB, cfg, name),
                             {k: 'N/A' if v is None else f'{v:g}%' for k, v in expected.items()})

    def test_inactive_metadata_is_reused_only_when_reselected_for_same_job(self):
        cfg = config(records={'Main': {'volume_basis': 'client_request', 'wet_stands': 10.}})
        self.assertNotIn('client request', generate(metadata=cfg))
        self.assertIn('based on client request with 10 stands wet', generate(active=('Main',), metadata=cfg))
        self.assertNotIn('client request', generate('LNR 7"', ('Main',), cfg))

    def test_no_optional_basis_or_context_fabricates_facts(self):
        for basis in ('auto', 'unspecified', 'excess', 'direct'):
            with self.subTest(basis=basis):
                cfg = config('CMT PLUG', {'Main': {'volume_basis': basis}}, placement_context_rows=['stale'])
                text = generate('CMT PLUG', ('Main',), cfg, excess=False)
                self.assertIn('inside casing / open hole', text)
                self.assertIn('The calculated cement volume is 731.0 bbl.', text)
                self.assertNotIn('client request', text)
                self.assertNotIn('excess', text)
                self.assertNotIn('stale', text)

    def test_config_trust_boundary_and_nullable_drafts(self):
        valid = config(records={'Main': {'volume_basis': 'auto', 'wet_stands': None, 'excess_csg_oh_pct': 0., 'liner_lap_host_row': 'stale'}})
        self.assertEqual(audit.round_trip({'job_type': JOB, 'executive_summary_config': valid})['executive_summary_config'], valid)
        self.assertNotIn('executive_summary_config', audit.round_trip({'job_type': JOB}))
        for malformed in (None, [], {'version': 2, 'job_type': JOB}, config(squeeze_scope='Unknown'),
                          config(squeeze_scope=2), config(placement_context_rows=[2]), config(squeeze_context_row=2),
                          config(per_slurry=[]), config(records={'Main': []}),
                          config(records={'Main': {'volume_basis': 'guess'}}),
                          config(records={'Main': {'wet_stands': float('inf')}}),
                          config(records={'Main': {'liner_lap_m': -1.}}),
                          config(records={'Main': {'excess_csg_oh_pct': True}}),
                          config(records={'Main': {'volume_basis': 2}}),
                          config(tie_back_origin='Unknown'),
                          {'version': True, 'job_type': JOB},
                          config(records={'Main': {'wet_inside_host_row': 3}}),
                          config(records={'Main': {'volume': 12.}})):
            with self.subTest(malformed=malformed):
                with self.assertRaises(ValueError):
                    audit.round_trip({'job_type': JOB, 'executive_summary_config': malformed})


class ExecutiveSummaryStateAndWord(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.case = audit.AuditRegressions()
        app = cls.case.configured_app(JOB)
        cls.case.export(app)
        cls.state = audit.round_trip(app.session_state.to_dict())

    def test_summary_metadata_changes_only_summary_source_and_document(self):
        state = deepcopy(self.state)
        with patch.object(report.st, 'session_state', state):
            before = report.synchronize_report_texts()
            state['exec_summary_text'] = 'Manually reviewed operational summary.'
            state['_compiled_doc_bytes'] = b'old Word'
            state['_test_summary'] = 'client_request'
            report._commit_summary_metadata('_test_summary', 'volume_basis', 'Main')
            after = report.synchronize_report_texts()
            self.assertTrue(after['exec_summary_text']['pending'])
            self.assertEqual(state['exec_summary_text'], 'Manually reviewed operational summary.')
            self.assertEqual(before['procedure_text'], after['procedure_text'])
            self.assertNotIn('_compiled_doc_bytes', state)
            spec = after['exec_summary_text']
            accept_report_text(state, 'exec_summary_text', spec['generated'], spec['signature'])
            self.assertFalse(report.synchronize_report_texts()['exec_summary_text']['pending'])
            accept_report_text(state, 'exec_summary_text', spec['generated'], spec['signature'], replace=True)
            self.assertEqual(state['exec_summary_text'], spec['generated'])
            restore_previous_text(state, 'exec_summary_text')
            self.assertEqual(state['exec_summary_text'], 'Manually reviewed operational summary.')
            self.assertTrue(report.synchronize_report_texts()['exec_summary_text']['pending'])

    def test_inactive_metadata_signature_and_phase_status_unchanged(self):
        state = deepcopy(self.state)
        with patch.object(report.st, 'session_state', state):
            before = report.synchronize_report_texts()
            status = compute_phase_status(state)
            state['executive_summary_config'] = config()
            empty = report.synchronize_report_texts()
            state['executive_summary_config']['per_slurry']['Tail'] = {'wet_stands': 17., 'volume_basis': 'client_request'}
            after = report.synchronize_report_texts()
            self.assertEqual(before, empty)
            self.assertEqual(empty, after)
            self.assertEqual(before['procedure_text'], after['procedure_text'])
            self.assertEqual(status, compute_phase_status(state))

    def test_metadata_first_edit_navigation_and_fresh_json_restore(self):
        app = self.case.app(deepcopy(self.state)); self.case.phase(app, 'phase10')
        next(w for w in app.selectbox if w.label == 'Volume basis — Main').set_value('client_request').run()
        self.case.healthy(app)
        self.assertIn('based on client request', app.session_state['exec_summary_text'])
        self.case.phase(app, 'phase4'); self.case.phase(app, 'phase10')
        self.assertEqual(next(w for w in app.selectbox if w.label == 'Volume basis — Main').value, 'client_request')
        next(w for w in app.number_input if w.label == 'Excess CSG-OH % override — Main').set_value(25.).run()
        self.assertEqual(app.session_state['executive_summary_config']['per_slurry']['Main']['excess_csg_oh_pct'], 25.)
        saved = audit.round_trip(app.session_state.to_dict())
        self.assertFalse(any(key.startswith('_summary') for key in saved))
        fresh = self.case.app(saved); self.case.phase(fresh, 'phase10')
        self.assertEqual(fresh.session_state['executive_summary_config'], saved['executive_summary_config'])
        self.assertEqual(next(w for w in fresh.number_input if w.label == 'Excess CSG-OH % override — Main').value, 25.)

    def test_context_callback_commits_first_edit_before_navigation(self):
        app = self.case.app(deepcopy(self.state)); self.case.phase(app, 'phase10')
        choices = hardware_choices(app.session_state['hardware_table'], {'Previous Casing'})
        token = next(iter(choices))
        app.multiselect[0].set_value([token])
        self.case.phase(app, 'phase4')
        self.assertEqual(app.session_state['executive_summary_config']['placement_context_rows'], [token])
        self.case.phase(app, 'phase10')
        self.assertEqual(app.multiselect[0].value, [token])
        self.assertEqual(app.session_state['exec_summary_text'],
                         app.session_state['report_text_state']['exec_summary_text']['generated'])

    def test_actual_word_summary_and_placement_consistent_without_engineering_change(self):
        state = deepcopy(self.state)
        with patch.object(report.st, 'session_state', state):
            old = report.build_master_context()
            state['executive_summary_config'] = config(records={'Main': {'volume_basis': 'excess',
                'excess_csg_oh_pct': 25., 'excess_csg_csg_pct': 10.5}})
            context = report.build_master_context()
            self.assertEqual(old['procedure'], context['procedure'])
            self.assertEqual((context['slurries'][0]['excess_oh'], context['slurries'][0]['excess_csg']), ('25%', '10.5%'))
            self.assertIn('25% excess for OH-CSG and 10.5% excess for CSG-CSG', context['exec_summary'])
            self.assertEqual({k:v for k,v in old['slurries'][0].items() if k not in ('excess_oh', 'excess_csg')},
                             {k:v for k,v in context['slurries'][0].items() if k not in ('excess_oh', 'excess_csg')})
            doc = DocxTemplate(report.TEMPLATE_PATH); doc.render(context)
            out = BytesIO(); doc.save(out); word = Document(BytesIO(out.getvalue()))
            self.assertIn(context['exec_summary'].split('\n\n')[2], '\n'.join(p.text for p in word.paragraphs))
            self.assertIn('25%', '\n'.join(c.text for t in word.tables for r in t.rows for c in r.cells))

    def test_three_slurry_word_excess_matches_each_summary_clause(self):
        state = deepcopy(self.state)
        names = ['Lead #1', 'Lead #2', 'Tail']
        state['fluids_config']['active'] = names + ['Displacement Fluid']
        for i, name in enumerate(names):
            state['fluids_config']['params'][name] = deepcopy(state['fluids_config']['params']['Main'])
            state['cement_params'][name] = dict(state['cement_params']['Main'],
                top_mode='Surface' if i == 0 else 'Depth (m MD)', top_depth=(0., 1500., 2850.)[i])
            state['cement_additives_dfs'][name] = state['cement_additives_dfs']['Main'].copy()
            state['lab_qc_params'][name] = dict(state['lab_qc_params']['Main'], reviewed=False)
        state['executive_summary_config'] = config(records={name: {'volume_basis': 'excess',
            'excess_csg_oh_pct': pct, 'excess_csg_csg_pct': pct / 10} for name, pct in zip(names, (100., 50., 25.))})
        app = self.case.app(audit.round_trip(state))
        for phase in ('phase4', 'phase5', 'phase7'):
            self.case.phase(app, phase)
        for button in list(app.button):
            if button.label == 'Sync with Phase V':
                app.button(key=button.key).click().run()
        for _ in names:
            next(b for b in app.button if b.label == 'Confirm measured lab results' and not b.disabled).click().run()
        self.case.export(app)
        word = Document(BytesIO(app.session_state['_compiled_doc_bytes']))
        with patch.object(report.st, 'session_state', deepcopy(app.session_state.to_dict())):
            context = report.build_master_context()
        intervals = slurry_intervals(app.session_state['hardware_table'], JOB, app.session_state['placement_config'],
                                     names, app.session_state['cement_params'])
        for payload, name, pct in zip(context['slurries'], names, (100., 50., 25.)):
            self.assertEqual((payload['excess_oh'], payload['excess_csg']), (f'{pct:g}%', f'{pct / 10:g}%'))
            self.assertIn(f'{pct:g}% excess for OH-CSG and {pct / 10:g}% excess for CSG-CSG', context['exec_summary'])
            self.assertEqual(payload['top'], intervals[name]['top'])
            table = next(t for t in word.tables if t.rows[0].cells[0].text == name + ' Cement Slurry Data')
            cells = '\n'.join(c.text for row in table.rows for c in row.cells)
            self.assertIn(f'{pct:g}%', cells)
            self.assertIn(f'{pct / 10:g}%', cells)

    def test_no_config_actual_export_still_valid(self):
        state = deepcopy(self.state); state.pop('executive_summary_config', None)
        with patch.object(report.st, 'session_state', state):
            context = report.build_master_context()
            self.assertEqual((context['slurries'][0]['excess_oh'], context['slurries'][0]['excess_csg']), ('N/A', 'N/A'))
            self.assertNotIn('client request', context['exec_summary'])


if __name__ == '__main__':
    unittest.main()
