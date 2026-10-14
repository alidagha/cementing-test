"""Round 9G consumes current calculations; the frozen Companion owns presentation."""
from copy import deepcopy
from io import BytesIO
import hashlib
import importlib
import importlib.util
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

import pandas as pd
import pytest
from docx import Document
from docx.oxml.ns import qn
from docxtpl import DocxTemplate

import centralizer_adapter as adapter
import engineering_tools
import materials_db
import phase_9_centralizer as phase9
import phase_10_procedure as report
import well_report
import test_audit_regressions as audit
from test_uat_round9f_phase9 import configured
from centralizer_fixtures import execute_centralizer

BODY = "Recommended Centralizer Performance Data is Based on API 10D Specifications."
IX = "IX. Centralizer Placement Recommendation"
X = "X. Cementing Procedure"
TEMPLATE_SHA = "0c3d911e471aae22c0b18fecc53ed99caa7e7479d94b22640bc19439da92e331"


def bridge():
    assert importlib.util.find_spec("centralizer_report"), "Current-result Word bridge is missing"
    return importlib.import_module("centralizer_report")


@pytest.fixture(scope="module")
def catalog():
    return adapter.load_official_catalog()


@pytest.fixture(scope="module")
def word_state():
    return audit.AuditRegressions().configured_app('CSG 9 5/8"').session_state.to_dict()


def context(state):
    with patch.object(report.st, "session_state", deepcopy(state)):
        return report.build_master_context(calculations_prepared=True)


def render(ctx):
    doc = DocxTemplate(report.TEMPLATE_PATH)
    doc.render(report._word_quantity_context(ctx))
    output = BytesIO(); doc.save(output)
    return well_report.postprocess_well_report(output.getvalue(), ctx["well_profile_report"],
        ctx["note_vertical"], ctx["cement_chart_report"], ctx["centralizer_report"])


def ix_region(doc):
    elements = list(doc._element.body)
    text = lambda element: "".join(
        t.text or "" for t in element.iter(qn("w:t")))
    start = next(i for i,e in enumerate(elements) if text(e) == IX)
    end = next(i for i,e in enumerate(elements) if text(e) == X)
    return elements[start:end]


@pytest.mark.parametrize("kind", ["casing", "liner", "tieback"])
def test_direct_counts_consume_exact_frozen_placement(catalog, kind):
    state = configured(catalog, kind)
    payload = adapter.build_static_payload(state, state["centralizer_config"]["zones"], catalog)
    frozen = adapter.centralizer_engine
    expected = frozen.place_zone_centralizers(payload["zones"][0], payload["casing"]["joint_length"],
        anchor=frozen._zone_anchor(payload["job_type"], payload["zones"], 0, []),
        shoe_md=payload["wellbore"]["string_shoe"])["count"]
    result = phase9.execute(state, catalog)
    assert result["zones"][0].get("count") == expected, "Direct count is absent or inferred"


def test_phase9_is_no_longer_word_advisory():
    assert "phase9" not in engineering_tools.WORD_ADVISORY_PHASES


@pytest.mark.parametrize("condition", ["never", "stale", "failed"])
def test_direct_context_cannot_bypass_current_result(word_state, catalog, condition):
    state = deepcopy(word_state)
    state.pop("_centralizer_result", None); state.pop("_centralizer_signature", None)
    if condition != "never":
        geometry = adapter.to_engine_geometry(state)
        config = phase9.new_config()
        product = adapter.filter_official_products(catalog, geometry["casing"]["od"] / .0254)[0]
        config["zones"] = [dict(top_md=0., bottom_md=3000., density="1/3", target_standoff_pct=70.,
            product_nr=product["nr"], custom_product_id=None)]
        state["centralizer_config"] = config
        phase9.execute(state, catalog)
        if condition == "stale": state["centralizer_config"]["zones"][0]["target_standoff_pct"] = 71.
        else:
            with patch.object(adapter, "solve_static", side_effect=ValueError("failed")):
                with pytest.raises(ValueError): phase9.execute(state, catalog)
    with pytest.raises(ValueError, match="Phase IX"):
        context(state)


@pytest.mark.parametrize("physics", ["static", "dynamic"])
@pytest.mark.parametrize("mode", ["direct", "optimize"])
@pytest.mark.parametrize("custom", [False, True])
def test_companion_view_owns_report_without_solver_rerun(catalog, physics, mode, custom):
    import cement_charts_engine as companion
    state = configured(catalog, physics=physics, mode=mode, use_custom=custom)
    phase9.execute(state, catalog)
    before = deepcopy(state)
    with patch.object(adapter, "solve_static", side_effect=AssertionError("report ran Static")), \
         patch.object(adapter, "solve_dynamic", side_effect=AssertionError("report ran Dynamic")), \
         patch.object(adapter.centralizer_engine, "compute_standoff_profile", side_effect=AssertionError("report ran physics")), \
         patch.object(adapter.centralizer_engine, "optimize_placement", side_effect=AssertionError("report optimized")):
        prepared = bridge().prepare_report(state)
    assert prepared["view"] == companion.build_view(prepared["analysis"])
    view = prepared["view"]
    assert view["n_rigid"] + view["n_spring"] == view["n_cent"]
    assert len(view["zone_rows"]) == len(state["_centralizer_result"]["zones"])
    assert prepared["analysis"]["between_md"] == state["_centralizer_result"]["profile"]["between_md"]
    assert prepared["analysis"]["between_stdf"] == state["_centralizer_result"]["profile"]["between_stdf"]
    pd.testing.assert_frame_equal(state["hardware_table"], before["hardware_table"])
    assert state["centralizer_config"] == before["centralizer_config"]
    assert state["_centralizer_result"] == before["_centralizer_result"]


@pytest.mark.parametrize("job", materials_db.JOB_TYPES)
def test_centralizer_note_is_sequential_and_applicable(job):
    notes = report.build_ordered_notes(job, {}, 40., [], False, is_vertical=True)
    assert [int(n.split(":")[0].split()[1]) for n in notes["all_notes"]] == list(range(1,len(notes["all_notes"])+1))
    if adapter.centralizer_applicable(job):
        assert notes["all_notes"][-1].endswith(BODY)
        assert notes["all_notes"][-2].endswith("The Well is Considered Vertical")
    else:
        assert not any(BODY in n for n in notes["all_notes"])


def test_template_is_byte_frozen():
    assert hashlib.sha256(Path(report.TEMPLATE_PATH).read_bytes()).hexdigest() == TEMPLATE_SHA


def test_companion_renderer_is_real_and_temp_is_removed(catalog):
    import cement_charts_engine as companion
    state = configured(catalog); phase9.execute(state, catalog)
    prepared = bridge().prepare_report(state)
    paths = []
    original = companion.render_pipe_standoff_chart
    def render_chart(analysis, path, **kwargs):
        paths.append(Path(path)); return original(analysis, path, **kwargs)
    with patch.object(companion, "render_pipe_standoff_chart", side_effect=render_chart):
        png = bridge().render_chart_png(prepared)
    assert png.startswith(b"\x89PNG\r\n\x1a\n")
    assert len(paths) == 1 and not paths[0].parent.exists()


@pytest.mark.parametrize("job", ['CSG 9 5/8"','LNR 7"','TIE BACK LNR 7"','CMT PLUG','CMT SQUEEZE'])
def test_actual_docx_tables_chart_note_and_page_transition(word_state, job):
    import cement_charts_engine as companion
    from project_state import prepare_calculations
    state = audit.AuditRegressions().configured_app(job).session_state.to_dict()
    assert prepare_calculations(state) == []
    execute_centralizer(state, zones=3, custom=job.startswith("LNR"))
    ctx = context(state)
    before = hashlib.sha256(Path(report.TEMPLATE_PATH).read_bytes()).hexdigest()
    png = None
    if adapter.centralizer_applicable(job):
        png = bridge().render_chart_png(ctx["centralizer_report"])
        with patch.object(bridge(), "render_chart_png", return_value=png) as call:
            result = render(ctx)
        assert call.call_count == 1
    else:
        with patch.object(companion, "render_pipe_standoff_chart", side_effect=AssertionError("N/A chart")):
            result = render(ctx)
    doc = Document(BytesIO(result))
    region = ix_region(doc)
    text = "\n".join("".join(t.text or "" for t in element.iter(qn("w:t"))) for element in region)
    tables = [table for table in doc.tables if table._tbl in region]
    drawings = [element for element in region if element.xpath('.//w:drawing')]
    assert len(region[-1].xpath('.//w:br[@w:type="page"]')) == 1
    assert '{{' not in text and '{%' not in text
    if png is None:
        assert not tables and not drawings and BODY not in text
        assert text.count("Not applicable (Plug/Squeeze).") == 1
    else:
        assert len(tables) == 2 and len(drawings) == 1
        view = companion.build_view(ctx["centralizer_report"]["analysis"])
        assert tables[0].cell(1,1).text == f"{view['top_cent_md']:.1f} m"
        assert tables[0].cell(1,3).text == str(view['n_cent'])
        assert tables[0].cell(2,1).text == f"{view['bottom_cent_md']:.1f} m"
        assert tables[0].cell(2,3).text == str(view['n_rigid'])
        assert tables[0].cell(3,1).text == f"{view['shoe']:.1f} m"
        assert tables[0].cell(3,3).text == str(view['n_spring'])
        assert len(tables[1].rows) == 2+len(view['zone_rows']) == 5
        for row, values in zip(tables[1].rows[2:],view['zone_rows']):
            assert [cell.text for cell in row.cells] == [f"{values['bottom_md']:.1f}",str(values['number']),
                values['cent_per_joint'],values['kind_display'],f"{values['min_stdf_pct']:.1f}",f"{values['min_depth']:.1f}"]
        assert text.count(ctx['note_centralizer']) == 1
        assert ctx['all_notes'][-1] == ctx['note_centralizer']
        blip=drawings[0].xpath('.//a:blip')[0]
        part=doc.part.related_parts[blip.get(qn('r:embed'))]
        assert part.blob == png
        inline=drawings[0].xpath('.//wp:inline')[0]
        assert inline.find(qn('wp:docPr')).get('descr')=='Pipe Standoff'
        extent=inline.find(qn('wp:extent'))
        assert int(extent.get('cx'))<=3381375 and int(extent.get('cy'))<=5000625
        assert int(extent.get('cx'))/int(extent.get('cy')) == pytest.approx(3.2/6.,abs=1e-6)
    assert hashlib.sha256(Path(report.TEMPLATE_PATH).read_bytes()).hexdigest() == before == TEMPLATE_SHA


@pytest.mark.parametrize("job", ['CSG 30"','CSG 24"','CSG 20"','CSG 18 5/8"'])
def test_shallow_csg_gets_centralizer_without_open_hole_notes(word_state,job):
    state=deepcopy(word_state); state.update(audit.chart_fixture(job),job_type=job)
    execute_centralizer(state)
    ctx=context(state)
    assert ctx['centralizer_report']['applicable']
    assert not ctx['has_open_hole_notes']
    assert not ctx['note_hydrostatic'] and not ctx['note_porefrac']


@pytest.mark.parametrize("physics",["static","dynamic"])
@pytest.mark.parametrize("mode",["direct","optimize"])
def test_all_four_modes_actual_word_without_execution_during_export(word_state,physics,mode):
    from project_state import prepare_calculations
    state=deepcopy(word_state)
    assert prepare_calculations(state)==[]
    if physics=='dynamic':
        from test_uat_round9e_dynamic_solver import readings
        state['lab_qc_params']['Main']['rheology']['bhct_down'].update(selected=True,readings=readings())
    execute_centralizer(state,physics=physics,mode=mode,custom=physics=='dynamic')
    with patch.object(adapter,'solve_static',side_effect=AssertionError('export executed Static')), \
         patch.object(adapter,'solve_dynamic',side_effect=AssertionError('export executed Dynamic')), \
         patch.object(adapter.centralizer_engine,'compute_standoff_profile',side_effect=AssertionError('export ran physics')), \
         patch.object(adapter.centralizer_engine,'optimize_placement',side_effect=AssertionError('export optimized')):
        ctx=context(state)
        result=render(ctx)
    assert len(ix_region(Document(BytesIO(result))))>4


def test_target_miss_exports_warning_once_before_sequential_note(word_state):
    from project_state import prepare_calculations
    state=deepcopy(word_state)
    state['well_auto_fields']['bhsp']=False
    state['well_geometry']={'version':1,'type':'Directional','survey':pd.DataFrame(
        [[0.,0.,0.],[1500.,65.,30.],[3300.,85.,80.]],columns=['MD','Deviation','Azimuth'])}
    assert prepare_calculations(state)==[]
    execute_centralizer(state,custom=True,target=100.)
    product=state['centralizer_config']['custom_products'][0]
    product.update(kind='spring',bows=6,min_od_in=9.625,restoring_force_lbf=50.,starting_force_lbf=10.,running_force_lbf=10.)
    phase9.execute(state)
    assert any(not z['meets_target'] for z in state['_centralizer_result']['zones'])
    assert engineering_tools.compute_phase_status(state)['phase9']['level']=='ok'
    ctx=context(state)
    doc=Document(BytesIO(render(ctx)))
    paragraphs=[p.text for p in doc.paragraphs]
    assert paragraphs.count(bridge().WARNING)==1
    assert paragraphs[paragraphs.index(ctx['note_centralizer'])-1]==bridge().WARNING


def test_met_target_has_no_warning_and_full_note_sequence(word_state):
    import re
    from project_state import prepare_calculations
    state=deepcopy(word_state);assert prepare_calculations(state)==[]
    execute_centralizer(state)
    ctx=context(state);doc=Document(BytesIO(render(ctx)))
    assert not ctx['centralizer_report']['warning']
    assert bridge().WARNING not in '\n'.join(p.text for p in doc.paragraphs)
    from html import unescape
    texts=[''.join(t.text or '' for t in p.iter(qn('w:t'))) for p in doc._element.body.xpath('.//w:p')]
    notes=[text.strip() for text in texts if re.match(r'^NOTE \d+:',text.strip())]
    assert notes==[unescape(n) for n in ctx['all_notes']]
    assert [int(n.split(':')[0].split()[1]) for n in notes]==list(range(1,len(notes)+1))


def test_manual_collision_uses_generated_numbers_not_permanent_twenty():
    specs={'exec_summary_text':{'generated':''},'procedure_text':{'generated':''}}
    report.validate_manual_numbering({'exec_summary_text':'NOTE 20: Extra'},specs,['NOTE 1: First'])
    with pytest.raises(ValueError,match='conflicts'):
        report.validate_manual_numbering({'exec_summary_text':'NOTE 2: Extra'},specs,['NOTE 2: '+BODY])


def test_signature_changes_with_current_result_and_restore_requires_execute(word_state):
    from project_state import prepare_calculations
    state=deepcopy(word_state);assert prepare_calculations(state)==[]
    execute_centralizer(state)
    before=report.document_signature(context(state),report.TEMPLATE_PATH)
    state['centralizer_config']['zones'][0]['density']='1/2'
    with pytest.raises(ValueError,match='Phase IX'):context(state)
    phase9.execute(state)
    assert report.document_signature(context(state),report.TEMPLATE_PATH)!=before
    restored=audit.round_trip(state)
    assert not any(k.startswith('_centralizer') for k in restored)
    with pytest.raises(ValueError,match='Phase IX'):context(restored)
    phase9.execute(restored)
    assert context(restored)['centralizer_report']['applicable']


def test_renderer_failure_cleans_temp_files_and_figures(catalog):
    import cement_charts_engine as companion
    state=configured(catalog);phase9.execute(state,catalog)
    prepared=bridge().prepare_report(state);paths=[]
    figures=companion.plt.get_fignums()
    def fail(analysis,path,**kwargs):
        paths.append(Path(path));Path(path).write_bytes(b'partial')
        companion.plt.figure()
        raise ValueError('controlled renderer failure')
    with patch.object(companion,'render_pipe_standoff_chart',side_effect=fail):
        with pytest.raises(ValueError,match='controlled'):bridge().render_chart_png(prepared)
    assert not paths[0].parent.exists()
    assert companion.plt.get_fignums()==figures


def test_concurrent_companion_renders_are_isolated(catalog):
    from concurrent.futures import ThreadPoolExecutor
    a=configured(catalog);b=deepcopy(a)
    b['centralizer_config']['zones'][0]['density']='1/1'
    phase9.execute(a,catalog);phase9.execute(b,catalog)
    reports=[bridge().prepare_report(s) for s in (a,b)]
    expected=[bridge().render_chart_png(p) for p in reports]
    with ThreadPoolExecutor(max_workers=2) as pool:
        actual=list(pool.map(bridge().render_chart_png,reports))
    assert actual==expected and actual[0]!=actual[1]


@pytest.mark.parametrize('condition',['never','stale','failed','restored'])
def test_phase_x_ui_blocks_noncurrent_result_and_keeps_phase9_deep_link(word_state,condition):
    state=deepcopy(word_state)
    if condition=='never':
        state.pop('_centralizer_result',None);state.pop('_centralizer_signature',None)
    elif condition=='stale':
        state['centralizer_config']['zones'][0]['density']='1/2'
    elif condition=='failed':
        with patch.object(adapter,'solve_static',side_effect=ValueError('Controlled failure')):
            with pytest.raises(ValueError):phase9.execute(state)
    else:
        state=audit.round_trip(state)
    case=audit.AuditRegressions();app=case.app(state);case.phase(app,'phase10')
    assert not app.exception
    assert next(b for b in app.button if b.label=='Build Word Document').disabled
    assert '_compiled_doc_bytes' not in app.session_state
    app.button(key='_rescue_phase9').click().run()
    assert not app.exception and app.radio(key='_app_mode_key').value=='phase9'


def test_phase_ix_edit_removes_actual_compiled_download(word_state):
    case=audit.AuditRegressions();app=case.app(deepcopy(word_state))
    case.export(app)
    assert '_compiled_doc_bytes' in app.session_state
    app.session_state['centralizer_config']['zones'][0]['density']='1/2'
    case.phase(app,'phase10')
    assert not app.exception
    assert next(b for b in app.button if b.label=='Build Word Document').disabled
    assert '_compiled_doc_bytes' not in app.session_state
    assert not app.get('download_button') or all(b.proto.label!='Download Final Document (.docx)' for b in app.get('download_button'))


@pytest.mark.parametrize('change',['table','note','drawing','header','page_break'])
def test_section_ix_unknown_structure_is_controlled_before_chart_render(word_state,change):
    from project_state import prepare_calculations
    state=deepcopy(word_state);assert prepare_calculations(state)==[]
    execute_centralizer(state);ctx=context(state)
    tpl=DocxTemplate(report.TEMPLATE_PATH);tpl.render(report._word_quantity_context(ctx))
    buf=BytesIO();tpl.save(buf);doc=Document(BytesIO(buf.getvalue()))
    region=ix_region(doc)
    tables=[table for table in doc.tables if table._tbl in region]
    note=next(p for p in doc.paragraphs if BODY in p.text)
    if change=='table':tables[0]._tbl.getparent().remove(tables[0]._tbl)
    elif change=='note':note._p.addnext(deepcopy(note._p))
    elif change=='drawing':
        drawing=region[-1].xpath('.//w:drawing')[0];drawing.addnext(deepcopy(drawing))
    elif change=='header':tables[1].cell(1,0).text='Unexpected column'
    else:
        br=region[-1].xpath('.//w:br[@w:type="page"]')[0];br.getparent().remove(br)
    with patch.object(bridge(),'render_chart_png',side_effect=AssertionError('unknown structure rendered a chart')):
        with pytest.raises(ValueError):well_report._postprocess_centralizer(doc,ctx['centralizer_report'])


def test_rendered_copy_outside_ix_and_row_properties_remain_identical(word_state):
    from lxml import etree
    from project_state import prepare_calculations
    state=deepcopy(word_state);assert prepare_calculations(state)==[]
    execute_centralizer(state,zones=4)
    ctx=context(state);tpl=DocxTemplate(report.TEMPLATE_PATH);tpl.render(report._word_quantity_context(ctx))
    buf=BytesIO();tpl.save(buf)
    base=well_report.postprocess_well_report(buf.getvalue(),ctx['well_profile_report'],ctx['note_vertical'],ctx['cement_chart_report'])
    doc=Document(BytesIO(base));before=deepcopy(doc._element)
    region=ix_region(doc);placement=next(t for t in doc.tables if t._tbl in region and t.cell(0,0).text.strip()=='Centralizer Placement')
    sample=deepcopy(placement.rows[2]._tr)
    well_report._postprocess_centralizer(doc,ctx['centralizer_report'])
    after=deepcopy(doc._element)
    def outside(root):
        body=root.find(qn('w:body'));elements=list(body)
        text=lambda e:''.join(t.text or '' for t in e.iter(qn('w:t')))
        start=next(i for i,e in enumerate(elements) if text(e)==IX)
        end=next(i for i,e in enumerate(elements) if text(e)==X)
        for element in elements[start+1:end]:body.remove(element)
        return etree.tostring(root,method='c14n')
    assert outside(before)==outside(after)
    def style(row):
        row=deepcopy(row)
        for node in row.iter(qn('w:t')):node.text=''
        return etree.tostring(row,method='c14n')
    placement=next(t for t in doc.tables if t._tbl in ix_region(doc) and t.cell(0,0).text.strip()=='Centralizer Placement')
    assert len(placement.rows)==6
    assert all(style(row._tr)==style(sample) for row in placement.rows[2:])
    output=BytesIO();doc.save(output)
    with ZipFile(BytesIO(base)) as old,ZipFile(BytesIO(output.getvalue())) as new:
        for name in old.namelist():
            if name.startswith(('word/styles','word/header','word/footer')):
                assert old.read(name)==new.read(name),name


@pytest.mark.parametrize('mode',['direct','optimize'])
def test_mixed_official_custom_kinds_bind_by_identity_and_counts(catalog,mode):
    state=configured(catalog,mode=mode,use_custom=True)
    spring=next(p for p in adapter.filter_official_products(catalog,9.625)
                if p['kind']=='spring' and p.get('restoring_force_lbf'))
    rigid=next(p for p in adapter.filter_official_products(catalog,9.625) if p['kind']=='rigid')
    state['centralizer_config']['custom_products'][0]['identification']=spring['identification']
    state['centralizer_config']['zones']=[
        dict(top_md=0.,bottom_md=300.,density='1/3',target_standoff_pct=70.,product_nr=spring['nr'],custom_product_id=None),
        dict(top_md=300.,bottom_md=600.,density='1/3',target_standoff_pct=70.,product_nr=None,custom_product_id='opaque-one'),
        dict(top_md=600.,bottom_md=1000.,density='1/3',target_standoff_pct=70.,product_nr=rigid['nr'],custom_product_id=None)]
    phase9.execute(state,catalog)
    prepared=bridge().prepare_report(state);view=prepared['view']
    assert view['n_rigid']>0 and view['n_spring']>0
    assert view['n_rigid']+view['n_spring']==view['n_cent']
    assert [r['kind_display'] for r in view['zone_rows']]==['Spring','Rigid','Rigid']
    assert [z['product_name'] for z in prepared['analysis']['zones']]==[
        f"OFFICIAL-CENT-{spring['nr']}",'CUSTOM-CENT-opaque-one',f"OFFICIAL-CENT-{rigid['nr']}"]
