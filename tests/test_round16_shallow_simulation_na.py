"""Owner-approved four-job N/A; real widgets/Word retain casing authority."""
from copy import deepcopy
from io import BytesIO
import json
import pickle
from unittest.mock import patch

from docx import Document
import pytest

import cement_chart_model as chart
import centralizer_adapter as adapter
import centralizer_report as bridge
import engineering_tools as tools
import materials_db
import phase_9_centralizer as phase9
import phase_10_procedure as report
import pumping_simulation as pumping
from project_state import fingerprint
import test_audit_regressions as audit
from test_uat_round9g_centralizer_word import BODY, IX, X, ix_region

# Independent approved identities: do not derive the expectation from production.
TARGETS = ('CSG 30"', 'CSG 24"', 'CSG 20"', 'CSG 18 5/8"')
PRIVATE = ('_centralizer_result', '_centralizer_signature', '_centralizer_error')
NA = 'Not applicable (shallow casing).'
UI_NA = 'Not applicable: Centralizer calculation is not required for this shallow casing job.'


@pytest.fixture(scope='module')
def applicable_state():
    # Existing explicitly synthetic measured prerequisites, genuine UI lab review
    # and real Centralizer execution; never mock a successful solver/review.
    return audit.AuditRegressions().configured_app('CSG 9 5/8"').session_state.to_dict()


@pytest.mark.parametrize('job', TARGETS)
def test_target_pumping_is_na_before_simulation_inputs(job):
    state = {'job_type': job, 'hardware_table': None, 'plastic_viscosity': 'bad',
             'yield_point': None, 'lab_qc_params': None, 'centralizer_config': None}
    before = pickle.dumps(state)
    with patch.object(pumping.fcm, 'build_three_charts', side_effect=AssertionError('N/A solver')):
        result = pumping.run_pumping_simulation(state)
    assert result['applicable'] is False
    assert result['inputs'] is result['result'] is result['figure'] is None
    assert pumping.source_signature(state) == fingerprint({'job': job})
    assert pumping.source_signature({'job_type': job}) == pumping.source_signature(state)
    assert chart.chart_job_kind(job) == 'casing'
    assert pickle.dumps(state) == before


@pytest.mark.parametrize('job', TARGETS)
def test_target_status_report_and_execute_ignore_private_results(job, applicable_state):
    state = deepcopy(applicable_state)
    state['job_type'] = job
    state.pop('centralizer_config', None)
    before = pickle.dumps(state)
    assert adapter.centralizer_applicable(job) is False
    assert adapter.to_engine_geometry(state) is None
    assert tools.compute_phase_status(state)['phase9'] == {'level': 'ok', 'message': NA}
    prepared = bridge.prepare_report(state)
    assert prepared == dict(applicable=False, analysis=None, view=None, source=None, warning=False, message=NA)
    assert pickle.dumps(state) == before
    with patch.object(adapter, 'solve_static', side_effect=AssertionError('N/A Static')), \
         patch.object(adapter, 'solve_dynamic', side_effect=AssertionError('N/A Dynamic')):
        assert phase9.execute(state) is None
    assert not any(key in state for key in PRIVATE)


@pytest.mark.parametrize('job', TARGETS)
def test_target_actual_ui_build_json_and_preserved_static_outputs(job, applicable_state, tmp_path):
    case = audit.AuditRegressions()
    state = deepcopy(applicable_state)
    state.update(audit.valid_export_chart_fixture(job), job_type=job)
    state.pop('centralizer_config', None)
    state = audit.round_trip(state)
    assert not any(key in state for key in PRIVATE)
    geometry_before = fingerprint(chart.resolve_chart_geometry(state))
    with patch.object(pumping.fcm, 'build_three_charts', side_effect=AssertionError('N/A Pumping')), \
         patch.object(adapter, 'solve_static', side_effect=AssertionError('N/A Static')), \
         patch.object(adapter, 'solve_dynamic', side_effect=AssertionError('N/A Dynamic')):
        app = case.app(state)
        case.phase(app, 'phase9')
        assert [item.value for item in app.info] == [UI_NA]
        assert not any(w.key and w.key.startswith('_centralizer_')
                       for group in (app.button, app.radio, app.selectbox, app.number_input, app.text_input)
                       for w in group)
        assert tools.compute_phase_status(app.session_state)['phase9'] == {'level': 'ok', 'message': NA}
        case.export(app)
    data = app.session_state['_compiled_doc_bytes']
    (tmp_path/'shallow.docx').write_bytes(data)
    doc = Document(BytesIO(data)); region = ix_region(doc)
    text = '\n'.join(p.text for p in doc.paragraphs)
    drawings = doc.element.body.xpath('.//wp:docPr/@descr')
    assert drawings.count('Static Well Security and Fluid Sequence') == 1
    assert 'Pumping Simulation Charts' not in drawings and 'Pipe Standoff' not in drawings
    assert text.count('Pumping simulation charts: Not applicable.') == 1
    assert text.count(NA) == 1 and text.count(IX) == text.count(X) == 1
    assert not [table for table in doc.tables if table._tbl in region]
    assert BODY not in text
    assert len(region[-1].xpath('.//w:br[@w:type="page"]')) == 1
    assert '{{' not in text and '{%' not in text
    with patch.object(report.st, 'session_state', app.session_state.to_dict()):
        ctx = report.build_master_context()
    assert ctx['has_centralizer'] is ctx['has_open_hole_notes'] is False
    assert not ctx['note_hydrostatic'] and not ctx['note_porefrac']
    final = app.session_state.to_dict()
    assert fingerprint(chart.resolve_chart_geometry(final)) == geometry_before
    restored = audit.round_trip(final)
    assert not any(key in restored for key in PRIVATE)
    assert fingerprint(chart.resolve_chart_geometry(restored)) == geometry_before
    # Actual output proves Schematic was retained, not merely a context flag.
    assert any('Schematic' in desc for desc in drawings)
    (tmp_path/'observation.json').write_text(json.dumps({'job': job, 'drawings': drawings,
        'geometry': geometry_before, 'phase9': tools.compute_phase_status(final)['phase9']}, indent=2))
    # Independent Lab admission is still required on the N/A target.
    final['lab_qc_params']['Main']['reviewed'] = False
    invalid = case.app(final)
    case.assert_export_blocked(invalid)


@pytest.mark.parametrize('job', TARGETS)
def test_actual_job_switch_retires_only_runtime_centralizer(job, applicable_state):
    case = audit.AuditRegressions(); app = case.app(deepcopy(applicable_state))
    case.phase(app, 'phase1')
    assert phase9.result_is_current(app.session_state)
    config = deepcopy(app.session_state['centralizer_config'])
    app.selectbox(key='_w_job_type').set_value(job).run(); case.healthy(app)
    assert not any(key in app.session_state for key in PRIVATE)
    assert app.session_state['centralizer_config'] == config
    assert not any(key.startswith('_compiled_doc_') for key in app.session_state.to_dict())
    app.selectbox(key='_w_job_type').set_value('CSG 9 5/8"').run(); case.healthy(app)
    assert not phase9.result_is_current(app.session_state)
    assert tools.compute_phase_status(app.session_state)['phase9']['level'] == 'warning'
    assert app.session_state['centralizer_config'] == config


@pytest.mark.parametrize('special', ('CMT PLUG', 'CMT SQUEEZE'))
def test_special_target_transitions_preserve_drafts_without_result_revival(special, applicable_state):
    case = audit.AuditRegressions(); app = case.app(deepcopy(applicable_state))
    case.phase(app, 'phase1')
    config = deepcopy(app.session_state['centralizer_config'])
    for job in (special, 'CSG 20"', special, 'CSG 9 5/8"'):
        app.selectbox(key='_w_job_type').set_value(job).run(); case.healthy(app)
    assert not any(key in app.session_state for key in PRIVATE)
    assert not phase9.result_is_current(app.session_state)
    assert app.session_state['centralizer_config'] == config


def test_unrelated_job_transition_keeps_existing_private_cache_policy(applicable_state):
    case = audit.AuditRegressions(); app = case.app(deepcopy(applicable_state)); case.phase(app, 'phase1')
    old = app.session_state['_centralizer_signature']
    app.selectbox(key='_w_job_type').set_value('CSG 13 3/8"').run(); case.healthy(app)
    assert app.session_state['_centralizer_signature'] == old
    assert app.session_state['_centralizer_result'] == applicable_state['_centralizer_result']


@pytest.mark.parametrize('job', [job for job in materials_db.JOB_TYPES if job not in TARGETS])
def test_nontarget_recognized_domain_and_old_na_shape_unchanged(job):
    special = job in ('CMT PLUG', 'CMT SQUEEZE', materials_db.CT_JOB)
    assert adapter.centralizer_applicable(job) is not special
    if special:
        assert bridge.prepare_report({'job_type': job}) == dict(
            applicable=False, analysis=None, view=None, source=None, warning=False)
        assert pumping.prepare_pumping_simulation({'job_type': job}) == dict(
            applicable=False, message='Not applicable (Plug/Squeeze).', inputs=None, warnings=[])
        assert tools.compute_phase_status({'job_type': job})['phase9'] == dict(
            level='ok', message='Not applicable (Plug/Squeeze).')
    else:
        with pytest.raises(ValueError, match='Phase IX'):
            bridge.prepare_report({'job_type': job})
        with pytest.raises(ValueError, match='Mud'):
            pumping.prepare_pumping_simulation({'job_type': job})


def test_unknown_job_stays_unsupported():
    state = {'job_type': 'CSG FUTURE 99"'}
    assert not adapter.centralizer_applicable(state['job_type'])
    assert tools.compute_phase_status(state)['phase9']['level'] == 'empty'
    for call in (pumping.prepare_pumping_simulation, pumping.source_signature, bridge.prepare_report):
        with pytest.raises(ValueError):
            call(state)


@pytest.mark.parametrize('job', TARGETS)
def test_target_valid_measured_inputs_still_do_not_enter_simulation(job, applicable_state):
    state = deepcopy(applicable_state)
    state.update(audit.valid_export_chart_fixture(job), job_type=job)
    assert chart.resolve_chart_geometry(state)['kind'] == 'casing'
    before = pickle.dumps(state)
    assert pumping.prepare_pumping_simulation(state)['applicable'] is False
    assert pickle.dumps(state) == before
