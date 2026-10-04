"""Functional UAT Round 2: global placement excess and Word Gradient display."""
from copy import deepcopy
from io import BytesIO
import json
import unittest
from unittest.mock import patch

import pandas as pd
from docx import Document
from docxtpl import DocxTemplate
import test_audit_regressions as audit
import phase_10_procedure as report
from engineering_tools import compute_phase_status
from placement import hardware_choices, target_descriptions, target_depth, host_label, slurry_intervals
from project_state import fingerprint

FIELDS = ('excess_csg_oh_pct', 'excess_csg_csg_pct')
LABELS = ('Excess CSG-OH %', 'Excess CSG-CSG %')


class Round2Phase23UAT(unittest.TestCase):
    app = audit.AuditRegressions.app
    healthy = audit.AuditRegressions.healthy
    phase = audit.AuditRegressions.phase
    configured_app = audit.AuditRegressions.configured_app
    export = audit.AuditRegressions.export
    assert_export_blocked = audit.AuditRegressions.assert_export_blocked

    def excess_input(self, app, field):
        return next(w for w in app.number_input if w.label == LABELS[FIELDS.index(field)])

    def context(self, app, prepared=False):
        with patch.object(report.st, 'session_state', deepcopy(app.session_state.to_dict())):
            return report.build_master_context(calculations_prepared=prepared)

    def confirm_lab(self, app):
        self.phase(app, 'phase7')
        keep = next((b for b in app.button if b.label == 'Keep reviewed lab entries'), None)
        if keep:
            keep.click().run()
        confirm = next((b for b in app.button if b.label == 'Confirm measured lab results'), None)
        if confirm:
            self.assertFalse(confirm.disabled)
            confirm.click().run()
        self.healthy(app)

    def word_excess(self, doc):
        return [[c.text for c in row.cells] for table in doc.tables for row in table.rows
                if row.cells[0].text == 'OH-CSG Excess']

    def test_exactly_two_global_nullable_inputs_replace_single_field(self):
        app = self.app()
        self.phase(app, 'phase2_3')
        self.assertFalse(any(w.label.startswith('Slurry volume basis') for w in app.text_input))
        self.assertEqual([w.label for w in app.number_input if w.label.startswith('Excess ')], list(LABELS))
        for field in FIELDS:
            self.assertIsNone(self.excess_input(app, field).value)
            self.assertIn(field, app.session_state['placement_config'])
            self.assertIsNone(app.session_state['placement_config'][field])
        self.assertNotIn('volume_basis', app.session_state['placement_config'])
        self.assertNotEqual(self.excess_input(app, FIELDS[0]).key, self.excess_input(app, FIELDS[1]).key)

    def test_independent_first_edits_commit_before_navigation(self):
        app = self.app()
        self.phase(app, 'phase2_3')
        for field, value in zip(FIELDS, (25.0, 10.5)):
            self.excess_input(app, field).set_value(value)
            self.phase(app, 'phase1')
            self.assertEqual(app.session_state['placement_config'][field], value)
            self.phase(app, 'phase2_3')
            self.assertEqual(self.excess_input(app, field).value, value)
        self.excess_input(app, FIELDS[0]).set_value(30.0).run()
        self.assertEqual(self.excess_input(app, FIELDS[1]).value, 10.5)
        self.excess_input(app, FIELDS[0]).set_value(None).run()
        self.assertIsNone(self.excess_input(app, FIELDS[0]).value)
        self.assertEqual(self.excess_input(app, FIELDS[1]).value, 10.5)
        self.excess_input(app, FIELDS[0]).set_value(0.0)
        self.phase(app, 'phase1')
        self.phase(app, 'phase2_3')
        self.assertEqual(self.excess_input(app, FIELDS[0]).value, 0.0)
        self.assertEqual(self.excess_input(app, FIELDS[1]).value, 10.5)

    def test_zero_only_excess_requires_replacement_confirmation_and_cancel_preserves_it(self):
        incoming = BytesIO(json.dumps({'well_name': 'Incoming Well', 'client': 'Incoming Client'}).encode())
        incoming.name = 'current-project.json'
        incoming.size = len(incoming.getvalue())
        for field in (None, *FIELDS):
            with self.subTest(field=field):
                app = self.app()
                self.phase(app, 'phase2_3')
                if field:
                    self.excess_input(app, field).set_value(0.0).run()
                self.phase(app, 'phase1')
                self.assertEqual(any('unsaved changes' in w.value for w in app.warning), field is not None)
                uploader_key = f"proj_uploader_{app.session_state.to_dict().get('_uploader_revision', 0)}"
                with patch('streamlit.delta_generator.DeltaGenerator.file_uploader',
                           side_effect=lambda *args, **kwargs: incoming if kwargs.get('key') == uploader_key else None):
                    app.run()
                    self.healthy(app)
                    self.assertEqual(any(b.key == '_confirm_upload_replacement' for b in app.button), field is not None)
                    if field:
                        next(b for b in app.button if b.key == '_cancel_upload_replacement').click().run()
                        self.healthy(app)
                        self.assertEqual(app.session_state['placement_config'][field], 0.0)
                        self.assertFalse(any(b.key == '_confirm_upload_replacement' for b in app.button))
                    else:
                        self.assertEqual(app.session_state['well_name'], 'Incoming Well')

    def test_current_json_preserves_blank_zero_decimal_and_percent_points(self):
        for values in ((None, None), (25.0, 10.5), (0.0, None), (None, 0.0), (30.0, 10.123456789123)):
            with self.subTest(values=values):
                app = self.app()
                self.phase(app, 'phase2_3')
                for field, value in zip(FIELDS, values):
                    self.excess_input(app, field).set_value(value).run()
                self.phase(app, 'phase1')
                saved = audit.round_trip(app.session_state.to_dict())
                self.assertEqual(tuple(saved['placement_config'][f] for f in FIELDS), values)
                restored = self.app(saved)
                self.phase(restored, 'phase2_3')
                self.assertEqual(tuple(self.excess_input(restored, f).value for f in FIELDS), values)
                self.assertEqual(restored.session_state['placement_config'], saved['placement_config'])

    def test_current_json_rejects_invalid_excess_without_state_replacement(self):
        for field in FIELDS:
            for invalid in (-1.0, float('nan'), float('inf'), float('-inf'), 'bad', '25', True, {}, []):
                with self.subTest(field=field, invalid=invalid):
                    project = {'job_type': 'CMT PLUG', 'placement_config': {'job_type': 'CMT PLUG', field: invalid}}
                    before = deepcopy(project)
                    raw = json.dumps(project).encode()
                    with self.assertRaises(ValueError):
                        audit.decode_project(raw, audit.namespace['deserialize_item'])
                    self.assertEqual(project['placement_config'].keys(), before['placement_config'].keys())

    def test_invalid_canonical_excess_keeps_repair_and_blocks_stale_word(self):
        app = self.configured_app('CSG 9 5/8"')
        for field in FIELDS:
            for invalid in (-1.0, float('nan'), 'bad'):
                with self.subTest(field=field, invalid=invalid):
                    self.export(app)
                    app.session_state['placement_config'][field] = invalid
                    self.assertEqual(compute_phase_status(app.session_state)['phase2_3']['level'], 'warning')
                    self.phase(app, 'phase2_3')
                    self.assertTrue(any('Correct Excess' in w.label for w in app.text_input))
                    self.assert_export_blocked(app)
                    app.session_state['placement_config'][field] = None

    def test_shared_payload_mapping_and_exact_percent_precision(self):
        hw = pd.DataFrame()
        job = 'CMT PLUG'
        params = {'top_mode': 'Surface', 'top_job_type': job}
        for values, expected in (((25.0, 10.5), ('25%', '10.5%')),
                                 ((None, 10.5), ('N/A', '10.5%')),
                                 ((0.0, None), ('0%', 'N/A')),
                                 ((10.123456789123, 1000.0), ('10.123456789123%', '1000%'))):
            with self.subTest(values=values):
                cfg = {'job_type': job, 'target_row': '__manual__', 'manual_depth_m': 3000.0,
                       **dict(zip(FIELDS, values))}
                payload = report._attach_placement_fields({'name': 'Main'}, params, hw, job, cfg)
                self.assertEqual((payload['excess_oh'], payload['excess_csg']), expected)
                self.assertEqual(payload['top'], 'surface')
                self.assertEqual(payload['bottom'], '3000.0 m MD')

    def test_current_target_selection_and_tieback_host_survive_excess_edits(self):
        for job in audit.BATCH1_JOBS:
            for manual in (False, True):
                with self.subTest(job=job, manual=manual):
                    app = self.configured_app(job)
                    if not manual:
                        description = 'Tie Back' if 'TIE BACK' in job else ('Liner' if job.startswith('LNR') else 'Casing')
                        table = app.session_state['hardware_table'].copy()
                        row = table.iloc[0].copy();row['Description'] = description;row['MD (m)'] = '2850'
                        app.session_state['hardware_table'] = pd.concat([table, pd.DataFrame([row])], ignore_index=True)
                    self.phase(app, 'phase2_3')
                    if not manual:
                        token = next(iter(hardware_choices(app.session_state['hardware_table'], target_descriptions(job))))
                        next(w for w in app.selectbox if w.label == 'Target shoe / treatment depth source').set_value(token).run()
                    old = deepcopy(app.session_state['placement_config'])
                    depth = target_depth(app.session_state['hardware_table'], job, old)
                    host = host_label(app.session_state['hardware_table'], job, old)
                    intervals = slurry_intervals(app.session_state['hardware_table'], job, old, ['Main'], app.session_state['cement_params'])
                    for field, value in zip(FIELDS, (25.0, 10.5)):
                        self.excess_input(app, field).set_value(value).run()
                    restored = self.app(audit.round_trip(app.session_state.to_dict()))
                    self.phase(restored, 'phase2_3')
                    cfg = restored.session_state['placement_config']
                    self.assertEqual({k: v for k, v in cfg.items() if k not in FIELDS}, {k: v for k, v in old.items() if k not in FIELDS})
                    self.assertEqual(target_depth(restored.session_state['hardware_table'], job, cfg), depth)
                    self.assertEqual(host_label(restored.session_state['hardware_table'], job, cfg), host)
                    self.assertEqual(slurry_intervals(restored.session_state['hardware_table'], job, cfg, ['Main'], restored.session_state['cement_params']), intervals)

    def test_chained_placement_unchanged_with_global_excess(self):
        for job in ('CSG 9 5/8"', 'LNR 7"'):
            with self.subTest(job=job):
                cfg = {'job_type': job, 'target_row': '__manual__', 'manual_depth_m': 3000.0}
                params = {s: {'top_mode': 'Depth (m MD)', 'top_depth': d, 'top_job_type': job}
                          for s, d in (('Lead', 1000.0), ('Tail', 1500.0), ('Main', 2000.0))}
                before = slurry_intervals(pd.DataFrame(), job, cfg, list(params), params)
                cfg.update(dict(zip(FIELDS, (25.0, 10.5))))
                after = slurry_intervals(pd.DataFrame(), job, cfg, list(params), params)
                self.assertEqual(after, before)
                self.assertEqual([(s, p['bottom_depth'], p['top_depth']) for s, p in after.items()],
                                 [('Main', 3000.0, 2000.0), ('Tail', 2000.0, 1500.0), ('Lead', 1500.0, 1000.0)])

    def test_actual_word_distinct_optional_zero_excess_for_five_workflows(self):
        for job in audit.BATCH1_JOBS:
            with self.subTest(job=job):
                app = self.configured_app(job)
                for values, expected in (((25.0, 10.5), ('25%', '10.5%')),
                                         ((None, 10.5), ('N/A', '10.5%')),
                                         ((0.0, None), ('0%', 'N/A')),
                                         ((None, 0.0), ('N/A', '0%'))):
                    with self.subTest(values=values):
                        self.phase(app, 'phase2_3')
                        protected = fingerprint({k: app.session_state[k] for k in
                                                  ('fluid_data', 'cement_params', 'cement_calc_Main', 'well_data', 'hardware_table')})
                        for field, value in zip(FIELDS, values):
                            self.excess_input(app, field).set_value(value).run()
                        self.assertEqual(compute_phase_status(app.session_state)['phase2_3']['level'], 'ok')
                        self.export(app)
                        self.assertEqual(fingerprint({k: app.session_state[k] for k in
                                                      ('fluid_data', 'cement_params', 'cement_calc_Main', 'well_data', 'hardware_table')}), protected)
                        doc = Document(BytesIO(app.session_state['_compiled_doc_bytes']))
                        rows = self.word_excess(doc)
                        self.assertEqual(len(rows), 1)
                        self.assertEqual((rows[0][1], rows[0][3]), expected)
                        for label, value in zip((rows[0][0], rows[0][2]), expected):
                            self.assertEqual((label + value).count('%'), 0 if value == 'N/A' else 1)
                        self.assertIn('CSG-OH excess: ' + expected[0], app.session_state['exec_summary_text'])
                        self.assertIn('CSG-CSG excess: ' + expected[1], app.session_state['exec_summary_text'])
                        self.assertNotIn('VOLUME BASIS', app.session_state['exec_summary_text'])

    def test_auto_gradient_word_only_formatting_preserves_formula_and_precision(self):
        app = self.configured_app('CSG 9 5/8"')
        app.session_state['well_auto_fields']['geo_gradient'] = True
        self.phase(app, 'phase2_3')
        expected = ((200 - 80) / (3000 * 3.28084)) * 100
        self.assertEqual(app.session_state['geo_gradient'], expected)
        self.confirm_lab(app)
        self.export(app)
        context = self.context(app)
        self.assertEqual(context['well_data']['geo_gradient'], '1.22')
        doc = Document(BytesIO(app.session_state['_compiled_doc_bytes']))
        self.assertEqual(doc.tables[5].rows[2].cells[3].text, '1.22')
        self.assertEqual(app.session_state['geo_gradient'], expected)
        self.assertEqual(app.session_state['well_data']['geo_gradient'], expected)
        self.assertTrue(app.session_state['well_auto_fields']['geo_gradient'])

    def test_manual_gradient_actual_word_two_decimals_and_provenance(self):
        app = self.configured_app('CSG 9 5/8"')
        for value, expected in ((1.23456789, '1.23'), (1.2, '1.20'), (0.9876, '0.99')):
            with self.subTest(value=value):
                self.phase(app, 'phase2_3')
                app.number_input(key='_w_geo_gradient').set_value(value)
                self.phase(app, 'phase1')
                restored = self.app(audit.round_trip(app.session_state.to_dict()))
                self.phase(restored, 'phase2_3')
                self.assertEqual(restored.session_state['geo_gradient'], value)
                self.assertFalse(restored.session_state['well_auto_fields']['geo_gradient'])
                self.confirm_lab(restored)
                self.export(restored)
                doc = Document(BytesIO(restored.session_state['_compiled_doc_bytes']))
                self.assertEqual(doc.tables[5].rows[2].cells[3].text, expected)
                self.assertEqual(restored.session_state['geo_gradient'], value)
                self.assertEqual(restored.session_state['well_data']['geo_gradient'], value)
                self.assertFalse(restored.session_state['well_auto_fields']['geo_gradient'])

    def test_missing_gradient_retains_missing_context_and_actual_word(self):
        app = self.configured_app('CSG 9 5/8"')
        self.phase(app, 'phase2_3')
        app.number_input(key='_w_geo_gradient').set_value(None).run()
        self.assertIsNone(app.session_state['geo_gradient'])
        # Inspect the presentation boundary independently of the normal required-input gate.
        context = self.context(app, prepared=True)
        self.assertIsNone(context['well_data']['geo_gradient'])
        template = DocxTemplate(str(audit.ROOT / 'master_template.docx'))
        template.render(context)
        data = BytesIO();template.save(data)
        doc = Document(BytesIO(data.getvalue()))
        self.assertEqual(doc.tables[5].rows[2].cells[3].text, 'None')
        self.export(app)
        actual = Document(BytesIO(app.session_state['_compiled_doc_bytes']))
        self.assertEqual(actual.tables[5].rows[2].cells[3].text, 'None')
        self.assertIsNone(app.session_state['geo_gradient'])
