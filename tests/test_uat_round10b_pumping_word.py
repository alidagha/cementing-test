"""Round 10B uses real Build operations and the existing Word pipeline."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
import hashlib
import pickle

from docx import Document
from docx.oxml.ns import qn
import pytest

import pumping_simulation as pumping
import phase_10_procedure as report
import project_state
import well_report
import test_audit_regressions as audit
from centralizer_fixtures import execute_centralizer
from test_uat_round10a_pumping_simulation import program_state
from test_uat_round9e_dynamic_solver import auxiliary

ANCHOR = "Pumping simulation charts: Not provided."
NA = "Pumping simulation charts: Not applicable."
DESCRIPTION = "Pumping Simulation Charts"
JOBS = {"casing": 'CSG 9 5/8"', "liner": 'LNR 7"', "tieback": 'TIE BACK LNR 7"'}


@pytest.fixture(scope="module", params=JOBS)
def project(request):
    kind = request.param
    case = audit.AuditRegressions()
    app = case.configured_app(JOBS[kind])
    geometry = program_state(kind)
    for column in ("Collapse (psi)", "Burst (psi)"):
        geometry["hardware_table"][column] = app.session_state["hardware_table"].iloc[0][column]
    for key in ("hardware_table", "placement_config"):
        app.session_state[key] = deepcopy(geometry[key])
    assert not app.exception and not app.error
    execute_centralizer(app.session_state)
    app.session_state["centralizer_config"]["aux_rheology"] = auxiliary()
    state = app.session_state.to_dict()
    assert project_state.prepare_calculations(state) == []
    pumping.prepare_pumping_simulation(state)
    return state


def run_build(state, *, processor=None, renderer=None, on_execute=None):
    case = audit.AuditRegressions()
    app = case.app(deepcopy(state))
    calls = []
    original = pumping.run_pumping_simulation
    def execute(current):
        result = original(current)
        if on_execute is not None:
            on_execute(result)
        calls.append((deepcopy(current.to_dict()), result))
        return result
    with patch.object(pumping, "run_pumping_simulation", side_effect=execute):
        case.phase(app, "phase10")
        assert not calls, "Simulation ran before Build"
        build = next(button for button in app.button if button.label == "Build Word Document")
        assert not build.disabled, [error.value for error in app.error]
        if processor is not None:
            with patch.object(report, "postprocess_well_report", side_effect=processor):
                build.click().run(timeout=120)
        elif renderer is not None:
            with patch.object(report.DocxTemplate, "render", side_effect=renderer):
                build.click().run(timeout=120)
        else:
            build.click().run(timeout=120)
    assert not app.exception
    return app, calls


def drawing_paragraph(doc, description):
    matches = [p for p in doc.paragraphs if p._p.xpath(f'.//wp:docPr[@descr="{description}"]')]
    assert len(matches) == 1
    return matches[0]


def image_bytes(doc, paragraph):
    rid = paragraph._p.xpath('.//a:blip/@r:embed')
    assert len(rid) == 1
    return doc.part.related_parts[rid[0]].blob


def test_source_signature_is_pure_and_has_no_simulation_or_result_dependency():
    signature = getattr(pumping, "source_signature", None)
    assert callable(signature), "Round 10B has no authoritative pumping source signature"
    state = program_state(fluids={"Main": 2., "Displacement Fluid": 2.})
    before = pickle.dumps(state)
    with patch.object(pumping.fcm, "build_three_charts", side_effect=AssertionError("Signature ran solver")):
        first = signature(state)
        state["yield_point"] = "2"
        assert signature(state) != first
        state["yield_point"] = "1"
        state.update(_unrelated_widget="draft", unrelated_report_text="presentation")
        assert signature(state) == first
        state.pop("_unrelated_widget"); state.pop("unrelated_report_text")
    assert pickle.dumps(state) == before


def test_preparation_failure_blocks_export_and_removes_old_document(project):
    state = deepcopy(project)
    state.update(yield_point="0", _compiled_doc_bytes=b"old", _compiled_doc_filename="old.docx",
                 _compiled_doc_signature="old")
    issues = project_state.prepare_calculations(state)
    assert any("Pumping Simulation" in issue and "YP > 0" in issue for issue in issues), issues
    assert not any(key.startswith("_compiled_doc_") for key in state)


def test_real_build_inserts_actual_three_chart_figure_in_section_iv(project):
    from cement_chart_report import render_chart_png
    baseline = pumping.fcm.plt.get_fignums()
    app, calls = run_build(project)
    assert not app.error, [error.value for error in app.error]
    assert len(calls) == 1, "Build did not execute a fresh Pumping Simulation"
    result = calls[0][1]
    assert len(result["figure"].axes) == 3
    doc = Document(BytesIO(app.session_state["_compiled_doc_bytes"]))
    pumping_paragraph = drawing_paragraph(doc, DESCRIPTION)
    diagnostic = drawing_paragraph(doc, "Static Well Security and Fluid Sequence")
    elements = list(doc._element.body)
    iv = next(p for p in doc.paragraphs if p.text.strip() == "IV. Fluids Sequence")
    v = next(p for p in doc.paragraphs if p.text.strip() == "V. Cementing Program")
    assert elements.index(iv._p) < elements.index(diagnostic._p) < elements.index(pumping_paragraph._p) < elements.index(v._p)
    assert ANCHOR not in [p.text for p in doc.paragraphs]
    png = BytesIO(); result["figure"].savefig(png, format="png", dpi=200, facecolor="white")
    assert image_bytes(doc, pumping_paragraph) == png.getvalue()
    assert image_bytes(doc, diagnostic) == render_chart_png(calls[0][0]["_cement_chart_model"])
    assert any(p._p.xpath('.//wp:docPr[@descr="Pipe Standoff"]') for p in doc.paragraphs)
    word_text = "\n".join(p.text for p in doc.paragraphs)
    assert all(warning not in word_text for warning in result["warnings"])
    assert len(pumping_paragraph._p.xpath('.//w:drawing')) == 1
    original = next(p for p in Document(report.TEMPLATE_PATH).paragraphs if p.text == ANCHOR)
    assert pumping_paragraph._p.pPr.xml == original._p.pPr.xml
    assert len(pumping_paragraph._p.xpath('.//w:lastRenderedPageBreak')) == len(original._p.xpath('.//w:lastRenderedPageBreak'))
    assert pumping.fcm.plt.get_fignums() == baseline
    assert not any(isinstance(value, pumping.fcm.SimResult) for value in app.session_state.to_dict().values())
    saved = audit.round_trip(app.session_state.to_dict())
    assert not any(key.startswith("_pumping") for key in saved)


@pytest.mark.parametrize("job", ["CMT PLUG", "CMT SQUEEZE"])
def test_special_jobs_do_not_execute_and_replace_only_the_pumping_anchor(job):
    case = audit.AuditRegressions()
    app = case.configured_app(job)
    with patch.object(pumping.fcm, "build_three_charts", side_effect=AssertionError("N/A executed solver")):
        case.export(app)
    doc = Document(BytesIO(app.session_state["_compiled_doc_bytes"]))
    text = [p.text for p in doc.paragraphs]
    assert text.count(NA) == 1 and ANCHOR not in text
    assert not any(p._p.xpath(f'.//wp:docPr[@descr="{DESCRIPTION}"]') for p in doc.paragraphs)


def test_protected_template_and_frozen_engine_remain_exact():
    assert hashlib.sha256(Path(report.TEMPLATE_PATH).read_bytes()).hexdigest() == "a1128b0e9352e1dd7c8768eecc9335c50858356cee80ed33c0ec275e26c4d08b"
    assert hashlib.sha256(Path(pumping.fcm.__file__).read_bytes()).hexdigest() == "ceaaff1311b241c8793fc65b0004a8ca28bef463d93db0957d9b650522bdc5ea"


@pytest.mark.parametrize("failure", ["postprocess", "template", "success"])
def test_owned_figure_is_closed_on_later_failure_without_closing_other_figures(project, failure):
    from streamlit.runtime.scriptrunner import script_runner
    original_cleanup = script_runner._clean_problem_modules
    figures, observed = {}, []
    def after_execute(result):
        figures.update(owned=result["figure"], other=pumping.fcm.plt.figure())
    def before_native_cleanup():
        if figures and not observed:
            active = pumping.fcm.plt.get_fignums()
            observed.append((figures["owned"].number in active, figures["other"].number in active))
        original_cleanup()
    try:
        error = ValueError("deliberate later Word failure")
        with patch.object(script_runner, "_clean_problem_modules", side_effect=before_native_cleanup):
            app, calls = run_build(project, on_execute=after_execute,
                **({"processor": error} if failure == "postprocess" else {"renderer": error} if failure == "template" else {}))
        assert len(calls) == 1
        if failure == "success":
            assert not app.error and "_compiled_doc_bytes" in app.session_state
        else:
            assert any("deliberate later Word failure" in item.value for item in app.error)
            assert "_compiled_doc_bytes" not in app.session_state
        assert observed == [(False, True)], "Build must close only its own figure before native Streamlit cleanup"
    finally:
        for figure in figures.values():
            pumping.fcm.plt.close(figure)


@pytest.mark.parametrize("field", ["plastic_viscosity", "yield_point", "bhsp", "mud_density"])
def test_canonical_mud_and_pressure_sources_change_signature(field):
    state = program_state(fluids={"Main": 2., "Displacement Fluid": 2.})
    old = pumping.source_signature(state)
    state[field] = "17"
    assert pumping.source_signature(state) != old


@pytest.mark.parametrize("edit", ["density", "volume", "pump_rate", "lab", "gel", "aux", "hardware", "target", "host"])
def test_other_consumed_sources_change_signature_without_execution(edit):
    state = program_state("tieback", fluids={"Main": 2., "Spacer Ahead": 1., "Displacement Fluid": 2.})
    old = pumping.source_signature(state)
    if edit in ("density", "volume", "pump_rate"):
        state["fluids_config"]["params"]["Main"][edit] = "117" if edit == "density" else 3.
    elif edit in ("lab", "gel"):
        row = state["lab_qc_params"]["Main"]["rheology"]["bhct_down"]
        if edit == "lab": row["readings"]["300"] = "169"
        else: row["gel_10_min"] = "11"
    elif edit == "aux": state["centralizer_config"]["aux_rheology"]["Spacer Ahead"]["300"] = "169"
    elif edit == "hardware": state["hardware_table"].at[0, "ID (in)"] = 6.1
    else: state["placement_config"][edit + "_row"] = "changed selection"
    with patch.object(pumping.fcm, "build_three_charts", side_effect=AssertionError("Signature ran solver")):
        assert pumping.source_signature(state) != old


def test_signature_ignores_unused_and_cached_values():
    state = program_state(fluids={"Main": 2., "Displacement Fluid": 2.})
    old = pumping.source_signature(state)
    state["effective_pv"] = state["effective_yp"] = 123.
    state["fluid_data"] = {"Main": {"min_rate": 999., "effective_density": 999.}}
    state["hardware_table"]["Grade"] = "informational"
    state["centralizer_config"]["aux_rheology"]["Spacer"]["300"] = 999.
    state["cement_params"] = {"Main": {"top_mode": "Depth (m MD)", "top_depth": 111., "report_bottom_depth": 222.}}
    state["lab_qc_params"]["Main"]["rheology"]["surface_down"] = {"selected": False}
    assert pumping.source_signature(state) == old


def test_two_builds_execute_twice_from_current_state_and_source_changes_remove_old_doc(project):
    case = audit.AuditRegressions()
    app = case.app(deepcopy(project))
    original = pumping.run_pumping_simulation
    values, figures = [], []
    def execute(current):
        values.append(current["yield_point"])
        result = original(current)
        figures.append(result["figure"])
        return result
    baseline = pumping.fcm.plt.get_fignums()
    with patch.object(pumping, "run_pumping_simulation", side_effect=execute):
        case.phase(app, "phase10")
        assert values == []
        next(button for button in app.button if button.label == "Build Word Document").click().run(timeout=120)
        assert len(values) == 1 and "_compiled_doc_bytes" in app.session_state
        app.session_state["yield_point"] = "16"
        app.session_state["well_data"]["yield_point"] = "16"
        app.run()
        assert len(values) == 1, "Source edit implicitly executed simulation"
        assert "_compiled_doc_bytes" not in app.session_state
        next(button for button in app.button if button.label == "Build Word Document").click().run(timeout=120)
        assert len(values) == 2 and values[-1] == "16"
        assert figures[0] is not figures[1]
        assert not app.exception and not app.error
        assert "_compiled_doc_bytes" in app.session_state
    assert pumping.fcm.plt.get_fignums() == baseline


def test_current_context_contains_only_pumping_source_signature(project):
    state = deepcopy(project)
    with patch.object(report.st, "session_state", state), \
         patch.object(pumping.fcm, "build_three_charts", side_effect=AssertionError("Context ran solver")):
        context = report.build_master_context()
    assert context["pumping_simulation_source"] == pumping.source_signature(state)
    changed = deepcopy(context)
    changed["pumping_simulation_source"] = "different current sources"
    assert report.document_signature(context, report.TEMPLATE_PATH) != report.document_signature(changed, report.TEMPLATE_PATH)


@pytest.mark.parametrize("fault", ["missing", "duplicate", "outside_iv"])
def test_pumping_anchor_guards_are_controlled_and_do_not_run_simulation(project, fault):
    from docxtpl import DocxTemplate
    state = deepcopy(project)
    with patch.object(report.st, "session_state", state):
        context = report.build_master_context()
    template = DocxTemplate(report.TEMPLATE_PATH)
    template.render(report._word_quantity_context(context))
    raw = BytesIO(); template.save(raw)
    doc = Document(BytesIO(raw.getvalue()))
    anchor = next(p for p in doc.paragraphs if p.text == ANCHOR)
    if fault == "missing": anchor._p.getparent().remove(anchor._p)
    elif fault == "duplicate": doc.add_paragraph(ANCHOR)
    else: doc._element.body.insert(len(doc._element.body)-1, anchor._p)
    changed = BytesIO(); doc.save(changed)
    with patch.object(pumping.fcm, "build_three_charts", side_effect=AssertionError("Postprocess ran solver")):
        with pytest.raises(ValueError, match="anchor"):
            well_report.postprocess_well_report(changed.getvalue(), context["well_profile_report"], context["note_vertical"],
                context["cement_chart_report"], context["centralizer_report"], pumping_report={"applicable": False})
