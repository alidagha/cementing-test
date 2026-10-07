"""Owner-approved D044 60 F dissolved-salt authority and review propagation."""
import unittest
from copy import deepcopy
from io import BytesIO
from unittest.mock import patch

import pandas as pd
from docx import Document
import engineering_tools as eng
import project_state as ps
import test_audit_regressions as audit
from phase_5_cement import build_components
from phase_7_lab import build_lab_df_from_phase5

NODES = ((2., .0360), (4., .0366), (6., .0371), (8., .0377), (10., .0383),
         (12., .0388), (14., .0392), (15., .03946), (16., .0397), (18., .04013),
         (20., .04052), (22., .0409), (25., .04154), (28., .0421), (30., .04245),
         (32., .0428), (35., .04320), (37.2, .0436))


def previous_signature(state, slurry):
    """Pre-D044 signature contract, to prove old reviewed sources become stale."""
    p = state.get('cement_params', {}).get(slurry, {})
    fluid = state.get('fluid_data', {}).get(slurry, {})
    well = state.get('well_data', {})
    return ps.fingerprint({'schema': 3, 'material_property_schema': state.get('material_property_schema', 0),
        'base_cement': p.get('base_cement', 'Cement G Delijan'), 'cmt_sg': eng.resolve_cement_sg(p),
        'density': fluid.get('density', '118.0'), 'effective_density': fluid.get('effective_density'),
        'bhst': well.get('bhst', state.get('bhst', '-')), 'bhsp': well.get('bhsp', ''),
        'additives': state.get('cement_additives_dfs', {}).get(slurry, pd.DataFrame())})


def salt_frame(pct):
    return pd.DataFrame([{'Material Type': 'Nacl', 'Name': 'SALT', 'Physical State': 'Powder',
                         'Mix Method': 'In Mix Water', 'User Input': pct, 'Density': 2.16}])


class D044SaltUAT(unittest.TestCase):
    app = audit.AuditRegressions.app
    healthy = audit.AuditRegressions.healthy
    phase = audit.AuditRegressions.phase
    configured_app = audit.AuditRegressions.configured_app
    export = audit.AuditRegressions.export
    assert_export_blocked = audit.AuditRegressions.assert_export_blocked

    def test_exact_direct_authority_nodes(self):
        self.assertEqual(eng.DISSOLVED_NACL_GAL_PER_LB, NODES)
        for pct, expected in NODES:
            with self.subTest(pct=pct):
                self.assertEqual(eng.dissolved_nacl_gal_per_lb(pct), expected)

    def test_interpolation_is_between_authority_nodes_not_measurement(self):
        for pct, expected in ((3., .0363), (23.5, .04122), (36.1, .0434)):
            with self.subTest(pct=pct):
                self.assertAlmostEqual(eng.dissolved_nacl_gal_per_lb(pct), expected, places=14)

    def test_domain_and_fresh_water_salt_mass_unchanged(self):
        self.assertEqual(eng.dissolved_nacl_gal_per_lb(0), 0.)
        for pct in (-1., .00001, 1., 1.999999, 37.200001, float('nan'), float('inf'), 'bad'):
            with self.subTest(pct=pct), self.assertRaises(ValueError):
                eng.dissolved_nacl_gal_per_lb(pct)
        for pct in (0., 2., 15., 18., 20., 25., 30., 35., 37.2):
            with self.subTest(pct=pct):
                field = eng.calculate_salt_field_amounts(pct, 123.456, 987.65, 31.)
                per_bbl = eng.WATER_DENSITY_PCF * eng.BBL_TO_CUFT * pct / 100
                self.assertEqual(field['lbs_per_bbl'], per_bbl)
                self.assertEqual(field['base_lb'], 123.456 * per_bbl)
                self.assertEqual(field['with_dead_lb'], (123.456 + 31.) * per_bbl)

    def test_mass_volume_and_phase5_phase7_consistency(self):
        from test_uat_round4_mass_balance import formulation
        from phase_5_cement import calculate_base_results
        for tail, av in ((False, .04013), (True, .03946)):
            with self.subTest(tail=tail):
                df = formulation(tail)
                powders, liquids, salt = build_components(df)
                sg = eng.catalog_cement_sg()
                density, volume = (123., 79.) if tail else (152., 731.)
                p = {'base_cement': 'Cement G Delijan', 'cmt_sg': sg, 'cmt_sg_source': 'catalog'}
                r = calculate_base_results(p, volume, density, powders, liquids, salt)
                water_mass = r['water_vol_per_sack'] * eng.WATER_DENSITY_PCF
                salt_mass = water_mass * salt / 100
                cement_av = 110 / (sg * eng.WATER_DENSITY_PCF)
                powder_av = sum(x['percent'] * 110 / 100 / x['density_pcf'] for x in powders)
                liquid_av = sum(x['gal_per_sk'] for x in liquids) / eng.GAL_PER_CUFT
                salt_av = salt_mass * av / eng.GAL_PER_CUFT
                total_av = cement_av + powder_av + liquid_av + r['water_vol_per_sack'] + salt_av
                total_mass = 110 + r['ma_additives'] + water_mass + salt_mass
                self.assertAlmostEqual(r['yield_ft3_per_sk'], total_av, places=12)
                self.assertAlmostEqual(total_mass / total_av, density, places=12)
                self.assertAlmostEqual(r['mix_water_gal_sk'], r['base_fluid_gal_sk'] + salt_mass * av, places=12)
                self.assertEqual(r['salt_lab_gr'], r['lab_water_gr'] * salt / 100)
                self.assertAlmostEqual(r['salt_field_lb'], r['field_water_bbl'] * eng.WATER_DENSITY_PCF * eng.BBL_TO_CUFT * salt / 100, places=10)
                grid = build_lab_df_from_phase5(df, slurry_weight_pcf=density, slurry_volume_bbl=volume, cmt_sg=sg, recalculate_mass=True)
                self.assertEqual(grid.iloc[0]['Mass'], f"{r['lab_cmt_gr']:.1f}")
                self.assertEqual(grid.loc[grid['Material']=='SALT', 'Mass'].iloc[0], f"{r['salt_lab_gr']:.1f}")
                for key in ('base_fluid_gal_sk', 'mix_water_gal_sk', 'mix_fluid_gal_sk'):
                    self.assertEqual(p[key], r[key])

    def test_signature_invalidates_only_positive_salt_without_material_migration(self):
        state = {'material_property_schema': 1, 'cement_additives_dfs': {'Main': salt_frame(18.)}}
        old = previous_signature(state, 'Main')
        self.assertNotEqual(ps.lab_source_signature(state, 'Main'), old)
        current = ps.lab_source_signature(state, 'Main')
        with patch.object(ps, 'SALT_MODEL_VERSION', 2):
            self.assertNotEqual(ps.lab_source_signature(state, 'Main'), current)
        for df in (pd.DataFrame(), salt_frame(0.), salt_frame(18.).assign(Name='NaCl-X', **{'Material Type':'Local Salt'})):
            state['cement_additives_dfs']['Main'] = df
            before = previous_signature(state, 'Main')
            self.assertEqual(ps.lab_source_signature(state, 'Main'), before)
            with patch.object(ps, 'SALT_MODEL_VERSION', 2):
                self.assertEqual(ps.lab_source_signature(state, 'Main'), before)
        self.assertEqual(state['material_property_schema'], 1)

    def old_review_app(self):
        app = self.configured_app(audit.BATCH1_JOBS[0], additives=True)
        app.session_state['cement_additives_dfs']['Main'] = pd.concat([
            app.session_state['cement_additives_dfs']['Main'], salt_frame(18.)], ignore_index=True)
        self.phase(app, 'phase5'); self.phase(app, 'phase7')
        next(b for b in app.button if b.label=='Sync with Phase V').click().run()
        next(b for b in app.button if b.label=='Confirm measured lab results').click().run()
        state = deepcopy(app.session_state.to_dict())
        state['lab_grid_dfs']['Main'].loc[0,'Mass'] = '777.7'
        state['lab_grid_dfs']['Main'].loc[0,'Lot No'] = 'D044-MANUAL-LOT'
        qc = state['lab_qc_params']['Main']
        qc['review_signature'] = eng.lab_review_signature(qc, state['lab_grid_dfs']['Main'])
        state['lab_source_signatures']['Main'] = previous_signature(state, 'Main')
        return state

    def test_old_salt_review_direct_restore_keep_confirm_word(self):
        state = self.old_review_app(); state['_compiled_doc_bytes'] = b'old compiled Word'
        self.assertTrue(any('Phase VII' in e for e in ps.prepare_calculations(state)))
        self.assertNotIn('_compiled_doc_bytes', state)
        self.assertEqual(state['material_property_schema'], 1)
        app = self.app(audit.round_trip(state)); self.assert_export_blocked(app)
        self.phase(app,'phase7')
        old = app.session_state['lab_grid_dfs']['Main'].copy(deep=True)
        next(b for b in app.button if b.label=='Keep reviewed lab entries').click().run()
        pd.testing.assert_frame_equal(app.session_state['lab_grid_dfs']['Main'], old)
        self.assertFalse(app.session_state['lab_qc_params']['Main']['reviewed'])
        self.assert_export_blocked(app); self.phase(app,'phase7')
        next(b for b in app.button if b.label=='Confirm measured lab results').click().run()
        self.export(app)
        text = '\n'.join(c.text for t in Document(BytesIO(app.session_state['_compiled_doc_bytes'])).tables for r in t.rows for c in r.cells)
        self.assertIn('777.7', text); self.assertIn('D044-MANUAL-LOT', text)
        self.export(self.app(audit.round_trip(app.session_state.to_dict())))

    def test_old_salt_review_sync_confirm_fresh_direct_word_and_full_precision(self):
        state = self.old_review_app(); app = self.app(audit.round_trip(state)); self.assert_export_blocked(app)
        self.phase(app,'phase7')
        next(b for b in app.button if b.label=='Sync with Phase V').click().run()
        s=app.session_state.to_dict(); p=s['cement_params']['Main']
        expected=build_lab_df_from_phase5(s['cement_additives_dfs']['Main'], slurry_weight_pcf=118., slurry_volume_bbl=50., cmt_sg=p['cmt_sg'], existing_df=state['lab_grid_dfs']['Main'], recalculate_mass=True)
        pd.testing.assert_frame_equal(s['lab_grid_dfs']['Main'], expected)
        self.assertEqual(expected.iloc[0]['Lot No'],'D044-MANUAL-LOT')
        self.assertNotEqual(expected.iloc[0]['Mass'],'777.7')
        self.assertFalse(s['lab_qc_params']['Main']['reviewed'])
        self.assert_export_blocked(app); self.phase(app,'phase7')
        next(b for b in app.button if b.label=='Confirm measured lab results').click().run()
        self.export(app)
        restored=self.app(audit.round_trip(app.session_state.to_dict())); self.export(restored)
        s=restored.session_state.to_dict()
        for field,key in [('base_fluid','base_fluid_gal_sk'),('mix_water','mix_water_gal_sk'),('mix_fluid','mix_fluid_gal_sk')]:
            self.assertEqual(s['lab_payload_Main'][field],s['cement_params']['Main'][key])
        text='\n'.join(c.text for t in Document(BytesIO(s['_compiled_doc_bytes'])).tables for r in t.rows for c in r.cells)
        for mass in expected['Mass']:self.assertIn(str(mass),text)
