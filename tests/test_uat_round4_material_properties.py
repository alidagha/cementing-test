"""Owner-approved exact properties and one-time material schema migration."""
import unittest
from copy import deepcopy
from io import BytesIO
from unittest.mock import patch
from docx import Document
import pandas as pd
import engineering_tools as eng
import materials_db as db
import project_state as ps
import test_audit_regressions as audit

EXACT = {**{n:84.90203857 for n in db.MATERIAL_TAXONOMY['F.L. Controller']['names']},
         **{n:89.27199554 for n in db.MATERIAL_TAXONOMY['Dispersant']['names']},
         **{n:76.78639984 for t in ('L.T. Retarder','H.T. Retarder') for n in db.MATERIAL_TAXONOMY[t]['names']},
         'Cacl2':109.2489471, 'Anti Settling':157.9427643, 'Micro Silica':137.3415375,
         'Silica Flour':165.4341278,'Light Weight':46.82097626,'Cenosphere':46.82097626,
         'Cement G Delijan':199.7695007,'Cement D Delijan':197.184}
OLD = {'O-uniFLC5':1.36,'O-CFR4':1.43,'O-R5':1.23,'Cacl2':1.75,
       'Anti Settling':2.53,'Micro Silica':2.20,'Silica Flour':2.65,'Light Weight':.75,
       'Cenosphere':.75,'Anti Foam':1.00}
def rows(values):
    return pd.DataFrame([{'Material Type':db.NAME_TO_MATERIAL[n][0], 'Name':n,
        'Physical State':db.NAME_TO_MATERIAL[n][1], 'Mix Method':'In Mix Water',
        'User Input':.2,'Density':v} for n,v in values])

class MaterialProperties(unittest.TestCase):
    app = audit.AuditRegressions.app
    phase = audit.AuditRegressions.phase
    healthy = audit.AuditRegressions.healthy
    configured_app = audit.AuditRegressions.configured_app
    export = audit.AuditRegressions.export
    assert_export_blocked = audit.AuditRegressions.assert_export_blocked

    def sg(self, app):
        return next(w for w in app.number_input if w.label == 'Cement SG - Main')

    def test_cement_catalog_switch_manual_edit_and_restore(self):
        app=self.app({'fluids_config':{'active':['Main'],'params':{'Main':{
            'volume':50.,'density':'118','pump_rate':'4'}}}})
        self.phase(app,'phase4');self.phase(app,'phase5')
        p=app.session_state['cement_params']['Main']
        self.assertEqual((p['cmt_sg_source'],p['cmt_sg']),('catalog',199.7695007/62.4))
        self.assertEqual(self.sg(app).proto.format,'%.2f')
        brand=next(w for w in app.selectbox if w.label=='Base Cement - Main')
        brand.set_value('Cement D Delijan').run()
        self.assertEqual(self.sg(app).value,3.16)
        self.sg(app).set_value(3.17).run()
        self.sg(app).set_value(3.16).run()  # A manual value equal to catalog stays manual.
        self.assertEqual(app.session_state['cement_params']['Main']['cmt_sg_source'],'manual')
        next(w for w in app.selectbox if w.label=='Base Cement - Main').set_value('Cement G Delijan').run()
        self.assertEqual(self.sg(app).value,3.16)
        restored=self.app(audit.round_trip(app.session_state.to_dict()))
        self.phase(restored,'phase5');self.assertEqual(self.sg(restored).value,3.16)
        self.assertEqual(restored.session_state['cement_params']['Main']['cmt_sg_source'],'manual')

    def test_live_mounted_editor_and_sg_cannot_undo_migration(self):
        app=self.configured_app(audit.BATCH1_JOBS[0]);self.phase(app,'phase5')
        app.session_state['cement_additives_dfs']['Main']=rows([('O-uniFLC5',1.36),('Micro Silica',2.4)])
        self.phase(app,'phase2_3');self.phase(app,'phase5')
        old=next(w for w in app.dataframe if 'Density' in w.proto.column_order).proto.id
        self.assertEqual(self.sg(app).value,3.20)
        app.session_state['material_property_schema']=0
        app.run();self.healthy(app)
        self.assertEqual(self.sg(app).value,199.7695007/62.4)
        self.assertNotEqual(old,next(w for w in app.dataframe if 'Density' in w.proto.column_order).proto.id)
        for _ in range(2):
            self.phase(app,'phase2_3');self.phase(app,'phase5')
            df=app.session_state['cement_additives_dfs']['Main']
            self.assertEqual(df.iloc[0]['Density'],84.90203857/62.4)
            self.assertEqual(df.iloc[1]['Density'],2.4)
            self.assertEqual(self.sg(app).value,199.7695007/62.4)

    def test_all_legacy_defaults_missing_exact_and_override(self):
        for name,old in eng.LEGACY_ADDITIVE_SG.items():
            with self.subTest(name=name):
                df=rows([(name,old),(name,None),(name,old+.012345),(name,old+1e-12)])
                state={'cement_additives_dfs':{'Lead':df}}
                ps.migrate_material_properties(state)
                got=state['cement_additives_dfs']['Lead']['Density'].tolist()
                new=eng.default_additive_density_gcm3(name,db.NAME_TO_MATERIAL[name][1])
                self.assertEqual(got,[new,new,old+.012345,old+1e-12])

    def test_inactive_migration_round_trip_reactivation(self):
        for name,sg,source,expected in [('Cement G Delijan',3.2,'catalog',199.7695007/62.4),
                                       ('Cement D Delijan',3.16,'catalog',3.16),
                                       ('Cement D Delijan',3.2,'manual',3.2)]:
            with self.subTest(name=name,sg=sg):
                state={'job_type':'CSG 20"','inactive_slurry_drafts':{'Tail':{
                    'cement_params':{'base_cement':name,'cmt_sg':sg},
                    'cement_additives_dfs':rows([('O-R12',1.23)]),
                    'lab_qc_params':{'reviewed':True}}}}
                current=audit.round_trip(state);restored=audit.round_trip(current)
                ps.purge_inactive_slurries(restored,['Tail'])
                p=restored['cement_params']['Tail']
                self.assertEqual((p['cmt_sg_source'],p['cmt_sg']),(source,expected))
                self.assertEqual(restored['cement_additives_dfs']['Tail'].iloc[0]['Density'],76.78639984/62.4)
                self.assertNotIn('Tail',restored.get('lab_source_signatures',{}))

    def test_cement_migration_reuses_existing_exact_catalog_lookup(self):
        for name,sg,source in [('CEMENT G DELIJAN',3.2,'catalog'),
                               ('cement d delijan',3.16,'catalog'),
                               ('CEMENT D DELIJAN',3.2,'manual')]:
            state={'cement_params':{'Main':{'base_cement':name,'cmt_sg':sg}}}
            ps.migrate_material_properties(state)
            self.assertEqual(state['cement_params']['Main']['cmt_sg_source'],source)
            self.assertEqual(state['cement_params']['Main']['base_cement'],name)

    def legacy_review(self):
        app=self.configured_app(audit.BATCH1_JOBS[0],additives=True)
        state=deepcopy(app.session_state.to_dict())
        state.pop('material_property_schema');state['cement_params']['Main'].pop('cmt_sg_source')
        state['cement_additives_dfs']['Main']=rows([('Anti Foam',1.)])
        state['lab_grid_dfs']['Main'].loc[0,'Mass']='777.7'
        state['lab_grid_dfs']['Main'].loc[0,'Lot No']='MANUAL-LOT'
        return state

    def test_direct_restore_export_stales_review_and_keep_preserves_entries(self):
        state=self.legacy_review();state['_compiled_doc_bytes']=b'old Word'
        oldgrid=state['lab_grid_dfs']['Main'].copy(deep=True)
        ps.migrate_material_properties(state)
        self.assertNotIn('_compiled_doc_bytes',state)
        self.assertTrue(any('Phase VII' in issue for issue in ps.prepare_calculations(state)))
        app=self.app(audit.round_trip(state));self.assert_export_blocked(app)
        self.phase(app,'phase7')
        next(b for b in app.button if b.label=='Keep reviewed lab entries').click().run()
        pd.testing.assert_frame_equal(app.session_state['lab_grid_dfs']['Main'],oldgrid)
        self.assertFalse(app.session_state['lab_qc_params']['Main']['reviewed'])
        self.assert_export_blocked(app);self.phase(app,'phase7')
        next(b for b in app.button if b.label=='Confirm measured lab results').click().run()
        self.export(app)
        text='\n'.join(c.text for t in Document(BytesIO(app.session_state['_compiled_doc_bytes'])).tables for r in t.rows for c in r.cells)
        self.assertIn('777.7',text);self.assertIn('MANUAL-LOT',text)

    def test_sync_corrects_masses_metadata_review_word_and_precision(self):
        import phase_10_procedure as report
        from phase_7_lab import build_lab_df_from_phase5
        state=self.legacy_review();app=self.app(audit.round_trip(state));self.phase(app,'phase5');self.phase(app,'phase7')
        next(b for b in app.button if b.label=='Sync with Phase V').click().run()
        s=app.session_state.to_dict();p=s['cement_params']['Main'];df=s['cement_additives_dfs']['Main']
        expected=build_lab_df_from_phase5(df,slurry_weight_pcf=118.,slurry_volume_bbl=50.,
            cmt_sg=p['cmt_sg'],existing_df=state['lab_grid_dfs']['Main'],recalculate_mass=True)
        pd.testing.assert_frame_equal(s['lab_grid_dfs']['Main'],expected)
        self.assertEqual(expected.iloc[0]['Lot No'],'MANUAL-LOT')
        self.assertNotEqual(expected.iloc[0]['Mass'],'777.7')
        self.assertFalse(s['lab_qc_params']['Main']['reviewed'])
        next(b for b in app.button if b.label=='Confirm measured lab results').click().run()
        self.export(app)
        restored=self.app(audit.round_trip(app.session_state.to_dict()))
        self.export(restored)  # Direct fresh restore/export, without opening Phase V.
        s=restored.session_state.to_dict()
        with patch.object(report.st,'session_state',s):ctx=report.build_master_context()
        for field,key in [('base_fluid','base_fluid_gal_sk'),('mix_water','mix_water_gal_sk'),('mix_fluid','mix_fluid_gal_sk')]:
            self.assertEqual(ctx['slurries'][0]['lab'][field],s['cement_params']['Main'][key])
        text='\n'.join(c.text for t in Document(BytesIO(s['_compiled_doc_bytes'])).tables for r in t.rows for c in r.cells)
        for mass in expected['Mass']:
            self.assertIn(str(mass),text)
    def test_exact_authorized_powders(self):
        for name,pcf in EXACT.items():
            with self.subTest(name=name):
                self.assertEqual(eng.resolve_additive_density(name,'Powder'),pcf)
                sg=eng.default_additive_density_gcm3(name)
                self.assertEqual(eng.resolve_additive_density(name,'Powder',sg),pcf)

    def test_exact_antifoam_properties(self):
        for name in ('Anti Foam','Defoamer','TA-47'):
            with self.subTest(name=name):
                ppg,factor=eng.resolve_additive_density(name,'Liquid')
                self.assertEqual(ppg,62.17825699/eng.GAL_PER_CUFT)
                self.assertEqual(factor,ppg/110.)

    def test_active_inactive_migration_and_idempotence(self):
        state={'job_type':'CSG 20"','cement_params':{'Main':{'base_cement':'Cement G Delijan','cmt_sg':3.20}},
               'cement_additives_dfs':{'Main':rows(list(OLD.items()))},
               'inactive_slurry_drafts':{'Tail':{'cement_params':{'base_cement':'Cement D Delijan','cmt_sg':3.20},
                  'cement_additives_dfs':rows([('O-R5',1.23),('Micro Silica',2.4)])}},
               '_compiled_doc_bytes':b'old'}
        self.assertTrue(ps.migrate_material_properties(state))
        self.assertEqual(state['material_property_schema'],1)
        self.assertEqual(state['cement_params']['Main']['cmt_sg_source'],'catalog')
        self.assertEqual(state['cement_params']['Main']['cmt_sg'],199.7695007/62.4)
        self.assertEqual(state['inactive_slurry_drafts']['Tail']['cement_params']['cmt_sg_source'],'manual')
        self.assertEqual(state['inactive_slurry_drafts']['Tail']['cement_params']['cmt_sg'],3.20)
        for _,r in state['cement_additives_dfs']['Main'].iterrows():
            self.assertEqual(r['Density'],eng.default_additive_density_gcm3(r['Name'],r['Physical State']))
        self.assertEqual(state['inactive_slurry_drafts']['Tail']['cement_additives_dfs'].iloc[1]['Density'],2.4)
        self.assertNotIn('_compiled_doc_bytes',state)
        before=ps.fingerprint(state)
        self.assertFalse(ps.migrate_material_properties(state))
        self.assertEqual(before,ps.fingerprint(state))

    def test_current_schema_historical_looking_manual_values_preserved(self):
        state={'job_type':'CSG 20"','material_property_schema':1,
               'cement_params':{'Main':{'cmt_sg':3.20,'cmt_sg_source':'manual'}},
               'cement_additives_dfs':{'Main':rows([('O-uniFLC5',1.36)])}}
        got=audit.round_trip(state)
        self.assertEqual(got['cement_params']['Main']['cmt_sg'],3.20)
        self.assertEqual(got['cement_additives_dfs']['Main'].iloc[0]['Density'],1.36)

    def test_json_rejects_unsupported_version_and_provenance(self):
        for version in (-1,2,True,'1',None):
            with self.subTest(version=version),self.assertRaises(ValueError):
                audit.round_trip({'job_type':'CSG 20"','material_property_schema':version})
        for source in (None,'automatic','',False):
            with self.subTest(source=source),self.assertRaises(ValueError):
                audit.round_trip({'job_type':'CSG 20"','material_property_schema':1,
                    'cement_params':{'Main':{'cmt_sg':3.2,'cmt_sg_source':source}}})
            with self.subTest(inactive_source=source),self.assertRaises(ValueError):
                audit.round_trip({'job_type':'CSG 20"','material_property_schema':1,
                    'inactive_slurry_drafts':{'Tail':{'cement_params':{'cmt_sg':3.2,'cmt_sg_source':source}}}})

    def test_explicit_catalog_and_manual_cement_authority(self):
        from phase_5_cement import calculate_base_results
        for name in ('Cement G Delijan','Cement D Delijan'):
            p={'base_cement':name,'cmt_sg':3.2,'cmt_sg_source':'catalog'}
            calculate_base_results(p,50.,118.,[],[],0.)
            self.assertEqual(p['cmt_sg'],eng.catalog_cement_sg(name))
            p['cmt_sg_source']='manual';p['cmt_sg']=3.2
            calculate_base_results(p,50.,118.,[],[],0.)
            self.assertEqual(p['cmt_sg'],3.2)

    def test_json_migration_before_open(self):
        got=audit.round_trip({'job_type':'CSG 20"','cement_params':{'Main':{'cmt_sg':3.2}},
                    'cement_additives_dfs':{'Main':rows([('O-uniFLC5',None),('O-CFR4',1.5)])}})
        self.assertEqual(got['material_property_schema'],1)
        self.assertEqual(got['cement_additives_dfs']['Main'].iloc[0]['Density'],84.90203857/62.4)
        self.assertEqual(got['cement_additives_dfs']['Main'].iloc[1]['Density'],1.5)

    def test_protected_properties(self):
        for n,sg in {'Bentonite':2.65,'Hidense':5.2,'Micro Max':4.8,'Boric Acid':1.43,'SALT':2.16}.items():
            self.assertEqual(eng.resolve_additive_density(n),sg*62.4)
        for n in db.MATERIAL_TAXONOMY['Anti Gas Migration']['names']:
            self.assertEqual(eng.resolve_additive_density(n,'Liquid')[0],1.05*eng.WATER_LB_PER_GAL)
        self.assertEqual(eng.resolve_additive_density('Micro Block','Liquid')[0],1.32*eng.WATER_LB_PER_GAL)
