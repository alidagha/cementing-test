"""Round 2 Batch 5: report order is independent of engineering traversal."""
from copy import deepcopy
from io import BytesIO
from html import unescape
import re
import unittest
from unittest.mock import patch

import pandas as pd
from docx import Document
from docx.oxml.ns import qn
from docxtpl import DocxTemplate
import test_audit_regressions as audit
import phase_10_procedure as report
from phase_5_cement import (build_components, calculate_base_results,
                            build_cement_tables, refresh_cement_calculations)
from phase_7_lab import build_lab_df_from_phase5
from placement import slurry_intervals, hardware_choices, HOST_DESCRIPTIONS

ORDER = ['Lead', 'Lead #1', 'Lead #2', 'Main', 'Tail']
HYDRAULIC_SLURRIES = ['Main', 'Lead', 'Lead #1', 'Lead #2', 'Tail']


class Round2ReportUAT(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        case = audit.AuditRegressions()
        app = case.configured_app('CSG 9 5/8"', optional_fluids=True)
        project = audit.round_trip(app.session_state.to_dict())
        project['fluids_config']['active'] = ['Pre Flush', 'Spacer', 'Spacer Ahead',
                                             *ORDER, 'Spacer Behind', 'Displacement Fluid']
        for i, slurry in enumerate(ORDER):
            project['fluids_config']['params'][slurry] = deepcopy(project['fluids_config']['params']['Main'])
            project['cement_params'][slurry] = deepcopy(project['cement_params']['Main'])
            placement_index = ['Main', 'Lead', 'Lead #1', 'Lead #2', 'Tail'].index(slurry)
            project['cement_params'][slurry].update(top_mode='Depth (m MD)', top_depth=500. + placement_index * 500,
                                                    top_job_type=project['job_type'])
            project['cement_additives_dfs'][slurry] = pd.DataFrame([
                {'Material Type': 'Extender', 'Name': 'Micro Silica', 'Physical State': 'Powder',
                 'Mix Method': 'Dry Blend', 'User Input': 5., 'Density': 2.3},
                {'Material Type': 'F.L. Controller', 'Name': 'O-uniFLC5', 'Physical State': 'Powder',
                 'Mix Method': 'In Mix Water', 'User Input': .5, 'Density': 1.5},
                {'Material Type': 'Anti Gas Migration', 'Name': 'O-GAS BLOCK', 'Physical State': 'Liquid',
                 'Mix Method': 'In Mix Water', 'User Input': .1, 'Density': 1.2}])
            project['lab_qc_params'][slurry] = {'free_water': float(i),
                'free_water_45': i + .5, 'surface_hardened_hours': 8. + i,
                'thickening_time': '03:30', 'thickening_endpoint': '70 Bc',
                'rheology': audit.rheology_fixture(), **audit.thickening_fixture(), **audit.compressive_fixture()}
        app = case.app(audit.round_trip(project))
        for phase in ('phase2_3', 'phase4', 'phase5', 'phase6', 'phase7'):
            case.phase(app, phase)
        for button in list(app.button):
            if button.label == 'Sync with Phase V':
                app.button(key=button.key).click().run()
        for _ in ORDER:
            next(b for b in app.button if b.label == 'Confirm measured lab results').click().run()
        case.export(app)
        cls.state = app.session_state.to_dict()

    def context(self, state=None):
        with patch.object(report.st, 'session_state', deepcopy(state or self.state)):
            return report.build_master_context()

    def word(self, context):
        doc = DocxTemplate(report.TEMPLATE_PATH)
        doc.render(context)
        output = BytesIO(); doc.save(output)
        return Document(BytesIO(output.getvalue()))

    def assert_word_order(self, doc, expected):
        program = [t.rows[0].cells[0].text.removesuffix(' Cement Slurry Data') for t in doc.tables
                   if t.rows[0].cells[0].text.endswith(' Cement Slurry Data')]
        lab = [p.text.removesuffix(' Cement Slurry Results') for p in doc.paragraphs
               if p.text.endswith(' Cement Slurry Results')]
        self.assertEqual(program, expected)
        self.assertEqual(lab, expected)

    def test_actual_word_sections_share_one_report_order(self):
        context = self.context()
        self.assertEqual([s['name'] for s in context['slurries']], ORDER)
        self.assert_word_order(self.word(context), ORDER)

    def test_subset_order_for_all_five_job_families(self):
        for job in audit.BATCH1_JOBS:
            for subset in (['Lead', 'Tail'], ['Lead #2', 'Main', 'Tail'], ['Main']):
                with self.subTest(job=job, subset=subset):
                    state = deepcopy(self.state)
                    state['job_type'] = state['doc_control']['job_type'] = job
                    state['placement_config'].update(job_type=job, host_row=next(iter(hardware_choices(state['hardware_table'], HOST_DESCRIPTIONS))))
                    state['fluids_config']['active'] = [*subset, 'Displacement Fluid']
                    for slurry in subset:
                        state['cement_params'][slurry].update(top_job_type=job, last_job_type=job)
                    # Current well/formulation inputs remain unchanged; report text is auto-generated.
                    state.pop('procedure_text', None); state.pop('exec_summary_text', None)
                    state.pop('report_text_state', None)
                    context = self.context(state)
                    self.assertEqual([s['name'] for s in context['slurries']], subset)
                    self.assert_word_order(self.word(context), subset)

    def test_report_order_ignores_interval_mapping_insertion_order(self):
        real_intervals = report.slurry_intervals
        def reverse_mapping(*args, **kwargs):
            return dict(reversed(list(real_intervals(*args, **kwargs).items())))
        original = self.context()
        with patch.object(report, 'slurry_intervals', side_effect=reverse_mapping):
            reordered = self.context()
        self.assertEqual(reordered['slurries'], original['slurries'])
        self.assertEqual([s['name'] for s in reordered['slurries']], ORDER)

    def test_chained_interval_values_and_summary_remain_engineering_order(self):
        for job in ('CSG 9 5/8"', 'LNR 7"'):
            with self.subTest(job=job):
                state = deepcopy(self.state)
                state['job_type'] = state['doc_control']['job_type'] = job
                state['placement_config']['job_type'] = job
                for p in state['cement_params'].values():
                    p.update(top_job_type=job, last_job_type=job)
                state.pop('procedure_text', None); state.pop('exec_summary_text', None)
                state.pop('report_text_state', None)
                intervals = slurry_intervals(state['hardware_table'], job, state['placement_config'],
                                             HYDRAULIC_SLURRIES, state['cement_params'])
                self.assertEqual(list(intervals), list(reversed(HYDRAULIC_SLURRIES)))
                context = self.context(state)
                for payload in context['slurries']:
                    self.assertEqual(payload['top'], intervals[payload['name']]['top'])
                    self.assertEqual(payload['bottom'], f"{intervals[payload['name']]['bottom_depth']:.1f} m MD")
                self.assertEqual([name.title() for name in re.findall(
                    r'(lead(?: #\d)?|main|tail) cement slurry will be cemented from', context['exec_summary'])], list(reversed(HYDRAULIC_SLURRIES)))
                self.assert_word_order(self.word(context), ORDER)

    def test_operational_fluids_procedure_and_timing_are_not_sorted_for_report(self):
        state = deepcopy(self.state)
        context = self.context(state)
        self.assertEqual([f['name'] for f in context['fluids_train']], state['fluids_config']['active'])
        self.assertEqual(unescape(context['procedure_text']), state['procedure_text'])
        positions = [context['procedure_text'].index(f'{s} cement slurry (') for s in HYDRAULIC_SLURRIES]
        self.assertEqual(positions, sorted(positions))
        self.assertEqual(state['total_pump_time_min'], 125.)
        self.assertIn('Safety Factor is 156 min', context['note_maxpump'])

    def test_existing_preflush_spacer_payload_and_document_flow(self):
        context = self.context()
        self.assertTrue(context['has_preflush']); self.assertTrue(context['has_spacer'])
        self.assertEqual([s['Name'] for s in context['spacers'] if s['Material Type'] == 'Spacer'],
                         ['Spacer', 'Spacer Ahead', 'Spacer Behind'])
        self.assertEqual(context['spacer_volume_bbl'], 150.)
        self.assertEqual(context['spacer_density_note'], '80.0')
        self.assertEqual(context['preflush']['volume'], 50.)
        doc = self.word(context)
        titles = [t.rows[0].cells[0].text for t in doc.tables]
        self.assertEqual(titles.count('Pre Flush Data'), 1)
        self.assertEqual(titles.count('Spacer Formulation Data'), 1)
        spacer = doc.tables[titles.index('Spacer Formulation Data')]
        self.assertEqual([r.cells[1].text for r in spacer.rows[3:] if r.cells[2].text == 'Spacer'],
                         ['Spacer', 'Spacer Ahead', 'Spacer Behind'])
        # XML body order includes paragraphs and tables, unlike separate python-docx lists.
        text = '\n'.join(''.join(t.text or '' for t in node.iter(qn('w:t'))) for node in doc.element.body)
        self.assertLess(text.index('Tail Cement Slurry Additives Data'), text.index('Pre Flush Data'))
        self.assertLess(text.index('Pre Flush Data'), text.index('Spacer Formulation Data'))
        self.assertLess(text.index('Spacer Formulation Data'), text.rindex('VII. Laboratory Report'))
        self.assertNotIn('Spacer Cement Slurry Results', text)
        self.assertNotIn('Pre Flush Cement Slurry Results', text)

    def test_both_mix_note_paths_keep_base_water_tank_and_remove_only_suffix(self):
        context = self.context()
        for s in context['slurries']:
            with self.subTest(slurry=s['name']):
                p = deepcopy(self.state['cement_params'][s['name']])
                df = self.state['cement_additives_dfs'][s['name']]
                result = calculate_base_results(p, s['volume_bbl'], float(s['density_pcf']), *build_components(df))
                blend, adds, note = build_cement_tables(p, df, result)
                self.assertEqual(note, f"Mix above Additives in **{p['mix_water'] + p['dead_vol']:.1f} bbl** Fresh Water at **{p['tank_name']}**.")
                self.assertRegex(s['note_mix'], rf"^NOTE \d+: Mix above Additives in {s['total_water_bbl']:.1f} bbl Fresh Water at {p['tank_name']}\.$")
                pd.testing.assert_frame_equal(blend, self.state['cement_blend_' + s['name']])
                pd.testing.assert_frame_equal(adds, self.state['cement_calc_' + s['name']])
                self.assertIn('Micro Silica', blend['Name'].tolist())
                self.assertNotIn('Micro Silica', adds['Name'].tolist())
        doc = self.word(context)
        text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertNotIn('pre-blended dry with bulk cement', text)
        self.assertEqual(text.count('Mix above Additives in '), 5)

    def test_note_counter_sequence_and_slurry_association(self):
        context = self.context()
        numbers = [int(re.match(r'NOTE (\d+):', n).group(1)) for n in context['all_notes']]
        self.assertEqual(numbers, [n for n in range(1, 28) if n != 20])
        self.assertEqual(len(numbers), len(set(numbers)))
        for i, s in enumerate(context['slurries']):
            self.assertTrue(s['note_mix'].startswith(f'NOTE {7 + i}:'))
            self.assertTrue(s['lab']['note_thickening'].endswith(f'After {8+i} Hours'))
        doc = self.word(context)
        actual = [int(re.match(r'NOTE (\d+):', p.text.strip()).group(1)) for p in doc.paragraphs
                  if re.match(r'NOTE \d+:', p.text.strip())]
        self.assertEqual(actual, numbers + [20])

    def test_density_override_engineering_and_lab_are_unchanged_by_notes(self):
        state = deepcopy(self.state)
        before = deepcopy(state['cement_params'])
        self.assertEqual(refresh_cement_calculations(state), [])
        self.assertEqual(state['cement_params'], before)
        for s in ORDER:
            df = state['cement_additives_dfs'][s]
            powders, liquids, _ = build_components(df)
            self.assertEqual(powders[0]['density_pcf'], 2.3 * 62.4)
            self.assertEqual(powders[1]['density_pcf'], 1.5 * 62.4)
            self.assertEqual(liquids[0]['density_ppg'], 1.2 * (62.4 / 7.48051945))
            expected = build_lab_df_from_phase5(df, base_cement=before[s]['base_cement'],
                slurry_weight_pcf=118., slurry_volume_bbl=50., cmt_sg=3.2, recalculate_mass=True)
            pd.testing.assert_frame_equal(expected, state['lab_grid_dfs'][s])


if __name__ == '__main__':
    unittest.main()
