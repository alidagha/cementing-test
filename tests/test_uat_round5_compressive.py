"""Round 5-3: independent measured UCA/Crush sources, gates and DOCX."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import pytest
from docx import Document
import test_audit_regressions as audit
from engineering_tools import compute_phase_status, lab_review_signature


def measured(uca=True, crush=True):
    return {'uca': {'selected': uca, 'cs_8': 100.25, 'cs_12': 200.5, 'cs_24': 300.125},
            'crush': {'selected': crush, 'force_1': 400.0, 'force_2': 401.0,
                      'force_3': 402.0, 'force_4': 403.0}}


def project(uca=True, crush=True):
    case = audit.AuditRegressions()
    app = case.configured_app('CSG 9 5/8"')
    state = audit.round_trip(app.session_state.to_dict())
    qc = state['lab_qc_params']['Main']
    qc.pop('comp_test', None)
    qc['compressive'] = measured(uca, crush)
    qc['review_signature'] = lab_review_signature(qc, state['lab_grid_dfs']['Main'])
    return state


def table(doc):
    return next(t for t in doc.tables if t.rows[0].cells[0].text == 'Compressive Strength Test')


def test_fresh_qc_has_no_implicit_crush():
    case = audit.AuditRegressions()
    state = project()
    state['lab_qc_params']['Main'].pop('compressive')
    app = case.app(state)
    case.phase(app, 'phase7')
    qc = app.session_state['lab_qc_params']['Main']
    assert qc.get('comp_test') != 'CRUSH'
    assert not qc['compressive']['uca']['selected']
    assert not qc['compressive']['crush']['selected']


def test_two_independent_checkboxes_replace_selectbox():
    case = audit.AuditRegressions()
    app = case.app(project())
    case.phase(app, 'phase7')
    assert {'UCA - Main', 'Crush Test - Main'} <= {w.label for w in app.checkbox}
    assert not any('Compressive Test' in w.label for w in app.selectbox)


def test_neither_selected_blocks_status_and_export():
    case = audit.AuditRegressions()
    app = case.app(project(False, False))
    assert compute_phase_status(app.session_state)['phase7']['level'] != 'ok'
    case.assert_export_blocked(app)


def test_actual_word_measurements_and_result():
    case = audit.AuditRegressions()
    app = case.app(project())
    case.export(app)
    data = app.session_state['_compiled_doc_bytes']
    Path('/tmp/r53-focused-word.docx').write_bytes(data)
    t = table(Document(BytesIO(data)))
    assert [t.rows[3].cells[i].text for i in (1, 2, 3)] == ['08:00', '12:00', '24:00']
    assert [t.rows[4].cells[i].text for i in (1, 2, 3)] == ['100.25', '200.5', '300.125']
    assert t.rows[5].cells[0].text == 'Compressive strength results: UCA = 300.125 psi; Crush Test = 101 psi.'

from engineering_tools import compressive_inputs, validate_compressive_test
from project_state import prepare_calculations, purge_inactive_slurries, refresh_fluids, lab_source_signature
from project_io import replace_project_state
from unittest.mock import patch
import phase_10_procedure as report
import re
import subprocess
import zipfile


@pytest.fixture(scope='module')
def valid_state():
    return project()


@pytest.mark.parametrize('uca,crush', [(True, False), (False, True), (True, True), (False, False)])
def test_selection_and_unselected_draft_exclusion(uca, crush):
    qc = {'compressive': measured(uca, crush)}
    if not uca and not crush:
        with pytest.raises(ValueError): validate_compressive_test(qc)
        return
    result = validate_compressive_test(qc)
    for test, selected in (('uca', uca), ('crush', crush)):
        assert result[test]['selected'] == selected
        assert ('result' in result[test]) == selected
        if not selected: assert result[test] == {'selected': False}
    assert 'result' not in qc['compressive']['uca']
    assert 'result' not in qc['compressive']['crush']


@pytest.mark.parametrize('forces,expected', [
    ([0, 0, 0, 0], 0), ([4, 4, 4, 4], 1), ([400, 400, 400, 400], 100),
    ([400, 400, 400, 400.0000001], 101), ([400, 401, 402, 403], 101),
    ([16, 0, 0, 1e-15], 2),
])
def test_exact_ceiling_without_intermediate_rounding(forces, expected):
    data = measured(False, True)
    data['crush'].update(dict(zip(('force_1', 'force_2', 'force_3', 'force_4'), forces)))
    result = validate_compressive_test({'compressive': data})['crush']['result']
    assert result == expected and isinstance(result, int)


@pytest.mark.parametrize('test,field', [('uca', f) for f in ('cs_8', 'cs_12', 'cs_24')] +
    [('crush', f'force_{i}') for i in range(1, 5)])
def test_each_measurement_required_finite_nonnegative_and_zero_valid(test, field):
    for value in (None, '', -1, True, '2', float('nan'), float('inf'), {}, []):
        data = measured()
        data[test][field] = value
        with pytest.raises(ValueError): validate_compressive_test({'compressive': data})
    data[test][field] = 0
    assert validate_compressive_test({'compressive': data})[test][field] == 0
    if test == 'uca' and field == 'cs_24':
        assert validate_compressive_test({'compressive': data})['uca']['result'] == 0


@pytest.mark.parametrize('mode,test', [('UCA', 'uca'), ('CRUSH', 'crush')])
def test_legacy_mapping_once_and_missing_measurements_still_block(mode, test):
    qc = {'comp_test': mode}
    data = compressive_inputs(qc)
    assert data[test]['selected'] and not data['crush' if test == 'uca' else 'uca']['selected']
    assert 'comp_test' not in qc
    with pytest.raises(ValueError): validate_compressive_test(qc)
    qc['comp_test'] = 'CRUSH' if test == 'uca' else 'UCA'
    assert compressive_inputs(qc) == data
    assert 'comp_test' not in qc
    restored = audit.round_trip({'job_type': 'CSG 9 5/8"', 'lab_qc_params': {'Main': {'comp_test': mode}}})
    assert restored['lab_qc_params']['Main']['compressive'] == data


@pytest.mark.parametrize('inactive', [False, True])
def test_trust_boundary_drafts_malformed_types_and_transactional_preservation(inactive):
    def state(data):
        qc = {'compressive': data}
        return {'job_type': 'CSG 9 5/8"', **(
            {'inactive_slurry_drafts': {'Tail': {'lab_qc_params': qc}}} if inactive else
            {'lab_qc_params': {'Main': qc}})}
    draft = measured()
    draft['uca']['cs_8'] = None
    draft['crush']['force_1'] = ''
    draft['crush']['force_2'] = -1  # Repairable semantic invalidity, not malformed JSON.
    restored = audit.round_trip(state(draft))
    assert {key: restored[key] for key in state(draft)} == state(draft)
    bad = [[], None, True, {'other': {}}, {'uca': []}, {'uca': {'selected': 'yes'}},
           {'uca': {'cs_8': True}}, {'crush': {'force_1': '10'}}, {'crush': {'area': 4}},
           {'uca': {'result': 300}}, {'crush': {'force_2': float('inf')}}]
    open_state = {'well_name': 'preserve', '_compiled_doc_bytes': b'old'}
    before = deepcopy(open_state)
    for value in bad:
        with pytest.raises(ValueError):
            parsed = audit.round_trip(state(value))
            replace_project_state(open_state, parsed, 'new')
        assert open_state == before


@pytest.mark.parametrize('uca,crush', [(True, False), (False, True), (True, True)])
def test_actual_docx_modes_conditions_and_direct_fresh_export(valid_state, uca, crush):
    state = deepcopy(valid_state)
    qc = state['lab_qc_params']['Main']
    qc['compressive'] = measured(uca, crush)
    qc['review_signature'] = lab_review_signature(qc, state['lab_grid_dfs']['Main'])
    # Prove stale payload values never survive the normal direct-export rebuild.
    state['lab_payload_Main']['compressive'] = measured(not uca, not crush)
    restored = audit.round_trip(state)
    assert prepare_calculations(restored) == []
    case = audit.AuditRegressions()
    app = case.app(restored)
    case.export(app)
    data = app.session_state['_compiled_doc_bytes']
    Path(f'/tmp/r53-word-{uca}-{crush}.docx').write_bytes(data)
    doc = Document(BytesIO(data)); t = table(doc)
    assert [t.rows[3].cells[i].text for i in (1,2,3)] == ['08:00', '12:00', '24:00']
    assert [t.rows[4].cells[i].text for i in (5,6,7,8)] == ['4'] * 4
    assert [t.rows[4].cells[i].text for i in (1,2,3)] == (['100.25','200.5','300.125'] if uca else ['-']*3)
    assert [t.rows[3].cells[i].text for i in (5,6,7,8)] == (['400.0','401.0','402.0','403.0'] if crush else ['-']*4)
    basic = next(t for t in doc.tables if t.rows[0].cells[0].text == 'Test Basic Data')
    bhst = next(r.cells[1].text for r in basic.rows if r.cells[0].text == 'BHST').removesuffix(' degF')
    bhsp = next(r.cells[3].text for r in basic.rows if r.cells[2].text == 'BHSP').removesuffix(' psi')
    for selected, col in ((uca,3),(crush,7)):
        assert t.rows[1].cells[col].text == (bhst if selected else '-')
        assert t.rows[2].cells[col].text == (bhsp if selected else '-')
    parts = (['UCA = 300.125 psi'] if uca else []) + (['Crush Test = 101 psi'] if crush else [])
    assert t.rows[5].cells[0].text == 'Compressive strength results: ' + '; '.join(parts) + '.'
    assert app.session_state['lab_qc_params']['Main']['compressive'] == measured(uca, crush)
    assert compute_phase_status(app.session_state)['phase7']['level'] == 'ok'


@pytest.mark.parametrize('test,field,label', [('uca','cs_8','CS @ 08:00'), ('uca','cs_12','CS @ 12:00'),
    ('uca','cs_24','CS @ 24:00')] + [('crush', f'force_{i}', f'Force {i} (lbf)') for i in range(1,5)])
def test_first_edit_navigation_review_document_and_fresh_json(valid_state, test, field, label):
    case = audit.AuditRegressions()
    app = case.app(deepcopy(valid_state)); case.phase(app,'phase7')
    before = lab_review_signature(app.session_state['lab_qc_params']['Main'], app.session_state['lab_grid_dfs']['Main'])
    app.session_state['_compiled_doc_bytes'] = b'stale'
    next(w for w in app.number_input if w.label == label+' - Main').set_value(123.456789)
    case.phase(app,'phase1')
    assert '_compiled_doc_bytes' not in app.session_state
    qc = app.session_state['lab_qc_params']['Main']
    assert qc['compressive'][test][field] == 123.456789
    assert lab_review_signature(qc,app.session_state['lab_grid_dfs']['Main']) != before
    assert compute_phase_status(app.session_state)['phase7']['level'] != 'ok'
    case.phase(app,'phase7')
    assert next(w for w in app.number_input if w.label == label+' - Main').value == 123.456789
    next(b for b in app.button if b.label == 'Confirm measured lab results').click().run()
    fresh = case.app(audit.round_trip(app.session_state.to_dict()))
    case.export(fresh)
    assert fresh.session_state['lab_qc_params']['Main']['compressive'][test][field] == 123.456789


def test_deselect_reselect_hidden_draft_and_selection_invalidates_review(valid_state):
    case = audit.AuditRegressions(); app = case.app(deepcopy(valid_state)); case.phase(app,'phase7')
    before = lab_review_signature(app.session_state['lab_qc_params']['Main'], app.session_state['lab_grid_dfs']['Main'])
    app.session_state['_compiled_doc_bytes'] = b'stale'
    next(w for w in app.checkbox if w.label == 'UCA - Main').uncheck().run()
    assert not any(w.label == 'CS @ 08:00 - Main' for w in app.number_input)
    assert app.session_state['lab_qc_params']['Main']['compressive']['uca']['cs_8'] == 100.25
    assert '_compiled_doc_bytes' not in app.session_state
    assert lab_review_signature(app.session_state['lab_qc_params']['Main'], app.session_state['lab_grid_dfs']['Main']) != before
    case.phase(app,'phase1'); case.phase(app,'phase7')
    next(w for w in app.checkbox if w.label == 'UCA - Main').check().run()
    assert next(w for w in app.number_input if w.label == 'CS @ 08:00 - Main').value == 100.25
    next(w for w in app.checkbox if w.label == 'UCA - Main').uncheck().run()
    next(w for w in app.checkbox if w.label == 'Crush Test - Main').uncheck().run()
    assert next(b for b in app.button if b.label == 'Confirm measured lab results').disabled
    case.assert_export_blocked(app)


@pytest.mark.parametrize('test,field', [('uca','cs_12'),('crush','force_3')])
@pytest.mark.parametrize('value', [None, -1])
def test_selected_partial_or_negative_blocks_matching_review_ui_export(valid_state,test,field,value):
    state=deepcopy(valid_state);qc=state['lab_qc_params']['Main'];qc['compressive'][test][field]=value
    qc['review_signature']=lab_review_signature(qc,state['lab_grid_dfs']['Main'])
    state=audit.round_trip(state);case=audit.AuditRegressions();app=case.app(state);case.phase(app,'phase7')
    assert next(b for b in app.button if b.label=='Confirm measured lab results').disabled
    assert compute_phase_status(app.session_state)['phase7']['level']!='ok'
    assert any('Compressive Test' in s for s in prepare_calculations(deepcopy(state)))
    case.assert_export_blocked(app)


def test_multislurry_inactive_archive_fresh_restore_and_word(valid_state):
    state=deepcopy(valid_state);state['fluids_config']['active']=['Lead','Tail','Displacement Fluid']
    for slurry,top in [('Lead',0),('Tail',2850)]:
        state['fluids_config']['params'][slurry]=deepcopy(state['fluids_config']['params']['Main'])
        for key in ('cement_params','cement_additives_dfs','lab_grid_dfs','lab_qc_params'):
            state[key][slurry]=deepcopy(state[key]['Main'])
        state['cement_params'][slurry].update(top_mode='Surface' if slurry=='Lead' else 'Depth (m MD)',top_depth=top)
        qc=state['lab_qc_params'][slurry];qc['compressive']=measured(slurry=='Lead',slurry=='Tail')
        qc['review_signature']=lab_review_signature(qc,state['lab_grid_dfs'][slurry])
    refresh_fluids(state)
    for slurry in ('Lead','Tail'):state['lab_source_signatures'][slurry]=lab_source_signature(state,slurry)
    state['_comp_tail_crush_force_1_default_project']=9999
    state['_comp_lead_uca_cs_8_default_project']=100.25
    purge_inactive_slurries(state,['Lead','Displacement Fluid'])
    assert '_comp_tail_crush_force_1_default_project' not in state
    assert '_comp_lead_uca_cs_8_default_project' in state
    archived=deepcopy(state['inactive_slurry_drafts']['Tail']['lab_qc_params']['compressive'])
    state=audit.round_trip(state)
    purge_inactive_slurries(state,['Lead','Tail','Displacement Fluid'])
    assert state['lab_qc_params']['Tail']['compressive']==archived
    assert 'Tail' not in state['lab_source_signatures']
    case=audit.AuditRegressions();app=case.app(state);case.phase(app,'phase7')
    assert next(w for w in app.number_input if w.label=='Force 1 (lbf) - Tail').value==400
    for b in list(app.button):
        if b.label=='Sync with Phase V': app.button(key=b.key).click().run()
    for _ in range(2):
        buttons = [b for b in app.button if b.label=='Confirm measured lab results']
        if not buttons: break
        assert not buttons[0].disabled
        buttons[0].click().run()
    case.export(app)
    data=app.session_state['_compiled_doc_bytes'];Path('/tmp/r53-multi.docx').write_bytes(data)
    tables=[t for t in Document(BytesIO(data)).tables if t.rows[0].cells[0].text=='Compressive Strength Test']
    assert [t.rows[5].cells[0].text for t in tables]==[
        'Compressive strength results: UCA = 300.125 psi.', 'Compressive strength results: Crush Test = 101 psi.']


def test_only_compressive_authorized_text_changes_template_package():
    baseline=BytesIO(subprocess.check_output(['git','show','a713cef:master_template.docx']))
    with zipfile.ZipFile(baseline) as old,zipfile.ZipFile('master_template.docx') as new:
        assert old.namelist()==new.namelist()
        for name in old.namelist():
            if name!='word/document.xml':assert old.read(name)==new.read(name)
        a,b=old.read('word/document.xml'),new.read('word/document.xml')
        locate=lambda x:next(m for m in re.finditer(rb'<w:tbl>.*?</w:tbl>',x,re.S) if b'Compressive Strength Test' in m.group())
        before,after=locate(a),locate(b)
        assert a[:before.start()]==b[:after.start()]
        assert a[before.end():]==b[after.end():]
        strip=lambda x:re.sub(rb'(<w:t(?:\s[^>]*)?>).*?(</w:t>)',rb'\1\2',x,flags=re.S)
        assert strip(before.group())==strip(after.group())
        rows=re.findall(rb'<w:tr[ >].*?</w:tr>',after.group(),re.S)
        assert len(rows)==6
        # Existing condition cells remain the same tags; only authorized data cells differ.
        old_rows=re.findall(rb'<w:tr[ >].*?</w:tr>',before.group(),re.S)
        assert rows[:3]==old_rows[:3]
