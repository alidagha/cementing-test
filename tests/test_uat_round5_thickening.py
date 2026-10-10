"""Round 5-2: canonical Thickening Test inputs and protected FLC/Free Water."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
import re
import subprocess
import zipfile
import pytest

from docx import Document
import test_audit_regressions as audit
from engineering_tools import compute_phase_status, lab_review_signature
from engineering_tools import validate_thickening_test, thickening_time_valid
from project_state import prepare_calculations, purge_inactive_slurries, lab_source_signature, refresh_fluids
import phase_10_procedure as report


FIELDS = {
    'thickening_test_start': '23:59',
    'thickening_enter_time': '00:00',
    'thickening_cell': 2.5,
    'thickening_30_bc': '01:10',
    'thickening_50_bc': '02:20',
    'thickening_time': '03:30',
}
LABELS = {
    'thickening_test_start': 'Test Start (HH:MM)',
    'thickening_enter_time': 'Enter Time (HH:MM)',
    'thickening_cell': 'Cell',
    'thickening_30_bc': '30 Bc (HH:MM)',
    'thickening_50_bc': '50 Bc (HH:MM)',
    'thickening_time': '70 Bc (HH:MM)',
}


def project():
    case = audit.AuditRegressions()
    app = case.configured_app('CSG 9 5/8"')
    state = audit.round_trip(app.session_state.to_dict())
    qc = state['lab_qc_params']['Main']
    qc.update(FIELDS)
    qc['review_signature'] = lab_review_signature(qc, state['lab_grid_dfs']['Main'])
    return state


def test_six_independent_thickening_inputs():
    case = audit.AuditRegressions()
    app = case.app(project())
    case.phase(app, 'phase7')
    for field, label in LABELS.items():
        widgets = app.number_input if field == 'thickening_cell' else app.text_input
        assert any(w.label == label + ' - Main' for w in widgets), label


def test_real_word_contains_all_six_values_without_endpoint_authority():
    case = audit.AuditRegressions()
    state = project()
    state['lab_qc_params']['Main']['thickening_endpoint'] = '100 Bc'
    app = case.app(state)
    case.export(app)
    data = app.session_state['_compiled_doc_bytes']
    Path('/tmp/r52-word.docx').write_bytes(data)
    table = next(t for t in Document(BytesIO(data)).tables
                 if t.rows[0].cells[0].text == 'Thickening Time Test')
    assert table.rows[2].cells[2].text == 'Enter Time\nhr:mm'
    assert [table.rows[3].cells[i].text for i in (4, 2, 5, 6, 7, 8)] == [
        '23:59', '00:00', '2.5', '01:10', '02:20', '03:30']
    assert 'thickening_endpoint' not in app.session_state['lab_payload_Main']


def test_missing_new_required_field_blocks_matching_review():
    case = audit.AuditRegressions()
    state = project()
    qc = state['lab_qc_params']['Main']
    qc['thickening_test_start'] = ''
    qc['review_signature'] = lab_review_signature(qc, state['lab_grid_dfs']['Main'])
    app = case.app(state)
    assert compute_phase_status(app.session_state)['phase7']['level'] != 'ok'
    case.assert_export_blocked(app)


def test_baseline_flc_temperature_already_uses_canonical_bhct():
    case = audit.AuditRegressions()
    state = project()
    qc = state['lab_qc_params']['Main']
    state['bhct'] = state['well_data']['bhct'] = 167
    state['lab_source_signatures']['Main'] = lab_source_signature(state, 'Main')
    qc['review_signature'] = lab_review_signature(qc, state['lab_grid_dfs']['Main'])
    app = case.app(state)
    case.export(app)
    doc = Document(BytesIO(app.session_state['_compiled_doc_bytes']))
    table = next(t for t in doc.tables if t.rows[0].cells[0].text == 'Free Water Test')
    assert table.rows[1].cells[3].text == '167 degF'
    assert table.rows[1].cells[1].text == '70 degF'
    assert table.rows[2].cells[1].text == '0.0 ml in 250 ml @ 2hr'


def test_free_water_docx_is_byte_identical_and_cell_retains_entered_precision():
    case = audit.AuditRegressions()
    state = project()
    qc = state['lab_qc_params']['Main']
    qc.update(free_water=1.5, free_water_45=2.5, thickening_cell=2.123456789)
    qc['review_signature'] = lab_review_signature(qc, state['lab_grid_dfs']['Main'])
    app = case.app(state)
    case.export(app)
    current = Document(BytesIO(app.session_state['_compiled_doc_bytes']))
    with patch.object(report.st, 'session_state', deepcopy(app.session_state.to_dict())):
        context = report.build_master_context()
    baseline_template = BytesIO(subprocess.check_output(['git', 'show', 'ecb0a6e:master_template.docx']))
    baseline = report.DocxTemplate(baseline_template)
    baseline.render(report._word_quantity_context(context))
    buf = BytesIO()
    baseline.save(buf)
    previous = Document(BytesIO(buf.getvalue()))
    flc = lambda doc: next(t for t in doc.tables if t.rows[0].cells[0].text == 'Free Water Test')
    assert flc(current)._tbl.xml == flc(previous)._tbl.xml
    assert flc(current).rows[2].cells[1].text == '1.5 ml in 250 ml @ 2hr'
    assert flc(current).rows[3].cells[1].text == '2.5 ml in 250 ml @ 2hr'
    tt = next(t for t in current.tables if t.rows[0].cells[0].text == 'Thickening Time Test')
    assert tt.rows[3].cells[5].text == '2.123456789'
    assert app.session_state['lab_qc_params']['Main']['thickening_cell'] == 2.123456789


def test_bhct_edit_blocks_old_document_then_reconfirmation_updates_flc():
    case = audit.AuditRegressions()
    app = case.app(project())
    case.export(app)
    assert '_compiled_doc_bytes' in app.session_state
    case.phase(app, 'phase2_3')
    next(w for w in app.number_input if w.label == 'BHCT (degF)').set_value(166.0)
    case.phase(app, 'phase10')
    assert compute_phase_status(app.session_state)['phase7']['level'] != 'ok'
    assert '_compiled_doc_bytes' not in app.session_state
    assert next(b for b in app.button if b.label == 'Build Word Document').disabled
    case.phase(app, 'phase7')
    next(b for b in app.button if b.label == 'Keep reviewed lab entries').click().run()
    next(b for b in app.button if b.label == 'Confirm measured lab results').click().run()
    # A cached old temperature must never win over the canonical source.
    app.session_state['lab_payload_Main']['bhct'] = 150
    case.export(app)
    doc = Document(BytesIO(app.session_state['_compiled_doc_bytes']))
    flc = next(t for t in doc.tables if t.rows[0].cells[0].text == 'Free Water Test')
    assert flc.rows[1].cells[3].text == f"{app.session_state['bhct']} degF"


@pytest.mark.parametrize('field', ['thickening_test_start', 'thickening_enter_time'])
@pytest.mark.parametrize('value,valid', [('00:00', True), ('23:59', True), ('24:00', True),
    ('24:01', False), ('25:00', False), ('01:60', False), ('1:00', False),
    ('00:0', False), ('01.00', False), ('', False), (None, False)])
def test_wall_clock_and_elapsed_boundaries(field, value, valid):
    qc = {**FIELDS, field: value}
    if valid:
        assert validate_thickening_test(qc)[field] == value
    else:
        with pytest.raises(ValueError): validate_thickening_test(qc)


def test_milestone_envelope_and_no_invented_order_or_cell_bounds():
    for field in ('thickening_30_bc', 'thickening_50_bc', 'thickening_time'):
        for value in ('00:01', '03:30', '3:30', '24:00'):
            assert validate_thickening_test({**FIELDS, field: value})[field] == value
        for value in ('00:00', '24:01', '99:59', '3:3', '', None):
            assert not thickening_time_valid(value)
            with pytest.raises(ValueError): validate_thickening_test({**FIELDS, field: value})
    for cell in (-5.25, 0.0, 2.5, 1e20):
        qc = {**FIELDS, 'thickening_cell': cell, 'thickening_30_bc': '24:00',
              'thickening_50_bc': '04:00', 'thickening_time': '00:01'}
        assert validate_thickening_test(qc) == qc
    for cell in (None, True, '2.5', float('nan'), float('inf')):
        with pytest.raises(ValueError): validate_thickening_test({**FIELDS, 'thickening_cell': cell})


@pytest.mark.parametrize('field', list(FIELDS))
def test_each_required_field_independently_blocks_ui_status_and_export(field):
    case = audit.AuditRegressions()
    state = project()
    qc = state['lab_qc_params']['Main']
    qc[field] = None if field == 'thickening_cell' else '24:01'
    qc['review_signature'] = lab_review_signature(qc, state['lab_grid_dfs']['Main'])
    restored = audit.round_trip(state)  # Semantically invalid text remains a repairable draft.
    app = case.app(restored)
    case.phase(app, 'phase7')
    assert next(b for b in app.button if b.label == 'Confirm measured lab results').disabled
    assert compute_phase_status(app.session_state)['phase7']['level'] != 'ok'
    assert any('Thickening Time Test' in issue for issue in prepare_calculations(deepcopy(restored)))
    case.assert_export_blocked(app)


def test_first_edit_immediate_navigation_review_invalidation_and_fresh_restore():
    case = audit.AuditRegressions()
    app = case.app(project())
    case.phase(app, 'phase7')
    edited = {**FIELDS, 'thickening_test_start': '24:00', 'thickening_enter_time': '01:25',
              'thickening_cell': -3.25, 'thickening_30_bc': '01:40',
              'thickening_50_bc': '02:40', 'thickening_time': '04:45'}
    for field, value in edited.items():
        qc = app.session_state['lab_qc_params']['Main']
        before = lab_review_signature(qc, app.session_state['lab_grid_dfs']['Main'])
        app.session_state['_compiled_doc_bytes'] = b'stale'
        widgets = app.number_input if field == 'thickening_cell' else app.text_input
        next(w for w in widgets if w.label == LABELS[field] + ' - Main').set_value(value)
        case.phase(app, 'phase1')
        assert app.session_state['lab_qc_params']['Main'][field] == value
        assert '_compiled_doc_bytes' not in app.session_state
        assert lab_review_signature(app.session_state['lab_qc_params']['Main'], app.session_state['lab_grid_dfs']['Main']) != before
        assert compute_phase_status(app.session_state)['phase7']['level'] != 'ok'
        case.phase(app, 'phase7')
        widgets = app.number_input if field == 'thickening_cell' else app.text_input
        assert next(w for w in widgets if w.label == LABELS[field] + ' - Main').value == value
    next(b for b in app.button if b.label == 'Confirm measured lab results').click().run()
    saved = audit.round_trip(app.session_state.to_dict())
    fresh = case.app(saved)
    case.export(fresh)  # Fresh JSON restore -> Phase X, without visiting Phase VII.
    qc = fresh.session_state['lab_qc_params']['Main']
    assert {k: qc[k] for k in FIELDS} == edited
    assert compute_phase_status(fresh.session_state)['phase7']['level'] == 'ok'
    case.phase(fresh, 'phase7')
    for field, value in edited.items():
        widgets = fresh.number_input if field == 'thickening_cell' else fresh.text_input
        assert next(w for w in widgets if w.label == LABELS[field] + ' - Main').value == value


def test_legacy_70bc_and_stale_endpoint_survive_as_draft_without_authority():
    state = project()
    qc = state['lab_qc_params']['Main']
    for field in FIELDS:
        if field != 'thickening_time': qc.pop(field)
    qc['thickening_endpoint'] = '100 Bc'
    restored = audit.round_trip(state)
    assert restored['lab_qc_params']['Main']['thickening_time'] == '03:30'
    assert restored['lab_qc_params']['Main']['thickening_endpoint'] == '100 Bc'
    assert prepare_calculations(restored)
    case = audit.AuditRegressions()
    app = case.app(restored)
    case.phase(app, 'phase7')
    assert next(w for w in app.text_input if w.label == '70 Bc (HH:MM) - Main').value == '03:30'
    assert next(b for b in app.button if b.label == 'Confirm measured lab results').disabled
    assert not any('Endpoint' in w.label for w in app.selectbox)
    case.assert_export_blocked(app)


def test_active_and_inactive_draft_json_and_malformed_type_rejection():
    for inactive in (False, True):
        qc = {**FIELDS, 'thickening_test_start': 'unfinished', 'thickening_enter_time': '',
              'thickening_cell': None, 'thickening_30_bc': None, 'thickening_50_bc': '24:01'}
        p = {'job_type': 'CSG 9 5/8"'}
        if inactive: p['inactive_slurry_drafts'] = {'Tail': {'lab_qc_params': qc}}
        else: p['lab_qc_params'] = {'Main': qc}
        assert audit.round_trip(p)['inactive_slurry_drafts' if inactive else 'lab_qc_params'] == p['inactive_slurry_drafts' if inactive else 'lab_qc_params']
        for bad in (True, 'bad', float('nan'), float('inf')):
            qc['thickening_cell'] = bad
            with pytest.raises(ValueError): audit.round_trip(p)
        qc['thickening_cell'] = None
        for field in FIELDS:
            if field in ('thickening_time', 'thickening_cell'): continue
            for bad in (True, 12, [], {}):
                qc[field] = bad
                with pytest.raises(ValueError): audit.round_trip(p)
            qc[field] = ''


def test_multislurry_archive_reactivation_flc_and_word():
    case = audit.AuditRegressions()
    state = project()
    state['fluids_config']['active'] = ['Lead', 'Tail', 'Displacement Fluid']
    state['bhct'] = state['well_data']['bhct'] = 155
    for slurry, top, start, cell in [('Lead', 0, '08:20', 1.25), ('Tail', 2850, '24:00', -2.5)]:
        state['fluids_config']['params'][slurry] = deepcopy(state['fluids_config']['params']['Main'])
        for key in ('cement_params', 'cement_additives_dfs', 'lab_grid_dfs', 'lab_qc_params'):
            state[key][slurry] = deepcopy(state[key]['Main'])
        state['cement_params'][slurry].update(top_mode='Surface' if slurry == 'Lead' else 'Depth (m MD)', top_depth=top)
        qc = state['lab_qc_params'][slurry]
        qc.update(thickening_test_start=start, thickening_cell=cell)
        qc['review_signature'] = lab_review_signature(qc, state['lab_grid_dfs'][slurry])
    refresh_fluids(state)
    for slurry in ('Lead', 'Tail'): state['lab_source_signatures'][slurry] = lab_source_signature(state, slurry)
    for field in FIELDS: state[f'_qc_{field}_in_tail_default_project'] = 'stale'
    state['_qc_thickening_time_in_lead_default_project'] = '03:30'
    purge_inactive_slurries(state, ['Lead', 'Displacement Fluid'])
    assert not any(k.startswith('_qc_thickening') and '_in_tail_' in k for k in state)
    assert '_qc_thickening_time_in_lead_default_project' in state
    archived = deepcopy(state['inactive_slurry_drafts']['Tail']['lab_qc_params'])
    restored = audit.round_trip(state)
    restored['fluids_config']['active'] = ['Lead', 'Tail', 'Displacement Fluid']
    purge_inactive_slurries(restored, restored['fluids_config']['active'])
    assert restored['lab_qc_params']['Tail'] == archived
    assert 'Tail' not in restored['lab_source_signatures']
    app = case.app(restored)
    case.phase(app, 'phase7')
    next(b for b in app.button if b.label == 'Keep reviewed lab entries').click().run()
    next(b for b in app.button if b.label == 'Confirm measured lab results' and not b.disabled).click().run()
    assert compute_phase_status(app.session_state)['phase7']['level'] == 'ok'
    case.export(app)
    data = app.session_state['_compiled_doc_bytes']
    Path('/tmp/r52-multi.docx').write_bytes(data)
    doc = Document(BytesIO(data))
    flc = [t for t in doc.tables if t.rows[0].cells[0].text == 'Free Water Test']
    tt = [t for t in doc.tables if t.rows[0].cells[0].text == 'Thickening Time Test']
    assert [t.rows[1].cells[3].text for t in flc] == ['155 degF', '155 degF']
    assert [t.rows[3].cells[4].text for t in tt] == ['08:20', '24:00']
    assert [t.rows[3].cells[5].text for t in tt] == ['1.25', '-2.5']
    assert all(t.rows[3].cells[8].text == '03:30' for t in tt)


def test_only_authorized_thickening_cell_text_changed_in_template():
    original = BytesIO(subprocess.check_output(['git', 'show', 'ecb0a6e:master_template.docx']))
    completed = BytesIO(subprocess.check_output(['git', 'show', 'a713cef:master_template.docx']))
    with zipfile.ZipFile(completed) as closed, zipfile.ZipFile('master_template.docx') as current:
        get_table = lambda x: next(m.group() for m in re.finditer(rb'<w:tbl>.*?</w:tbl>', x, re.S) if b'Thickening Time Test' in m.group())
        assert get_table(closed.read('word/document.xml')) == get_table(current.read('word/document.xml'))
    with zipfile.ZipFile(original) as old, zipfile.ZipFile(completed) as new:
        assert old.namelist() == new.namelist()
        for name in old.namelist():
            if name != 'word/document.xml': assert old.read(name) == new.read(name)
        a, b = old.read('word/document.xml'), new.read('word/document.xml')
        table = lambda x: next(m for m in re.finditer(rb'<w:tbl>.*?</w:tbl>', x, re.S) if b'Thickening Time Test' in m.group())
        before, after = table(a), table(b)
        assert a[:before.start()] == b[:after.start()]
        assert a[before.end():] == b[after.end():]
        clean = after.group().replace(b'70 Bc', b'{{ s.lab.thickening_endpoint_label }}')
        clean = re.sub(rb'{{ s.lab.thickening_(?:enter_time|test_start|cell|30_bc|50_bc) }}', b'-', clean)
        # Reverse just the Enter Time heading's four existing text runs.
        heading = next(c for c in re.findall(rb'<w:tc>.*?</w:tc>', clean, re.S) if b'Enter Time' in c and b'hr:mm' in c)
        old_heading = heading.replace(b'Enter Time', b'Initial Temp')
        texts = list(re.finditer(rb'<w:t(?:\s[^>]*)?>(.*?)</w:t>', old_heading))
        for index, value in reversed(list(enumerate([b'Initial Temp', b' (', b'degF', b')']))):
            match = texts[index]
            old_heading = old_heading[:match.start(1)] + value + old_heading[match.end(1):]
        clean = clean.replace(heading, old_heading, 1)
        assert clean == before.group()
