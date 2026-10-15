"""Final Round-9 certification through real execution, restore and Word boundaries."""
from copy import deepcopy
from io import BytesIO
import math
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

import pandas as pd
import pytest
from docx import Document
from docx.oxml.ns import qn
from streamlit.testing.v1 import AppTest

import centralizer_adapter as adapter
import centralizer_optimizer_engine as engine
import centralizer_report as bridge
import engineering_tools as tools
import materials_db
import phase_9_centralizer as phase9
import phase_10_procedure as report
from project_io import replace_project_state
from project_state import fingerprint, is_project_key
import test_audit_regressions as audit
from centralizer_fixtures import execute_centralizer
from test_uat_round9e_dynamic_solver import readings
from test_uat_round9f_phase9 import configured
from test_uat_round9g_centralizer_word import ix_region


@pytest.fixture(scope="module")
def catalog():
    return adapter.load_official_catalog()


def canonical(state):
    return fingerprint({key: value for key, value in state.items() if is_project_key(key)})


def directional(state):
    state["well_geometry"] = {"version": 1, "type": "Directional", "survey": pd.DataFrame(
        [[0., 0., 359.], [400., 20., 1.], [800., 45., 30.], [1200., 65., 90.]],
        columns=["MD", "Deviation", "Azimuth"])}


@pytest.mark.parametrize("job", materials_db.JOB_TYPES)
def test_all_jobs_navigate_without_implicit_execution_or_false_completion(job):
    assert len(materials_db.JOB_TYPES) == 13
    state = {"job_type": job, "mud_density": "80", **audit.well_profile_fixture(),
             **audit.chart_fixture(job)}
    app = AppTest.from_file(str(audit.ROOT / "main.py"), default_timeout=30)
    for key, value in state.items():
        app.session_state[key] = value
    app.run()
    with patch.object(adapter, "solve_static", side_effect=AssertionError("Implicit Static execution")), \
         patch.object(adapter, "solve_dynamic", side_effect=AssertionError("Implicit Dynamic execution")):
        app.radio(key="_app_mode_key").set_value("phase9").run()
        assert not app.exception
        assert app.radio(key="_app_mode_key").value == "phase9"
        status = tools.compute_phase_status(app.session_state)["phase9"]
        if job in ("CMT PLUG", "CMT SQUEEZE"):
            assert not adapter.centralizer_applicable(job)
            assert status == {"level": "ok", "message": "Not applicable (Plug/Squeeze)."}
            assert not any(button.label == "Execute Centralizer" for button in app.button)
        else:
            assert adapter.centralizer_applicable(job)
            assert status["level"] == "warning"
            assert any(button.label == "Execute Centralizer" for button in app.button)


@pytest.mark.parametrize("kind", ["casing", "liner", "tieback"])
@pytest.mark.parametrize("physics", ["static", "dynamic"])
@pytest.mark.parametrize("mode", ["direct", "optimize"])
@pytest.mark.parametrize("use_custom", [False, True], ids=["official", "custom"])
def test_real_24_case_matrix_reaches_current_companion_view(catalog, kind, physics, mode, use_custom):
    state = configured(catalog, kind, physics, mode, use_custom)
    geometry = adapter.to_engine_geometry(state)
    if kind == "tieback" and physics == "dynamic":
        directional(state)
    if use_custom and physics == "dynamic":
        product = state["centralizer_config"]["custom_products"][0]
        product.update(kind="spring", bows=6, min_od_in=product["pipe_size_in"],
                       max_od_in=geometry["wellbore"]["hole_size"] / adapter.INCH_TO_M,
                       restoring_force_lbf=50., starting_force_lbf=10., running_force_lbf=10.)
    before = canonical(state)
    route = "solve_static" if physics == "static" else "solve_dynamic"
    other = "solve_dynamic" if physics == "static" else "solve_static"
    with patch.object(adapter, route, wraps=getattr(adapter, route)) as call, \
         patch.object(adapter, other, side_effect=AssertionError("Wrong physics route")), \
         patch.object(engine, "compute_standoff_profile", wraps=engine.compute_standoff_profile) as direct, \
         patch.object(engine, "optimize_placement", wraps=engine.optimize_placement) as optimize, \
         patch.object(engine, "run_full_analysis", side_effect=AssertionError("Wrong solver API")):
        result = phase9.execute(state, catalog)
        assert call.call_count == 1 and call.call_args.kwargs["mode"] == mode
        assert (optimize.call_count == 1) if mode == "optimize" else (optimize.call_count == 0)
        assert direct.call_count >= 1
    assert canonical(state) == before
    assert result is state["_centralizer_result"]
    assert phase9.result_is_current(state, catalog)
    assert tools.compute_phase_status(state)["phase9"]["level"] == "ok"
    for name in ("centralizer_md", "stdf_at", "between_stdf", "tension_force", "normal_force"):
        assert result["profile"][name]
        assert all(math.isfinite(value) for value in result["profile"][name])
    for actual, source in zip(result["zones"], state["centralizer_config"]["zones"]):
        assert (actual["top_md"], actual["bottom_md"]) == (source["top_md"], source["bottom_md"])
        assert actual["product_nr"] == source["product_nr"]
        assert actual.get("custom_product_id") == source["custom_product_id"]
        assert actual["product_name"] == adapter.resolve_zone_product(
            source, catalog, state["centralizer_config"]["custom_products"])["name"]
        assert type(actual["meets_target"]) is bool and math.isfinite(actual["min_stdf"])
    with patch.object(adapter, "solve_static", side_effect=AssertionError("Report solved again")), \
         patch.object(adapter, "solve_dynamic", side_effect=AssertionError("Report solved again")):
        prepared = bridge.prepare_report(state)
    assert prepared["view"]["n_rigid"] + prepared["view"]["n_spring"] == prepared["view"]["n_cent"]
    assert len(prepared["view"]["zone_rows"]) == len(result["zones"])
    assert prepared["warning"] == any(not zone["meets_target"] for zone in result["zones"])
    if kind == "liner":
        assert geometry["wellbore"]["string_top"] == 400.
        assert adapter.build_fluid_snapshot(state)["inside_intervals"][0]["top_md"] == 0.
    elif kind == "tieback":
        assert geometry["wellbore"]["prev_shoe_depth"] == geometry["wellbore"]["string_shoe"]
        assert geometry["wellbore"]["hole_size"] == geometry["wellbore"]["prev_casing_id"]


@pytest.mark.parametrize("physics", ["static", "dynamic"])
@pytest.mark.parametrize("use_custom", [False, True])
def test_real_project_restore_requires_execute_before_report(catalog, physics, use_custom):
    state = configured(catalog, physics=physics, use_custom=use_custom)
    state["hardware_table"] = state["hardware_table"].assign(**{"Collapse (psi)": None, "Burst (psi)": None})
    phase9.execute(state, catalog)
    saved = audit.round_trip(state)
    assert saved["centralizer_config"] == state["centralizer_config"]
    assert not any(key.startswith("_centralizer") for key in saved)
    restored = {}
    replace_project_state(restored, saved, "round9h-current-project")
    assert not phase9.result_is_current(restored, catalog)
    assert tools.compute_phase_status(restored)["phase9"]["level"] == "warning"
    with pytest.raises(ValueError, match="Phase IX"):
        bridge.prepare_report(restored)
    phase9.execute(restored, catalog)
    assert restored["centralizer_config"] == saved["centralizer_config"]
    assert tools.compute_phase_status(restored)["phase9"]["level"] == "ok"
    assert bridge.prepare_report(restored)["applicable"]


@pytest.mark.parametrize("edit", ["geometry", "fluid", "product", "zone", "mud", "lab", "aux", "pump"])
def test_consumed_source_edit_blocks_report_without_implicit_reexecution(catalog, edit):
    state = configured(catalog, physics="dynamic", use_custom=True)
    phase9.execute(state, catalog)
    if edit == "geometry": state["well_geometry"]["td_m"] = 1300.
    elif edit == "fluid": state["fluids_config"]["params"]["Tail"]["volume"] += 1.
    elif edit == "product": state["centralizer_config"]["custom_products"][0]["description"] = "Selected edited"
    elif edit == "zone": state["centralizer_config"]["zones"][0]["target_standoff_pct"] = 71.
    elif edit == "mud": state["plastic_viscosity"] = "18"
    elif edit == "lab": state["lab_qc_params"]["Tail"]["rheology"]["bhct_down"]["readings"]["300"] += 1.
    elif edit == "aux": state["centralizer_config"]["aux_rheology"]["Spacer"]["300"] += 1.
    else: state["fluids_config"]["params"]["Displacement Fluid"]["pump_rate"] = "5-7"
    with patch.object(adapter, "solve_dynamic", side_effect=AssertionError("Implicit re-execution")):
        assert not phase9.result_is_current(state, catalog)
        assert tools.compute_phase_status(state)["phase9"]["level"] == "warning"
        with pytest.raises(ValueError, match="Phase IX"):
            bridge.prepare_report(state)
    phase9.execute(state, catalog)
    assert phase9.result_is_current(state, catalog) and bridge.prepare_report(state)["applicable"]


@pytest.mark.parametrize("physics", ["static", "dynamic"])
def test_report_only_changes_do_not_stale_engineering(catalog, physics):
    state = configured(catalog, physics=physics)
    phase9.execute(state, catalog)
    before = bridge.prepare_report(state)
    state["procedure_text"] = "Reviewed report-only text"
    state["cement_params"] = {"Tail": {"top_depth": 999., "report_bottom_depth": 1.}}
    assert phase9.result_is_current(state, catalog)
    assert bridge.prepare_report(state) == before


def test_all_active_dynamic_sources_and_excluded_spacer_behind(catalog):
    state = configured(catalog, physics="dynamic")
    for name in ("Scavenger", "Pre Flush", "Spacer Ahead", "Spacer Behind"):
        state["fluids_config"]["active"].append(name)
        state["fluids_config"]["params"][name] = {"volume": 20., "density": "95", "pump_rate": "4"}
    state["lab_qc_params"]["Scavenger"] = {"rheology": {"bhct_down": {"selected": True, "readings": readings()}}}
    result = phase9.execute(state, catalog)
    diagnostics = result["diagnostics"]
    assert diagnostics["pump_rate_bpm"] == tools.require_positive_pump_rate("4-6")
    assert diagnostics["rheology"]["Displacement Fluid"]["source"] == "phase_ii_mud"
    assert diagnostics["rheology"]["Scavenger"]["source"] == "lab_bhct_down"
    assert diagnostics["rheology"]["Spacer Behind"]["source"] == "aux_rheology"
    snapshot = adapter.build_fluid_snapshot(state)
    assert [(row["name"], row["top_md"], row["bottom_md"]) for row in diagnostics["stages_rheology"]] == [
        (row["fluid"], row["top_md"], row["bottom_md"]) for row in snapshot["annulus_intervals"]
        if row["fluid"] not in ("Mud", "Displacement Fluid")]
    assert all(row["name"] != "Spacer Behind" for row in diagnostics["stages_rheology"])
    del state["centralizer_config"]["aux_rheology"]["Spacer Behind"]
    with pytest.raises(ValueError, match="MissingRheology.*Spacer Behind"):
        phase9.execute(state, catalog)
    assert not phase9.result_is_current(state, catalog)
    state["centralizer_config"]["aux_rheology"]["Spacer Behind"] = readings()
    phase9.execute(state, catalog)
    assert bridge.prepare_report(state)["applicable"]


def test_static_tier2_remains_real_and_dynamic_rejects_before_solver(catalog):
    state = configured(catalog)
    before = canonical(state)
    config = state["centralizer_config"]
    override = {"inside_fluids": [[0., 1300.], [100., 1550.]], "annulus_fluids": [[0., 1500.]]}
    original = deepcopy(override)
    result = adapter.solve_static(state, config["zones"], catalog, mode="direct", override=override)
    assert result["profile"]["centralizer_md"]
    with patch.object(engine, "compute_standoff_profile", side_effect=AssertionError("Tier-2 Dynamic executed")), \
         pytest.raises(ValueError, match="Tier-2"):
        adapter.solve_dynamic(state, config["zones"], catalog, mode="direct", override=override)
    assert override == original and canonical(state) == before


@pytest.mark.parametrize("kind", ["casing", "liner", "tieback"])
def test_auto_populated_zones_remain_explicit_editable_and_reportable(catalog, kind):
    from cement_chart_model import resolve_chart_geometry, interval_volume
    state = configured(catalog, kind)
    geometry = resolve_chart_geometry(state)
    toc = 600. if kind == "tieback" else 450.
    state["fluids_config"]["params"]["Tail"]["volume"] = (
        interval_volume(geometry["capacity_segments"], toc, 1000.) + geometry["shoe_track_bbl"])
    state["cement_params"] = {"Tail": {"top_depth": 999., "report_bottom_depth": 1.}}
    generated = phase9.auto_populate(state)
    assert all(zone["product_nr"] is None and zone["custom_product_id"] is None for zone in generated)
    product = state["centralizer_config"]["zones"][0]["product_nr"]
    for zone in generated:
        zone["product_nr"] = product
    generated[0]["target_standoff_pct"] = 68.
    state["centralizer_config"]["zones"] = generated
    phase9.execute(state, catalog)
    assert state["centralizer_config"]["zones"] == generated
    assert len(bridge.prepare_report(state)["view"]["zone_rows"]) == len(generated)


@pytest.fixture(scope="module")
def word_projects():
    case = audit.AuditRegressions()
    return {job: case.configured_app(job).session_state.to_dict() for job in
            ('CSG 9 5/8"', 'LNR 7"', 'TIE BACK LNR 7"', 'CMT PLUG', 'CMT SQUEEZE')}


DOCUMENT_CASES = [
    ('CSG 9 5/8"', "static", "direct", False, 2, False),
    ('LNR 7"', "static", "optimize", True, 4, False),
    ('TIE BACK LNR 7"', "dynamic", "direct", False, 2, False),
    ('CSG 9 5/8"', "dynamic", "optimize", True, 3, True),
    ('CMT PLUG', "static", "direct", False, 0, False),
    ('CMT SQUEEZE', "static", "direct", False, 0, False),
]


@pytest.mark.parametrize("job,physics,mode,custom,zones,miss", DOCUMENT_CASES)
def test_real_phase_x_build_and_section_ix_contract(word_projects, job, physics, mode, custom, zones, miss):
    state = deepcopy(word_projects[job])
    if miss:
        state["well_auto_fields"]["bhsp"] = False
        state["well_geometry"] = {"version": 1, "type": "Directional", "survey": pd.DataFrame(
            [[0., 0., 0.], [1500., 65., 30.], [3300., 85., 80.]], columns=["MD", "Deviation", "Azimuth"])}
    case = audit.AuditRegressions()
    app = case.app(state)
    if physics == "dynamic":
        app.session_state["lab_qc_params"]["Main"]["rheology"]["bhct_down"] = {
            "selected": True, "readings": readings(), "gel_10_sec": "8", "gel_10_min": "12"}
        case.phase(app, "phase7")
        next(button for button in app.button if button.label == "Confirm measured lab results").click().run()
        assert not app.exception and not app.error
    state = app.session_state.to_dict()
    if zones:
        execute_centralizer(state, zones=zones, physics=physics, mode=mode, custom=custom,
                            target=100. if miss else 70.)
        if miss:
            state["centralizer_config"]["custom_products"][0].update(kind="spring", bows=6,
                min_od_in=9.625, restoring_force_lbf=50., starting_force_lbf=10., running_force_lbf=10.)
            phase9.execute(state)
        elif job == 'LNR 7"':
            official = next(product for product in adapter.filter_official_products(adapter.load_official_catalog(), 7.)
                            if product["kind"] == "rigid")
            state["centralizer_config"]["custom_products"][0]["identification"] = official["identification"]
            for index in (1, 3):
                state["centralizer_config"]["zones"][index].update(product_nr=official["nr"], custom_product_id=None)
            phase9.execute(state)
            assert len({row["product_name"] for row in state["_centralizer_result"]["zones"]}) == 2
        assert phase9.result_is_current(state)
    for key, value in state.items():
        app.session_state[key] = value
    with patch.object(adapter, "solve_static", side_effect=AssertionError("Word executed Static")), \
         patch.object(adapter, "solve_dynamic", side_effect=AssertionError("Word executed Dynamic")):
        case.phase(app, "phase10")
        build = next(button for button in app.button if button.label == "Build Word Document")
        assert not build.disabled, [error.value for error in app.error]
        build.click().run()
        assert not app.exception and not app.error
    data = app.session_state["_compiled_doc_bytes"]
    with patch.object(report.st, "session_state", app.session_state.to_dict()):
        context = report.build_master_context()
    doc = Document(BytesIO(data))
    region = ix_region(doc)
    text = "\n".join("".join(t.text or "" for t in element.iter(qn("w:t"))) for element in region)
    tables = [table for table in doc.tables if table._tbl in region]
    drawings = [element for element in region if element.xpath('.//w:drawing')]
    assert len(region[-1].xpath('.//w:br[@w:type="page"]')) == 1
    with ZipFile(BytesIO(data)) as package:
        xml = package.read("word/document.xml").decode()
        assert "{{" not in xml and "{%" not in xml
        assert "Pumping simulation charts: Not provided." not in xml
    pumping = [p for p in doc.paragraphs if p._p.xpath('.//wp:docPr[@descr="Pumping Simulation Charts"]')]
    if zones:
        assert len(pumping) == 1 and len(pumping[0]._p.xpath('.//w:drawing')) == 1
        diagnostic = [p for p in doc.paragraphs if p._p.xpath('.//wp:docPr[@descr="Static Well Security and Fluid Sequence"]')]
        assert len(diagnostic) == 1
        elements = list(doc._element.body)
        iv = next(p for p in doc.paragraphs if p.text.strip() == 'IV. Fluids Sequence')
        v = next(p for p in doc.paragraphs if p.text.strip() == 'V. Cementing Program')
        assert elements.index(iv._p) < elements.index(diagnostic[0]._p) < elements.index(pumping[0]._p) < elements.index(v._p)
    else:
        assert not pumping
        assert [p.text for p in doc.paragraphs].count('Pumping simulation charts: Not applicable.') == 1
    if not zones:
        assert not tables and not drawings and bridge.NOTE_BODY not in text
        assert text.count(bridge.NOT_APPLICABLE) == 1
        return
    view = context["centralizer_report"]["view"]
    assert len(tables) == 2 and len(drawings) == 1
    assert len(tables[1].rows) == 2 + zones
    assert [tables[0].cell(i, j).text for i, j in ((1,1),(1,3),(2,1),(2,3),(3,1),(3,3))] == [
        f"{view['top_cent_md']:.1f} m", str(view["n_cent"]), f"{view['bottom_cent_md']:.1f} m",
        str(view["n_rigid"]), f"{view['shoe']:.1f} m", str(view["n_spring"])]
    for row, expected in zip(tables[1].rows[2:], view["zone_rows"]):
        assert [cell.text for cell in row.cells] == [f"{expected['bottom_md']:.1f}", str(expected["number"]),
            expected["cent_per_joint"], expected["kind_display"], f"{expected['min_stdf_pct']:.1f}", f"{expected['min_depth']:.1f}"]
    rid = drawings[0].xpath('.//a:blip')[0].get(qn("r:embed"))
    assert doc.part.related_parts[rid].blob == bridge.render_chart_png(context["centralizer_report"])
    assert text.count(context["note_centralizer"]) == 1
    assert context["all_notes"][-1] == context["note_centralizer"]
    assert text.count(bridge.WARNING) == int(miss)
    if miss:
        assert text.index(bridge.WARNING) < text.index(context["note_centralizer"])
    assert not list(Path("/tmp").glob("round9g-standoff-*"))
