"""Round 12B: register one standalone liner size through existing authorities."""
from copy import deepcopy
from io import BytesIO
import json
import math
import re
from unittest.mock import patch

from docx import Document
from docxtpl import DocxTemplate
import pandas as pd
import pytest

import centralizer_adapter as centralizer
import cement_chart_model as model
import engineering_tools as tools
import materials_db as db
import phase_10_procedure as report
import phase_9_centralizer as phase9
import placement
import project_state as ps
import pumping_simulation as pumping
import well_schematic_adapter as schematic
from project_io import decode_project
import test_audit_regressions as audit
from test_uat_round9b_geometry import hardware
from test_uat_round10a_pumping_simulation import program_state


JOB = 'LNR 4 1/2"'


def liner_state(job=JOB, size='4 1/2', diesel=False):
    fluids = {'Main': 2., 'Displacement Fluid': 2.}
    if diesel:
        fluids = {'Diesel': 1., **fluids}
    state = program_state('liner', fluids=fluids)
    state['job_type'] = job
    state['hardware_table'] = hardware(
        ('Drill Pipes', '400', '3 1/2', 2.75),
        ('Liner', '1000', size, 4.), ('Liner Hanger', '400', size, 4.),
        ('Previous Casing', '500', '7', 6.276), ('Open Hole Size', '1000', '6', 6.))
    state['placement_config'].update(job_type=job, target_row=next(iter(
        placement.hardware_choices(state['hardware_table'], {'Liner'}))))
    state['cement_params'] = {'Main': {'top_mode': 'Surface', 'top_job_type': job,
                                      'mix_water': 20., 'dead_vol': 2., 'tank_name': 'Mud Reserve Tanks'}}
    state.update(pore_gradient_emw_pcf=70., frac_gradient_emw_pcf=160., mud_type='OBM' if diesel else 'WBM')
    if diesel:
        state['fluids_config']['params']['Diesel'].update(density='52.5', material_name='Diesel')
        state['lab_qc_params']['Diesel'] = deepcopy(state['lab_qc_params']['Main'])
    return state


def zones(catalog, state):
    size = next(row['Size (in)'] for row in state['hardware_table'].to_dict('records') if row['Description'] == 'Liner')
    # Explicit fixture selection from real candidates; the application never auto-selects.
    product = next(p for p in centralizer.filter_official_products(catalog, tools.parse_fractional_size(size))
                   if p['kind'] == 'rigid')
    return [{'top_md': 400., 'bottom_md': 1000., 'density': '1/3',
             'target_standoff_pct': 70., 'product_nr': product['nr']}]


def procedure(state):
    ps.refresh_fluids(state)
    with patch.object(report.st, 'session_state', state):
        return report.generate_official_procedure(state['job_type'], state['fluid_data'],
            state['fluids_config']['active'], state['cement_params'], {}, None,
            1000., state['total_pump_time_min'], state['placement_config'])


def register_roots(monkeypatch):
    """Isolate baseline downstream GREEN probes by supplying only future root entries.

    No solver, geometry, persistence, UI or renderer function is replaced.
    """
    monkeypatch.setattr(db, 'JOB_TYPES', [*db.JOB_TYPES, JOB] if JOB not in db.JOB_TYPES else db.JOB_TYPES)
    monkeypatch.setitem(db.DEFAULT_HOLE_SIZES, JOB, '6"')
    monkeypatch.setitem(db.TANK_OPTIONS_BY_JOB, JOB, db.TANK_OPTIONS_BY_JOB['LNR 5"'])
    monkeypatch.setitem(model.CHART_JOBS, JOB, 'liner')


def test_exact_taxonomy_and_position():
    assert JOB in db.JOB_TYPES
    assert db.JOB_TYPES[db.JOB_TYPES.index('LNR 5"') + 1] == JOB
    assert 'TIE BACK LNR 4 1/2"' not in db.JOB_TYPES


def test_default_hole():
    assert db.DEFAULT_HOLE_SIZES.get(JOB) == '6"'


def test_explicit_tank_mapping_matches_lnr5():
    assert JOB in db.TANK_OPTIONS_BY_JOB
    assert db.get_available_tanks(JOB) == db.get_available_tanks('LNR 5"') == [
        'Mud Reserve Tanks', 'Batch Mixer Tanks', 'Cement Unit Tanks']


def test_existing_chart_kind_admits_exact_job():
    assert model.chart_job_kind(JOB) == 'liner'


def test_saved_project_trust_boundary_admits_exact_job():
    assert decode_project(json.dumps({'job_type': JOB}).encode(), lambda value: value)['job_type'] == JOB


def test_centralizer_admits_actual_liner_hardware():
    assert centralizer.centralizer_applicable(JOB)
    payload = centralizer.to_engine_geometry(liner_state())
    assert payload['job_type'] == 'liner'
    assert payload['casing']['od'] == 4.5 * centralizer.INCH_TO_M


def test_existing_pumping_admits_chart_liner_path():
    prepared = pumping.prepare_pumping_simulation(liner_state())
    assert prepared['applicable']
    assert prepared['inputs']['geometry'].pipe_sections[-1].od_in == 4.5


def test_existing_schematic_admits_fractional_liner():
    payload = schematic.prepare_schematic(liner_state())
    assert any(row['size_in'] == 4.5 for row in payload['casings'])


def test_procedure_has_fractional_operation_size():
    text = procedure(liner_state())
    assert 'running 4 1/2 in Liner' in text
    assert 'running 5 in Liner' not in text


@pytest.mark.parametrize('size', ['4 1/2', '4 1/2"'])
def test_existing_fractional_parser_and_placement_family_are_green(size):
    assert tools.parse_fractional_size(size) == 4.5
    assert placement.target_descriptions(JOB) == {'Liner'}
    assert db.get_cementing_method(JOB) == 'Primary Cementing'
    assert db.get_procedure_group(JOB) == 'Group A'


def test_real_catalog_has_fractional_pipe_and_six_inch_family_without_hole_filter():
    catalog = centralizer.load_official_catalog()
    candidates = centralizer.filter_official_products(catalog, tools.parse_fractional_size('4 1/2"'))
    assert candidates and all(p['pipe_size_in'] == 4.5 for p in candidates)
    assert any(p['hole_size_in'] == 6. for p in candidates)
    assert len({p['hole_size_in'] for p in candidates}) > 1


def test_generic_ui_trust_and_metadata_need_only_root_registration(monkeypatch):
    register_roots(monkeypatch)
    assert decode_project(json.dumps({'job_type': JOB}).encode(), lambda value: value)['job_type'] == JOB
    app = audit.AuditRegressions().app()
    app.selectbox(key='_w_job_type').set_value(JOB).run()
    assert not app.exception
    assert app.session_state['job_type'] == JOB
    assert app.session_state['hole_size'] == '6"'
    assert app.session_state['cementing_method'] == 'Primary Cementing'
    assert ps.request_description_for_job(JOB) == '4 1/2" LNR Cementing'


def test_real_downstream_contracts_need_only_root_registration(monkeypatch):
    register_roots(monkeypatch)
    state = liner_state(diesel=True)
    catalog = centralizer.load_official_catalog()
    zone = zones(catalog, state)
    geometry = model.resolve_chart_geometry(state)
    assert geometry['kind'] == 'liner'
    payload = centralizer.to_engine_geometry(state)
    assert payload['casing'] == {'od': 4.5 * .0254, 'id': 4. * .0254,
        'weight_kg_m': 50. * .45359237 / .3048, 'joint_length': 12.2}
    result = centralizer.solve_dynamic(state, zone, catalog, mode='direct')
    assert result['profile']['centralizer_md']
    assert result['diagnostics']['rheology']['Diesel']['source'] == 'lab_bhct_down'
    prepared = pumping.prepare_pumping_simulation(state)
    assert prepared['applicable'] and prepared['inputs']['geometry'].pipe_sections[-1].od_in == 4.5
    drawing = schematic.prepare_schematic(state)
    assert schematic._job_size(JOB) == 4.5
    assert any(row['size_in'] == 4.5 for row in drawing['casings'])


@pytest.mark.parametrize('alias', ['LNR 4.5"', 'LNR 4-1/2"', 'LNR 4½"', 'TIE BACK LNR 4 1/2"'])
def test_only_exact_canonical_job_is_admitted(alias):
    assert alias not in db.JOB_TYPES
    with pytest.raises(ValueError, match='not supported'):
        decode_project(json.dumps({'job_type': alias}).encode(), lambda value: value)
    with pytest.raises(ValueError, match='Unsupported chart job'):
        model.chart_job_kind(alias)


def test_closed_diesel_taxonomy_and_order_are_unchanged():
    assert db.HYDRAULIC_EXECUTION_ORDER == ('Diesel', 'Pre Flush', 'Spacer', 'Spacer Ahead',
        'Scavenger', 'Main', 'Lead', 'Lead #1', 'Lead #2', 'Tail', 'Spacer Behind', 'Displacement Fluid')
    assert db.PHASE7_LAB_FLUIDS == (*db.CEMENT_FORMULATION_FLUIDS, *db.RHEOLOGY_ONLY_FLUIDS)
    assert db.RHEOLOGY_ONLY_FLUIDS == ('Diesel', 'Pre Flush')
    assert db.RHEOLOGY_LAB_FLUIDS == db.CEMENT_FORMULATION_FLUIDS


@pytest.mark.parametrize('size', ['4 1/2', '4 1/2"'])
def test_actual_hardware_dimensions_joint_and_weight_remain_authoritative(size):
    state = liner_state(size=size)
    frame = state['hardware_table']
    frame.loc[frame['Description'] == 'Liner', ['ID (in)', 'Joint (m)', 'Weight (ppf)']] = [4.05, 9.5, 23.5]
    translated = centralizer.to_engine_geometry(state)
    assert translated['casing'] == {'od': 4.5 * .0254, 'id': 4.05 * .0254,
        'weight_kg_m': 23.5 * .45359237 / .3048, 'joint_length': 9.5}
    resolved = model.resolve_chart_geometry(state)
    liner = next(row for row in resolved['active_string'] if row['description'] == 'Liner')
    assert (liner['od_in'], liner['id_in']) == (4.5, 4.05)
    simulation = pumping.prepare_pumping_simulation(state)['inputs']['geometry']
    assert (simulation.pipe_sections[-1].od_in, simulation.pipe_sections[-1].id_in) == (4.5, 4.05)
    drawing = schematic.prepare_schematic(state)
    target = drawing['casings'][drawing['current_job']['target_index']]
    assert (target['size_in'], target['id_in']) == (4.5, 4.05)


@pytest.mark.parametrize('formation', ['basic', 'advanced'])
def test_existing_liner_placement_hydrostatic_and_security_are_size_agnostic(formation):
    state = liner_state(diesel=True)
    if formation == 'advanced':
        state['formation_pressure_config'] = {'mode': 'advanced', 'profile': [
            {'md_m': 500., 'pore_psi': 400., 'frac_psi': 2000.},
            {'md_m': 1000., 'pore_psi': 800., 'frac_psi': 3000.}]}
    resolved = model.build_chart_model(state)
    assert resolved['kind'] == 'liner' and resolved['static_security_ready']
    assert resolved['static_security']['domain']['top_md'] == 500.
    assert resolved['static_security']['domain']['bottom_md'] == 1000.
    intervals = placement.slurry_intervals(state['hardware_table'], JOB, state['placement_config'],
        ['Main'], state['cement_params'], state['fluids_config']['params'])
    official = model.resolve_cement_placement(state)
    assert list(intervals) == [row['fluid'] for row in official]
    for row in official:
        interval = intervals[row['fluid']]
        assert interval['top_depth'] == row['top_md']
        assert interval['bottom_depth'] == row['bottom_md']
        assert interval['returned_bbl'] == row['returned_bbl']
    # Same actual Hardware, using an existing registered liner identity, must
    # produce the same engineering values rather than guessed Job-Type geometry.
    reference = deepcopy(state)
    reference['job_type'] = reference['placement_config']['job_type'] = 'LNR 5"'
    reference['cement_params']['Main']['top_job_type'] = 'LNR 5"'
    expected = model.build_chart_model(reference)
    for key in ('hydrostatic', 'pore_frac', 'static_security', 'annular_intervals', 'cement_placement'):
        assert resolved[key] == expected[key]
    assert all(math.isfinite(value) for value in resolved['hydrostatic']['pressure_psi'])


@pytest.mark.parametrize('physics', ['static', 'dynamic'])
@pytest.mark.parametrize('mode', ['direct', 'optimize'])
@pytest.mark.parametrize('diesel', [False, True])
def test_real_centralizer_runs_existing_modes_with_explicit_catalog_identity(physics, mode, diesel):
    state = liner_state(diesel=diesel)
    catalog = centralizer.load_official_catalog()
    config = phase9.new_config()
    config.update(physics_mode=physics, placement_mode=mode, zones=zones(catalog, state))
    state['centralizer_config'] = config
    result = phase9.execute(state, catalog)
    assert result['profile']['centralizer_md']
    assert all(math.isfinite(value) for value in result['profile']['tension_force'])
    assert phase9.result_is_current(state, catalog)
    if physics == 'dynamic' and diesel:
        assert result['diagnostics']['rheology']['Diesel']['source'] == 'lab_bhct_down'
    assert result['zones'][0]['product_nr'] == config['zones'][0]['product_nr']


@pytest.mark.parametrize('diesel', [False, True])
def test_real_frozen_pumping_executes_without_new_job_model(diesel):
    state = liner_state(diesel=diesel)
    inputs = pumping.prepare_pumping_simulation(state)['inputs']
    result = pumping.run_pumping_simulation(state)
    try:
        assert result['applicable'] and result['result'] is not None
        assert len(result['figure'].axes) == 3
        assert inputs['geometry'].track_depth_m == 1000.
        assert inputs['geometry'].pipe_sections[-1].od_in == 4.5
        if diesel:
            assert inputs['rheology']['Diesel']['source'] == 'phase7:bhct_down'
    finally:
        pumping.fcm.plt.close(result['figure'])


def test_real_fractional_schematic_keeps_identity_and_renders(tmp_path):
    import well_schematic_engine as engine
    state = liner_state()
    payload = schematic.prepare_schematic(state)
    assert payload['current_job']['type'] == 'liner'
    target = payload['casings'][payload['current_job']['target_index']]
    assert target['size_in'] == 4.5
    assert schematic.format_size(target['size_in']) == '4 1/2"'
    assert engine.check_well_design(payload) == []
    path = tmp_path / 'lnr-4-1-2.png'
    engine.render_well_schematic(payload, str(path), dpi=90)
    assert path.read_bytes().startswith(b'\x89PNG\r\n\x1a\n')


@pytest.mark.parametrize('job,size,pressure,rotation', [(JOB, '4 1/2', '150', False),
    ('LNR 5"', '5', '150', False), ('LNR 7"', '7', '100', True)])
def test_procedure_display_is_separate_from_exact_lnr7_variant(job, size, pressure, rotation):
    text = procedure(liner_state(job, size, diesel=True))
    assert f'running {size} in Liner' in text
    assert f'hold {pressure} psi pressure' in text
    assert ('raised and lowered with a 7 in liner' in text) == rotation
    assert 'If plug does not bump, do not displace more than %25 of the calculated shoe track.' in text
    assert 'If observed bumping, increase pressure to 3000 psi and hold for 15 min due to liner integrity test.(Optional)' in text
    local_notes = [int(value) for value in re.findall(r'(?m)^Note(\d+):', text)]
    assert local_notes == list(range(1, (3 if rotation else 2) + 1))
    steps = [int(value) for value in re.findall(r'(?m)^(\d+)\. ', text)]
    assert steps == list(range(1, len(steps) + 1))
    assert text.count('Pump 1.0 bbl Diesel as pre flush.(52.5 pcf)') == 1


@pytest.fixture(scope='module')
def word_project():
    from test_uat_round12a5_diesel_word_e2e import projects as diesel_projects
    from centralizer_fixtures import execute_centralizer
    # Start with reviewed Phase-IV–VII data, then explicitly supply the new
    # operation's Hardware. Historical audit fixture size dictionaries are not
    # authorities for new-job geometry.
    state = diesel_projects.__wrapped__()('LNR 5"')
    state['job_type'] = state['doc_control']['job_type'] = JOB
    state['hole_size'] = state['doc_control']['hole_size'] = '6"'
    for key in ('request_description', 'revision_description'):
        state[key] = state['doc_control'][key] = ps.request_description_for_job(JOB)
    frame = state['hardware_table']
    frame.loc[frame['Description'].isin(['Liner', 'Liner Hanger']), 'Size (in)'] = '4 1/2"'
    frame.loc[frame['Description'].isin(['Liner', 'Liner Hanger']), 'ID (in)'] = 4.
    frame.loc[frame['Description'] == 'Open Hole Size', 'Size (in)'] = '6'
    frame.loc[frame['Description'] == 'Open Hole Size', 'ID (in)'] = 6.
    state['placement_config'].update(job_type=JOB, target_row=next(iter(
        placement.hardware_choices(frame, {'Liner'}))))
    for params in state['cement_params'].values():
        params.update(top_job_type=JOB, last_job_type=JOB)
    for name, params in state['fluids_config']['params'].items():
        params['volume'] = 20. if name == 'Diesel' else 4.
    assert ps.prepare_calculations(state) == []
    assert execute_centralizer(state, physics='dynamic') is not None
    return state


def word_context(state):
    with patch.object(report.st, 'session_state', state):
        return report.build_master_context()


@pytest.mark.parametrize('diesel', [False, True])
def test_real_word_context_and_protected_template_preserve_fractional_identity(word_project, diesel):
    from centralizer_fixtures import execute_centralizer
    state = deepcopy(word_project)
    if not diesel:
        state['fluids_config']['active'].remove('Diesel')
        assert ps.prepare_calculations(state) == []
        execute_centralizer(state, physics='dynamic')
    ctx = word_context(state)
    assert ctx['job_type'] == JOB
    assert '4 1/2 in Liner' in ctx['procedure_text']
    assert 'running 5 in Liner' not in ctx['procedure_text']
    assert 'Tie Back' not in ctx['procedure_text']
    assert ctx['centralizer_report']['applicable']
    assert ctx['pumping_simulation_source'] == pumping.source_signature(state)
    template = DocxTemplate(report.TEMPLATE_PATH)
    template.render(report._word_quantity_context(ctx))
    stream = BytesIO(); template.save(stream)
    doc = Document(BytesIO(stream.getvalue()))
    paragraphs = '\n'.join(p.text for p in doc.paragraphs)
    tables = '\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert '4 1/2 in Liner' in paragraphs and 'running 5 in Liner' not in paragraphs
    assert '4 1/2"' in tables
    assert ('Pump 20.0 bbl Diesel as pre flush.(52.5 pcf)' in paragraphs) == diesel
    assert [int(number) for number in re.findall(r'(?m)^NOTE (\d+):', '\n'.join(ctx['all_notes']))] == list(
        range(1, len(ctx['all_notes']) + 1))


@pytest.mark.parametrize('diesel', [False, True])
def test_exact_job_save_restore_and_full_build_use_existing_consumers(word_project, diesel):
    from types import SimpleNamespace
    from centralizer_fixtures import execute_centralizer
    from test_uat_round10c_integration import save_project, restore_project, build
    state = deepcopy(word_project)
    if not diesel:
        state['fluids_config']['active'].remove('Diesel')
        state['mud_type'] = 'WBM'
        assert ps.prepare_calculations(state) == []
        execute_centralizer(state, physics='dynamic')
    state['schematic_config'] = {'version': 1}
    raw = save_project(SimpleNamespace(session_state=state))
    app = restore_project(raw)
    restored = app.session_state.to_dict()
    assert restored['job_type'] == JOB
    pd.testing.assert_frame_equal(restored['hardware_table'], state['hardware_table'])
    for key in ('placement_config', 'fluids_config', 'cement_params', 'lab_qc_params',
                'centralizer_config', 'schematic_config'):
        assert ps.fingerprint(restored[key]) == ps.fingerprint(state[key]), key
    assert '_centralizer_result' not in restored
    assert not any('liner_variant' in key for key in restored)
    assert execute_centralizer(app.session_state, physics='dynamic') is not None
    data = build(app)
    doc = Document(BytesIO(data))
    paragraphs = '\n'.join(p.text for p in doc.paragraphs)
    tables = '\n'.join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert '4 1/2 in Liner' in paragraphs and 'running 5 in Liner' not in paragraphs
    assert '4 1/2"' in tables
    for description in ('Well Schematic', 'Static Well Security and Fluid Sequence', 'Pipe Standoff', 'Pumping Simulation Charts'):
        assert any(p._p.xpath(f'.//wp:docPr[@descr="{description}"]') for p in doc.paragraphs), description
    assert phase9.result_is_current(app.session_state)
    assert app.session_state['_pumping_simulation_source'] == pumping.source_signature(app.session_state)
    assert ('Pump 20.0 bbl Diesel as pre flush.(52.5 pcf)' in paragraphs) == diesel


def test_real_new_job_phase_navigation_reuses_existing_iv_through_vii_ui(word_project):
    case = audit.AuditRegressions()
    app = case.app(deepcopy(word_project))
    case.phase(app, 'phase1')
    assert app.selectbox(key='_w_job_type').value == JOB
    assert app.session_state['hole_size'] == '6"'
    assert app.session_state['cementing_method'] == 'Primary Cementing'
    case.phase(app, 'phase2_3')
    assert app.session_state['placement_config']['job_type'] == JOB
    case.phase(app, 'phase4')
    assert app.checkbox(key='chk_fluid_Diesel').value
    assert app.number_input(key='vol_Diesel').value == 20.
    case.phase(app, 'phase5')
    tank = next(w for w in app.selectbox if w.label == 'Mixing Tank - Main')
    assert tank.options == db.get_available_tanks('LNR 5"')
    case.phase(app, 'phase6')
    assert 'Diesel' not in app.session_state.get('spacer_dfs', {})
    case.phase(app, 'phase7')
    assert {tab.label for tab in app.tabs} == {'Main', 'Diesel', 'Pre Flush'}
    assert 'Diesel' not in app.session_state['cement_params']
    assert phase9.result_is_current(app.session_state)
