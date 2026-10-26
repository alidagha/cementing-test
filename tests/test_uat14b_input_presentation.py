"""14B UI projections preserve authoritative engineering inputs."""
from copy import deepcopy
import json
import math
from io import BytesIO
from zipfile import ZipFile
import xml.etree.ElementTree as ET

import pandas as pd
import pytest
from docx import Document
import test_audit_regressions as audit
from phase_2_3_well_data import HARDWARE_COLUMNS


CASE = audit.AuditRegressions()


def hardware(description="Open Hole Size", md="1500"):
    return pd.DataFrame([{
        "Description": description, "MD (m)": md, "Size (in)": "8.5", "ID (in)": 8.5 if description == "Open Hole Size" else 7.5,
        "Joint (m)": 0. if description == "Open Hole Size" else 12.2,
        "Weight (ppf)": 0., "Grade": "-", "Collapse (psi)": 0., "Burst (psi)": 0.,
    }], columns=HARDWARE_COLUMNS)


def editor(app):
    return next(e for e in app.dataframe if "_editor_hardware_" in e.proto.id)


def test_no_unapproved_hardware_table_button_or_caption():
    app = CASE.app({"hardware_table": hardware()})
    CASE.phase(app, "phase2_3")
    assert not [e for e in app.dataframe if not e.proto.id and list(e.value.columns) == HARDWARE_COLUMNS]
    assert not [b for b in app.button if b.key == "_apply_open_hole_td"]
    assert not [c for c in app.caption if "Read-only hardware display" in c.value or "Automatic BHSP (ceiling)" in c.value]


def test_fresh_job_choices_omit_csg7_but_not_backend_support():
    app = CASE.app()
    assert 'CSG 7"' not in app.selectbox(key="_w_job_type").options
    assert 'CSG 7"' in audit.materials_db.JOB_TYPES


def test_fresh_hardware_choices_omit_tubing():
    app = CASE.app()
    CASE.phase(app, "phase2_3")
    options = json.loads(editor(app).proto.columns)["Description"]["type_config"]["options"]
    assert "Tubing" not in options
    assert "Tubing" in audit.materials_db.HARDWARE_DESCRIPTIONS


@pytest.mark.parametrize("md,shown", [("1500", "1500.0"), ("1500.25000001", "1500.25000001"), ("1000.0-1500.25", "1000.0-1500.25")])
def test_md_presentation_never_rounds_or_rewrites_authoritative_text(md, shown):
    app = CASE.app({"hardware_table": hardware(md=md)})
    CASE.phase(app, "phase2_3")
    assert editor(app).value.at[0, "MD (m)"] == md  # Native text grid; Owner Word fallback is checked below.
    assert app.session_state["hardware_table"].at[0, "MD (m)"] == md
    restored = audit.round_trip(app.session_state.to_dict())
    assert restored["hardware_table"].at[0, "MD (m)"] == md


def test_open_hole_presentation_is_not_numeric_canonical_data():
    app = CASE.app({"hardware_table": hardware()})
    CASE.phase(app, "phase2_3")
    fields = ("Joint (m)", "Weight (ppf)", "Grade", "Collapse (psi)", "Burst (psi)")
    assert editor(app).value.at[0, "Joint (m)"] == 0.0  # Native row locking unavailable; Word-only dashes below.
    canonical = app.session_state["hardware_table"]
    assert canonical.at[0, "Grade"] == "-"
    assert all(canonical.at[0, key] == 0.0 for key in fields if key != "Grade")


def test_vertical_td_automatically_seeds_original_control_and_preserves_edit():
    app = CASE.app({"hardware_table": hardware(md="1500.25"),
                    "well_geometry": {"version": 1, "type": "Vertical", "td_m": None}})
    CASE.phase(app, "phase2_3")
    assert app.number_input(key="_w_vertical_td").value == 1500.25
    assert not [b for b in app.button if b.key == "_apply_open_hole_td"]
    assert app.session_state["well_geometry"]["td_m"] == 1500.25
    app.number_input(key="_w_vertical_td").set_value(1600.125).run()
    CASE.phase(app, "phase1")
    CASE.phase(app, "phase2_3")
    assert app.session_state["well_geometry"]["td_m"] == 1600.125
    assert not [b for b in app.button if b.key == "_apply_open_hole_td"]
    CASE.hardware_event(app, {"edited_rows": {"0": {"MD (m)": "1800"}}})
    assert app.session_state["hardware_table"].at[0, "MD (m)"] == "1800"
    assert app.session_state["well_geometry"]["td_m"] == 1600.125
    assert audit.round_trip(app.session_state.to_dict())["well_geometry"]["td_m"] == 1600.125


def test_automatic_bhsp_ceiling_is_display_only_and_manual_remains_exact():
    app = CASE.app({**audit.well_profile_fixture(1500.25), "mud_density": "80-82", "well_auto_fields": {"bhsp": True}})
    CASE.phase(app, "phase2_3")
    expected = 1500.25 * 82 * 0.02278
    assert float(app.session_state["bhsp"]) == expected
    assert app.session_state["well_data"]["bhsp"] == app.session_state["bhsp"]
    assert app.text_input(key="_w_bhsp").value == str(math.ceil(expected))
    app.text_input(key="_w_bhsp").set_value(str(math.ceil(expected))).run()
    CASE.phase(app, "phase1")
    CASE.phase(app, "phase2_3")
    assert float(app.session_state["bhsp"]) == expected
    assert app.session_state["well_auto_fields"]["bhsp"] is True
    assert float(audit.round_trip(app.session_state.to_dict())["bhsp"]) == expected
    app.text_input(key="_w_bhsp").set_value("1234.125").run()
    assert app.session_state["bhsp"] == "1234.125"
    assert app.session_state["well_auto_fields"]["bhsp"] is False
    assert audit.round_trip(app.session_state.to_dict())["well_data"]["bhsp"] == "1234.125"


@pytest.mark.parametrize("job", audit.materials_db.JOB_TYPES)
def test_every_supported_saved_job_is_representable_and_round_trips(job):
    app = CASE.app({"job_type": job, "hole_size": audit.materials_db.DEFAULT_HOLE_SIZES[job]})
    assert app.selectbox(key="_w_job_type").value == job
    assert job in app.selectbox(key="_w_job_type").options
    assert audit.round_trip(app.session_state.to_dict())["job_type"] == job
    if job == 'CSG 7"':
        app.selectbox(key="_w_job_type").set_value('LNR 7"').run()
        CASE.healthy(app)
        assert 'CSG 7"' not in app.selectbox(key="_w_job_type").options
        assert app.session_state["job_type"] == 'LNR 7"'


def test_saved_tubing_is_preserved_in_editor_and_json():
    app = CASE.app({"hardware_table": hardware("Tubing", "1500.25")})
    CASE.phase(app, "phase2_3")
    options = json.loads(editor(app).proto.columns)["Description"]["type_config"]["options"]
    assert "Tubing" in options
    CASE.phase(app, "phase1")
    restored = CASE.app(audit.round_trip(app.session_state.to_dict()))
    CASE.phase(restored, "phase2_3")
    assert restored.session_state["hardware_table"].at[0, "Description"] == "Tubing"
    assert restored.session_state["hardware_table"].at[0, "MD (m)"] == "1500.25"


@pytest.mark.parametrize("md,td,kind,count", [
    ("1000-1500", None, "Vertical", 1), ("bad", None, "Vertical", 1),
    ("0", None, "Vertical", 1), ("-1", None, "Vertical", 1),
    ("1500", None, "Vertical", 2), ("1500", 0.0, "Vertical", 1),
    ("1500", 1600.125, "Vertical", 1), ("1500", None, "Directional", 1),
])
def test_td_suggestion_never_resolves_ambiguous_invalid_or_explicit_sources(md, td, kind, count):
    table = pd.concat([hardware(md=md)] * count, ignore_index=True)
    app = CASE.app({"hardware_table": table,
                    "well_geometry": {"version": 1, "type": kind, "td_m": td}})
    CASE.phase(app, "phase2_3")
    assert not [b for b in app.button if b.key == "_apply_open_hole_td"]
    assert app.session_state["well_geometry"]["td_m"] == td


def test_pending_hardware_cannot_supply_a_td_suggestion():
    table = hardware()
    pending = pd.concat([table, hardware("", "1800")], ignore_index=True)
    app = CASE.app({"hardware_table": table, "hardware_editor_draft": pending,
                    "well_geometry": {"version": 1, "type": "Vertical", "td_m": None}})
    CASE.phase(app, "phase2_3")
    assert not [b for b in app.button if b.key == "_apply_open_hole_td"]
    assert app.session_state["well_geometry"]["td_m"] is None
    assert len(app.session_state["hardware_editor_draft"]) == 2




def test_csg7_valid_fixture_still_builds_real_word(tmp_path):
    app = CASE.configured_app('CSG 7"')
    assert app.session_state["job_type"] == 'CSG 7"'
    assert app.session_state["hole_size"] == audit.materials_db.DEFAULT_HOLE_SIZES['CSG 7"']
    # Match B2's computed-placement fixture contract, retaining all real Lab prerequisites.
    for params in app.session_state["cement_params"].values():
        params.update(top_mode="Not entered", top_depth=None, report_bottom_depth=None)
    CASE.export(app)
    data = app.session_state["_compiled_doc_bytes"]
    (tmp_path / "supported-csg7.docx").write_bytes(data)
    doc = Document(BytesIO(data))
    assert doc.tables
    with ZipFile(BytesIO(data)) as package:
        xml = ET.fromstring(package.read("word/document.xml"))
    text = "\n".join(node.text or "" for node in xml.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))
    assert "Regression Well" in text and "Regression Client" in text
    assert '7"' in text
    table = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == "Tubular and Casing Hardware")
    rows = [[c.text.strip() for c in row.cells] for row in table.rows]
    state = deepcopy(app.session_state.to_dict())
    original = state["hardware_table"].copy(deep=True)
    from unittest.mock import patch
    import phase_10_procedure as report
    with patch.object(report.st, "session_state", state):
        ctx = report.build_master_context(calculations_prepared=True)
    pd.testing.assert_frame_equal(original, state["hardware_table"], check_exact=True)
    from test_pre_checkpoint_b2_cross_module_integration import assert_shared_geometry
    assert_shared_geometry(state, ctx)  # Real CSG consumer, sharing this required DOCX build.
    for raw, displayed in zip(original.to_dict("records"), ctx["hardware"]):
        assert displayed["MD (m)"] == f'{float(raw["MD (m)"]):.1f}'
        actual = next(row for row in rows if raw["Description"] in row)
        assert displayed["MD (m)"] in actual
        if raw["Description"] == "Open Hole Size":
            for field in ("Joint (m)", "Weight (ppf)", "Grade", "Collapse (psi)", "Burst (psi)"):
                assert displayed[field] == "-"
            assert actual.count("-") >= 5


def test_blank_automatic_bhsp_callback_preserves_automatic_provenance():
    from unittest.mock import patch
    import phase_2_3_well_data as well
    app = CASE.app({"well_auto_fields": {"bhsp": True}})
    CASE.phase(app, "phase2_3")
    assert app.text_input(key="_w_bhsp").value == ""
    state = deepcopy(app.session_state.to_dict())
    state["_w_bhsp"] = ""
    with patch.object(well.st, "session_state", state):
        well._commit_well_widget("bhsp")  # Same unchanged value delivered by a queued blur callback.
    assert state["well_auto_fields"]["bhsp"] is True
    assert state["bhsp"] == state["well_data"]["bhsp"] == ""


def test_automatic_bhsp_source_change_reseeds_display_without_manual_drift():
    app = CASE.app({**audit.well_profile_fixture(1500.25), "mud_density": "80-82", "well_auto_fields": {"bhsp": True}})
    CASE.phase(app, "phase2_3")
    app.text_input(key="_w_mud_density").set_value("81").run()
    CASE.healthy(app)
    precise = 1500.25 * 81 * 0.02278
    assert float(app.session_state["bhsp"]) == precise
    assert app.text_input(key="_w_bhsp").value == str(math.ceil(precise))
    assert app.session_state["well_auto_fields"]["bhsp"] is True
    restored = CASE.app(audit.round_trip(app.session_state.to_dict()))
    CASE.phase(restored, "phase2_3")
    assert float(restored.session_state["bhsp"]) == precise
    assert restored.text_input(key="_w_bhsp").value == str(math.ceil(precise))
    assert restored.session_state["well_auto_fields"]["bhsp"] is True


@pytest.mark.parametrize("raw", ["1e10000", "1.50025e3", "broken", "", "1500.25000000000001"])
def test_primary_grid_preserves_nonstandard_and_precise_md_without_display_coercion(raw):
    app = CASE.app({"hardware_table": hardware(md=raw)})
    CASE.phase(app, "phase2_3")
    assert editor(app).value.at[0, "MD (m)"] == raw
    assert app.session_state["hardware_table"].at[0, "MD (m)"] == raw
    assert audit.round_trip(app.session_state.to_dict())["hardware_table"].at[0, "MD (m)"] == raw


from test_pre14_sh_batch_e import reviewed_main


@pytest.mark.parametrize("md,expected", [("1500.25000001", "1500.3"), ("1000.0-1500.25", "1000.0-1500.25")])
@pytest.mark.parametrize("prepared", [False, True])
def test_report_fallback_preserves_exact_source_and_interval(reviewed_main, md, expected, prepared, tmp_path):
    from unittest.mock import patch
    import project_state as ps
    import phase_10_procedure as report
    state = deepcopy(reviewed_main)
    pipe = hardware("Drill Pipes", md)
    pipe.loc[0, ["Size (in)", "ID (in)"]] = ["5", 4.276]
    state["hardware_table"] = pd.concat([state["hardware_table"], pipe], ignore_index=True)
    before = state["hardware_table"].copy(deep=True)
    assert ps.prepare_calculations(state) == []
    with patch.object(report.st, "session_state", state):
        ctx = report.build_master_context(calculations_prepared=prepared)
    assert ctx["hardware"][-1]["MD (m)"] == expected
    pd.testing.assert_frame_equal(state["hardware_table"], before, check_exact=True)
    assert audit.round_trip(state)["hardware_table"].iloc[-1]["MD (m)"] == md
    from docxtpl import DocxTemplate
    template = DocxTemplate(report.TEMPLATE_PATH)
    template.render(report._word_quantity_context(ctx))
    output = tmp_path / "hardware-projection.docx"
    template.save(output)
    document = Document(output)
    table = next(t for t in document.tables if t.rows[0].cells[0].text.strip() == "Tubular and Casing Hardware")
    row = next(r for r in table.rows if r.cells[0].text.strip() == "Drill Pipes")
    assert row.cells[1].text.strip() == expected
    pd.testing.assert_frame_equal(state["hardware_table"], before, check_exact=True)
