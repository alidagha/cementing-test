"""B3: adversarial edits between real Word builds, reviews and checkpoints."""
from copy import deepcopy
from io import BytesIO
from html import unescape
import hashlib
import re
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
import pytest

import engineering_tools as eng
import materials_db as db
import project_state as ps
import phase_10_procedure as report
import phase_9_centralizer as phase9
import cement_chart_model as chart
import pumping_simulation as pumping
import test_audit_regressions as audit
from centralizer_fixtures import execute_centralizer
from test_uat_round10c_integration import build, save_project, restore_project
from test_uat_round12a5_diesel_word_e2e import numbered_steps
from test_pre_checkpoint_b2_cross_module_integration import (
    projects as integrated_projects, JOBS, CSG, LNR, select_target, signatures,
)


CASE = audit.AuditRegressions()


@pytest.fixture(scope="module")
def projects():
    return integrated_projects.__wrapped__()


def context(state):
    with patch.object(report.st, "session_state", state):
        return report.build_master_context()


def checkpoint(state):
    return restore_project(save_project(SimpleNamespace(session_state=state)))


def no_document(state):
    assert not any(key in state for key in (
        "_compiled_doc_bytes", "_compiled_doc_filename", "_compiled_doc_signature"))


def blocked(app):
    CASE.phase(app, "phase10")
    assert next(b for b in app.button if b.label == "Build Word Document").disabled
    no_document(app.session_state)


def review_main(app, sync=False):
    CASE.phase(app, "phase7")
    if sync:
        next(b for b in app.button if b.key.startswith("_sync_btn_main_")).click().run()
        assert not app.session_state["lab_qc_params"]["Main"]["reviewed"]
    button = next(b for b in app.button if b.key.startswith("_confirm_lab_btn_main_"))
    assert not button.disabled
    button.click().run()
    CASE.healthy(app)
    assert eng.compute_phase_status(app.session_state)["phase7"]["level"] == "ok"


def rebuild(app):
    CASE.phase(app, "phase10")  # Run the real rebuild against Streamlit's proxy.
    assert ps.prepare_calculations(app.session_state.to_dict()) == []
    if not phase9.result_is_current(app.session_state):
        execute_centralizer(app.session_state, physics="dynamic")
    return build(app)


def semantics(data):
    doc = Document(BytesIO(data))
    return ([p.text for p in doc.paragraphs],
            [[[c.text for c in row.cells] for row in table.rows] for table in doc.tables],
            [node.get("descr") for node in doc.element.body.iter(qn("wp:docPr"))])


def media_signatures(data):
    with ZipFile(BytesIO(data)) as archive:
        return {hashlib.sha256(archive.read(name)).hexdigest() for name in archive.namelist()
                if name.startswith("word/media/")}


def sequence_table(doc):
    matches = [t for t in doc.tables if t.rows[0].cells[0].text.strip() == "Fluids"]
    assert len(matches) == 1
    table = matches[0]
    body = list(doc.element.body)
    headings = {p.text.strip(): body.index(p._p) for p in doc.paragraphs}
    assert headings["IV. Fluids Sequence"] < body.index(table._tbl) < headings["V. Cementing Program"]
    return table


def assert_document(data, state):
    doc = Document(BytesIO(data))
    text = "\n".join(node.text or "" for node in doc.element.body.iter(qn("w:t")))
    with ZipFile(BytesIO(data)) as archive:
        for name in archive.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                xml = archive.read(name).decode()
                assert "{{" not in xml and "{%" not in xml, name
    for heading in ("II. Well Data Summary", "III. Geothermal Temperature Profile",
                    "IV. Fluids Sequence", "V. Cementing Program",
                    "IX. Centralizer Placement Recommendation"):
        assert [p.text.strip() for p in doc.paragraphs].count(heading) == 1
    rows = [[c.text.strip() for c in row.cells] for row in sequence_table(doc).rows]
    names = [r[0] for r in rows if r[0] in db.HYDRAULIC_EXECUTION_ORDER]
    active = state["fluids_config"]["active"]
    assert names == [n for n in db.HYDRAULIC_EXECUTION_ORDER if n in active]
    assert len(names) == len(set(names))
    ctx = context(state)
    for source in ctx["fluids_train"]:
        row = next(r for r in rows if r[0] == source["name"])
        assert row[2:5] == [str(source["density"]), f'{source["volume"]:.1f}', str(source["pump_rate"])]
    actual_steps = [p.text.strip() for p in doc.paragraphs if re.match(r"^\d+\. ", p.text.strip())]
    steps = numbered_steps(unescape(ctx["procedure_text"]))
    assert actual_steps == [f"{number}. {body}" for number, body in steps]
    assert [n for n, _ in steps] == list(range(1, len(steps) + 1))
    notes = [p.text.strip() for p in doc.paragraphs if re.match(r"^NOTE \d+:", p.text.strip())]
    assert notes == [unescape(note) for note in ctx["all_notes"]]
    if "Diesel" in active:
        source = state["fluid_data"]["Diesel"]
        sentence = f'Pump {source["volume"]:.1f} bbl Diesel as pre flush.({source["density"]} pcf)'
        assert [body for _, body in steps].count(sentence) == 1
        assert not any("Diesel" in note for note in notes)
    else:
        assert "Diesel" not in text
    hardware = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == "Tubular and Casing Hardware")
    hardware_text = "\n".join(c.text for r in hardware.rows for c in r.cells)
    for row in state["hardware_table"].to_dict("records"):
        assert row["Description"] in hardware_text
        assert str(row["MD (m)"]) in hardware_text
    return doc, ctx


@pytest.mark.parametrize("job", JOBS)
def test_actual_word_structure_and_sequence_for_eight_job_families(projects, job):
    state = projects(job)
    doc, ctx = assert_document(state["_compiled_doc_bytes"], state)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert ctx["job_type"] == job
    if job.startswith("LNR "):
        size = job[len("LNR "):-1]
        assert f"running {size} in Liner" in text
        assert ("raised and lowered with a 7 in liner" in text) == (job == 'LNR 7"')
        pressure = 100 if job == 'LNR 7"' else 150
        assert f'hold {pressure} psi pressure' in text
    if job in ("CMT PLUG", "CMT SQUEEZE"):
        assert not ctx["has_centralizer"]
        assert not pumping.prepare_pumping_simulation(state)["applicable"]


@pytest.mark.parametrize("job", [CSG, LNR, "CMT PLUG"])
def test_repeated_real_build_and_checkpoint_are_semantically_idempotent(projects, job):
    app = CASE.app(projects(job))
    before = save_project(app)
    signature = app.session_state["_compiled_doc_signature"]
    first = build(app)
    second = build(app)
    assert semantics(first) == semantics(second)
    assert app.session_state["_compiled_doc_signature"] == signature
    assert save_project(app) == before
    restored = restore_project(before)
    no_document(restored.session_state)
    third = rebuild(restored)
    assert semantics(third) == semantics(second)
    assert save_project(restored) == before


@pytest.mark.parametrize("change", ["hardware", "volume", "rate", "density", "Diesel rheology", "Formation"])
def test_source_edit_navigation_gate_and_new_actual_docx(projects, change):
    app = CASE.app(projects())
    old_signature = app.session_state["_compiled_doc_signature"]
    old_bytes = app.session_state["_compiled_doc_bytes"]
    state = app.session_state
    if change == "hardware":
        state["hardware_table"].loc[state["hardware_table"]["Description"] == "Casing", "ID (in)"] = 8.4
    elif change in ("volume", "rate", "density"):
        CASE.phase(app, "phase4")
        if change == "volume":
            app.number_input(key="vol_Main").set_value(51.)
        else:
            app.text_input(key=("rate_Main" if change == "rate" else "den_Main")).set_value("5" if change == "rate" else "120")
    elif change == "Diesel rheology":
        CASE.phase(app, "phase7")
        next(w for w in app.text_input if w.key.startswith("_rheo_diesel_bhct_down_300_")).set_value("169")
    else:
        state[chart.PORE_FIELD] = state["well_data"][chart.PORE_FIELD] = 71.
    CASE.phase(app, "phase1")
    CASE.phase(app, "phase10")
    no_document(state)
    if change in ("hardware", "volume", "density", "Diesel rheology"):
        assert not phase9.result_is_current(state)
        assert next(b for b in app.button if b.label == "Build Word Document").disabled
    else:
        assert phase9.result_is_current(state)
    if change == "density":
        review_main(app, sync=True)
    data = rebuild(app)
    assert app.session_state["_compiled_doc_signature"] != old_signature
    assert semantics(data) != semantics(old_bytes) or media_signatures(data) != media_signatures(old_bytes)
    assert_document(data, app.session_state.to_dict())


def test_report_override_clear_restore_and_engineering_change_remain_separate(projects):
    app = CASE.app(projects())
    baseline = signatures(app.session_state)
    source = app.session_state["cement_params"]["Main"]
    source.update(top_mode="Depth (m MD)", top_depth=2100., report_bottom_depth=2900.)
    overridden = build(app)
    assert signatures(app.session_state) == baseline
    doc, ctx = assert_document(overridden, app.session_state.to_dict())
    assert ctx["slurries"][0]["top"] == "2100.0 m"
    assert ctx["slurries"][0]["bottom"] == "2900.0 m MD"
    assert "2100.0 m" in "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    app = checkpoint(app.session_state.to_dict())
    assert app.session_state["cement_params"]["Main"]["top_depth"] == 2100.
    assert semantics(rebuild(app)) == semantics(overridden)
    # Explicit report overrides remain intentional even as engineering changes.
    app.session_state["fluids_config"]["params"]["Main"]["volume"] = 51.
    blocked(app)
    assert not phase9.result_is_current(app.session_state)
    rebuild(app)
    app.session_state["cement_params"]["Main"].update(top_mode="Not entered", top_depth=None, report_bottom_depth=None)
    data = build(app)
    _, ctx = assert_document(data, app.session_state.to_dict())
    interval = next(r for r in chart.resolve_cement_placement(app.session_state) if r["fluid"] == "Main")
    assert ctx["slurries"][0]["top"] == ("surface" if interval["top_md"] == 0 else f'{interval["top_md"]:.1f} m')
    assert ctx["slurries"][0]["bottom"] == f'{interval["bottom_md"]:.1f} m MD'


@pytest.mark.parametrize("change", ["rheology", "thickening", "compressive", "lab mass"])
def test_measured_lab_edit_saved_before_review_blocks_then_exports_current_values(projects, change):
    state = projects()
    qc = state["lab_qc_params"]["Main"]
    if change == "rheology":
        qc["rheology"]["bhct_down"]["readings"]["300"] = "169"
    elif change == "thickening":
        qc["thickening_time"] = "03:45"
    elif change == "compressive":
        qc["compressive"]["crush"]["force_1"] = 450.
    else:
        grid = state["lab_grid_dfs"]["Main"]
        grid.at[0, "Mass"] = str(float(grid.at[0, "Mass"]) + 1.)
    assert qc["review_signature"] != eng.lab_review_signature(qc, state["lab_grid_dfs"]["Main"])
    app = checkpoint(state)
    blocked(app)
    review_main(app)
    data = rebuild(app)
    doc, ctx = assert_document(data, app.session_state.to_dict())
    assert eng.compute_phase_status(app.session_state)["phase7"]["level"] == "ok"
    assert semantics(data) != semantics(projects()["_compiled_doc_bytes"])
    assert ctx["slurries"][0]["lab"]["grid"] == app.session_state["lab_grid_dfs"]["Main"].to_dict("records")


@pytest.mark.parametrize("density,expected", [(None, 1.50), (1.47, 1.47)])
def test_sodium_edit_checkpoint_sync_resets_review_and_word_uses_current_additive(projects, density, expected):
    state = projects()
    old_mass = state["lab_grid_dfs"]["Main"].loc[lambda d: d.Material == "Sodium Silicate", "Mass"].iloc[0]
    row = state["cement_additives_dfs"]["Main"]
    row.at[0, "User Input"] = .75
    row.at[0, "Density"] = density
    app = checkpoint(state)
    blocked(app)
    CASE.phase(app, "phase5")
    assert app.session_state["cement_additives_dfs"]["Main"].at[0, "Density"] == expected
    CASE.phase(app, "phase7")
    next(b for b in app.button if b.key.startswith("_sync_btn_main_")).click().run()
    assert not app.session_state["lab_qc_params"]["Main"]["reviewed"]
    blocked(app)
    review_main(app)
    data = rebuild(app)
    doc, ctx = assert_document(data, app.session_state.to_dict())
    slurry = ctx["slurries"][0]
    assert "Sodium Silicate" not in [r["Name"] for r in slurry["blends"]]
    lab = next(r for r in slurry["lab_additives"] if r["Code"] == "Sodium Silicate")
    assert lab["Mass"] != old_mass
    rows = [" | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows]
    sodium = [r for r in rows if "Sodium Silicate" in r]
    assert len(sodium) == 2 and all("Liquid Extender" in r for r in sodium)
    assert any("0.75 gal/sk" in r for r in sodium)
    assert any(str(lab["Mass"]) in r and "gal/sk" in r for r in sodium)
    assert "Sodium Silicate" not in [r["name"] for r in ctx["fluids_train"]]


@pytest.mark.parametrize("invalid", ["target", "geometry", "density", "lab", "Formation"])
def test_export_invalid_valid_invalid_cannot_resurrect_old_docx(projects, invalid):
    app = CASE.app(projects())
    state = app.session_state
    original = deepcopy(state.to_dict())
    def corrupt():
        if invalid == "target":
            state["placement_config"]["target_row"] = "missing-target"
        elif invalid == "geometry":
            state["hardware_table"].loc[state["hardware_table"]["Description"] == "Casing", "ID (in)"] = 10.
        elif invalid == "density":
            state["fluids_config"]["params"]["Diesel"]["density"] = ""
        elif invalid == "lab":
            state["lab_qc_params"]["Main"]["rheology"]["bhct_down"]["selected"] = False
        else:
            domain = chart.current_open_hole_domain(state)
            state["formation_pressure_config"] = {"mode": "advanced", "profile": [
                {"md_m": domain["top_md"], "pore_psi": None, "frac_psi": None},
                {"md_m": domain["bottom_md"], "pore_psi": 4000., "frac_psi": 10000.}]}
    corrupt()
    blocked(app)
    for key in ("placement_config", "hardware_table", "fluids_config", "lab_qc_params", "formation_pressure_config"):
        if key in original:
            state[key] = deepcopy(original[key])
        elif key in state:
            del state[key]
    data = rebuild(app)
    assert_document(data, state.to_dict())
    corrupt()
    blocked(app)


@pytest.mark.parametrize("diesel,preflush,mud", [(True, False, "OBM"), (False, True, "OBM"), (False, False, "WBM")])
def test_independent_diesel_preflush_selection_never_leaks_inactive_drafts(projects, diesel, preflush, mud):
    state = projects()
    active = state["fluids_config"]["active"]
    if not diesel:
        active.remove("Diesel")
    if not preflush:
        active.remove("Pre Flush")
    state["mud_type"] = state["well_data"]["mud_type"] = mud
    app = checkpoint(state)
    data = rebuild(app)
    doc, ctx = assert_document(data, app.session_state.to_dict())
    assert ctx["has_preflush"] == preflush
    assert "Diesel" in app.session_state["lab_qc_params"]  # Preserved draft is not report authority.
    assert ("Diesel" in ctx["procedure_text"]) == diesel
    if not preflush:
        pumps = [body for _, body in numbered_steps(ctx["procedure_text"]) if body.startswith("Pump ")]
        assert all("preflush" not in line and ("pre flush" not in line or "Diesel" in line) for line in pumps)
    if diesel and preflush:
        names = [r["name"] for r in ctx["fluids_train"]]
        assert names.index("Diesel") + 1 == names.index("Pre Flush")


def test_manual_report_edits_xml_escape_review_and_regeneration_preserve_current_sources(projects):
    app = CASE.app(projects())
    CASE.phase(app, "phase10")
    summary = app.session_state["exec_summary_text"] + "\nOwner review: A & B < C > D."
    procedure = app.session_state["procedure_text"] + "\nOwner instruction: A & B < C > D."
    next(w for w in app.text_area if w.key.startswith("_exec_summary_editor_area_")).set_value(summary).run()
    next(w for w in app.text_area if w.key.startswith("_proc_editor_area_")).set_value(procedure).run()
    first = build(app)
    assert "Owner review: A & B < C > D." in "\n".join(semantics(first)[0])
    CASE.phase(app, "phase4")
    app.number_input(key="vol_Diesel").set_value(24.)
    CASE.phase(app, "phase1")
    blocked(app)
    assert app.session_state["exec_summary_text"] == summary
    assert app.session_state["procedure_text"] == procedure
    with patch.object(report.st, "session_state", app.session_state.to_dict()):
        specs = report.synchronize_report_texts()
    assert not specs["exec_summary_text"]["pending"]
    assert specs["procedure_text"]["pending"]
    next(b for b in app.button if b.key == "_regenerate_procedure_text").click().run()
    assert app.session_state["procedure_text"] != procedure
    assert app.session_state["report_text_state"]["procedure_text"]["history"][-1] == procedure
    current = rebuild(app)
    assert_document(current, app.session_state.to_dict())
    assert app.session_state["exec_summary_text"] == summary
    next(b for b in app.button if b.key == "_restore_procedure_text").click().run()
    blocked(app)
    assert app.session_state["procedure_text"] == procedure
    # The existing Keep action is the owner's explicit approval of manual text.
    next(b for b in app.button if b.key == "_keep_procedure_text").click().run()
    reviewed = build(app)
    assert "Owner instruction: A & B < C > D." in semantics(reviewed)[0]
    assert app.session_state["exec_summary_text"] == summary


def test_unused_diesel_surface_source_does_not_drift_current_word_or_engineering(projects):
    app = CASE.app(projects())
    old_doc = app.session_state["_compiled_doc_bytes"]
    old_signature = app.session_state["_compiled_doc_signature"]
    physics = signatures(app.session_state)
    app.session_state["lab_qc_params"]["Diesel"]["rheology"]["surface_down"]["readings"]["300"] = "169"
    CASE.phase(app, "phase1")
    CASE.phase(app, "phase10")
    assert signatures(app.session_state) == physics
    assert app.session_state["_compiled_doc_signature"] == old_signature
    assert app.session_state["_compiled_doc_bytes"] == old_doc
    assert semantics(build(app)) == semantics(old_doc)


@pytest.fixture(scope="module")
def multi_slurry(projects):
    app = CASE.app(projects())
    CASE.phase(app, "phase4")
    for name in ("Scavenger", "Lead", "Tail"):
        app.checkbox(key="chk_fluid_" + name).check().run()
        app.number_input(key="vol_" + name).set_value(50.).run()
        app.text_input(key="den_" + name).set_value("118").run()
        app.text_input(key="rate_" + name).set_value("4").run()
        app.session_state["cement_params"][name] = deepcopy(app.session_state["cement_params"]["Main"])
        app.session_state["lab_qc_params"][name] = deepcopy(app.session_state["lab_qc_params"]["Main"])
    CASE.phase(app, "phase5")
    CASE.phase(app, "phase7")
    for name in ("Scavenger", "Lead", "Tail"):
        sync = next((b for b in app.button if (b.key or "").startswith(f"_sync_btn_{name.lower()}_")), None)
        assert sync is not None, [error.value for error in app.error]
        sync.click().run()
        next(b for b in app.button if (b.key or "").startswith(f"_confirm_lab_btn_{name.lower()}_")).click().run()
    rebuild(app)
    return app.session_state.to_dict()


def test_multislurry_scavenger_and_legacy_spacers_render_once_in_their_own_orders(multi_slurry):
    state = deepcopy(multi_slurry)
    doc, ctx = assert_document(state["_compiled_doc_bytes"], state)
    assert [r["name"] for r in ctx["fluids_train"]] == ["Diesel", "Pre Flush", "Spacer Ahead", "Scavenger",
                                                     "Main", "Lead", "Tail", "Spacer Behind", "Displacement Fluid"]
    assert ctx["has_scavenger"]
    assert [r["name"] for r in ctx["slurries"]] == ["Lead", "Main", "Tail"]
    actual = [p.text for p in doc.paragraphs if p.text.endswith(" Cement Slurry Results")]
    assert actual == [f"{name} Cement Slurry Results" for name in ("Lead", "Main", "Tail")]
    steps = [body for _, body in numbered_steps(unescape(ctx["procedure_text"]))]
    selected = [line for line in steps if line.startswith(("Pump ", "Mix and pump "))]
    expected_names = [r["name"] for r in ctx["fluids_train"] if r["name"] not in ("Pre Flush", "Displacement Fluid")]
    indices = [next(i for i, line in enumerate(selected) if name in line) for name in expected_names]
    assert indices == sorted(indices) and len(indices) == len(set(indices))


def test_inactive_sodium_archive_survives_restore_without_ghost_word_content(multi_slurry):
    app = CASE.app(deepcopy(multi_slurry))
    CASE.phase(app, "phase4")
    app.checkbox(key="chk_fluid_Main").uncheck().run()
    app = checkpoint(app.session_state.to_dict())
    archived = app.session_state["inactive_slurry_drafts"]["Main"]["cement_additives_dfs"]
    assert archived.at[0, "Name"] == "Sodium Silicate"
    data = rebuild(app)
    doc, ctx = assert_document(data, app.session_state.to_dict())
    assert "Main" not in [r["name"] for r in ctx["slurries"]]
    assert "Sodium Silicate" not in "\n".join(node.text or "" for node in doc.element.body.iter(qn("w:t")))
    assert app.session_state["inactive_slurry_drafts"]["Main"]["cement_additives_dfs"].equals(archived)


@pytest.mark.parametrize("change", ["hardware", "placement", "Diesel", "Sodium", "Phase VII"])
def test_checkpoint_after_edit_before_any_rebuild_exports_new_canonical_truth(projects, change):
    state = projects()
    if change == "hardware":
        state["hardware_table"].loc[state["hardware_table"]["Description"] == "Casing", "MD (m)"] = "2900"
        select_target(state)
    elif change == "placement":
        state["cement_params"]["Main"].update(top_mode="Depth (m MD)", top_depth=2100., report_bottom_depth=2900.)
    elif change == "Diesel":
        state["fluids_config"]["params"]["Diesel"].update(density="54", volume=24.)
        state["lab_qc_params"]["Diesel"]["rheology"]["bhct_down"]["readings"]["300"] = "169"
    elif change == "Sodium":
        state["cement_additives_dfs"]["Main"].at[0, "User Input"] = .75
    else:
        state["lab_qc_params"]["Main"]["api_fl"] = .5
    app = checkpoint(state)
    blocked(app)  # Restored results cannot replace explicit Centralizer execution.
    if change in ("Sodium", "Phase VII"):
        review_main(app, sync=change == "Sodium")
    data = rebuild(app)
    doc, ctx = assert_document(data, app.session_state.to_dict())
    if change == "hardware":
        assert any(r["Description"] == "Casing" and r["MD (m)"] == "2900.0" for r in ctx["hardware"])
        frame = app.session_state["hardware_table"]
        assert frame.loc[frame.Description == "Casing", "MD (m)"].tolist() == ["2900"]
        hardware_table = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == "Tubular and Casing Hardware")
        assert any("Casing" in [c.text.strip() for c in row.cells] and
                   "2900.0" in [c.text.strip() for c in row.cells] for row in hardware_table.rows)
    elif change == "placement":
        assert ctx["slurries"][0]["top"] == "2100.0 m"
        assert ctx["slurries"][0]["bottom"] == "2900.0 m MD"
    elif change == "Diesel":
        assert "Pump 24.0 bbl Diesel as pre flush.(54 pcf)" in ctx["procedure_text"]
        assert app.session_state["lab_qc_params"]["Diesel"]["rheology"]["bhct_down"]["readings"]["300"] == "169"
    elif change == "Sodium":
        assert ctx["slurries"][0]["additives"][0]["(lbs or gal)/sk"] == "0.75 gal/sk"
    else:
        assert ctx["slurries"][0]["lab"]["api_fl_collected"] == .5
        assert ctx["slurries"][0]["lab"]["fluid_loss"] == 1.


def test_advanced_formation_and_incomplete_endpoint_cycle_replaces_real_word_chart(projects):
    app = CASE.app(projects())
    state = app.session_state
    domain = chart.current_open_hole_domain(state)
    state["formation_pressure_config"] = {"mode": "advanced", "profile": [
        {"md_m": domain["top_md"], "pore_psi": 2000., "frac_psi": 6000.},
        {"md_m": domain["bottom_md"], "pore_psi": 4000., "frac_psi": 10000.}]}
    assert phase9.result_is_current(state)
    first = build(app)
    assert_document(first, state.to_dict())
    state["hardware_table"].loc[state["hardware_table"]["Description"] == "Previous Casing", "MD (m)"] = "1600"
    draft = state.to_dict()
    ps.reconcile_formation_profile(draft)
    state["formation_pressure_config"] = draft["formation_pressure_config"]
    assert state["formation_pressure_config"]["profile"][0]["pore_psi"] is None
    blocked(app)
    app = checkpoint(state.to_dict())
    blocked(app)
    app.session_state["formation_pressure_config"]["profile"][0].update(pore_psi=2200., frac_psi=6500.)
    second = rebuild(app)
    doc, ctx = assert_document(second, app.session_state.to_dict())
    formation = ctx["cement_chart_report"]["model"]["formation_pressure"]
    assert formation["domain"]["top_md"] == 1600.
    assert formation["mode"] == "advanced"
    assert app.session_state["formation_pressure_config"]["profile"][0]["pore_psi"] == 2200.
    assert app.session_state["_compiled_doc_signature"] != projects()["_compiled_doc_signature"]


@pytest.mark.parametrize("dataset,field,value", [
    ("surface_down", "300", "169"), ("bhct_up", "300", "169"),
    ("surface_down", "gel_10_min", "14"), ("bhct_up", "selected", False),
])
def test_unused_diesel_widget_does_not_retire_a_current_real_document(projects, dataset, field, value):
    state = projects()
    if dataset == "bhct_up":
        state["lab_qc_params"]["Diesel"]["rheology"][dataset] = deepcopy(
            state["lab_qc_params"]["Diesel"]["rheology"]["surface_down"])
    app = CASE.app(state)
    current_bytes = app.session_state["_compiled_doc_bytes"]
    current_signature = app.session_state["_compiled_doc_signature"]
    before = signatures(app.session_state)
    CASE.phase(app, "phase7")
    widgets = app.checkbox if field == "selected" else app.text_input
    next(w for w in widgets if (w.key or "").startswith(f"_rheo_diesel_{dataset}_{field}_")).set_value(value).run()
    row = app.session_state["lab_qc_params"]["Diesel"]["rheology"][dataset]
    assert (row["readings"]["300"] if field == "300" else row[field]) == value
    assert eng.compute_phase_status(app.session_state)["phase7"]["level"] == "ok"
    assert signatures(app.session_state) == before
    CASE.phase(app, "phase1")
    CASE.phase(app, "phase10")
    ctx = context(app.session_state.to_dict())
    assert report.document_signature(ctx, str(report.TEMPLATE_PATH)) == current_signature
    assert "_compiled_doc_bytes" in app.session_state, "Unused Diesel dataset edit retired a current real DOCX"
    assert app.session_state["_compiled_doc_signature"] == current_signature
    assert app.session_state["_compiled_doc_bytes"] == current_bytes


def test_invalid_unused_diesel_dataset_still_blocks_recovered_export(projects):
    app = CASE.app(projects())
    for invalid in (True, False, True):
        CASE.phase(app, "phase7")
        next(w for w in app.text_input if w.key.startswith("_rheo_diesel_surface_down_300_")).set_value("bad" if invalid else "168").run()
        if invalid:
            blocked(app)
        else:
            assert phase9.result_is_current(app.session_state)
            assert_document(build(app), app.session_state.to_dict())


def test_reported_cement_surface_widget_still_retires_word_and_requires_review(projects):
    app = CASE.app(projects())
    CASE.phase(app, "phase7")
    next(w for w in app.text_input if w.key.startswith("_rheo_main_surface_down_300_")).set_value("169").run()
    no_document(app.session_state)
    blocked(app)
    assert phase9.result_is_current(app.session_state)
    review_main(app)
    data = build(app)
    _, ctx = assert_document(data, app.session_state.to_dict())
    assert ctx["slurries"][0]["lab"]["rheology"]["surface_down"]["readings"]["300"] == 169.


def test_consumed_diesel_bhct_gel_retires_word_without_false_centralizer_staleness(projects):
    app = CASE.app(projects())
    previous = pumping.source_signature(app.session_state)
    CASE.phase(app, "phase7")
    next(w for w in app.text_input if (w.key or "").startswith("_rheo_diesel_bhct_down_gel_10_min_")).set_value("14").run()
    no_document(app.session_state)
    assert phase9.result_is_current(app.session_state)
    assert pumping.source_signature(app.session_state) != previous
    data = build(app)
    assert_document(data, app.session_state.to_dict())
