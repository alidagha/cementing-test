"""Whole-psi Word presentation without changing measured/automatic state."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

import test_audit_regressions as audit
from engineering_tools import lab_review_signature
from phase_10_procedure import _word_quantity_context
from project_state import prepare_calculations


@pytest.mark.parametrize("raw,expected", [
    (5467.200000000001, "5467"), (5467.5, "5468"),
    ("5467.200000000001", "5467"), ("5467.5", "5468"),
    (7300, "7300"), ("7300+1000", "7300+1000"), ("", ""), ("-", "-"),
    ("NaN", "NaN"), ("Infinity", "Infinity"),
])
def test_terminal_pressure_copy_preserves_raw_context(raw, expected):
    context = {"slurries": [{"lab": {"bhsp": raw, "uca_pressure": raw}}]}
    before = deepcopy(context)
    lab = _word_quantity_context(context)["slurries"][0]["lab"]
    assert lab["bhsp"] == expected
    assert lab["uca_pressure"] == expected
    assert context == before


def test_inactive_uca_placeholder_stays_dash():
    context = {"slurries": [{"lab": {"bhsp": "5467.5", "uca_pressure": "-"}}]}
    lab = _word_quantity_context(context)["slurries"][0]["lab"]
    assert lab == {"bhsp": "5468", "uca_pressure": "-"}


def pressure_cells(doc):
    basic = next(r.cells[-1].text for t in doc.tables for r in t.rows
                 if len(r.cells) == 4 and r.cells[2].text == "BHSP")
    thick = next(t for t in doc.tables if t.rows[0].cells[0].text == "Thickening Time Test")
    uca = next(r.cells[3].text for t in doc.tables for r in t.rows
               if len(r.cells) >= 4 and r.cells[0].text == "UCA" and r.cells[2].text == "Pressure")
    return basic, thick.rows[3].cells[1].text, uca


@pytest.mark.parametrize("manual,expected", [
    (None, "5467"), ("5467.5", "5468"), ("7300+1000", "7300+1000"),
])
def test_normal_review_fresh_restore_direct_export_all_pressure_consumers(manual, expected):
    case = audit.AuditRegressions()
    app = case.configured_app('CSG 9 5/8"')
    if manual is not None:
        case.phase(app, "phase2_3")
        app.text_input(key="_w_bhsp").set_value(manual).run()
        case.phase(app, "phase7")
        next(b for b in app.button if b.label == "Sync with Phase V").click().run()
    case.phase(app, "phase7")
    next(w for w in app.selectbox if "Compressive Test" in w.label).set_value("UCA").run()
    next(b for b in app.button if b.label == "Confirm measured lab results").click().run()
    saved = audit.round_trip(app.session_state.to_dict())
    raw = saved["well_data"]["bhsp"]
    assert raw == (manual if manual is not None else "5467.200000000001")
    qc = deepcopy(saved["lab_qc_params"]["Main"])
    review = lab_review_signature(qc, saved["lab_grid_dfs"]["Main"])
    # Rebuild the payload without visiting Phase VII, through the real export gate.
    saved.pop("lab_payload_Main", None)
    assert prepare_calculations(saved) == []
    assert saved["lab_payload_Main"]["bhsp"] == raw
    restored = case.app(saved)
    case.export(restored)
    doc = Document(BytesIO(restored.session_state["_compiled_doc_bytes"]))
    label = "automatic" if manual is None else "half-up" if manual == "5467.5" else "manual"
    Path(f"/tmp/bhsp-{label}.docx").write_bytes(
        restored.session_state["_compiled_doc_bytes"])
    assert pressure_cells(doc) == (expected + " psi", expected, expected)
    assert restored.session_state["well_data"]["bhsp"] == raw
    assert restored.session_state["lab_payload_Main"]["bhsp"] == raw
    assert restored.session_state["lab_qc_params"]["Main"] == qc
    assert lab_review_signature(qc, restored.session_state["lab_grid_dfs"]["Main"]) == review
    assert restored.session_state["well_auto_fields"]["bhsp"] == (manual is None)
    # A normal well edit still invalidates the compiled document via existing gates.
    case.phase(restored, "phase2_3")
    restored.text_input(key="_w_bhsp").set_value("8000+500").run()
    case.phase(restored, "phase10")
    assert "_compiled_doc_bytes" not in restored.session_state
