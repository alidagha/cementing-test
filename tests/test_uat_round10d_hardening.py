"""Round 10D attacks cross-boundary freshness, failure and isolation lifecycles."""
from copy import deepcopy
from io import BytesIO
import json
import pickle
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from lxml import etree
import pytest

import phase_7_lab as phase7
import phase_9_centralizer as phase9
import phase_10_procedure as report
import pumping_simulation as pumping
from project_io import content_signature, decode_project, replace_project_state
from project_state import fingerprint
import test_audit_regressions as audit
from centralizer_fixtures import execute_centralizer
from test_uat_round10a_pumping_simulation import program_state
from test_uat_round9e_dynamic_solver import reviewed_spacer
from test_uat_round10b_pumping_word import ANCHOR, run_build
from test_uat_round10c_integration import inspect_word, save_project


@pytest.fixture(scope="module")
def built_state():
    case = audit.AuditRegressions()
    app = case.configured_app('CSG 9 5/8"', optional_fluids=True, include_standalone_spacer=False)
    case.phase(app, "phase4")
    for fluid in ("Pre Flush", "Spacer Ahead", "Spacer Behind"):
        app.number_input(key="vol_" + fluid).set_value(2.).run()
    execute_centralizer(app.session_state)
    app, calls = run_build(app.session_state.to_dict())
    assert len(calls) == 1 and not app.error
    inspect_word(app.session_state["_compiled_doc_bytes"])
    return app.session_state.to_dict()


@pytest.mark.parametrize("source", ["phase2", "phase4", "phase7", "phase9"])
def test_valid_source_edits_retire_download_without_implicit_simulation(built_state, source):
    state = deepcopy(built_state)
    old = pumping.source_signature(state)
    # These do not change Static Centralizer physics, but all affect Pumping.
    if source == "phase2":
        state["yield_point"] = state["well_data"]["yield_point"] = "2"
    elif source == "phase4":
        state["fluids_config"]["params"]["Displacement Fluid"]["pump_rate"] = "5-6"
    elif source == "phase7":
        state["_edit"] = "169"
        with patch.object(phase7.st, "session_state", state):
            phase7._commit_rheology("Main", "bhct_down", "readings", "_edit", rpm=300)
    else:
        # Spacer Behind is excluded from Static Centralizer, not from Pumping.
        assert phase9.result_is_current(state)
        old = pumping.source_signature(state)
        state["_edit"] = 169.
        with patch.object(phase9.st, "session_state", state):
            phase9._commit("aux", "_edit", fluid="Spacer Behind", rpm=300)
    assert phase9.result_is_current(state)
    assert pumping.source_signature(state) != old
    with patch.object(pumping.fcm, "build_three_charts", side_effect=AssertionError("Rerun executed solver")):
        app = audit.AuditRegressions().app(state)
        audit.AuditRegressions().phase(app, "phase10")
        app.run()
    assert not app.exception
    if source == "phase9":
        assert not app.error
        assert not next(b for b in app.button if b.label == "Build Word Document").disabled
    assert "_compiled_doc_bytes" not in app.session_state
    assert not any(button.proto.label == "Download Final Document (.docx)" for button in app.get("download_button"))


@pytest.mark.parametrize("fluid", pumping.fcm.SPACER_TYPES)
def test_auxiliary_activation_failure_repair_cycle_never_uses_stale_rheology(fluid):
    state = program_state(fluids={"Main": 2., "Displacement Fluid": 2.})
    initial = pumping.source_signature(state)
    with patch.object(pumping.fcm, "build_three_charts", side_effect=AssertionError("Preflight executed solver")):
        for _ in range(2):
            state["fluids_config"]["active"].append(fluid)
            state["fluids_config"]["params"][fluid] = {
                "volume": 1., "density": "80", "pump_rate": "4", "material_name": "WATER"}
            if fluid == "Spacer":
                reviewed_spacer(state)
                sources = state["lab_qc_params"]
                missing = "Phase VII Spacer.*bhct_down"
            else:
                sources = state["centralizer_config"]["aux_rheology"]
                missing = f"Phase IX {fluid}.*seven-RPM"
            readings = sources.pop(fluid)
            with pytest.raises(ValueError, match=missing):
                pumping.prepare_pumping_simulation(state)
            sources[fluid] = readings
            prepared = pumping.prepare_pumping_simulation(state)
            assert fluid in prepared["inputs"]["fluid_data"]
            assert pumping.source_signature(state) != initial
            state["fluids_config"]["active"].remove(fluid)
            # Inactive saved readings can be unfinished and must not participate.
            sources[fluid] = {}
            assert pumping.source_signature(state) == initial
            pumping.prepare_pumping_simulation(state)
            sources[fluid] = readings


def test_two_sessions_and_prepared_outputs_are_independent(built_state):
    case = audit.AuditRegressions()
    left = case.app(deepcopy(built_state))
    right = case.app(deepcopy(built_state))
    before = fingerprint(right.session_state.to_dict())
    left.session_state["fluids_config"]["params"]["Displacement Fluid"]["pump_rate"] = "7"
    with patch.object(pumping.fcm, "build_three_charts", side_effect=AssertionError("Preparation executed solver")):
        one = pumping.prepare_pumping_simulation(left.session_state)
        two = pumping.prepare_pumping_simulation(right.session_state)
    assert one["inputs"]["fluid_data"]["Displacement Fluid"]["min_rate"] == 7.
    assert two["inputs"]["fluid_data"]["Displacement Fluid"]["min_rate"] == 4.
    one["inputs"]["geometry"].pipe_sections[0].id_in = 1.
    one["inputs"]["fluid_data"]["Main"]["volume"] = 999.
    assert two["inputs"]["geometry"].pipe_sections[0].id_in != 1.
    assert two["inputs"]["fluid_data"]["Main"]["volume"] != 999.
    assert fingerprint(right.session_state.to_dict()) == before
    assert left.session_state["fluids_config"]["params"]["Main"]["volume"] != 999.


def test_load_cannot_resurrect_pumping_output_or_download_and_bad_load_is_atomic(built_state):
    app = audit.AuditRegressions().app(deepcopy(built_state))
    raw = save_project(app)
    incoming = json.loads(raw)
    incoming.update(_pumping_simulation_source="forged", _pumping_result={"fake": True},
                    _compiled_doc_bytes="forged DOCX", _compiled_doc_signature="forged")
    encoded = json.dumps(incoming, allow_nan=False).encode()
    state = deepcopy(built_state)
    decoded = decode_project(encoded, audit.namespace["deserialize_item"])
    replace_project_state(state, decoded, content_signature(encoded))
    assert not any(key.startswith(("_pumping", "_compiled_doc_", "_centralizer_")) for key in state)
    assert pumping.source_signature(state) == pumping.source_signature(built_state)
    assert not phase9.result_is_current(state)
    # A hostile load must not partially replace even a currently compiled project.
    current = deepcopy(built_state)
    before = pickle.dumps(current)
    incoming["centralizer_config"]["aux_rheology"]["Spacer"] = {"300": True}
    with pytest.raises(ValueError):
        replace_project_state(current, decode_project(json.dumps(incoming).encode(),
                              audit.namespace["deserialize_item"]), "rejected")
    assert pickle.dumps(current) == before


def test_solver_failure_after_prior_success_retires_document_and_leaves_no_owned_figure(built_state):
    before = pumping.fcm.plt.get_fignums()
    with patch.object(pumping.fcm, "simulate", side_effect=ValueError("deliberate simulation failure")) as solver:
        app, calls = run_build(built_state)
    assert solver.call_count == 1 and calls == []
    assert any("deliberate simulation failure" in error.value for error in app.error)
    assert not any(key.startswith("_compiled_doc_") for key in app.session_state.to_dict())
    assert pumping.fcm.plt.get_fignums() == before


def test_png_failure_closes_only_build_owned_figure_and_retires_old_download(built_state):
    from streamlit.runtime.scriptrunner import script_runner
    original_cleanup = script_runner._clean_problem_modules
    figures, observed = {}, []

    def after_execute(result):
        figures.update(owned=result["figure"], other=pumping.fcm.plt.figure())
        def fail_png(*args, **kwargs):
            raise OSError("deliberate PNG failure")
        result["figure"].savefig = fail_png

    def cleanup():
        if figures and not observed:
            active = pumping.fcm.plt.get_fignums()
            observed.append((figures["owned"].number in active, figures["other"].number in active))
        original_cleanup()

    try:
        with patch.object(script_runner, "_clean_problem_modules", side_effect=cleanup):
            app, calls = run_build(built_state, on_execute=after_execute)
        assert len(calls) == 1
        assert any("deliberate PNG failure" in error.value for error in app.error)
        assert not any(key.startswith("_compiled_doc_") for key in app.session_state.to_dict())
        assert observed == [(False, True)]
    finally:
        for figure in figures.values():
            pumping.fcm.plt.close(figure)


@pytest.mark.parametrize("job", ["CMT PLUG", "CMT SQUEEZE"])
def test_na_replacement_preserves_every_other_body_element_and_anchor_formatting(job):
    from docxtpl import DocxTemplate
    import well_report
    case = audit.AuditRegressions()
    app = case.configured_app(job)
    with patch.object(report.st, "session_state", app.session_state.to_dict()):
        context = report.build_master_context()
    template = DocxTemplate(report.TEMPLATE_PATH)
    template.render(report._word_quantity_context(context))
    rendered = BytesIO()
    template.save(rendered)
    args = (rendered.getvalue(), context["well_profile_report"], context["note_vertical"],
            context["cement_chart_report"], context["centralizer_report"])
    with patch.object(pumping.fcm, "build_three_charts", side_effect=AssertionError("N/A executed solver")):
        baseline_bytes = well_report.postprocess_well_report(*args)
        final_bytes = well_report.postprocess_well_report(*args, pumping_report={"applicable": False})
    baseline, final = Document(BytesIO(baseline_bytes)), Document(BytesIO(final_bytes))
    anchor = next(p for p in baseline.paragraphs if p.text == ANCHOR)
    replacement = next(p for p in final.paragraphs if p.text == "Pumping simulation charts: Not applicable.")
    old, new = list(baseline._element.body), list(final._element.body)
    index = old.index(anchor._p)
    assert len(old) == len(new)
    assert [i for i, (a, b) in enumerate(zip(old, new)) if etree.tostring(a) != etree.tostring(b)] == [index]
    assert anchor._p.pPr.xml == replacement._p.pPr.xml
    assert [r._r.rPr.xml if r._r.rPr is not None else None for r in anchor.runs] == [
        r._r.rPr.xml if r._r.rPr is not None else None for r in replacement.runs]
    assert len(final._element.xpath('.//wp:docPr[@descr="Pumping Simulation Charts"]')) == 0
    assert baseline.sections[-1]._sectPr.xml == final.sections[-1]._sectPr.xml
    with ZipFile(BytesIO(baseline_bytes)) as a, ZipFile(BytesIO(final_bytes)) as b:
        assert a.namelist() == b.namelist()
        assert all(a.read(name) == b.read(name) for name in a.namelist() if name != "word/document.xml")
