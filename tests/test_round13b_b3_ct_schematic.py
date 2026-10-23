"""Truthful CT schematic presentation reuses the frozen Plug renderer."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import hashlib
import importlib.util
import json
import pickle
import subprocess
from unittest.mock import patch
from zipfile import ZipFile

import pandas as pd
import pytest
from PIL import Image
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

import materials_db as db
import placement
import project_state as ps
import engineering_tools as eng
import well_schematic_adapter as schematic
import well_schematic_engine as engine
import well_report
import phase_10_procedure as report
import test_audit_regressions as audit
from test_round13b_b2_ct_stages import state as stage_state
from test_round13b_b1_ct_foundation import hardware as ct_hardware
from phase_2_3_well_data import HARDWARE_COLUMNS

PLACEHOLDER = 'Well Schematic Not Available – Incomplete Well Geometry'
CASE = audit.AuditRegressions()


def row(description, md=None, size=None, inside=None):
    item = dict.fromkeys(HARDWARE_COLUMNS, None)
    item.update({'Description':description,'MD (m)':md,'Size (in)':size,'ID (in)':inside})
    return item


def state(kind='complete', bottom=3527., top=3470.):
    s = stage_state(20,2)
    history = [row('Previous Casing','3000','13 3/8',12.415),
               row('Previous Liner','2500-3800','7',6.276)]
    if kind == 'casing': history = [row('Previous Casing','3800','7',6.276)]
    if kind == 'hole': history = [row('Previous Casing','2500','13 3/8',12.415),
                                 row('Open Hole Size','2500-4000','12 1/4',12.25)]
    if kind == 'partial': history = [row('Previous Casing'),row('Previous Liner',size='7'),row('Previous Liner',size='5')]
    if kind == 'none': history = []
    s.update(well_name='B3 engineering UAT',hardware_table=pd.DataFrame(history+[ct_hardware(str(bottom)).iloc[0].to_dict()],columns=HARDWARE_COLUMNS))
    s['cement_params']['Main'].update(top_depth=top,top_job_type=db.CT_JOB,top_mode='Depth (m MD)')
    ps.refresh_well_derived(s)
    return s


def legacy_reference(s):
    """Test-only ordinary Plug operation using the same measured interval."""
    out = deepcopy(s)
    out['hardware_table'] = out['hardware_table'][out['hardware_table']['Description'] != 'CT'].copy()
    out['job_type'] = 'CMT PLUG'
    out['placement_config'] = {'job_type':'CMT PLUG','target_row':'__manual__',
                               'manual_depth_m':placement.ct_setting_depth(s['hardware_table'])}
    out['cement_params']['Main']['top_job_type'] = 'CMT PLUG'
    return out


@pytest.mark.parametrize('kind',['complete','casing','hole'])
def test_complete_geometry_is_exact_plug_payload_and_real_png(kind,tmp_path):
    s = state(kind); before=pickle.dumps(s)
    result = schematic.prepare_schematic(s)
    assert result['available'] is True and result['job_type'] == db.CT_JOB
    assert not result['missing_geometry']
    payload=result['payload']
    assert payload == schematic.prepare_schematic(legacy_reference(s))
    assert payload['cement_plugs'] == [{'top_md':3470.,'bottom_md':3527.}]
    assert payload['squeezes'] == [] and payload['current_job']['type'] == 'cmt_plug'
    assert set(payload) <= {'well_name','show_mid_depth','casings','open_holes','cement_plugs','squeezes','current_job','final_relevant_depth_m'}
    blob=well_report.render_schematic_png(result)
    image=Image.open(BytesIO(blob)); assert image.format == 'PNG' and min(image.size)>1000
    (tmp_path/f'ct-{kind}.png').write_bytes(blob)
    assert blob == well_report.render_schematic_png(payload)
    assert pickle.dumps(s) == before


@pytest.mark.parametrize('field',['MD (m)','Size (in)','ID (in)'])
def test_missing_history_is_valid_unavailable_without_fabricated_payload(field):
    s=state(); s['hardware_table'].loc[0,field]=None; before=pickle.dumps(s)
    result=schematic.prepare_schematic(s)
    assert result['available'] is False and result['message'] == PLACEHOLDER
    assert result['payload'] is None
    assert any(field in reason for reason in result['missing_geometry'])
    assert eng.compute_phase_status(s)['phase2_3']['level'] == 'ok'
    assert PLACEHOLDER in eng.compute_phase_status(s)['phase2_3']['message']
    assert pickle.dumps(s)==before


@pytest.mark.parametrize('kind',['none','partial'])
def test_unknown_rows_remain_separate_and_restore_deterministically(kind):
    s=state(kind); table=s['hardware_table'].copy(deep=True)
    result=schematic.prepare_schematic(s)
    assert not result['available'] and result['message']==PLACEHOLDER
    restored=audit.round_trip(s)
    assert schematic.prepare_schematic(restored)==result
    assert schematic.source_signature(restored)==schematic.source_signature(s)
    pd.testing.assert_frame_equal(restored['hardware_table'],table,check_dtype=False)
    assert len(restored['hardware_table'])==len(table)
    assert not any(k.startswith('_well_schematic') for k in restored)


@pytest.mark.parametrize('field,value,match',[
    ('ID (in)',14.,'ID'),('ID (in)',-1.,'ID'),('ID (in)',float('inf'),'finite'),
    ('Size (in)','bad','Size'),('Size (in)','-7','Size'),
    ('MD (m)','3100-3000','interval'),('MD (m)','-5','MD'),('MD (m)','4100','well'),
])
def test_known_invalid_geometry_is_not_hidden_by_another_missing_row(field,value,match):
    s=state(); s['hardware_table'].loc[0,field]=value
    s['hardware_table']=pd.DataFrame([*s['hardware_table'].to_dict('records'),row('Previous Casing')],columns=HARDWARE_COLUMNS)
    with pytest.raises(ValueError,match=match): schematic.prepare_schematic(s)
    assert eng.compute_phase_status(s)['phase2_3']['level']!='ok'


def test_known_casing_clearance_fails_even_if_other_history_is_missing():
    s=state('casing'); outer=row('Previous Casing',None,'9 5/8',6.5)
    s['hardware_table']=pd.concat([pd.DataFrame([outer]),s['hardware_table']],ignore_index=True)
    with pytest.raises(ValueError,match='fit|pass|ID'): schematic.prepare_schematic(s)


def test_missing_lap_is_advisory_but_supplied_lap_conflict_is_invalid():
    s=state();s['hardware_table'].loc[1,'MD (m)']='3800'
    assert schematic.prepare_schematic(s)['message']==PLACEHOLDER
    s['hardware_table'].loc[1,'MD (m)']='2500-3800'
    s['hardware_table']=pd.DataFrame([*s['hardware_table'].to_dict('records'),row('Previous Liner Hanger','2600','7')],columns=HARDWARE_COLUMNS)
    with pytest.raises(ValueError,match='contradict|match'):schematic.prepare_schematic(s)


def test_missing_bore_never_consumes_legacy_manual_diameter():
    s=state('casing');s['hardware_table'].loc[0,'MD (m)']='3400'
    s['schematic_config']={'version':1,'plug_bore_sections':[{'top_md':3400.,'bottom_md':3527.,'diameter_in':8.5}]}
    result=schematic.prepare_schematic(s)
    assert not result['available'] and result['message']==PLACEHOLDER
    assert result['payload'] is None
    assert s['schematic_config']['plug_bore_sections'][0]['diameter_in']==8.5


@pytest.mark.parametrize('fault',['missing','duplicate','negative','interval','nonfinite','outside_well','top_missing','top_bottom','top_deeper','stale_top','no_main','extra_slurry'])
def test_operation_failures_cannot_become_missing_geometry_placeholder(fault):
    s=state('partial');table=s['hardware_table'];idx=table.index[table['Description']=='CT'][0]
    if fault=='missing':s['hardware_table']=table.drop(idx)
    elif fault=='duplicate':s['hardware_table']=pd.concat([table,table.loc[[idx]]],ignore_index=True)
    elif fault in ('negative','interval','nonfinite','outside_well'):
        table.loc[idx,'MD (m)']={'negative':-1,'interval':'3500-3527','nonfinite':float('inf'),'outside_well':4001}[fault]
    elif fault=='no_main':s['fluids_config']['active'].remove('Main')
    elif fault=='extra_slurry':s['fluids_config']['active'].append('Tail')
    else:
        s['cement_params']['Main']['top_depth']={'top_missing':None,'top_bottom':3527.,'top_deeper':3528.,'stale_top':3470.}[fault]
        if fault=='stale_top':s['cement_params']['Main']['top_job_type']='CMT PLUG'
    with pytest.raises(ValueError):schematic.prepare_schematic(s)
    assert eng.compute_phase_status(s)['phase2_3']['level']!='ok'


@pytest.mark.parametrize('bottom,top',[(3527,3470),(2750.25,2700.125)])
def test_dynamic_depth_authority_never_uses_manual_target_or_historical_shoe(bottom,top):
    s=state('casing',bottom,top)
    s['placement_config']={'job_type':db.CT_JOB,'target_row':'__manual__','manual_depth_m':123.}
    assert schematic.prepare_schematic(s)['payload']['cement_plugs']==[{'top_md':top,'bottom_md':bottom}]
    assert placement.ct_setting_depth(s['hardware_table'])==bottom


def test_explicit_surface_and_profile_requirements_remain():
    s=state('casing');s['cement_params']['Main'].update(top_mode='Surface',top_depth=0.)
    assert schematic.prepare_schematic(s)['payload']['cement_plugs']==[{'top_md':0.,'bottom_md':3527.}]
    s['geothermal_config']['value']=None
    assert eng.compute_phase_status(s)['phase2_3']['level']!='ok'


@pytest.mark.parametrize('change',['ct_md','top','row_count','row_order','size','id','well_name','well_bound'])
def test_consumed_sources_invalidate_image_signature_and_documents(change):
    s=state();ps.prepare_calculations(s);original=schematic.source_signature(s)
    s['_compiled_doc_bytes']=b'old'
    if change=='ct_md':s['hardware_table'].loc[2,'MD (m)']=3550.
    elif change=='top':s['cement_params']['Main']['top_depth']=3460.
    elif change=='row_count':s['hardware_table']=pd.DataFrame([*s['hardware_table'].to_dict('records'),row('Previous Casing')],columns=HARDWARE_COLUMNS)
    elif change=='row_order':s['hardware_table']=s['hardware_table'].iloc[::-1].reset_index(drop=True)
    elif change=='size':s['hardware_table'].loc[0,'Size (in)']='13.5'
    elif change=='id':s['hardware_table'].loc[0,'ID (in)']=12.4
    elif change=='well_name':s['well_name']='Changed & well'
    else:s['well_geometry']['td_m']=4100.
    assert schematic.source_signature(s)!=original
    issues=ps.prepare_calculations(s)
    assert '_compiled_doc_bytes' not in s
    assert any('Phase VII' in issue for issue in issues)
    assert s['_well_schematic_source']==schematic.source_signature(s)


def test_stage_choices_and_cement_quantities_are_not_schematic_sources():
    s=state();from ct_stages import ct_stage_projection
    stage=ct_stage_projection(s);before=pickle.dumps(s['cement_params']);sig=schematic.source_signature(s);result=schematic.prepare_schematic(s)
    s['ct_stage_config']['count']=3
    assert schematic.source_signature(s)==sig and schematic.prepare_schematic(s)==result
    assert ct_stage_projection(s)['cement_bbl']!=stage['cement_bbl']
    assert pickle.dumps(s['cement_params'])==before


@pytest.mark.parametrize('kind',['partial','complete'])
def test_phase2_navigation_preserves_raw_hardware_and_keeps_export_blocked(kind):
    import well_report as output
    s=state(kind);original=s['hardware_table'].copy(deep=True)
    app=CASE.app(s);CASE.phase(app,'phase2_3')
    pd.testing.assert_frame_equal(app.session_state['hardware_table'],original,check_dtype=False)
    assert eng.compute_phase_status(app.session_state)['phase2_3']['level']=='ok'
    if kind=='partial':assert any(PLACEHOLDER in x.value for x in app.info)
    else:
        button=next(b for b in app.button if b.label=='Preview Well Schematic')
        button.click().run();CASE.healthy(app)
    CASE.phase(app,'phase5');CASE.phase(app,'phase2_3')
    restored=CASE.app(audit.round_trip(app.session_state.to_dict()));CASE.phase(restored,'phase2_3')
    pd.testing.assert_frame_equal(restored.session_state['hardware_table'],original,check_dtype=False)
    CASE.phase(restored,'phase10')
    assert not any(b.label=='Build Word Document' for b in restored.button)
    assert any('Phase VII' in w.value for w in restored.warning)


@pytest.mark.parametrize('prepared',[False,True])
def test_direct_ct_word_gate_is_unchanged(prepared):
    s=state();s['_compiled_doc_bytes']=b'old'
    with patch.object(report.st,'session_state',s),pytest.raises(ValueError,match='Phase VII'):
        report.build_master_context(calculations_prepared=prepared)
    assert '_compiled_doc_bytes' not in s


def test_unknown_hanger_size_cannot_assign_a_measured_liner_relationship():
    s=state()
    s['hardware_table']=pd.DataFrame([*s['hardware_table'].to_dict('records'),row('Previous Liner Hanger','2500')],columns=HARDWARE_COLUMNS)
    result=schematic.prepare_schematic(s)
    assert not result['available'] and result['message']==PLACEHOLDER
    assert any('Hanger' in x and 'Size (in)' in x for x in result['missing_geometry'])


def test_known_later_section_clearance_is_validated_before_missing_fallback():
    s=state()
    s['hardware_table']=pd.DataFrame([
        row('Previous Casing','2000','13 3/8',12.415),
        row('Previous Casing','3000','13 3/8',6.5),
        row('Previous Liner','2500-3800','7',6.276),
        row('Previous Casing'),*s['hardware_table'].tail(1).to_dict('records')],columns=HARDWARE_COLUMNS)
    with pytest.raises(ValueError,match='fit|ID'):
        schematic.prepare_schematic(s)


@pytest.fixture(scope='module')
def approved_adapter(tmp_path_factory):
    """Execute the exact approved parent adapter against identical legacy inputs."""
    import sys
    path=tmp_path_factory.mktemp('approved-adapter')/'adapter.py'
    path.write_bytes(subprocess.check_output(['git','show','184cc888552f54df6a1d02fd0bdca16b77e358a4:well_schematic_adapter.py']))
    spec=importlib.util.spec_from_file_location('round13b_b3_parent_adapter',path)
    module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module)
    return module


@pytest.mark.parametrize('job',[job for job in db.JOB_TYPES if job!=db.CT_JOB])
def test_all_fourteen_original_payloads_signatures_and_pngs_equal_parent(job,approved_adapter,tmp_path):
    s={'job_type':job,**audit.valid_export_chart_fixture(job)}
    s['fluids_config']={'active':['Main','Displacement Fluid']}
    s['cement_params']={'Main':{'top_mode':'Depth (m MD)','top_depth':2000.,'top_job_type':job}}
    before=pickle.dumps(s)
    expected=approved_adapter.prepare_schematic(s);actual=schematic.prepare_schematic(s)
    assert actual==expected
    assert schematic.source_signature(s)==approved_adapter.source_signature(s)
    engine.render_well_schematic(expected,str(tmp_path/'before.png'),dpi=90,strict=True)
    engine.render_well_schematic(actual,str(tmp_path/'after.png'),dpi=90,strict=True)
    assert (tmp_path/'before.png').read_bytes()==(tmp_path/'after.png').read_bytes()
    assert pickle.dumps(s)==before


@pytest.fixture(scope='module')
def ordinary_render_context():
    # Controlled anchor harness uses an ordinary, normally admitted report.
    # No CT export gate is bypassed and this is not a final CT document.
    from test_uat_round7b_report import rendered
    app=CASE.configured_app('CMT PLUG')
    with patch.object(report.st,'session_state',app.session_state.to_dict()):
        context=report.build_master_context()
    return rendered(context),context['well_profile_report'],context['note_vertical']


@pytest.mark.parametrize('kind',['complete','partial'])
def test_section_ii_projection_changes_only_existing_anchor(ordinary_render_context,kind,tmp_path):
    args=ordinary_render_context
    projection=schematic.prepare_schematic(state(kind))
    baseline=well_report.postprocess_well_report(*args)
    result=well_report.postprocess_well_report(*args,schematic_report=projection)
    before,after=Document(BytesIO(baseline)),Document(BytesIO(result))
    old,new=list(before._element.body),list(after._element.body)
    assert len(old)==len(new)
    expected=next(i for i,e in enumerate(old) if ''.join(t.text or '' for t in e.iter(qn('w:t')))==well_report.SCHEMATIC_ANCHOR)
    assert [i for i,(a,b) in enumerate(zip(old,new)) if etree.tostring(a)!=etree.tostring(b)]==[expected]
    oldp=next(p for p in before.paragraphs if p.text==well_report.SCHEMATIC_ANCHOR)
    if kind=='partial':
        newp=next(p for p in after.paragraphs if p.text==PLACEHOLDER)
        assert not newp._p.xpath('.//pic:pic')
    else:
        from test_uat_round11a_schematic import schematic_picture
        newp,png=schematic_picture(after)
        assert png==well_report.render_schematic_png(projection)
        (tmp_path/'ct-anchor.png').write_bytes(png)
    assert oldp._p.pPr.xml==newp._p.pPr.xml
    assert oldp.runs[0]._r.rPr.xml==newp.runs[0]._r.rPr.xml
    body=list(after._element.body)
    assert body.index(next(p._p for p in after.paragraphs if p.text.strip()=='II. Well Data Summary'))<body.index(newp._p)<body.index(next(p._p for p in after.paragraphs if p.text.strip()=='III. Geothermal Temperature Profile'))
    with ZipFile(BytesIO(baseline)) as a,ZipFile(BytesIO(result)) as b:
        added=set(b.namelist())-set(a.namelist())
        assert not set(a.namelist())-set(b.namelist())
        if kind=='partial':
            assert not added
            allowed={'word/document.xml'}
        else:
            assert len(added)==1 and next(iter(added)).startswith('word/media/')
            allowed={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
        assert all(a.read(name)==b.read(name) for name in set(a.namelist())-allowed)
    (tmp_path/f'controlled-anchor-{kind}.docx').write_bytes(result)


def test_incomplete_schematic_is_advisory_but_procedure_gate_remains():
    s=state('partial')
    issues=ps.prepare_calculations(s)
    assert issues==['Main: open Phase VII to enter and review the internal Lab/Rheology data.',
                    'Phase VII: Lab QC or formulation rows incomplete for: Main.']
    assert s['_well_schematic_source']==schematic.source_signature(s)
    assert eng.compute_phase_status(s)['phase2_3']['level']=='ok'
    assert not any(isinstance(v,bytes) for k,v in s.items() if k.startswith('_well_schematic'))


def test_frozen_renderer_template_and_previous_ct_authorities_match_parent():
    for name in ['well_schematic_engine.py','placement.py','ct_stages.py','phase_5_cement.py','standalone_spacer.py']:
        reference='fdcfeb14a6246e59261c89775635ca9574674d42' if name=='phase_5_cement.py' else '184cc888552f54df6a1d02fd0bdca16b77e358a4'
        assert Path(name).read_bytes()==subprocess.check_output(['git','show',f'{reference}:{name}'])
    from test_round13b_post_b4_word_quality import test_template_changes_only_approved_sample_text_payloads
    test_template_changes_only_approved_sample_text_payloads()


def test_conflicting_measured_open_hole_sections_cannot_hide_behind_missing_history():
    s=state('hole')
    records=s['hardware_table'].to_dict('records')
    records[1]['MD (m)']='2500-3800'
    records.insert(2,row('Open Hole Size','3800-4000','8 1/2',8.5))
    records.insert(0,row('Previous Casing'))
    s['hardware_table']=pd.DataFrame(records,columns=HARDWARE_COLUMNS)
    with pytest.raises(ValueError,match='Open Hole.*diameter|Open Hole.*Size'):
        schematic.prepare_schematic(s)


@pytest.mark.parametrize('grade',[None,'N-80'])
def test_explicit_retained_ct_edits_survive_navigation_and_real_project_restore(grade):
    from test_uat_round10c_integration import save_project,restore_project
    app=CASE.app(state('none'));CASE.phase(app,'phase1')
    app.selectbox(key='_w_job_type').set_value('CMT PLUG').run()
    CASE.phase(app,'phase2_3')
    edits={'MD (m)':'3550','Size (in)':'2 3/8','ID (in)':1.9,
           'Joint (m)':0.,'Weight (ppf)':2.,'Grade':grade,
           'Collapse (psi)':0.,'Burst (psi)':0.}
    CASE.hardware_event(app,{'edited_rows':{'0':edits}},navigate='phase5')
    assert ps.hardware_draft_pending(app.session_state)
    CASE.phase(app,'phase2_3')
    expected=app.session_state['hardware_table'].copy(deep=True)
    for key,value in edits.items():
        assert expected.iloc[0][key]==value
    CASE.phase(app,'phase5')
    raw=save_project(app)
    assert b'_well_schematic' not in raw and b'PNG' not in raw
    restored=restore_project(raw);CASE.phase(restored,'phase2_3')
    pd.testing.assert_frame_equal(restored.session_state['hardware_table'],expected,check_dtype=False)
    assert restored.session_state['well_name']=='B3 engineering UAT'
    CASE.phase(restored,'phase1')
    restored.selectbox(key='_w_job_type').set_value(db.CT_JOB).run()
    CASE.phase(restored,'phase2_3')
    pd.testing.assert_frame_equal(restored.session_state['hardware_table'],expected,check_dtype=False)
    assert placement.ct_setting_depth(restored.session_state['hardware_table'])==3550.
