"""F1: truthful nullable metadata and job-owned excess widget lifetimes."""
from copy import deepcopy
from datetime import date
from io import BytesIO
from html import unescape
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docxtpl import DocxTemplate
import pytest

import materials_db as db
import engineering_tools as eng
import phase_1_doc_control as p1
import phase_2_3_well_data as p2
import phase_4_fluids as p4
import phase_10_procedure as report
import project_state as ps
from placement import EXCESS_FIELDS
import test_audit_regressions as audit

CASE = audit.AuditRegressions()
DOC_KEYS = ('_compiled_doc_bytes', '_compiled_doc_filename', '_compiled_doc_signature')


@pytest.fixture(scope='module')
def reviewed():
    app = CASE.configured_app('CMT PLUG')
    state = deepcopy(app.session_state.to_dict())
    assert eng.compute_phase_status(state)['phase7']['level'] == 'ok'
    assert ps.prepare_calculations(state) == []
    return state


def context(state, prepared):
    with patch.object(report.st, 'session_state', state):
        return report.build_master_context(calculations_prepared=prepared)


def stale(state):
    for key, value in zip(DOC_KEYS, (b'old word', 'old.docx', 'old signature')):
        state[key] = value


@pytest.mark.parametrize('fluid', ['Main', 'Displacement Fluid'])
def test_m4_null_mount_is_blank_without_canonical_normalization(reviewed, fluid):
    state = deepcopy(reviewed)
    state['fluids_config']['params'][fluid]['material_name'] = None
    app = CASE.app(audit.round_trip(state))
    CASE.phase(app, 'phase4')
    assert app.text_input(key='matname_'+fluid).value == ''
    assert app.session_state['fluids_config']['params'][fluid]['material_name'] is None
    assert app.session_state['fluid_data'][fluid]['material_name'] == ''
    CASE.phase(app, 'phase1')
    restored = audit.round_trip(app.session_state.to_dict())
    assert restored['fluids_config']['params'][fluid]['material_name'] is None
    assert restored['fluid_data'][fluid]['material_name'] == ''


@pytest.mark.parametrize('value', [None, '', 'None', 'Field & <wash>'])
def test_m4_record_factory_does_not_stringify_null(value):
    row = p4.build_fluid_record('Main', value, 50., '118', 118., '4', 4., 12.5, '00:12', 12.5, '00:12')
    assert row['material_name'] == ('' if value is None else value)
    assert (row['volume'], row['effective_density'], row['duration_min']) == (50., 118., 12.5)


@pytest.mark.parametrize('value', [None, '', 'None', 'Field & <wash>'])
def test_m4_canonical_to_derived_projection_preserves_source_and_quantities(reviewed, value):
    state = deepcopy(reviewed)
    state['fluids_config']['params']['Main']['material_name'] = value
    original = deepcopy(state['fluids_config'])
    old = deepcopy(state['fluid_data']['Main'])
    ps.refresh_fluids(state)
    assert state['fluids_config'] == original
    assert state['fluid_data']['Main']['material_name'] == ('' if value is None else value)
    assert {k:v for k,v in old.items() if k!='material_name'} == {
        k:v for k,v in state['fluid_data']['Main'].items() if k!='material_name'}


@pytest.mark.parametrize('canonical,cached', [(None, 'None'), (None, 'Cement Slurry'),
                                            ('Current name', 'Old name')])
def test_m4_prepared_stale_identity_rejects_and_retires_word(reviewed, canonical, cached):
    state = deepcopy(reviewed)
    state['fluids_config']['params']['Main']['material_name'] = canonical
    state['fluid_data']['Main']['material_name'] = cached
    before = deepcopy(state['fluids_config'])
    stale(state)
    with pytest.raises(ValueError, match='Phase IV.*Main.*(?:identity|name)'):
        context(state, True)
    assert not any(k in state for k in DOC_KEYS)
    assert state['fluids_config'] == before
    assert state['fluid_data']['Main']['material_name'] == cached


@pytest.mark.parametrize('prepared', [False, True])
@pytest.mark.parametrize('value', [None, '', 'None', 'Field & <wash>'])
def test_m4_current_nullable_and_literal_names_have_truthful_context(reviewed, prepared, value):
    state = deepcopy(reviewed)
    state['fluids_config']['params']['Main']['material_name'] = value
    ps.refresh_fluids(state)
    ctx = context(state, prepared)
    assert unescape(next(r for r in ctx['fluids_train'] if r['name']=='Main')['material_name']) == ('' if value is None else value)
    assert state['fluids_config']['params']['Main']['material_name'] == value


@pytest.mark.parametrize('prepared', [False, True])
def test_m4_legacy_absent_params_uses_its_genuine_record_name(reviewed, prepared):
    state = deepcopy(reviewed)
    state['fluids_config']['params'].pop('Main')
    state['fluid_data']['Main']['material_name'] = 'Legacy field name'
    assert next(r for r in context(state, prepared)['fluids_train'] if r['name']=='Main')['material_name'] == 'Legacy field name'


@pytest.mark.parametrize('fluid', ['Pre Flush', 'Spacer Ahead', 'Spacer Behind'])
@pytest.mark.parametrize('prepared', [False, True])
def test_m4_required_null_name_cannot_use_an_old_valid_record(reviewed, fluid, prepared):
    state = deepcopy(reviewed)
    state['fluids_config']['active'].insert(0, fluid)
    state['fluids_config']['params'][fluid] = {'volume': 25., 'density': '80', 'pump_rate': '4', 'material_name': None}
    state['fluid_data'][fluid] = dict(state['fluid_data']['Main'], material_name='Old valid name', name=fluid)
    stale(state)
    with pytest.raises(ValueError, match='Phase IV.*'+fluid):
        context(state, prepared)
    assert not any(k in state for k in DOC_KEYS)


def test_m4_f3_volume_error_precedes_a_simultaneous_name_error(reviewed):
    state = deepcopy(reviewed)
    state['fluids_config']['params']['Main']['material_name'] = None
    state['fluid_data']['Main']['volume'] = 'abc'
    stale(state)
    with pytest.raises(ValueError, match='Unreadable numeric volume.*Main.*volume'):
        context(state, True)
    assert not any(k in state for k in DOC_KEYS)


@pytest.mark.parametrize('field,label,_', EXCESS_FIELDS)
@pytest.mark.parametrize('value', [None, 0., 25.])
@pytest.mark.parametrize('other_job', ['CMT PLUG', 'LNR 7"'])
def test_m5_phase1_round_trip_jobs_clears_both_excess_without_phase2(field, label, _, value, other_job):
    app = CASE.app({'job_type':'CSG 9 5/8"'})
    CASE.phase(app, 'phase2_3')
    widget = next(w for w in app.number_input if w.label==label)
    old_key = widget.key
    widget.set_value(25.).run()
    next(w for w in app.number_input if w.label==label).set_value(value)
    CASE.phase(app, 'phase1')
    assert app.session_state['placement_config'][field] == value
    app.selectbox(key='_w_job_type').set_value(other_job).run()
    app.selectbox(key='_w_job_type').set_value('CSG 9 5/8"').run()
    assert app.session_state['placement_config'].get(field) is None
    CASE.phase(app, 'phase2_3')
    assert next(w for w in app.number_input if w.label==label).value is None
    assert all(app.session_state['placement_config'].get(f) is None for f,_,_ in EXCESS_FIELDS)
    # Exercise the real registered callback with its retired key/late value.
    state = deepcopy(app.session_state.to_dict())
    before = deepcopy(state['placement_config'])
    state[old_key] = 25.
    with patch.object(p2.st, 'session_state', state):
        p2._commit_placement(field, old_key)
    assert state['placement_config'] == before
    restored = audit.round_trip(app.session_state.to_dict())
    assert all(restored['placement_config'].get(f) is None for f,_,_ in EXCESS_FIELDS)
    assert not any(k.startswith('_placement_') for k in restored)


@pytest.mark.parametrize('field,label,_', EXCESS_FIELDS)
@pytest.mark.parametrize('value', [None, 0., 25.])
def test_m5_same_job_navigation_and_json_preserve_explicit_values(field, label, _, value):
    app = CASE.app({'job_type':'CSG 9 5/8"'})
    CASE.phase(app, 'phase2_3')
    next(w for w in app.number_input if w.label==label).set_value(value)
    CASE.phase(app, 'phase1')
    app.selectbox(key='_w_job_type').set_value('CSG 9 5/8"').run()
    CASE.phase(app, 'phase2_3')
    assert next(w for w in app.number_input if w.label==label).value == value
    restored = CASE.app(audit.round_trip(app.session_state.to_dict()))
    CASE.phase(restored, 'phase2_3')
    assert restored.session_state['placement_config'][field] == value
    assert next(w for w in restored.number_input if w.label==label).value == value


@pytest.mark.parametrize('other_job', ['CMT PLUG', 'LNR 7"'])
def test_m5_visited_intermediate_job_cannot_reuse_old_excess(other_job):
    app = CASE.app({'job_type':'CSG 9 5/8"'})
    CASE.phase(app, 'phase2_3')
    for _,label,_ in EXCESS_FIELDS:
        next(w for w in app.number_input if w.label==label).set_value(25.).run()
    CASE.phase(app, 'phase1')
    app.selectbox(key='_w_job_type').set_value(other_job).run()
    CASE.phase(app, 'phase2_3')
    assert all(app.session_state['placement_config'].get(f) is None for f,_,_ in EXCESS_FIELDS)
    CASE.phase(app, 'phase1')
    app.selectbox(key='_w_job_type').set_value('CSG 9 5/8"').run()
    CASE.phase(app, 'phase2_3')
    assert all(app.session_state['placement_config'].get(f) is None for f,_,_ in EXCESS_FIELDS)


def test_m10_explicit_flat_null_survives_restore_mount_and_date_clear():
    state = audit.round_trip({'job_type':'CMT PLUG', 'report_date':None,
                             'doc_control':{'date':'2025-01-02'}})
    assert state['report_date'] is None and state['doc_control']['date'] == ''
    app = CASE.app(state)
    assert app.date_input(key='_w_report_date').value is None
    app.date_input(key='_w_report_date').set_value(date(2025,1,2)).run()
    app.date_input(key='_w_report_date').set_value(None)
    CASE.phase(app, 'phase2_3')
    assert app.session_state['report_date'] is None
    assert app.session_state['doc_control']['date'] == ''
    restored = audit.round_trip(app.session_state.to_dict())
    assert restored['report_date'] is None and restored['doc_control']['date'] == ''


@pytest.mark.parametrize('nested', [None, ''])
def test_m10_legacy_nested_empty_decode_contract_stays_intact(nested):
    state = audit.round_trip({'job_type':'CMT PLUG', 'doc_control':{'date':nested}})
    assert 'report_date' not in state
    assert state['doc_control']['date'] == nested
    if nested is None:
        app = CASE.app(deepcopy(state))
        assert app.session_state['report_date'] is None
        assert app.date_input(key='_w_report_date').value is None
        assert app.session_state['doc_control']['date'] == ''


@pytest.mark.parametrize('prepared', [False, True])
def test_m10_legacy_null_context_is_blank_without_mutating_date_provenance(reviewed, prepared):
    state = deepcopy(reviewed)
    state.pop('report_date')
    state['doc_control']['date'] = None
    state = audit.round_trip(state)
    assert 'report_date' not in state and state['doc_control']['date'] is None
    assert context(state, prepared)['date'] == ''
    assert 'report_date' not in state and state['doc_control']['date'] is None


def test_m10_fresh_default_and_selected_iso_date_keep_existing_meaning():
    app = CASE.app()
    assert app.session_state['report_date'] == date.today()
    app.date_input(key='_w_report_date').set_value(date(2025,1,2))
    CASE.phase(app, 'phase2_3')
    restored = audit.round_trip(app.session_state.to_dict())
    assert restored['report_date'] == date(2025,1,2)
    assert restored['doc_control']['date'] == '2025-01-02'


def test_m4_m10_null_and_repaired_identity_dates_in_real_word(reviewed, tmp_path):
    state = deepcopy(reviewed)
    state['fluids_config']['params']['Main']['material_name'] = None
    state['report_date'] = None
    state = audit.round_trip(state)
    before = deepcopy(state['cement_params'])
    ctx = context(state, False)
    doc = DocxTemplate('master_template.docx')
    doc.render(report._word_quantity_context(ctx))
    buf = BytesIO();doc.save(buf)
    out = Document(BytesIO(buf.getvalue()))
    main = next([c.text for c in r.cells] for t in out.tables for r in t.rows
                if r.cells[0].text=='Main')
    assert main[1] == '' and ctx['date'] == ''
    texts = [e.text or '' for e in out.element.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')]
    assert texts[texts.index('Date:')+1] == ''
    (tmp_path/'f1-null.docx').write_bytes(buf.getvalue())
    app = CASE.app(deepcopy(state));CASE.phase(app, 'phase4')
    stale(app.session_state)
    app.text_input(key='matname_Main').set_value('F1 field & <mix>')
    CASE.phase(app, 'phase1')
    assert not any(k in app.session_state for k in DOC_KEYS)
    assert app.session_state['fluids_config']['params']['Main']['material_name'] == 'F1 field & <mix>'
    app.date_input(key='_w_report_date').set_value(date(2025,1,2))
    CASE.phase(app, 'phase7')
    restored = audit.round_trip(app.session_state.to_dict())
    assert restored['cement_params'] == before
    assert restored['report_date'] == date(2025,1,2)
    for prepared in (False, True):
        current = context(deepcopy(restored), prepared)
        assert current['date'] == '2025-01-02'
        assert unescape(next(r for r in current['fluids_train'] if r['name']=='Main')['material_name']) == 'F1 field & <mix>'
    CASE.export(app)
    data = app.session_state['_compiled_doc_bytes']
    out = Document(BytesIO(data))
    main = next([c.text for c in r.cells] for t in out.tables for r in t.rows if r.cells[0].text=='Main')
    assert main[1] == 'F1 field & <mix>' and main[3] == '50.0'
    xml = ZipFile(BytesIO(data)).read('word/document.xml').decode()
    assert '2025-01-02' in xml and 'F1 field &amp; &lt;mix&gt;' in xml
    (tmp_path/'f1-repaired.docx').write_bytes(data)


@pytest.mark.parametrize('name', ['', 'None', 'Field & <wash>'])
def test_m4_real_name_edit_before_navigation_preserves_exact_user_text(reviewed, name):
    state = deepcopy(reviewed)
    state['fluids_config']['params']['Main']['material_name'] = None
    app = CASE.app(audit.round_trip(state))
    CASE.phase(app, 'phase4')
    # Start with a genuine nonblank edit so clearing to empty is an event.
    app.text_input(key='matname_Main').set_value('Temporary field name').run()
    stale(app.session_state)
    app.text_input(key='matname_Main').set_value(name)
    CASE.phase(app, 'phase1')
    assert not any(k in app.session_state for k in DOC_KEYS)
    restored = audit.round_trip(app.session_state.to_dict())
    assert restored['fluids_config']['params']['Main']['material_name'] == name
    assert restored['fluid_data']['Main']['material_name'] == name
    assert unescape(next(r for r in context(restored, True)['fluids_train'] if r['name']=='Main')['material_name']) == name


@pytest.mark.parametrize('legacy', [False, True])
def test_m4_missing_name_keeps_default_or_real_legacy_value(reviewed, legacy):
    state = deepcopy(reviewed)
    state['fluids_config']['params']['Main'].pop('material_name')
    if legacy:
        state['fluid_data']['Main']['material_name'] = 'Actual legacy name'
        expected = 'Actual legacy name'
    else:
        state['fluid_data']['Main'].pop('material_name')
        expected = db.DEFAULT_MATERIAL_NAMES['Main']
    app = CASE.app(audit.round_trip(state))
    CASE.phase(app, 'phase4')
    assert app.text_input(key='matname_Main').value == expected
    assert app.session_state['fluids_config']['params']['Main']['material_name'] == expected


def test_m4_diesel_fixed_authority_and_blank_standalone_spacer_are_unchanged(reviewed):
    state = deepcopy(reviewed)
    state['mud_type'] = state['well_data']['mud_type'] = 'OBM'
    state['fluids_config']['active'] += ['Diesel', 'Spacer']
    state['fluids_config']['params'].update({
        'Diesel': {'volume': 10., 'density': '52', 'pump_rate': '4', 'material_name': None},
        'Spacer': {'volume': 20., 'density': '110', 'pump_rate': '4', 'material_name': ''}})
    ps.refresh_fluids(state)
    assert state['fluid_data']['Diesel']['material_name'] == 'Diesel'
    assert state['fluids_config']['params']['Diesel']['material_name'] is None
    assert state['fluid_data']['Spacer']['material_name'] == ''
    app = CASE.app(audit.round_trip(state))
    CASE.phase(app, 'phase4')
    assert not any(w.label=='Material Name - Diesel' for w in app.text_input)
    assert app.session_state['fluids_config']['params']['Diesel']['material_name'] == 'Diesel'
    assert app.text_input(key='matname_Spacer').value == ''


@pytest.mark.parametrize('date_fields', [
    {'report_date': None},
    {'doc_control': {'date': None}},
    {'doc_control': {'date': '2025-01-02'}},
    {'doc_control': {'date': ''}},
    {'report_date': '2025-01-02'},
    {},
])
def test_m10_real_project_load_keeps_unchanged_date_checkpoint_clean(date_fields):
    import json
    from test_pre_round14_ui02_navigation_persistence import upload, dirty
    app = CASE.app()
    incoming = {'well_name': 'F1 loaded checkpoint', 'client': 'Synthetic client', **date_fields}
    with upload(app, json.dumps(incoming).encode()):
        app.run()
        CASE.healthy(app)
        assert app.session_state['well_name'] == 'F1 loaded checkpoint'
        expected = (incoming['report_date'] if 'report_date' in incoming else
                    None if incoming.get('doc_control', {}).get('date', 'absent') is None else
                    date(2025,1,2) if incoming.get('doc_control', {}).get('date') else date.today())
        assert app.session_state['report_date'] == expected
        assert not dirty(app)
        CASE.phase(app, 'phase2_3')
        CASE.phase(app, 'phase1')
        assert app.session_state['report_date'] == expected
        assert not dirty(app)


@pytest.mark.parametrize('prepared', [False, True])
def test_m10_flat_null_without_document_mirror_is_blank_in_context(reviewed, prepared):
    state = deepcopy(reviewed)
    state['report_date'] = None
    state.pop('doc_control')
    restored = audit.round_trip(state)
    assert restored['report_date'] is None
    assert restored.get('doc_control', {}).get('date') == ''
    assert context(restored, prepared)['date'] == ''
    assert restored['report_date'] is None


def test_m10_bare_iso_mount_preserves_canonical_string_until_real_edit():
    app = CASE.app(audit.round_trip({'job_type': 'CMT PLUG', 'report_date': '2025-01-02'}))
    assert app.date_input(key='_w_report_date').value == date(2025,1,2)
    assert app.session_state['report_date'] == '2025-01-02'
    CASE.phase(app, 'phase2_3')
    CASE.phase(app, 'phase1')
    assert app.session_state['report_date'] == '2025-01-02'
    app.date_input(key='_w_report_date').set_value(date(2025,2,3))
    CASE.phase(app, 'phase2_3')
    assert app.session_state['report_date'] == date(2025,2,3)
    assert audit.round_trip(app.session_state.to_dict())['doc_control']['date'] == '2025-02-03'


@pytest.mark.parametrize('edit', ['date', 'clear_date', 'well'])
def test_m10_loaded_checkpoint_real_edit_requires_confirmation_and_cancel_is_lossless(edit):
    import json
    from test_pre_round14_ui02_navigation_persistence import upload, dirty
    from test_uat_round10c_integration import save_project
    app = CASE.app()
    with upload(app, b'{"well_name":"F1 loaded","report_date":"2025-01-02"}'):
        app.run()
        CASE.healthy(app)
        assert not dirty(app)
    if edit == 'well':
        app.text_input(key='_w_well_name').set_value('F1 genuine edit')
    else:
        app.date_input(key='_w_report_date').set_value(None if edit=='clear_date' else date(2025,2,3))
    CASE.phase(app, 'phase2_3')
    assert dirty(app)
    if edit != 'well':
        assert app.session_state['report_date'] == (None if edit=='clear_date' else date(2025,2,3))
    before = save_project(app)
    with upload(app, json.dumps({'well_name':'F1 replacement'}).encode()):
        app.run()
        CASE.healthy(app)
        assert any(b.key=='_confirm_upload_replacement' for b in app.button)
        assert save_project(app) == before
        app.button(key='_cancel_upload_replacement').click().run()
        CASE.healthy(app)
        assert save_project(app) == before and dirty(app)
    with upload(app, b'{"well_name":"F1 replacement"}'):
        app.run()
        app.button(key='_confirm_upload_replacement').click().run()
        CASE.healthy(app)
        assert app.session_state['well_name'] == 'F1 replacement'
        assert not dirty(app)
