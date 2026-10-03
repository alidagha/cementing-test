"""Regression checks for project restoration, defaults and Word export.

Run from the repository root with: python -m unittest discover -s tests -v
"""
import ast
from datetime import date, datetime
from copy import deepcopy
from io import BytesIO
import json
import math
import re
from pathlib import Path
import sys
import unittest
from unittest.mock import patch

import pandas as pd
from docx import Document
from streamlit.testing.v1 import AppTest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import materials_db
from engineering_tools import compute_phase_status, lab_review_signature, round_half_up
from project_io import decode_project
from project_state import is_project_key, prepare_calculations, hardware_draft_pending

# Use the application's serializers without running its top-level UI.
tree = ast.parse((ROOT / "main.py").read_text())
namespace = {"pd": pd, "date": date, "datetime": datetime, "math": math}
exec(compile(ast.Module(body=[node for node in tree.body
                             if isinstance(node, ast.FunctionDef)
                             and node.name in ("serialize_item", "deserialize_item")],
                        type_ignores=[]), str(ROOT / "main.py"), "exec"), namespace)


def round_trip(state):
    saved = {key: namespace["serialize_item"](value)
             for key, value in state.items() if is_project_key(key)}
    return decode_project(json.dumps(saved, allow_nan=False).encode(),
                          namespace["deserialize_item"])


BATCH1_JOBS = ('CSG 9 5/8"', 'CMT PLUG', 'CMT SQUEEZE', 'LNR 7"', 'TIE BACK LNR 7"')


class AuditRegressions(unittest.TestCase):
    def app(self, project=None):
        app = AppTest.from_file(str(ROOT / "main.py"), default_timeout=30)
        for key, value in (project or {}).items():
            app.session_state[key] = value
        app.run()
        self.healthy(app)
        return app

    def healthy(self, app):
        self.assertFalse(app.exception, [e.message for e in app.exception])

    def phase(self, app, name):
        app.radio(key="_app_mode_key").set_value(name).run()
        self.healthy(app)
        self.assertEqual(app.radio(key="_app_mode_key").value, name)

    def test_all_pages(self):
        app = self.app()
        for phase in ("phase1", "phase2_3", "phase4", "phase5", "phase6", "phase7", "phase10"):
            with self.subTest(phase=phase):
                self.phase(app, phase)

    def test_hole_size_save_load_and_legacy_defaults(self):
        for hole, customized, expected in (
            ('Custom approved hole', None, 'Custom approved hole'),
            ('26"', None, '12 1/4"'),
            ('26"', True, '26"'),
        ):
            with self.subTest(hole=hole, customized=customized):
                project = {"job_type": 'CSG 20"', "hole_size": hole}
                if customized is not None:
                    project["hole_size_customized"] = customized
                app = self.app(round_trip(project))
                app.selectbox(key="_w_job_type").set_value('CSG 9 5/8"').run()
                self.healthy(app)
                self.assertEqual(app.text_input(key="_w_hole_size").value, expected)

        app = self.app()
        app.text_input(key="_w_hole_size").set_value('Custom approved hole').run()
        app.text_input(key="_w_hole_size").set_value('26"').run()
        restored = round_trip(app.session_state.to_dict())
        self.assertTrue(restored["hole_size_customized"])
        app = self.app(restored)
        app.selectbox(key="_w_job_type").set_value('CSG 9 5/8"').run()
        self.healthy(app)
        self.assertEqual(app.text_input(key="_w_hole_size").value, '26"')

    def test_numeric_lab_tables_restore_and_render(self):
        grid = pd.DataFrame([
            {"Material": "CEMENT G DELIJAN", "Concentration": 100.0,
             "Unit": "BWOB", "Mass": 600.0, "Lot No": "sample"},
            {"Material": "Base Water", "Concentration": None,
             "Unit": "BWOW", "Mass": None, "Lot No": None},
        ])
        for archived in (False, True):
            with self.subTest(archived=archived):
                project = {
                    "job_type": 'CSG 20"',
                    "fluids_config": {"active": ["Main"], "params": {
                        "Main": {"volume": 50.0, "density": "118", "pump_rate": "4"}}},
                }
                if archived:
                    project["inactive_slurry_drafts"] = {"Main": {"lab_grid_dfs": grid}}
                else:
                    project["lab_grid_dfs"] = {"Main": grid}
                restored = round_trip(project)
                normalized = (restored["inactive_slurry_drafts"]["Main"]["lab_grid_dfs"]
                              if archived else restored["lab_grid_dfs"]["Main"])
                self.assertEqual(normalized.iloc[0]["Concentration"], "100.0")
                self.assertTrue(pd.isna(normalized.iloc[1]["Concentration"]))
                app = self.app(restored)
                self.phase(app, "phase4")
                self.phase(app, "phase7")
                self.assertIn("Main", app.session_state["lab_grid_dfs"])

    def test_spacer_defaults_and_saved_densities(self):
        spacers = ("Spacer", "Spacer Ahead", "Spacer Behind")
        for params in (None, {}, {s: {"volume": 0.0, "density": "101.0", "pump_rate": "4"}
                                 for s in spacers}):
            with self.subTest(params=params):
                expected = "101.0" if params else "80.0"
                app = self.app(None if params is None else {
                    "job_type": 'CSG 20"',
                    "fluids_config": {"active": list(spacers), "params": params},
                })
                self.phase(app, "phase4")
                if params is None:
                    for name in spacers:
                        app.checkbox(key="chk_fluid_" + name).check().run()
                self.healthy(app)
                for name in spacers:
                    self.assertEqual(app.text_input(key="den_" + name).value, expected)

    def configured_app(self, job, additives=False, optional_fluids=False):
        # Valid export fixtures explicitly supply the former well defaults.
        app = self.app({"mud_density": "80.0", "plastic_viscosity": "45", "yield_point": "15",
                        "geo_md": 3000.0, "geo_tvd": 3000.0, "bhst": 200, "geo_gradient": 1.25})
        app.selectbox(key="_w_job_type").set_value(job).run()
        app.text_input(key="_w_well_name").set_value("Regression Well").run()
        app.text_input(key="_w_client").set_value("Regression Client").run()
        columns = ["Description", "MD (m)", "Size (in)", "ID (in)", "Joint (m)",
                   "Weight (ppf)", "Grade", "Collapse (psi)", "Burst (psi)"]
        app.session_state["hardware_table"] = pd.DataFrame([
            ["Previous Casing", "3000", "20", 18.0, 12.0, 100.0, "K-55", 1000.0, 2000.0],
        ], columns=columns)
        self.phase(app, "phase2_3")
        next(w for w in app.selectbox if w.label == "Target shoe / treatment depth source").set_value("__manual__").run()
        next(w for w in app.number_input if w.label == "Measured target depth (m MD)").set_value(3000.0).run()
        if "TIE BACK" in job:
            host = next(w for w in app.selectbox if w.label.startswith("Tie-back host"))
            host.select_index(1).run()
        self.phase(app, "phase4")
        selected = ["Main"]
        if optional_fluids:
            selected += ["Pre Flush", "Spacer", "Spacer Ahead", "Spacer Behind"]
        for name in selected:
            app.checkbox(key="chk_fluid_" + name).check().run()
        for name in app.session_state["fluids_config"]["active"]:
            app.number_input(key="vol_" + name).set_value(50.0).run()
            if name in ("Main", "Pre Flush"):
                app.text_input(key="den_" + name).set_value("118.0" if name == "Main" else "80.0").run()
        if additives:
            app.session_state["cement_additives_dfs"] = {"Main": pd.DataFrame([
                {"Material Type": "Anti Foam", "Name": "Anti Foam", "Physical State": "Liquid",
                 "Mix Method": "In Mix Water", "User Input": 0.01},
            ])}
        self.phase(app, "phase5")
        next(w for w in app.number_input if w.label == "Dead Vol (bbl) - Main").set_value(31.0).run()
        next(w for w in app.selectbox if w.label == "Top of cement - Main").set_value("Surface").run()
        if optional_fluids:
            app.session_state["spacer_dfs"] = {s: pd.DataFrame([
                {"Chemical": "Spacer", "User Input (% or gal)": 1.0, "Weighting Agent Type": "-"},
            ]) for s in ("Spacer", "Spacer Ahead", "Spacer Behind")}
        self.phase(app, "phase6")
        self.phase(app, "phase7")
        next(b for b in app.button if b.label == "Confirm measured lab results").click().run()
        self.healthy(app)
        return app

    def export(self, app):
        self.phase(app, "phase10")
        self.assertFalse(app.error, [e.value for e in app.error])
        build = next(b for b in app.button if b.label == "Build Word Document")
        self.assertFalse(build.disabled)
        build.click().run()
        self.healthy(app)
        self.assertFalse(app.error, [e.value for e in app.error])
        doc = Document(BytesIO(app.session_state["_compiled_doc_bytes"]))
        self.assertTrue(doc.tables)
        text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        self.assertIn("Regression Client", text)

    def test_neat_cement_export_for_every_job_type(self):
        for job in materials_db.JOB_TYPES:
            with self.subTest(job=job):
                app = self.configured_app(job)
                self.assertTrue(app.session_state["cement_additives_dfs"]["Main"].empty)
                self.assertEqual(compute_phase_status(app.session_state)["phase5"]["level"], "ok")
                self.export(app)

    def test_additives_optional_fluids_and_saved_project_export(self):
        app = self.configured_app('CSG 20"', additives=True, optional_fluids=True)
        self.export(app)
        app = self.app(round_trip(app.session_state.to_dict()))
        for phase in ("phase2_3", "phase4", "phase5", "phase6", "phase7"):
            self.phase(app, phase)
        self.export(app)
        app.session_state["cement_additives_dfs"]["Main"].loc[0, "User Input"] = None
        self.assertEqual(compute_phase_status(app.session_state)["phase5"]["level"], "warning")
        self.phase(app, "phase10")
        self.assertTrue(next(b for b in app.button if b.label == "Build Word Document").disabled)
        self.assertNotIn("_compiled_doc_bytes", app.session_state)


    def assert_export_blocked(self, app):
        self.phase(app, "phase10")
        self.assertTrue(next(b for b in app.button if b.label == "Build Word Document").disabled)
        for key in ("_compiled_doc_bytes", "_compiled_doc_filename", "_compiled_doc_signature"):
            self.assertNotIn(key, app.session_state)
        self.assertFalse(any(b.proto.label == "Download Final Document (.docx)"
                             for b in app.get("download_button")))
        readiness = next(e for e in app.expander if "Export readiness checklist" in e.label)
        self.assertFalse(readiness.label.startswith("✅"))
        self.assertEqual(len(app.text_area), 2)  # Repairable report editors stay mounted.

    def test_batch1_calculation_errors_block_export(self):
        for job in BATCH1_JOBS:
            with self.subTest(job=job):
                app = self.configured_app(job, optional_fluids=True)
                self.export(app)
                app.session_state["preflush_config"]["wash_multiplier"] = "abc"
                # The stale Phase VI cache still looks complete: the calculation
                # failure itself must block export, independently of status.
                self.assertTrue(all(s["level"] == "ok"
                                    for s in compute_phase_status(app.session_state).values()))
                self.assert_export_blocked(app)
                self.assertTrue(app.error)
                self.assertTrue(prepare_calculations(deepcopy(app.session_state.to_dict())))
                self.assertEqual(round_trip(app.session_state.to_dict())["preflush_config"]["wash_multiplier"], "abc")
                self.phase(app, "phase6")
                next(w for w in app.text_input if w.label.startswith("Correct Preflush wash ratio")).set_value("3").run()
                next(b for b in app.button if b.label == "Apply corrected values").click().run()
                self.export(app)

    def test_batch1_restored_cement_bounds_block_export(self):
        invalid = (("yield", -1.0), ("yield", 0.0), ("yield", 0.0005),
                   ("mix_water", -1.0), ("dead_vol", -31.0),
                   ("cmt_sg", 2.49), ("cmt_sg", 3.51))
        for job in BATCH1_JOBS:
            valid = self.configured_app(job)
            self.export(valid)
            saved = round_trip(valid.session_state.to_dict())
            compiled = {k: valid.session_state[k] for k in
                        ("_compiled_doc_bytes", "_compiled_doc_filename", "_compiled_doc_signature")}
            for field, value in invalid:
                with self.subTest(job=job, field=field, value=value):
                    project = deepcopy(saved)
                    project["cement_params"]["Main"].update(auto_calc=False)
                    project["cement_params"]["Main"][field] = value
                    app = self.app({**round_trip(project), **compiled})
                    self.assertEqual(compute_phase_status(app.session_state)["phase5"]["level"], "warning")
                    self.assertNotIn("_compiled_doc_bytes", app.session_state)
                    self.assert_export_blocked(app)
                    self.assertTrue(prepare_calculations(deepcopy(app.session_state.to_dict())))
                    self.assertEqual(app.session_state["cement_params"]["Main"][field], value)
                    self.phase(app, "phase5")
                    self.assertTrue(any(w.label.startswith("Correct ") for w in app.text_input))
            with self.subTest(job=job, valid_manual=True):
                project = deepcopy(saved)
                project["cement_params"]["Main"].update(auto_calc=False, mix_water=30.0, dead_vol=0.0)
                project["cement_params"]["Main"]["yield"] = 1.5
                app = self.app(round_trip(project))
                self.export(app)
                self.assertEqual(app.session_state["cement_params"]["Main"]["yield"], 1.5)
                self.assertEqual(app.session_state["cement_params"]["Main"]["total_sacks"], 187.2)

    def test_batch1_lab_masses_block_confirmation_and_export(self):
        for job in BATCH1_JOBS:
            app = self.configured_app(job)
            self.export(app)
            for mass in ("-10", "abc", "nan", "inf", "", None):
                with self.subTest(job=job, mass=mass):
                    self.phase(app, "phase7")
                    editor = next(e for e in app.dataframe if "_editor_lab_tbl_main_" in e.proto.id)
                    widgets = app._tree.get_widget_states()
                    event = widgets.widgets.add()
                    event.id = editor.proto.id
                    event.string_value = json.dumps({"edited_rows": {"0": {"Mass": mass}},
                                                     "added_rows": [], "deleted_rows": []})
                    app._run(widgets)
                    self.healthy(app)
                    confirm = next(b for b in app.button if b.label == "Confirm measured lab results")
                    self.assertTrue(confirm.disabled)
                    self.assertEqual(compute_phase_status(app.session_state)["phase7"]["level"], "warning")
                    self.assertNotIn("_compiled_doc_bytes", app.session_state)
                    # Even matching restored review metadata cannot bless a bad mass.
                    qc = app.session_state["lab_qc_params"]["Main"]
                    qc["reviewed"] = True
                    qc["review_signature"] = lab_review_signature(qc, app.session_state["lab_grid_dfs"]["Main"])
                    restored = round_trip(app.session_state.to_dict())
                    self.assertEqual(compute_phase_status(restored)["phase7"]["level"], "warning")
                    self.assert_export_blocked(app)
                    self.assertTrue(prepare_calculations(deepcopy(app.session_state.to_dict())))
                    stored = restored["lab_grid_dfs"]["Main"].iloc[0]["Mass"]
                    self.assertTrue(pd.isna(stored) if mass is None else stored == mass)
            # Repaired data can be confirmed and exported without syncing away edits.
            self.phase(app, "phase7")
            editor = next(e for e in app.dataframe if "_editor_lab_tbl_main_" in e.proto.id)
            widgets = app._tree.get_widget_states()
            event = widgets.widgets.add()
            event.id = editor.proto.id
            event.string_value = json.dumps({"edited_rows": {"0": {"Mass": "0"}},
                                             "added_rows": [], "deleted_rows": []})
            app._run(widgets)
            next(b for b in app.button if b.label == "Confirm measured lab results").click().run()
            self.export(app)
            self.assertEqual(app.session_state["lab_grid_dfs"]["Main"].iloc[0]["Mass"], "0")

    def test_batch1_reviewed_thickening_time_is_revalidated(self):
        for job in BATCH1_JOBS:
            valid = self.configured_app(job)
            self.export(valid)
            saved = round_trip(valid.session_state.to_dict())
            compiled = {k: valid.session_state[k] for k in
                        ("_compiled_doc_bytes", "_compiled_doc_filename", "_compiled_doc_signature")}
            for time in ("00:00", "24:01", "99:59", "03.30", None, "00:01", "24:00"):
                with self.subTest(job=job, time=time):
                    project = deepcopy(saved)
                    qc = project["lab_qc_params"]["Main"]
                    qc["thickening_time"] = time
                    qc["review_signature"] = lab_review_signature(qc, project["lab_grid_dfs"]["Main"])
                    app = self.app({**round_trip(project), **compiled})
                    if time in ("00:01", "24:00"):
                        self.export(app)
                        doc = Document(BytesIO(app.session_state["_compiled_doc_bytes"]))
                        self.assertTrue(any(time in c.text for t in doc.tables for row in t.rows for c in row.cells))
                    else:
                        self.assertEqual(compute_phase_status(app.session_state)["phase7"]["level"], "warning")
                        self.assertNotIn("_compiled_doc_bytes", app.session_state)
                        self.assert_export_blocked(app)
                        self.assertTrue(prepare_calculations(deepcopy(app.session_state.to_dict())))
                        self.assertEqual(app.session_state["lab_qc_params"]["Main"]["thickening_time"], time)
                        self.phase(app, "phase7")
                        self.assertTrue(next(b for b in app.button if b.label == "Confirm measured lab results").disabled)


    def hardware_event(self, app, edits, navigate=None):
        if navigate:
            app.radio(key="_app_mode_key").set_value(navigate)
        editor = next(e for e in app.dataframe if "_editor_hardware_" in e.proto.id)
        widgets = app._tree.get_widget_states()
        event = widgets.widgets.add()
        event.id = editor.proto.id
        event.string_value = json.dumps({"edited_rows": {}, "added_rows": [], "deleted_rows": [], **edits})
        app._run(widgets)
        self.healthy(app)

    def test_batch2_final_edits_survive_navigation(self):
        import phase_10_procedure
        for index, job in enumerate(BATCH1_JOBS):
            with self.subTest(job=job):
                app = self.configured_app(job)
                self.phase(app, "phase1")
                app.selectbox(key="_w_job_type").set_value('CSG 20"').run()
                name = f"Navigation Well {index}"
                app.text_input(key="_w_well_name").set_value(name)
                app.text_input(key="_w_prepared_by").set_value("Navigation Engineer")
                app.text_input(key="_w_request_number").set_value("REQ-7")
                app.date_input(key="_w_report_date").set_value(date(2026, 9, 20))
                app.selectbox(key="_w_job_type").set_value(job)
                app.radio(key="_app_mode_key").set_value("phase2_3").run()
                self.healthy(app)
                self.assertEqual(app.session_state["well_name"], name)
                self.assertEqual(app.session_state["doc_control"]["well_name"], name)
                self.assertEqual(app.session_state["job_type"], job)
                self.assertEqual(app.session_state["doc_control"]["job_type"], job)
                self.assertEqual(app.session_state["doc_control"]["date"], "2026-09-20")
                next(w for w in app.selectbox if w.label == "Target shoe / treatment depth source").set_value("__manual__")
                app.radio(key="_app_mode_key").set_value("phase4").run()
                self.assertEqual(app.session_state["placement_config"]["target_row"], "__manual__")
                self.phase(app, "phase2_3")
                app.text_input(key="_w_mud_density").set_value("84-86")
                app.text_input(key="_w_plastic_viscosity").set_value("46-48")
                app.text_input(key="_w_yield_point").set_value("16-18")
                app.text_input(key="_w_bhsp").set_value("7300+1000")
                next(w for w in app.number_input if w.label == "Measured target depth (m MD)").set_value(2850.0)
                next(w for w in app.text_input if w.label.startswith("Slurry volume basis")).set_value("Approved field basis")
                if "TIE BACK" in job:
                    next(w for w in app.selectbox if w.label.startswith("Tie-back host")).select_index(1)
                app.radio(key="_app_mode_key").set_value("phase4").run()
                self.healthy(app)
                well = app.session_state["well_data"]
                self.assertEqual(well["mud_density"], "84-86")
                self.assertEqual(well["effective_mud_density"], 85.0)
                self.assertEqual(well["effective_pv"], 47.0)
                self.assertEqual(well["effective_yp"], 17.0)
                self.assertEqual(well["bhsp"], "7300+1000")
                placement = app.session_state["placement_config"]
                self.assertEqual(placement["manual_depth_m"], 2850.0)
                self.assertEqual(placement["volume_basis"], "Approved field basis")
                saved = round_trip(app.session_state.to_dict())
                self.assertEqual(saved["well_name"], name)
                self.assertEqual(saved["doc_control"]["job_type"], job)
                self.assertEqual(saved["well_data"]["mud_density"], "84-86")
                app = self.app(saved)
                self.assertEqual(app.text_input(key="_w_well_name").value, name)
                self.assertEqual(app.text_input(key="_w_prepared_by").value, "Navigation Engineer")
                self.assertEqual(app.text_input(key="_w_request_number").value, "REQ-7")
                self.assertEqual(app.selectbox(key="_w_job_type").value, job)
                self.phase(app, "phase2_3")
                self.assertEqual(app.text_input(key="_w_mud_density").value, "84-86")
                self.assertEqual(next(w for w in app.number_input if w.label == "Measured target depth (m MD)").value, 2850.0)
                for phase in ("phase4", "phase5", "phase6", "phase7"):
                    self.phase(app, phase)
                next(b for b in app.button if b.label == "Keep reviewed lab entries").click().run()
                next(b for b in app.button if b.label == "Confirm measured lab results").click().run()
                self.export(app)
                with patch.object(phase_10_procedure.st, "session_state", deepcopy(app.session_state.to_dict())):
                    context = phase_10_procedure.build_master_context()
                self.assertEqual(context["well_name"], name)
                self.assertEqual(context["job_type"], job)
                self.assertEqual(context["well_data"]["mud_density"], "84-86")
                self.assertIn("2850.0", context["exec_summary"])
                doc = Document(BytesIO(app.session_state["_compiled_doc_bytes"]))
                self.assertIn(name, doc._element.xml)  # Metadata also appears inside nested Word tables.

    def test_batch2_pending_hardware_blocks_stale_export(self):
        actions = ({"edited_rows": {"0": {"MD (m)": "2950"}}},
                   {"added_rows": [{"MD (m)": "1200"}]},
                   {"deleted_rows": [0]})
        for job in BATCH1_JOBS:
            valid = self.configured_app(job)
            self.export(valid)
            saved = round_trip(valid.session_state.to_dict())
            compiled = {key: valid.session_state[key] for key in
                        ("_compiled_doc_bytes", "_compiled_doc_filename", "_compiled_doc_signature")}
            for action in actions:
                with self.subTest(job=job, action=action):
                    app = self.app({**deepcopy(saved), **compiled})
                    self.phase(app, "phase2_3")
                    original = app.session_state["hardware_table"].copy(deep=True)
                    self.hardware_event(app, action, navigate="phase10")
                    self.assertTrue(app.session_state["hardware_table"].equals(original))
                    self.assertTrue(hardware_draft_pending(app.session_state))
                    self.assertEqual(compute_phase_status(app.session_state)["phase2_3"]["level"], "warning")
                    self.assert_export_blocked(app)
                    self.assertTrue(any("pending hardware" in issue for issue in
                                        prepare_calculations(deepcopy(app.session_state.to_dict()))))
                    restored = round_trip(app.session_state.to_dict())
                    self.assertTrue(hardware_draft_pending(restored))
                    app = self.app(restored)
                    self.assert_export_blocked(app)
                    self.phase(app, "phase2_3")
                    if "added_rows" in action:
                        self.assertTrue(hardware_draft_pending(app.session_state))
                        self.hardware_event(app, {"edited_rows": {"1": {
                            "Description": "Open Hole Size", "Size (in)": "12.25", "ID (in)": 12.25,
                            "Joint (m)": 3.0, "Weight (ppf)": 30.0, "Grade": "L-80",
                            "Collapse (psi)": 100.0, "Burst (psi)": 200.0}}})
                        open_hole = app.session_state["hardware_table"].iloc[1]
                        for field in ("Joint (m)", "Weight (ppf)", "Collapse (psi)", "Burst (psi)"):
                            self.assertEqual(open_hole[field], 0.0)
                        self.assertEqual(open_hole["Grade"], "-")
                    elif "deleted_rows" in action:
                        self.assertTrue(app.session_state["hardware_table"].empty)
                        self.assertNotEqual(compute_phase_status(app.session_state)["phase2_3"]["level"], "ok")
                        self.hardware_event(app, {"added_rows": original.to_dict("records")})
                    else:
                        self.assertEqual(app.session_state["hardware_table"].iloc[0]["MD (m)"], "2950")
                    self.assertFalse(hardware_draft_pending(app.session_state))
                    if "TIE BACK" in job:
                        # A changed/deleted host must be selected again after review.
                        next(w for w in app.selectbox if w.label.startswith("Tie-back host")).select_index(1).run()
                    self.assertEqual(compute_phase_status(app.session_state)["phase2_3"]["level"], "ok")
                    app = self.app(round_trip(app.session_state.to_dict()))
                    self.phase(app, "phase2_3")
                    self.export(app)

    def test_batch2_nested_restoration_and_flat_precedence(self):
        nested = {"doc_control": {"job_type": 'CSG 9 5/8"', "well_name": "Nested Well",
                                  "client": "Nested Client", "hole_size": "Approved hole",
                                  "date": "2026-09-20", "prepared_by": "Nested Engineer"},
                  "well_data": {"mud_type": "OBM", "mud_density": "90-92", "plastic_viscosity": "50",
                                "yield_point": "20", "geo_md": 3300.0, "geo_tvd": 3200.0,
                                "bhst": 210, "geo_gradient": 1.40, "bhsp": "6000+500"}}
        for flat in ({}, {"well_name": "Flat Well", "client": "", "job_type": 'CMT PLUG',
                         "mud_density": "84", "geo_md": 3400.0, "report_date": date(2026, 9, 21)}):
            with self.subTest(flat=flat):
                restored = round_trip({**deepcopy(nested), **flat})
                self.assertEqual(restored["well_name"], flat.get("well_name", "Nested Well"))
                self.assertEqual(restored["job_type"], flat.get("job_type", 'CSG 9 5/8"'))
                self.assertEqual(restored["client"], flat.get("client", "Nested Client"))
                self.assertEqual(restored["mud_density"], flat.get("mud_density", "90-92"))
                self.assertEqual(restored["geo_md"], flat.get("geo_md", 3300.0))
                self.assertEqual(restored["effective_mud_density"], 84.0 if flat else 91.0)
                app = self.app(restored)
                self.assertEqual(app.text_input(key="_w_well_name").value, restored["well_name"])
                self.phase(app, "phase2_3")
                self.assertEqual(app.text_input(key="_w_mud_density").value, restored["mud_density"])
                self.assertEqual(app.number_input(key="_w_geo_md").value, restored["geo_md"])
                self.phase(app, "phase4")
                self.phase(app, "phase1")
                saved = round_trip(app.session_state.to_dict())
                for key in ("well_name", "job_type", "client", "mud_density", "geo_md", "report_date"):
                    self.assertEqual(saved[key], restored[key])
                self.assertEqual(saved["doc_control"]["well_name"], saved["well_name"])
                self.assertEqual(saved["doc_control"]["job_type"], saved["job_type"])
                self.assertEqual(saved["well_data"]["mud_density"], saved["mud_density"])
        with self.assertRaises(ValueError):
            round_trip({"doc_control": {"job_type": "Unsupported job"}})
        with self.assertRaises(ValueError):
            round_trip({"well_data": {"geo_md": "invalid"}})

    def test_batch2_well_only_replacement_requires_confirmation(self):
        incoming = BytesIO(json.dumps({"well_name": "Incoming Well", "client": "Incoming Client"}).encode())
        incoming.name = "replacement.json"
        incoming.size = len(incoming.getvalue())
        for field, value in (("mud_type", "OBM"), ("mud_density", "90"), ("plastic_viscosity", "50"),
                             ("yield_point", "20"), ("geo_md", 3300.0), ("geo_tvd", 2800.0),
                             ("bhst", 210), ("geo_gradient", 1.40), ("bhsp", "6000")):
            with self.subTest(field=field):
                app = self.app()
                self.phase(app, "phase2_3")
                widget = next(w for w in [*app.text_input, *app.number_input, *app.selectbox]
                              if w.key == f"_w_{field}")
                widget.set_value(value)
                app.radio(key="_app_mode_key").set_value("phase1").run()
                self.assertTrue(any("unsaved changes" in w.value for w in app.warning))
                uploader_key = f"proj_uploader_{app.session_state.to_dict().get('_uploader_revision', 0)}"
                with patch("streamlit.delta_generator.DeltaGenerator.file_uploader",
                           side_effect=lambda *args, **kwargs: incoming if kwargs.get("key") == uploader_key else None):
                    app.run()
                    self.healthy(app)
                    self.assertTrue(any(b.key == "_confirm_upload_replacement" for b in app.button))
                    self.assertEqual(app.session_state[field], value)
                    next(b for b in app.button if b.key == "_cancel_upload_replacement").click().run()
                    self.healthy(app)
                    self.assertEqual(app.session_state[field], value)
                    self.assertEqual(app.session_state["well_data"][field], value)
                    self.assertFalse(any(b.key == "_confirm_upload_replacement" for b in app.button))
        # Visiting untouched initialization pages must not count as user edits.
        app = self.app()
        for phase in ("phase2_3", "phase4", "phase5", "phase6", "phase7", "phase1"):
            self.phase(app, phase)
        self.assertFalse(any("unsaved changes" in w.value for w in app.warning))
        with patch("streamlit.delta_generator.DeltaGenerator.file_uploader", return_value=incoming):
            app.run()
            self.healthy(app)
            self.assertEqual(app.session_state["well_name"], "Incoming Well")
            self.assertFalse(any(b.key == "_confirm_upload_replacement" for b in app.button))
        # Explicit confirmation uses the same replacement path after a well-only edit.
        app = self.app()
        self.phase(app, "phase2_3")
        app.text_input(key="_w_mud_density").set_value("91").run()
        with patch("streamlit.delta_generator.DeltaGenerator.file_uploader", return_value=incoming):
            app.run()
            next(b for b in app.button if b.key == "_confirm_upload_replacement").click().run()
            self.healthy(app)
            self.assertEqual(app.session_state["well_name"], "Incoming Well")
            self.assertNotIn("mud_density", app.session_state)

    def test_batch2_numeric_hardware_is_normalized_before_mount(self):
        from phase_2_3_well_data import HARDWARE_COLUMNS
        grid = pd.DataFrame([
            ["Previous Casing", 3000.0, 20.0, 18.0, 12.0, 100.0, "K-55", 1000.0, 2000.0],
            [None, None, None, None, None, None, None, None, None],
        ], columns=HARDWARE_COLUMNS)
        for draft in (False, True):
            with self.subTest(draft=draft):
                key = "hardware_editor_draft" if draft else "hardware_table"
                restored = round_trip({"job_type": 'CSG 20"', key: grid.copy(deep=True)})
                self.assertEqual(restored[key].iloc[0]["MD (m)"], "3000.0")
                self.assertEqual(restored[key].iloc[0]["Size (in)"], "20.0")
                self.assertTrue(pd.isna(restored[key].iloc[1]["MD (m)"]))
                self.assertTrue(pd.isna(restored[key].iloc[1]["Size (in)"]))
                app = self.app(restored)
                self.phase(app, "phase2_3")
                source = app.session_state["_editor_source__editor_hardware_0"]
                self.assertEqual(source.iloc[0]["MD (m)"], "3000.0")
                self.assertEqual(source.iloc[0]["Size (in)"], "20.0")
                self.assertTrue(pd.isna(source.iloc[1]["MD (m)"]))
                self.assertTrue(pd.isna(source.iloc[1]["Size (in)"]))
                self.phase(app, "phase1")
                app = self.app(round_trip(app.session_state.to_dict()))
                self.phase(app, "phase2_3")
                # In-memory/legacy state also normalizes before the first mount.
                app = self.app({"job_type": 'CSG 20"', key: grid.copy(deep=True)})
                self.phase(app, "phase2_3")
        from project_io import normalize_hardware_text_columns
        missing = pd.DataFrame({"MD (m)": pd.Series([pd.NA, None], dtype=object),
                                "Size (in)": [None, float("nan")]})
        normalize_hardware_text_columns(missing)
        self.assertTrue(missing.isna().all().all())
        self.assertEqual(missing["MD (m)"].dtype, object)
        self.assertEqual(missing["Size (in)"].dtype, object)

    def test_batch3_first_top_depth_and_restored_repairs(self):
        valid = self.configured_app('CSG 9 5/8"')
        saved = round_trip(valid.session_state.to_dict())
        project = deepcopy(saved)
        project.pop("cement_params")  # First Phase V entry for the slurry.
        app = self.app(project)
        self.phase(app, "phase5")
        next(w for w in app.number_input if w.label == "Dead Vol (bbl) - Main").set_value(31.0).run()
        mode = next(w for w in app.selectbox if w.label == "Top of cement - Main")
        self.assertEqual(mode.value, "Not entered")
        self.assertIsNone(app.session_state["cement_params"]["Main"]["top_depth"])
        self.assertEqual(compute_phase_status(app.session_state)["phase5"]["level"], "warning")
        mode.set_value("Depth (m MD)").run()
        self.healthy(app)
        depth = next(w for w in app.number_input if w.label == "Top of cement depth (m MD) - Main")
        self.assertEqual(depth.value, 0.0)
        self.assertFalse(any(w.label.startswith("Correct Top depth") for w in app.text_input))
        self.assertEqual(compute_phase_status(app.session_state)["phase5"]["level"], "warning")
        self.assert_export_blocked(app)
        self.phase(app, "phase5")
        next(w for w in app.number_input if w.label == "Top of cement depth (m MD) - Main").set_value(2100.0)
        app.radio(key="_app_mode_key").set_value("phase10").run()
        self.assertEqual(app.session_state["cement_params"]["Main"]["top_depth"], 2100.0)
        self.assertEqual(compute_phase_status(app.session_state)["phase5"]["level"], "ok")
        self.export(app)
        saved_depth = round_trip(app.session_state.to_dict())
        for value in (None, "bad depth", -10.0, 2100.0):
            with self.subTest(restored_depth=value):
                project = deepcopy(saved_depth)
                project["cement_params"]["Main"]["top_depth"] = value
                project["cement_params"]["Main"].pop("draft_top_depth", None)
                app = self.app(round_trip(project))
                self.phase(app, "phase5")
                if value is None:
                    self.assertEqual(next(w for w in app.number_input if w.label.startswith("Top of cement depth")).value, 0.0)
                    self.assertNotEqual(compute_phase_status(app.session_state)["phase5"]["level"], "ok")
                elif value == 2100.0:
                    self.assertEqual(next(w for w in app.number_input if w.label.startswith("Top of cement depth")).value, value)
                    self.assertEqual(app.session_state["cement_params"]["Main"]["top_depth"], value)
                    self.export(app)
                else:
                    repair = next(w for w in app.text_input if w.label.startswith("Correct Top depth"))
                    self.assertEqual(app.session_state["cement_params"]["Main"]["top_depth"], value)
                    self.assertNotEqual(compute_phase_status(app.session_state)["phase5"]["level"], "ok")
                    repair.set_value("2100").run()
                    next(b for b in app.button if b.label == "Apply corrected values").click().run()
                    self.healthy(app)
                    self.assertEqual(app.session_state["cement_params"]["Main"]["top_depth"], 2100.0)
                    self.export(app)

    def test_batch3_top_ownership_after_job_changes(self):
        import phase_10_procedure as report
        for index, job in enumerate(BATCH1_JOBS):
            with self.subTest(job=job):
                app = self.configured_app(job)
                self.export(app)
                next_job = BATCH1_JOBS[(index + 2) % len(BATCH1_JOBS)]
                for marker in ("absent", None):
                    with self.subTest(legacy_owner=marker):
                        legacy = round_trip(app.session_state.to_dict())
                        legacy["cement_params"]["Main"].pop("top_job_type", None)
                        if marker is None:
                            legacy["cement_params"]["Main"]["top_job_type"] = None
                        old = self.app(legacy)
                        self.assertEqual(compute_phase_status(old.session_state)["phase5"]["level"], "ok")
                        old.selectbox(key="_w_job_type").set_value(next_job)
                        old.radio(key="_app_mode_key").set_value("phase10").run()
                        self.assertEqual(old.session_state["cement_params"]["Main"]["top_job_type"], job)
                        old.session_state["placement_config"]["job_type"] = next_job
                        self.assertEqual(compute_phase_status(old.session_state)["phase5"]["level"], "warning")
                        self.assert_export_blocked(old)
                self.phase(app, "phase1")
                app.selectbox(key="_w_job_type").set_value(next_job)
                app.radio(key="_app_mode_key").set_value("phase10").run()
                self.assertEqual(app.session_state["cement_params"]["Main"]["top_job_type"], job)
                self.assertEqual(compute_phase_status(app.session_state)["phase5"]["level"], "warning")
                self.assert_export_blocked(app)
                # Reviewing the new job target alone cannot bless the old top.
                self.phase(app, "phase2_3")
                next(w for w in app.selectbox if w.label == "Target shoe / treatment depth source").set_value("__manual__").run()
                next(w for w in app.number_input if w.label == "Measured target depth (m MD)").set_value(3000.0).run()
                if "TIE BACK" in next_job:
                    next(w for w in app.selectbox if w.label.startswith("Tie-back host")).select_index(1).run()
                self.assertEqual(compute_phase_status(app.session_state)["phase2_3"]["level"], "ok")
                self.assertEqual(compute_phase_status(app.session_state)["phase5"]["level"], "warning")
                restored = round_trip(app.session_state.to_dict())
                app = self.app(restored)
                self.assert_export_blocked(app)
                with patch.object(report.st, "session_state", deepcopy(restored)):
                    with self.assertRaises(ValueError):
                        report.build_master_context()
                self.phase(app, "phase5")
                mode = next(w for w in app.selectbox if w.label == "Top of cement - Main")
                self.assertEqual(mode.value, "Not entered")
                mode.set_value("Depth (m MD)").run()
                next(w for w in app.number_input if w.label.startswith("Top of cement depth")).set_value(2100.0).run()
                self.assertEqual(app.session_state["cement_params"]["Main"]["top_job_type"], next_job)
                self.assertEqual(compute_phase_status(app.session_state)["phase5"]["level"], "ok")
                self.export(app)

    def test_batch3_top_limits_match_export_readiness(self):
        import phase_10_procedure as report
        for job in BATCH1_JOBS:
            saved = round_trip(self.configured_app(job).session_state.to_dict())
            for depth, accepted in ((2999.0, True), (3000.0, False), (3001.0, False), (0.0, False),
                                    (-10.0, False), ("2000-2100", False), ("NaN", False)):
                with self.subTest(job=job, top=depth):
                    project = deepcopy(saved)
                    project["cement_params"]["Main"].update(top_mode="Depth (m MD)", top_depth=depth, top_job_type=job)
                    app = self.app(round_trip(project))
                    self.assertEqual(compute_phase_status(app.session_state)["phase5"]["level"], "ok" if accepted else "warning")
                    if accepted:
                        self.export(app)
                        with patch.object(report.st, "session_state", deepcopy(app.session_state.to_dict())):
                            context = report.build_master_context()
                        self.assertEqual(context["slurries"][0]["top"], "2999.0 m")
                        self.assertEqual(context["slurries"][0]["bottom"], "3000.0 m MD")
                    else:
                        self.assert_export_blocked(app)
                        with patch.object(report.st, "session_state", deepcopy(app.session_state.to_dict())):
                            with self.assertRaises(ValueError):
                                report.build_master_context()
            # Legacy approved depths without top_mode/ownership remain supported.
            project = deepcopy(saved)
            top = project["cement_params"]["Main"]
            top.pop("top_mode", None)
            top.pop("top_job_type", None)
            top["top_depth"] = 2100.0
            app = self.app(round_trip(project))
            self.assertEqual(compute_phase_status(app.session_state)["phase5"]["level"], "ok")
            self.export(app)
            project["placement_config"]["manual_depth_m"] = 0.0
            self.assertEqual(compute_phase_status(project)["phase5"]["level"], "warning")

    def test_batch3_chained_word_placement_matches_summary(self):
        import phase_10_procedure as report
        for job in ('CSG 9 5/8"', 'LNR 7"'):
            valid = self.configured_app(job)
            self.export(valid)  # Single-slurry behavior remains valid.
            saved = round_trip(valid.session_state.to_dict())
            for slurries, depths in ((["Lead", "Main"], {"Main": 2000.0, "Lead": 1000.0}),
                                     (["Lead #1", "Lead #2", "Main"], {"Main": 2000.0, "Lead #2": 1000.0, "Lead #1": 500.0})):
                with self.subTest(job=job, slurries=slurries):
                    project = deepcopy(saved)
                    project["fluids_config"]["active"] = [*slurries, "Displacement Fluid"]
                    for slurry in slurries:
                        project["fluids_config"]["params"][slurry] = deepcopy(project["fluids_config"]["params"]["Main"])
                        project["cement_params"][slurry] = deepcopy(saved["cement_params"]["Main"])
                        project["cement_params"][slurry].update(top_mode="Depth (m MD)", top_depth=depths[slurry], top_job_type=job)
                        project["cement_additives_dfs"][slurry] = project["cement_additives_dfs"]["Main"].copy(deep=True)
                    app = self.app(round_trip(project))
                    for phase in ("phase2_3", "phase4", "phase5", "phase7"):
                        self.phase(app, phase)
                    for _ in slurries:
                        confirms = [b for b in app.button if b.label == "Confirm measured lab results"]
                        if confirms:
                            confirms[0].click().run()
                    self.assertEqual(compute_phase_status(app.session_state)["phase5"]["level"], "ok")
                    app = self.app(round_trip(app.session_state.to_dict()))
                    for phase in ("phase2_3", "phase4", "phase5", "phase7"):
                        self.phase(app, phase)
                    self.export(app)
                    with patch.object(report.st, "session_state", deepcopy(app.session_state.to_dict())):
                        context = report.build_master_context()
                    expected, bottom = [], 3000.0
                    for slurry in reversed(slurries):
                        expected.append((slurry, bottom, depths[slurry]))
                        bottom = depths[slurry]
                    payload = [(p["name"], float(p["bottom"].split()[0]), float(p["top"].split()[0]))
                               for p in context["slurries"]]
                    # Chained upper bottoms use a top label, without the MD suffix.
                    summary = [(name.title(), float(bottom), float(top)) for name, bottom, top in re.findall(
                        r"(main|lead(?: #\d)?) cement slurry from ([\d.]+) m(?: MD)? to ([\d.]+) m", context["exec_summary"])]
                    self.assertEqual(payload, expected)
                    self.assertEqual(summary, expected)
                    self.assertIn("target depth is 3000.0 m MD", context["exec_summary"])
                    doc = Document(BytesIO(app.session_state["_compiled_doc_bytes"]))
                    word = []
                    for table in doc.tables:
                        title = table.rows[0].cells[0].text
                        if title.endswith(" Cement Slurry Data"):
                            row = next(r for r in table.rows if r.cells[0].text == "Bottom")
                            word.append((title.removesuffix(" Cement Slurry Data"),
                                         float(row.cells[1].text.split()[0]), float(row.cells[3].text.split()[0])))
                    self.assertEqual(word, expected)
                    # A top below the previous slurry's top must fail both gates.
                    invalid = round_trip(app.session_state.to_dict())
                    invalid["cement_params"][slurries[0]]["top_depth"] = 2500.0
                    bad = self.app(invalid)
                    self.assertEqual(compute_phase_status(bad.session_state)["phase5"]["level"], "warning")
                    self.assert_export_blocked(bad)
        # Sibling jobs retain independent target-based bottoms, without chaining.
        from placement import slurry_intervals
        for job in ('CMT PLUG', 'CMT SQUEEZE', 'TIE BACK LNR 7"'):
            intervals = slurry_intervals(None, job, {"job_type": job, "target_row": "__manual__", "manual_depth_m": 3000.0},
                                         ["Lead", "Main"], {"Lead": {"top_mode": "Depth (m MD)", "top_depth": 1000.0},
                                                            "Main": {"top_mode": "Depth (m MD)", "top_depth": 2000.0}})
            self.assertEqual(list(intervals), ["Lead", "Main"])
            self.assertTrue(all(p["bottom_depth"] == 3000.0 for p in intervals.values()))

    def test_batch3_note_and_procedure_use_total_water(self):
        import phase_10_procedure as report
        from phase_5_cement import build_components, calculate_base_results
        for job in BATCH1_JOBS:
            valid = self.configured_app(job, additives=True)
            saved = round_trip(valid.session_state.to_dict())
            for auto, dead in ((True, 20.0), (False, 20.0), (False, 0.0)):
                with self.subTest(job=job, automatic=auto, dead_volume=dead):
                    project = deepcopy(saved)
                    p = project["cement_params"]["Main"]
                    p.update(auto_calc=auto, dead_vol=dead)
                    if not auto:
                        p.update(mix_water=30.0, manual_mix_water=30.0, manual_yield=1.5)
                        p["yield"] = 1.5
                        project["cement_additives_dfs"]["Main"].loc[0, "User Input"] = 60.0 / 187.2
                    app = self.app(round_trip(project))
                    self.phase(app, "phase5")
                    self.phase(app, "phase7")
                    keeps = [b for b in app.button if b.label == "Keep reviewed lab entries"]
                    if keeps:
                        keeps[0].click().run()
                        next(b for b in app.button if b.label == "Confirm measured lab results").click().run()
                    p = deepcopy(app.session_state["cement_params"]["Main"])
                    adds = app.session_state["cement_calc_Main"].copy(deep=True)
                    formulation = app.session_state["cement_additives_dfs"]["Main"]
                    fluid = app.session_state["fluid_data"]["Main"]
                    engine = calculate_base_results(deepcopy(p), fluid["volume"], float(fluid["effective_density"]), *build_components(formulation))
                    if auto:
                        self.assertEqual(p["mix_water"], round_half_up(engine["field_water_bbl"], 1))
                    else:
                        self.assertEqual(p["mix_water"], 30.0)
                        self.assertEqual(p["yield"], 1.5)
                        self.assertEqual(p["total_sacks"], 187.2)
                    self.assertEqual(p["dead_vol"], dead)
                    total = round_half_up(p["mix_water"] + dead, 1)
                    concentration = float(adds.iloc[0]["(lbs or gal)/bbl"].split()[0])
                    amount = float(adds.iloc[0]["lbs or gal (with dead Vol.)"].split()[0])
                    self.assertEqual(amount, round_half_up((p["mix_water"] + dead) * concentration, 2))
                    if not auto:
                        self.assertEqual(concentration, 2.0)
                        self.assertEqual(amount, 100.0 if dead else 60.0)
                    self.assertIn(f"**{total:.1f} bbl**", app.session_state["cement_note_Main"])
                    self.export(app)
                    with patch.object(report.st, "session_state", deepcopy(app.session_state.to_dict())):
                        context = report.build_master_context()
                    payload = context["slurries"][0]
                    self.assertEqual(payload["mix_water_bbl"], p["mix_water"])
                    self.assertEqual(payload["dead_vol_bbl"], dead)
                    self.assertEqual(payload["total_water_bbl"], total)
                    self.assertIn(f"{total:.1f} bbl Fresh Water", payload["note_mix"])
                    water_line = next(line for line in context["procedure_text"].splitlines() if "Fill up" in line)
                    self.assertIn(f"{total:.1f} bbl", water_line)
                    self.assertEqual(app.session_state["cement_params"]["Main"], p)
                    pd.testing.assert_frame_equal(app.session_state["cement_calc_Main"], adds)

    def test_batch4_word_units_and_surface(self):
        import phase_10_procedure as report
        for job in BATCH1_JOBS:
            with self.subTest(job=job):
                app = self.configured_app(job, optional_fluids=True)
                self.phase(app, "phase5")
                next(w for w in app.selectbox if w.label == "Top of cement - Main").set_value("Depth (m MD)").run()
                next(w for w in app.number_input if w.label == "Top of cement depth (m MD) - Main").set_value(1000.0).run()
                self.export(app)
                doc = Document(BytesIO(app.session_state["_compiled_doc_bytes"]))
                slurry = next(t for t in doc.tables if t.rows[0].cells[0].text == "Main Cement Slurry Data")
                placement = next(r for r in slurry.rows if r.cells[0].text == "Bottom")
                self.assertEqual(placement.cells[1].text, "3000.0 m MD")
                self.assertEqual(placement.cells[3].text, "1000.0 m")
                self.assertEqual(slurry.rows[1].cells[1].text, "50.0 bbl")
                self.assertEqual(slurry.rows[1].cells[3].text, "118.0 pcf")
                self.assertTrue(slurry.rows[4].cells[1].text.endswith(" cuft/sk"))
                spacer = next(t for t in doc.tables if t.rows[0].cells[0].text == "Spacer Formulation Data")
                self.assertEqual(spacer.rows[1].cells[4].text, "80.0 pcf")
                self.assertEqual(spacer.rows[3].cells[3].text, "50.0 bbl @ 80.0 pcf")
                preflush = next(t for t in doc.tables if t.rows[0].cells[0].text == "Pre Flush Data")
                self.assertEqual(preflush.rows[1].cells[2].text, "50.0 bbl")
                self.assertEqual(preflush.rows[1].cells[4].text.count("pcf"), 1)
                word_text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
                self.assertNotRegex(word_text, r"\bm MD m\b|\bm m\b|\bpcf pcf\b")
                with patch.object(report.st, "session_state", deepcopy(app.session_state.to_dict())):
                    context = report.build_master_context()
                self.assertIn("3000.0 m MD", context["exec_summary"])
                self.assertEqual(context["slurries"][0]["top"], "1000.0 m")
                # Surface is a representation, not a numeric depth with a unit.
                project = round_trip(app.session_state.to_dict())
                project["cement_params"]["Main"].update(top_mode="Surface", top_depth=0.0)
                app = self.app(round_trip(project))
                for phase in ("phase2_3", "phase4", "phase5", "phase6", "phase7"):
                    self.phase(app, phase)
                self.export(app)
                doc = Document(BytesIO(app.session_state["_compiled_doc_bytes"]))
                slurry = next(t for t in doc.tables if t.rows[0].cells[0].text == "Main Cement Slurry Data")
                placement = next(r for r in slurry.rows if r.cells[0].text == "Bottom")
                self.assertEqual(placement.cells[1].text, "3000.0 m MD")
                self.assertEqual(placement.cells[3].text, "surface")

    def test_batch4_single_slurry_water_label(self):
        import phase_10_procedure as report
        for job in BATCH1_JOBS:
            saved = round_trip(self.configured_app(job).session_state.to_dict())
            for slurry in ("Main", "Lead", "Tail"):
                with self.subTest(job=job, slurry=slurry):
                    project = deepcopy(saved)
                    project["fluids_config"]["active"] = [slurry, "Displacement Fluid"]
                    project["fluids_config"]["params"][slurry] = deepcopy(saved["fluids_config"]["params"]["Main"])
                    for key in ("cement_params", "cement_additives_dfs"):
                        project[key][slurry] = deepcopy(saved[key]["Main"])
                    app = self.app(round_trip(project))
                    for phase in ("phase2_3", "phase4", "phase5", "phase7"):
                        self.phase(app, phase)
                    confirms = [b for b in app.button if b.label == "Confirm measured lab results"]
                    if confirms:
                        confirms[0].click().run()
                        self.healthy(app)
                    self.export(app)
                    with patch.object(report.st, "session_state", deepcopy(app.session_state.to_dict())):
                        context = report.build_master_context()
                    self.assertEqual([p["name"] for p in context["slurries"]], [slurry])
                    water_line = next(line for line in context["procedure_text"].splitlines() if "Fill up" in line)
                    prefix = "PLUG" not in job and not ("LNR" in job and "TIE BACK" not in job)
                    total = context["slurries"][0]["total_water_bbl"]
                    expected = (f"with {total:.1f} bbl fresh water" if "PLUG" in job else
                                f"Fresh water ({slurry + ': ' if prefix else ''}{total:.1f} bbl)")
                    self.assertIn(expected, water_line)
                    if slurry != "Main":
                        self.assertNotIn("Main:", water_line)
                    doc = Document(BytesIO(app.session_state["_compiled_doc_bytes"]))
                    word_line = next(p.text for p in doc.paragraphs if "Fill up" in p.text)
                    self.assertIn(expected, word_line)

    def test_batch4_spacer_zero_positive_and_missing_dosage(self):
        for job in BATCH1_JOBS:
            saved = round_trip(self.configured_app(job, optional_fluids=True).session_state.to_dict())
            for dosage in (0.0, 1.25, None):
                with self.subTest(job=job, dosage=dosage):
                    project = deepcopy(saved)
                    project["spacer_dfs"]["Spacer"].loc[0, "User Input (% or gal)"] = dosage
                    app = self.app(round_trip(project))
                    for phase in ("phase2_3", "phase4", "phase5", "phase6", "phase7"):
                        self.phase(app, phase)
                    amount = app.session_state["spacer_dfs"]["Spacer"].iloc[0]["User Input (% or gal)"]
                    if dosage is None:
                        self.assertTrue(pd.isna(amount))
                        self.assertEqual(compute_phase_status(app.session_state)["phase6"]["level"], "warning")
                        self.assert_export_blocked(app)
                        continue
                    self.assertEqual(amount, dosage)
                    self.assertEqual(compute_phase_status(app.session_state)["phase6"]["level"], "ok")
                    before = app.session_state["spacer_dfs"]["Spacer"].copy(deep=True)
                    self.export(app)
                    doc = Document(BytesIO(app.session_state["_compiled_doc_bytes"]))
                    spacer = next(t for t in doc.tables if t.rows[0].cells[0].text == "Spacer Formulation Data")
                    chemical = next(r for r in spacer.rows if r.cells[1].text == "Spacer" and r.cells[2].text == "-")
                    self.assertEqual(chemical.cells[3].text, str(dosage))
                    pd.testing.assert_frame_equal(app.session_state["spacer_dfs"]["Spacer"], before)

if __name__ == "__main__":
    unittest.main()
