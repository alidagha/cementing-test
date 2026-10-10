"""Round 6: fluid roles and the single well-level BHCT authority."""
from copy import deepcopy
from unittest.mock import patch
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile
import ast
import json
import re
import subprocess

import pandas as pd
import pytest
from docx import Document

import materials_db
import phase_10_procedure as report
from engineering_tools import (lab_review_signature, catalog_cement_sg,
                               compute_phase_status, validate_lab_rheology)
from phase_5_cement import REQUIRED_ADDITIVE_COLUMNS, refresh_cement_calculations
from phase_7_lab import build_lab_df_from_phase5
from project_io import decode_project
from project_state import (refresh_well_derived, lab_source_signature, refresh_fluids,
                           refresh_lab_payloads, prepare_calculations, purge_inactive_slurries)
import test_audit_regressions as audit

PLACEMENT = ("Main", "Lead", "Lead #1", "Lead #2", "Tail")
ORDER = ("Pre Flush", "Spacer", "Spacer Ahead", "Scavenger", *PLACEMENT,
         "Spacer Behind", "Displacement Fluid")


def state_for(*names):
    """Explicit current-state prerequisites; no production defaults or migration."""
    state = {
        **audit.well_profile_fixture(),
        "job_type": 'CSG 9 5/8"', "well_name": "Round 6", "client": "NIDC",
        "geo_md": 3000.0, "geo_tvd": 3000.0, "bhst": 200, "bhct": 150,
        "bhsp": "5603.88", "mud_density": "80-82", "plastic_viscosity": "45", "yield_point": "15",
        "well_data": {"geo_md": 3000.0, "geo_tvd": 3000.0, "bhst": 200,
                      "bhct": 150, "bhsp": "5603.88", "mud_density": "80-82", "geo_gradient": 1.2},
        "material_property_schema": 1,
        "placement_config": {"job_type": 'CSG 9 5/8"', "target_row": "__manual__", "manual_depth_m": 3000.0},
        "fluids_config": {"active": list(names), "params": {}},
        "cement_params": {}, "cement_additives_dfs": {},
        "lab_qc_params": {}, "lab_grid_dfs": {}, "lab_source_signatures": {},
    }
    for name in names:
        state["fluids_config"]["params"][name] = {
            "volume": 30.0, "density": "118", "pump_rate": "3-5", "material_name": "Cement Slurry"}
        params = {"base_cement": "Cement G Delijan", "cmt_sg": catalog_cement_sg(),
                  "cmt_sg_source": "catalog", "auto_calc": True,
                  "dead_vol": 2.0, "tank_name": "Cement Unit Tanks"}
        if name in PLACEMENT:
            params.update(top_mode="Surface", top_depth=None)
        state["cement_params"][name] = params
        frame = pd.DataFrame(columns=REQUIRED_ADDITIVE_COLUMNS + ["Density"])
        state["cement_additives_dfs"][name] = frame
        state["lab_grid_dfs"][name] = build_lab_df_from_phase5(
            frame, base_cement=params["base_cement"], slurry_weight_pcf=118,
            slurry_volume_bbl=30, cmt_sg=params["cmt_sg"], recalculate_mass=True)
        qc = {"rheology": {"bhct_down": deepcopy(audit.rheology_fixture()["surface_down"])}}
        if name in PLACEMENT:
            qc.update(api_fl=4.0, free_water=0.0, free_water_45=0.0,
                      surface_hardened_hours=8.0, thickening_time="03:30",
                      **audit.thickening_fixture(), **audit.compressive_fixture())
        state["lab_qc_params"][name] = qc
    refresh_fluids(state)
    for name in names:
        qc = state["lab_qc_params"][name]
        qc["reviewed"] = True
        qc["review_signature"] = lab_review_signature(qc, state["lab_grid_dfs"][name])
        state["lab_source_signatures"][name] = lab_source_signature(state, name)
    return state


def attach_scavenger(project):
    extra = state_for("Scavenger")
    project["fluids_config"]["active"].append("Scavenger")
    project["fluids_config"]["params"]["Scavenger"] = extra["fluids_config"]["params"]["Scavenger"]
    for key in ("cement_params", "cement_additives_dfs", "lab_grid_dfs", "lab_qc_params"):
        project[key]["Scavenger"] = deepcopy(extra[key]["Scavenger"])
    refresh_fluids(project)
    project["lab_source_signatures"]["Scavenger"] = lab_source_signature(project, "Scavenger")
    return project


def test_central_roles_and_hydraulic_order():
    assert tuple(materials_db.FLUID_TYPES) == ORDER
    assert tuple(getattr(materials_db, "PLACEMENT_SLURRIES", ())) == PLACEMENT
    assert tuple(getattr(materials_db, "CEMENT_FORMULATION_FLUIDS", ())) == ("Scavenger", *PLACEMENT)
    assert materials_db.DEFAULT_MATERIAL_NAMES["Scavenger"] == "Cement Slurry"


def test_one_bhct_editor_in_well_phase_and_none_in_lab():
    case = audit.AuditRegressions()
    app = case.app(state_for("Main"))
    case.phase(app, "phase2_3")
    assert len([w for w in app.number_input if "BHCT" in w.label]) == 1
    case.phase(app, "phase7")
    assert not [w for w in app.number_input if "BHCT" in w.label]


@pytest.mark.parametrize("inactive", [False, True])
def test_current_json_rejects_qc_bhct(inactive):
    project = {"job_type": 'CSG 9 5/8"', "bhct": 150,
               "well_data": {"bhct": 150}, "lab_qc_params": {"Lead": {"bhct": 150}}}
    if inactive:
        project["inactive_slurry_drafts"] = {"Lead": {"lab_qc_params": project.pop("lab_qc_params")["Lead"]}}
    with pytest.raises(ValueError, match="BHCT|bhct"):
        decode_project(json.dumps(project).encode(), audit.namespace["deserialize_item"])


def test_well_bhct_participates_in_every_lab_source():
    state = state_for("Main", "Scavenger")
    signatures = {name: lab_source_signature(state, name) for name in ("Main", "Scavenger")}
    state["bhct"] = state["well_data"]["bhct"] = 155
    for name, previous in signatures.items():
        assert lab_source_signature(state, name) != previous


def test_scavenger_formulation_uses_existing_engine_without_placement():
    state = state_for("Main", "Scavenger")
    assert refresh_cement_calculations(state) == []
    for key in ("yield", "mix_water", "solution", "total_sacks", "base_fluid_gal_sk",
                "mix_water_gal_sk", "mix_fluid_gal_sk"):
        assert state["cement_params"]["Scavenger"][key] == state["cement_params"]["Main"][key]
    assert not ({"top_mode", "top_depth", "bottom", "excess_oh", "excess_csg"}
                & state["cement_params"]["Scavenger"].keys())


def test_scavenger_internal_lab_needs_only_shared_conditions_and_rheology():
    state = state_for("Scavenger")
    assert refresh_cement_calculations(state) == []
    assert refresh_lab_payloads(state) == []
    payload = state["lab_payload_Scavenger"]
    assert (payload["bhct"], payload["bhst"], payload["bhsp"]) == (150, 200, "5603.88")
    assert not ({"compressive", "thickening_time", "api_fl", "free_water", "surface_hardened_hours"}
                & state["lab_qc_params"]["Scavenger"].keys())


def test_scavenger_requires_bhct_ramp_down():
    state = state_for("Scavenger")
    qc = state["lab_qc_params"]["Scavenger"]
    qc["rheology"] = audit.rheology_fixture()
    qc["review_signature"] = lab_review_signature(qc, state["lab_grid_dfs"]["Scavenger"])
    assert any("BHCT" in issue and "Ramp-Down" in issue for issue in refresh_lab_payloads(state))


def test_direct_rebuild_uses_well_bhct():
    state = state_for("Main")
    assert prepare_calculations(state) == []
    assert state["lab_payload_Main"]["bhct"] == 150
    assert "bhct" not in state["lab_qc_params"]["Main"]


def test_template_has_exact_three_scavenger_program_tables():
    doc = Document(report.TEMPLATE_PATH)
    tables = [table for table in doc.tables if "Scavenger" in table.rows[0].cells[0].text]
    assert len(tables) == 3
    assert len(tables[0].rows) == 3
    assert not any(text in cell.text for row in tables[0].rows for cell in row.cells
                   for text in ("Excess", "Bottom", "Top"))


def test_scavenger_does_not_change_summary_or_its_source_signature():
    state = state_for("Main")
    with patch.object(report.st, "session_state", state):
        before = report.synchronize_report_texts()["exec_summary_text"]
        extra = state_for("Scavenger")
        state["fluids_config"]["active"].insert(0, "Scavenger")
        for key in ("fluid_data", "cement_params", "cement_additives_dfs"):
            state[key]["Scavenger"] = extra[key]["Scavenger"]
        state["fluids_config"]["params"]["Scavenger"] = extra["fluids_config"]["params"]["Scavenger"]
        refresh_fluids(state)
        after = report.synchronize_report_texts()["exec_summary_text"]
    assert after["generated"] == before["generated"]
    assert after["signature"] == before["signature"]


def test_scavenger_is_unselected_by_default_and_has_normal_fluid_widgets():
    case = audit.AuditRegressions()
    app = case.app({"mud_density": "80", "effective_mud_density": 80.0})
    case.phase(app, "phase4")
    assert not app.checkbox(key="chk_fluid_Scavenger").value
    app.checkbox(key="chk_fluid_Scavenger").check().run()
    assert any(w.label == "Volume (bbl) - Scavenger" for w in app.number_input)
    assert any(w.label == "Rate (bpm) - Scavenger" for w in app.text_input)
    assert any(w.label == "Density (pcf) - Scavenger" for w in app.text_input)
    assert app.session_state["fluids_config"]["params"]["Scavenger"]["material_name"] == "Cement Slurry"


def test_canonical_refresh_sorts_full_hydraulic_order_and_uses_minimum_rate():
    state = state_for(*reversed(materials_db.CEMENT_FORMULATION_FLUIDS))
    for name in ORDER:
        state["fluids_config"]["params"].setdefault(name, {
            "volume": 30.0, "density": "80", "pump_rate": "3-5",
            "material_name": materials_db.DEFAULT_MATERIAL_NAMES[name]})
    state["fluids_config"]["active"] = list(reversed(ORDER))
    refresh_fluids(state)
    assert tuple(state["fluids_config"]["active"]) == tuple(state["fluid_data"]) == ORDER
    for i, name in enumerate(ORDER, 1):
        fluid = state["fluid_data"][name]
        assert fluid["min_rate"] == 3.0
        assert fluid["duration_min"] == 10.0
        assert fluid["cumul_time_min"] == i * 10.0


@pytest.mark.parametrize("rows", [
    [{"Material Type": "Local Powder", "Name": "P-X", "Physical State": "Powder", "Mix Method": "Dry Blend", "User Input": 5.0, "Density": 2.2}],
    [{"Material Type": "Local Powder", "Name": "P-X", "Physical State": "Powder", "Mix Method": "In Mix Water", "User Input": 2.0, "Density": 2.2}],
    [{"Material Type": "Local Liquid", "Name": "L-X", "Physical State": "Liquid", "Mix Method": "Dry Blend", "User Input": 0.1, "Density": 1.15}],
    [{"Material Type": "Nacl", "Name": "SALT", "Physical State": "Powder", "Mix Method": "In Mix Water", "User Input": 18.0, "Density": None}],
])
def test_scavenger_engine_and_lab_parity_for_dry_wet_liquid_and_salt(rows):
    state = state_for("Main", "Scavenger")
    for name in ("Main", "Scavenger"):
        state["cement_additives_dfs"][name] = pd.DataFrame(rows)
    assert refresh_cement_calculations(state) == []
    for key in ("yield", "mix_water", "solution", "total_sacks", "base_fluid_gal_sk",
                "mix_water_gal_sk", "mix_fluid_gal_sk"):
        assert state["cement_params"]["Scavenger"][key] == state["cement_params"]["Main"][key]
    assert state["cement_calc_Main"].equals(state["cement_calc_Scavenger"])
    assert state["cement_blend_Main"].equals(state["cement_blend_Scavenger"])
    grids = [build_lab_df_from_phase5(state["cement_additives_dfs"][name],
                                    base_cement="Cement G Delijan", slurry_weight_pcf=118,
                                    slurry_volume_bbl=30, cmt_sg=catalog_cement_sg(),
                                    recalculate_mass=True) for name in ("Main", "Scavenger")]
    assert grids[0].equals(grids[1])


@pytest.mark.parametrize("job", audit.BATCH1_JOBS)
def test_scavenger_phase5_readiness_needs_formulation_not_placement(job):
    state = state_for("Scavenger")
    state["job_type"] = job
    assert refresh_cement_calculations(state) == []
    assert compute_phase_status(state)["phase5"]["level"] == "ok"
    assert not state["cement_params"]["Scavenger"].get("top_mode")


@pytest.mark.parametrize("optional", ["surface_down", "bhct_up"])
def test_scavenger_optional_rheology_draft_only_blocks_when_selected(optional):
    state = state_for("Scavenger")
    qc = state["lab_qc_params"]["Scavenger"]
    draft = {"selected": False, "readings": {"300": "168"}, "gel_10_sec": "", "gel_10_min": ""}
    qc["rheology"][optional] = draft
    assert validate_lab_rheology(qc, "Scavenger")["bhct_down"]["is_valid_bingham"]
    draft["selected"] = True
    with pytest.raises(ValueError): validate_lab_rheology(qc, "Scavenger")
    draft["selected"] = False
    assert validate_lab_rheology(qc, "Scavenger")["bhct_down"]["is_valid_bingham"]
    assert draft["readings"] == {"300": "168"}


def test_scavenger_ui_has_no_placement_or_full_lab_controls():
    case = audit.AuditRegressions()
    app = case.app(state_for("Main", "Scavenger"))
    case.phase(app, "phase5")
    assert any(w.label == "Cement SG - Scavenger" for w in app.number_input)
    assert any(w.label == "Top of cement - Main" for w in app.selectbox)
    assert not [w for w in app.selectbox if w.label == "Top of cement - Scavenger"]
    assert not [w for w in app.number_input if "Top of cement" in w.label and "Scavenger" in w.label]
    case.phase(app, "phase7")
    for widget in (*app.number_input, *app.text_input, *app.checkbox):
        if "Scavenger" in widget.label:
            assert not any(label in widget.label for label in ("Free Water", "Filtrate", "Surface Sample", "70 Bc", "Cell", "CS @", "Force", "Crush Test", "UCA"))
    assert "compressive" not in app.session_state["lab_qc_params"]["Scavenger"]


@pytest.mark.parametrize("field,value", [("bhct", 155), ("bhst", 205), ("bhsp", "7300+1000")])
def test_shared_well_source_edit_stales_all_six_fluids(field, value):
    state = state_for(*materials_db.CEMENT_FORMULATION_FLUIDS)
    previous = deepcopy(state["lab_source_signatures"])
    state[field] = state["well_data"][field] = value
    if field == "bhst":
        state["geothermal_config"]["value"] = value
        refresh_well_derived(state)
    for name in materials_db.CEMENT_FORMULATION_FLUIDS:
        assert lab_source_signature(state, name) != previous[name]
    assert len(refresh_lab_payloads(state)) == 6


@pytest.mark.parametrize("field", ["flat", "nested"])
@pytest.mark.parametrize("value", [True, "150", float("inf"), float("nan"), [], {}])
def test_well_bhct_trust_boundary_rejects_malformed_values(field, value):
    project = {"job_type": 'CSG 9 5/8"'}
    if field == "flat": project["bhct"] = value
    else: project["well_data"] = {"bhct": value}
    with pytest.raises(ValueError):
        decode_project(json.dumps(project).encode(), audit.namespace["deserialize_item"])


def test_scavenger_inactive_draft_round_trip_and_transient_cleanup():
    state = state_for("Main", "Scavenger")
    draft = deepcopy(state["lab_qc_params"]["Scavenger"])
    frame = state["cement_additives_dfs"]["Scavenger"].copy()
    state["_rheo_scavenger_bhct_down_300_old"] = "STALE"
    state["_rheo_main_bhct_down_300_old"] = "MAIN"
    state["_editor_source__editor_additives_scavenger_old"] = frame
    state["fluids_config"]["active"] = ["Main"]
    purge_inactive_slurries(state, ["Main"])
    assert "Scavenger" not in state["lab_qc_params"]
    assert "_rheo_scavenger_bhct_down_300_old" not in state
    assert "_editor_source__editor_additives_scavenger_old" not in state
    assert state["_rheo_main_bhct_down_300_old"] == "MAIN"
    restored = audit.round_trip(state)
    assert restored["inactive_slurry_drafts"]["Scavenger"]["lab_qc_params"] == draft
    restored["fluids_config"]["active"] = ["Scavenger", "Main"]
    purge_inactive_slurries(restored, restored["fluids_config"]["active"])
    assert restored["lab_qc_params"]["Scavenger"] == draft
    assert restored["cement_additives_dfs"]["Scavenger"].equals(frame)
    assert "Scavenger" not in restored["lab_source_signatures"]
    assert any("Scavenger" in issue for issue in prepare_calculations(restored))


def test_first_well_bhct_edit_navigation_restore_and_document_invalidation():
    case = audit.AuditRegressions()
    project = state_for("Main", "Scavenger")
    app = case.app(project)
    app.session_state["_compiled_doc_bytes"] = b"obsolete"
    case.phase(app, "phase2_3")
    app.number_input(key="_w_bhct").set_value(155.0)
    case.phase(app, "phase1")
    assert app.session_state["well_data"]["bhct"] == app.session_state["bhct"] == 155
    assert "_compiled_doc_bytes" not in app.session_state
    assert compute_phase_status(app.session_state)["phase7"]["level"] != "ok"
    restored = case.app(audit.round_trip(app.session_state.to_dict()))
    case.phase(restored, "phase2_3")
    assert restored.number_input(key="_w_bhct").value == 155
    case.phase(restored, "phase7")
    assert not [w for w in restored.number_input if "BHCT" in w.label]
    assert not any("bhct" in qc for qc in restored.session_state["lab_qc_params"].values())


def test_suggestion_cannot_apply_a_non_bhct_regime():
    import phase_2_3_well_data as well
    state = state_for("Main")
    with patch.object(well.st, "session_state", state), patch.object(well, "suggest_bhct", return_value={"kind": "SQUEEZE", "temp_degF": 999}):
        assert well._bhct_suggestion() is None
        well._apply_bhct_suggestion()
    assert state["bhct"] == state["well_data"]["bhct"] == 150


def test_template_package_is_unchanged_outside_exact_section_vi_insertion():
    original = BytesIO(subprocess.check_output(["git", "show", "80353e4:master_template.docx"]))
    with ZipFile(original) as old, ZipFile(report.TEMPLATE_PATH) as new:
        assert old.namelist() == new.namelist()
        for name in old.namelist():
            if name != "word/document.xml": assert old.read(name) == new.read(name)
        before, after = old.read("word/document.xml"), new.read("word/document.xml")
    paragraphs = list(re.finditer(rb'<w:p[ >].*?</w:p>', after, re.S))
    start = next(p.start() for p in paragraphs if b'{%p if has_scavenger %}' in p.group())
    end = next(p.end() for p in paragraphs if p.start() > start and b'{%p endif %}' in p.group())
    fragment = after[start:end]
    assert fragment.count(b'<w:tbl>') == 3
    assert not any(value in fragment for value in (b'excess_oh', b'excess_csg', b'.top', b'.bottom', b'.lab'))
    assert after[:start] + after[end:] == before


def test_protected_algorithms_and_properties_are_byte_or_ast_identical():
    for filename in ("rheology.py", "bhct_helper.py", "placement.py"):
        assert Path(filename).read_bytes() == subprocess.check_output(["git", "show", "80353e4:" + filename])
    before = ast.parse(subprocess.check_output(["git", "show", "80353e4:engineering_tools.py"]))
    after = ast.parse(Path("engineering_tools.py").read_text())
    functions = ("calculate_slurry_from_components", "calculate_salt_field_amounts", "resolve_additive_density",
                 "normalize_additive_mix", "lab_temperature_valid", "validate_thickening_test", "validate_compressive_test")
    for name in functions:
        original = next(node for node in before.body if isinstance(node, ast.FunctionDef) and node.name == name)
        current = next(node for node in after.body if isinstance(node, ast.FunctionDef) and node.name == name)
        assert ast.dump(current) == ast.dump(original), name


def test_real_docx_direct_fresh_restore_roles_procedure_and_notes():
    case = audit.AuditRegressions()
    project = audit.round_trip(case.configured_app('CSG 9 5/8"').session_state.to_dict())
    project = attach_scavenger(project)
    restored = case.app(audit.round_trip(project))
    # Export directly; no Phase IV/V/VII visit after fresh restore.
    case.export(restored)
    doc = Document(BytesIO(restored.session_state['_compiled_doc_bytes']))
    with patch.object(report.st, "session_state", deepcopy(restored.session_state.to_dict())):
        context = report.build_master_context()
    assert [row["name"] for row in context["slurries"]] == ["Main"]
    assert context["has_scavenger"] and context["scavenger"]["name"] == "Scavenger"
    assert not any(key in context["scavenger"] for key in ("top", "bottom", "excess_oh", "excess_csg", "lab"))
    assert "Scavenger" not in context["executive_summary"]
    assert "Scavenger:" in context["procedure"]
    assert "Mix and pump 30.0 bbl Scavenger cement slurry (118 pcf)." in context["procedure"]
    assert context["procedure"].index("Scavenger cement slurry") < context["procedure"].index("Main cement slurry")
    names = [table.rows[0].cells[0].text for table in doc.tables]
    assert names.count("Scavenger Cement Slurry Data") == 1
    assert names.count("Scavenger Cement Slurry Blend Data") == 1
    assert names.count("Scavenger Cement Slurry Additives Data") == 1
    assert names.count("Test Basic Data") == 1 and names.count("Rheology Test") == 1
    table = next(table for table in doc.tables if table.rows[0].cells[0].text == "Scavenger Cement Slurry Data")
    assert len(table.rows) == 3 and table.rows[1].cells[1].text == "30.0 bbl"
    numbers = [int(re.match(r"NOTE (\d+):", note)[1]) for note in context["all_notes"]]
    assert numbers == [n for n in range(1, len(numbers) + 2) if n != 20][:len(numbers)]
    main_note = context["slurries"][0]["note_mix"]
    scavenger_note = context["scavenger"]["note_mix"]
    lab_note = context["slurries"][0]["lab"]["note_freshwater"]
    assert context["all_notes"].index(main_note) < context["all_notes"].index(scavenger_note) < context["all_notes"].index(lab_note)
    Path('/tmp/round6-direct.docx').write_bytes(restored.session_state['_compiled_doc_bytes'])


@pytest.mark.parametrize("defect", ["missing_down", "partial", "gel", "ty_zero"])
def test_scavenger_invalid_required_dataset_blocks_confirmation_and_direct_export(defect):
    case = audit.AuditRegressions()
    project = attach_scavenger(audit.round_trip(case.configured_app('CSG 9 5/8"').session_state.to_dict()))
    row = project["lab_qc_params"]["Scavenger"]["rheology"]["bhct_down"]
    if defect == "missing_down": row["selected"] = False
    elif defect == "partial": row["readings"] = {"300": "168"}
    elif defect == "gel": row["gel_10_sec"] = ""
    else: row["readings"] = dict(zip(map(str, (300, 200, 100, 60, 30, 6, 3)), map(str, (150, 100, 50, 30, 15, 2, 1))))
    app = case.app(project)
    case.phase(app, "phase7")
    assert next(button for button in app.button if "scavenger" in button.key and button.label == "Confirm measured lab results").disabled
    case.phase(app, "phase10")
    assert any("Scavenger" in error.value for error in app.error)
    assert "_compiled_doc_bytes" not in app.session_state
    assert not app.exception


def test_scavenger_first_formulation_and_rheology_edits_survive_navigation_and_fresh_restore():
    case = audit.AuditRegressions()
    app = case.app(state_for("Main", "Scavenger"))
    case.phase(app, "phase5")
    next(w for w in app.number_input if w.label == "Dead Vol (bbl) - Scavenger").set_value(3.25)
    case.phase(app, "phase7")
    reading = next(w for w in app.text_input if w.label == "300 RPM - BHCT — Ramp-Down - Scavenger")
    reading.set_value("170")
    case.phase(app, "phase1")
    assert app.session_state["cement_params"]["Scavenger"]["dead_vol"] == 3.25
    assert app.session_state["lab_qc_params"]["Scavenger"]["rheology"]["bhct_down"]["readings"]["300"] == "170"
    assert app.session_state["lab_qc_params"]["Main"]["rheology"]["bhct_down"]["readings"]["300"] == "168"
    assert compute_phase_status(app.session_state)["phase7"]["level"] != "ok"
    restored = case.app(audit.round_trip(app.session_state.to_dict()))
    case.phase(restored, "phase7")
    assert next(w for w in restored.text_input if w.label == reading.label).value == "170"
    case.phase(restored, "phase5")
    assert next(w for w in restored.number_input if w.label == "Dead Vol (bbl) - Scavenger").value == 3.25


def test_normal_placement_chaining_ignores_scavenger_in_new_hydraulic_order():
    from placement import slurry_intervals
    state = state_for("Tail", "Scavenger", "Lead")
    state["cement_params"]["Tail"].update(top_mode="Depth (m MD)", top_depth=2850.0)
    refresh_fluids(state)
    assert state["fluids_config"]["active"] == ["Scavenger", "Lead", "Tail"]
    intervals = slurry_intervals(state.get("hardware_table"), state["job_type"], state["placement_config"],
                                ["Lead", "Tail"], state["cement_params"])
    assert (intervals["Tail"]["top_depth"], intervals["Tail"]["bottom_depth"]) == (2850.0, 3000.0)
    assert (intervals["Lead"]["top_depth"], intervals["Lead"]["bottom_depth"]) == (0.0, 2850.0)
    assert compute_phase_status(state)["phase5"]["level"] == "ok"


def test_incomplete_scavenger_direct_phase_x_has_controlled_phase_iv_feedback():
    case = audit.AuditRegressions()
    app = case.app({"job_type": 'CSG 9 5/8"', "fluids_config": {
        "active": ["Scavenger", "Main"], "params": {"Scavenger": {"volume": None}}}})
    case.phase(app, "phase10")
    assert any("Phase IV" in error.value and "Scavenger" in error.value for error in app.error)
    assert not app.exception
    assert "_compiled_doc_bytes" not in app.session_state


def test_all_five_normal_lab_sections_share_well_conditions_and_exclude_scavenger():
    case = audit.AuditRegressions()
    project = audit.round_trip(case.configured_app('CSG 9 5/8"').session_state.to_dict())
    for i, name in enumerate(PLACEMENT):
        if name != "Main":
            project["fluids_config"]["params"][name] = deepcopy(project["fluids_config"]["params"]["Main"])
            for key in ("cement_params", "cement_additives_dfs", "lab_grid_dfs", "lab_qc_params"):
                project[key][name] = deepcopy(project[key]["Main"])
        project["cement_params"][name].update(top_mode="Surface" if i == 0 else "Depth (m MD)", top_depth=None if i == 0 else 500.0 * i)
    project["fluids_config"]["active"] = [*PLACEMENT, "Displacement Fluid"]
    attach_scavenger(project)
    for name in materials_db.CEMENT_FORMULATION_FLUIDS:
        project["lab_source_signatures"][name] = lab_source_signature(project, name)
    restored = case.app(audit.round_trip(project))
    case.export(restored)
    doc = Document(BytesIO(restored.session_state["_compiled_doc_bytes"]))
    assert sum(t.rows[0].cells[0].text == "Test Basic Data" for t in doc.tables) == 5
    assert sum(t.rows[0].cells[0].text == "Rheology Test" for t in doc.tables) == 5
    for name in materials_db.CEMENT_FORMULATION_FLUIDS:
        payload = restored.session_state["lab_payload_" + name]
        assert (payload["bhct"], payload["bhst"], payload["bhsp"]) == (150, 200, project["well_data"]["bhsp"])
    with patch.object(report.st, "session_state", restored.session_state.to_dict()):
        context = report.build_master_context()
    assert [s["name"] for s in context["slurries"]] == ["Lead", "Lead #1", "Lead #2", "Main", "Tail"]
    for table in (t for t in doc.tables if t.rows[0].cells[0].text == "Rheology Test"):
        assert table.rows[1].cells[2].text == table.rows[1].cells[3].text == "Tstart: 150 degF"
    Path('/tmp/round6-all-normal.docx').write_bytes(restored.session_state["_compiled_doc_bytes"])


def test_scavenger_manual_override_uses_existing_raw_engine_and_survives_restore():
    from engineering_tools import BBL_TO_CUFT
    case = audit.AuditRegressions()
    app = case.app(state_for("Main", "Scavenger"))
    case.phase(app, "phase5")
    next(w for w in app.checkbox if "scavenger" in w.key and w.label.startswith("Manual Override")).check().run()
    next(w for w in app.number_input if "scavenger" in w.key and w.label == "Custom Yield (cuft/sk)").set_value(1.234).run()
    next(w for w in app.number_input if "scavenger" in w.key and w.label == "Custom Fresh Water (bbl)").set_value(19.125)
    case.phase(app, "phase1")
    params = app.session_state["cement_params"]["Scavenger"]
    assert (params["yield"], params["mix_water"], params["manual_yield"], params["manual_mix_water"]) == (1.234, 19.125, 1.234, 19.125)
    restored = audit.round_trip(app.session_state.to_dict())
    assert prepare_calculations(restored) == []
    assert restored["cement_params"]["Scavenger"]["total_sacks"] == 30 * BBL_TO_CUFT / 1.234
    assert restored["cement_params"]["Scavenger"]["mix_water"] == 19.125
    assert restored["cement_params"]["Main"]["auto_calc"]
