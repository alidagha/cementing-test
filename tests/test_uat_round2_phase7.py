"""Round 2 Batch 6: measured lab inputs, fixed endpoint and report taxonomy."""
from copy import deepcopy
from io import BytesIO
import unittest
from unittest.mock import patch

import pandas as pd
from docx import Document
import test_audit_regressions as audit
import phase_10_procedure as report
from engineering_tools import compute_phase_status, lab_review_signature, thickening_time_valid
from project_state import prepare_calculations, lab_source_signature
from phase_2_3_well_data import _bhct_suggestion

SLURRIES = ['Lead', 'Lead #1', 'Lead #2', 'Main', 'Tail']


class Round2LabUAT(unittest.TestCase):
    app = audit.AuditRegressions.app
    phase = audit.AuditRegressions.phase
    healthy = audit.AuditRegressions.healthy
    configured_app = audit.AuditRegressions.configured_app
    export = audit.AuditRegressions.export
    assert_export_blocked = audit.AuditRegressions.assert_export_blocked

    @classmethod
    def setUpClass(cls):
        case = audit.AuditRegressions()
        cls.valid = audit.round_trip(case.configured_app(audit.BATCH1_JOBS[0]).session_state.to_dict())

    def fresh_lab(self):
        p = deepcopy(self.valid)
        p['bhct'] = p['well_data']['bhct'] = None
        p['fluids_config']['active'] = [*SLURRIES, 'Displacement Fluid']
        for i, s in enumerate(SLURRIES):
            p['fluids_config']['params'][s] = deepcopy(p['fluids_config']['params']['Main'])
            p['cement_params'][s] = deepcopy(p['cement_params']['Main'])
            p['cement_params'][s].update(top_mode='Depth (m MD)', top_depth=500.+500*i)
            p['cement_additives_dfs'][s] = p['cement_additives_dfs']['Main'].copy(deep=True)
        p.update(lab_qc_params={}, lab_grid_dfs={}, lab_source_signatures={}, lab_initialized_slurries=[])
        app = self.app(p)
        for phase in ('phase2_3','phase4','phase5','phase7'):
            self.phase(app, phase)
        return app

    def number(self, app, label):
        return next(w for w in app.number_input if w.label == label)

    def time(self, app, s):
        return next(w for w in app.text_input if w.label == '70 Bc (HH:MM) - '+s)

    def context(self, state, prepared=False):
        with patch.object(report.st, 'session_state', deepcopy(state)):
            return report.build_master_context(calculations_prepared=prepared)

    def test_all_archetypes_blank_no_endpoint_and_review_disabled(self):
        app = self.fresh_lab()
        for s in SLURRIES:
            with self.subTest(slurry=s):
                self.assertIsNone(app.session_state['bhct'])
                self.assertNotIn('bhct',app.session_state['lab_qc_params'][s])
                self.assertEqual(self.time(app,s).value,'')
                self.assertEqual(app.session_state['lab_qc_params'][s]['thickening_time'],'')
        self.assertFalse(any('Thickening Time Endpoint' in w.label for w in app.selectbox))
        self.assertTrue(all(b.disabled for b in app.button if b.label=='Confirm measured lab results'))
        self.assertNotEqual(compute_phase_status(app.session_state)['phase7']['level'],'ok')
        self.assert_export_blocked(app)
        self.phase(app,'phase7')
        for s in SLURRIES:
            self.assertIsNone(app.session_state['bhct'])
            self.assertNotIn('bhct',app.session_state['lab_qc_params'][s])
            self.assertEqual(self.time(app,s).value,'')

    def test_current_nullable_active_and_inactive_drafts_and_numeric_rejections(self):
        for s in SLURRIES:
            for inactive in (False, True):
                with self.subTest(slurry=s, inactive=inactive):
                    qc={'thickening_time':''}
                    p={'job_type':audit.BATCH1_JOBS[0], 'material_property_schema':1, 'bhct':None}
                    if inactive:p['inactive_slurry_drafts']={s:{'lab_qc_params':qc}}
                    else:p['lab_qc_params']={s:qc}
                    self.assertEqual(audit.round_trip(p),p)
                    for bad in (True,'bad',float('inf'),float('nan')):
                        p['bhct']=bad
                        with self.assertRaises(ValueError):audit.round_trip(p)

    def test_first_entries_navigation_and_fresh_restore_preserve_blank_and_entered(self):
        app=self.fresh_lab()
        self.phase(app,'phase2_3')
        self.number(app,'BHCT (degF)').set_value(155.0)
        self.phase(app,'phase1');self.phase(app,'phase2_3')
        self.assertEqual(self.number(app,'BHCT (degF)').value,155)
        self.phase(app,'phase7')
        for s in SLURRIES[:-1]:
            self.assertEqual(app.session_state[f'lab_payload_{s}']['bhct'],155)
            self.time(app,s).set_value('04:45')
            self.phase(app,'phase1');self.phase(app,'phase7')
            self.assertEqual(self.time(app,s).value,'04:45')
        restored=self.app(audit.round_trip(app.session_state.to_dict()))
        self.phase(restored,'phase7')
        for s in SLURRIES:
            with self.subTest(slurry=s):
                self.assertEqual(restored.session_state[f'lab_payload_{s}']['bhct'],155)
                self.assertEqual(self.time(restored,s).value,'' if s=='Tail' else '04:45')

    def test_blank_bhct_or_time_blocks_even_matching_review_and_invalidates_word(self):
        for s in SLURRIES:
            for field, value in (('bhct',None),('thickening_time','')):
                with self.subTest(slurry=s, field=field):
                    p=deepcopy(self.valid)
                    if s!='Main':
                        p['fluids_config']['active']=[s,'Displacement Fluid']
                        for key in ('cement_params','cement_additives_dfs','lab_grid_dfs','lab_qc_params','lab_source_signatures'):
                            p[key][s]=deepcopy(p[key]['Main'])
                        p['fluids_config']['params'][s]=deepcopy(p['fluids_config']['params']['Main'])
                    qc=p['lab_qc_params'][s]
                    if field=='bhct':
                        p['bhct']=p['well_data']['bhct']=value
                        p['lab_source_signatures'][s]=lab_source_signature(p,s)
                    else:qc[field]=value
                    qc['review_signature']=lab_review_signature(qc,p['lab_grid_dfs'][s])
                    p['_compiled_doc_bytes']=b'stale'
                    app=self.app(audit.round_trip(p));self.phase(app,'phase4');self.phase(app,'phase7')
                    self.assertTrue(next(b for b in app.button if b.label=='Confirm measured lab results').disabled)
                    self.assertNotEqual(compute_phase_status(app.session_state)['phase7']['level'],'ok')
                    self.assert_export_blocked(app)
                    self.assertTrue(prepare_calculations(deepcopy(app.session_state.to_dict())))

    def test_thickening_time_validation_envelope_unchanged(self):
        for value in ('03:30','04:45','3:30','00:01','24:00'):
            self.assertTrue(thickening_time_valid(value))
        for value in ('',None,'00:00','24:01','99:59','3:3'):
            self.assertFalse(thickening_time_valid(value))

    def test_obsolete_endpoint_is_ignored_but_all_measured_fields_and_grid_still_invalidate(self):
        qc=deepcopy(self.valid['lab_qc_params']['Main']);grid=self.valid['lab_grid_dfs']['Main']
        signature=lab_review_signature(qc,grid)
        for endpoint in ('100 Bc','Not specified','arbitrary',None):
            with self.subTest(endpoint=endpoint):
                self.assertEqual(lab_review_signature({**qc,'thickening_endpoint':endpoint},grid),signature)
        for field in ('thickening_time','free_water','free_water_45','surface_hardened_hours','api_fl'):
            with self.subTest(measured=field):
                self.assertNotEqual(lab_review_signature({**qc,field:'changed'},grid),signature)
        changed_well=deepcopy(self.valid)
        changed_well['bhct']=changed_well['well_data']['bhct']=155
        self.assertNotEqual(lab_source_signature(changed_well,'Main'),lab_source_signature(self.valid,'Main'))
        changed_qc = deepcopy(qc)
        changed_qc['compressive']['crush']['force_1'] += 1
        self.assertNotEqual(lab_review_signature(changed_qc, grid), signature)
        for column in grid.columns:
            changed=grid.copy();changed.loc[0,column]='changed'
            self.assertNotEqual(lab_review_signature(qc,changed),signature)

    def test_stale_endpoint_ignored_live_prepared_and_cached_context_word_all_jobs(self):
        for job in audit.BATCH1_JOBS:
            valid=self.configured_app(job)
            for endpoint in ('100 Bc','Not specified'):
                with self.subTest(job=job, endpoint=endpoint):
                    p=audit.round_trip(valid.session_state.to_dict())
                    p['lab_qc_params']['Main']['thickening_endpoint']=endpoint
                    app=self.app(p);self.phase(app,'phase7')
                    self.assertEqual(app.session_state['lab_qc_params']['Main']['thickening_endpoint'],endpoint)
                    self.assertFalse(any('Thickening Time Endpoint' in w.label for w in app.selectbox))
                    self.assertNotIn('thickening_endpoint',app.session_state['lab_payload_Main'])
                    self.assertEqual(compute_phase_status(app.session_state)['phase7']['level'],'ok')
                    state=deepcopy(app.session_state.to_dict())
                    self.assertEqual(prepare_calculations(state),[])
                    self.assertNotIn('thickening_endpoint',state['lab_payload_Main'])
                    state['lab_payload_Main']['thickening_endpoint']=endpoint
                    context=self.context(state,prepared=True)
                    self.assertNotIn('thickening_endpoint',context['slurries'][0]['lab'])
                    self.assertNotIn('thickening_endpoint_label',context['slurries'][0]['lab'])
                    self.export(app)
                    doc=Document(BytesIO(app.session_state['_compiled_doc_bytes']))
                    table=next(t for t in doc.tables if t.rows[0].cells[0].text=='Thickening Time Test')
                    self.assertEqual(table.rows[2].cells[-1].text,'70 Bc\nhr:mm')
                    self.assertEqual(table.rows[3].cells[-1].text,'03:30')

    def test_bhct_advisory_does_not_autofill_and_apply_is_explicit_for_all_slurries(self):
        app=self.fresh_lab()
        self.assertFalse(any('Suggested SQUEEZE' in c.value for c in app.caption))
        self.assertFalse(any(b.label.startswith('Apply ') for b in app.button))
        p=audit.round_trip(app.session_state.to_dict())
        p['well_geometry']['td_m']=3100.
        p['geo_md']=p['geo_tvd']=3100.
        p['well_data']['geo_md']=p['well_data']['geo_tvd']=3100.
        app=self.app(p);self.phase(app,'phase2_3')
        with patch('phase_2_3_well_data.st.session_state',app.session_state):suggestion=_bhct_suggestion()[1]
        expected=max(60,min(400,round(suggestion['temp_degF'])))
        self.assertIsNone(self.number(app,'BHCT (degF)').value)
        next(b for b in app.button if b.key=='_apply_well_bhct').click().run()
        self.assertEqual(self.number(app,'BHCT (degF)').value,expected)
        self.phase(app,'phase7')
        for s in SLURRIES:
            with self.subTest(slurry=s):
                self.assertEqual(app.session_state[f'lab_payload_{s}']['bhct'],expected)
                self.assertNotIn('bhct',app.session_state['lab_qc_params'][s])

    def test_material_types_occurrences_custom_fallback_and_unchanged_membership(self):
        names=['Micro Silica','Silica Flour','Hidense','O-uniFLC5','O-GAS BLOCK','Micro Block']
        expected=['Extender','CS Stabilizer','Weighting Agent','F.L. Controller','Anti Gas Migration','Liquid Extender']
        phase5=pd.DataFrame([{'Name':n,'Material Type':t} for n,t in zip(names,expected)]+[
            {'Name':'Repeated Custom','Material Type':'Custom Type A'},
            {'Name':'Repeated Custom','Material Type':'Custom Type B'},
            {'Name':'Other (Custom)','Material Type':'Owner Powder'}])
        materials=['CEMENT G DELIJAN',*names,'Repeated Custom','Repeated Custom','Owner Powder','Base Water','Unknown manual']
        lab=pd.DataFrame([{'Material':n,'Concentration':str(i),'Unit':'BWOB' if i==0 else '% BWOC',
                           'Mass':str(i+1),'Lot No':'LOT'+str(i)} for i,n in enumerate(materials)])
        before=report.separate_lab_tables(lab)
        after=report.separate_lab_tables(lab,phase5)
        for old_rows,new_rows in zip(before,after):
            self.assertEqual(len(old_rows),len(new_rows))
            for old,new in zip(old_rows,new_rows):
                self.assertEqual({k:v for k,v in old.items() if k!='Material Type'},
                                 {k:v for k,v in new.items() if k!='Material Type'})
        types={r['Name']:r['Material Type'] for rows in after for r in rows if r['Name']}
        for name,kind in zip(names,expected):self.assertEqual(types[name],kind)
        self.assertEqual(types['CEMENT G DELIJAN'],'Cement')
        self.assertEqual(types['Owner Powder'],'Owner Powder')
        self.assertEqual(types['Base Water'],'Base Water')
        self.assertEqual(types['Unknown manual'],'Unknown manual')
        repeated=[r['Material Type'] for rows in after for r in rows if r['Name']=='Repeated Custom']
        self.assertEqual(repeated,['Custom Type A','Custom Type B'])
        fallback=report.separate_lab_tables(lab)
        lookup={r['Name']:r['Material Type'] for rows in fallback for r in rows if r['Name']}
        for n,t in zip(names,expected):self.assertEqual(lookup[n],t)

    def test_actual_word_material_types_use_same_slurry_phase5_rows(self):
        app=self.configured_app(audit.BATCH1_JOBS[0]);self.phase(app,'phase5')
        source=pd.DataFrame([
            {'Name':'Micro Silica','Material Type':'Extender','Physical State':'Powder','Mix Method':'Dry Blend','User Input':5.,'Density':2.3},
            {'Name':'O-uniFLC5','Material Type':'F.L. Controller','Physical State':'Powder','Mix Method':'In Mix Water','User Input':.5,'Density':1.5},
            {'Name':'O-GAS BLOCK','Material Type':'Anti Gas Migration','Physical State':'Liquid','Mix Method':'In Mix Water','User Input':.1,'Density':1.2}])
        app.session_state['cement_additives_dfs']={'Main':source}
        self.phase(app,'phase7')
        next(b for b in app.button if b.label=='Sync with Phase V').click().run()
        before=deepcopy(app.session_state['lab_grid_dfs']['Main'])
        next(b for b in app.button if b.label=='Confirm measured lab results').click().run()
        self.export(app)
        pd.testing.assert_frame_equal(before,app.session_state['lab_grid_dfs']['Main'])
        context=self.context(app.session_state.to_dict())
        rows=context['slurries'][0]['lab']['conventional']+context['slurries'][0]['lab']['additives']
        types={r['Name']:r['Material Type'] for r in rows}
        self.assertEqual(types['Micro Silica'],'Extender')
        self.assertEqual(types['O-uniFLC5'],'F.L. Controller')
        self.assertEqual(types['O-GAS BLOCK'],'Anti Gas Migration')
        doc=Document(BytesIO(app.session_state['_compiled_doc_bytes']))
        cells=[[c.text for c in r.cells] for t in doc.tables for r in t.rows]
        for name,kind in (('Micro Silica','Extender'),('O-uniFLC5','F.L. Controller'),('O-GAS BLOCK','Anti Gas Migration')):
            self.assertTrue(any(c[0]==kind and c[1]==name and len(c)==6 for c in cells))


if __name__=='__main__':unittest.main()
