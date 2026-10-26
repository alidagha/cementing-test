"""B2: change engineering sources after real execution, then cross consumers."""
from copy import deepcopy
from io import BytesIO
import math
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
import pandas as pd
import pytest

import test_audit_regressions as audit
import materials_db as db
import placement
import cement_chart_model as chart
import cement_chart_report as chart_report
import centralizer_adapter as centralizer
import phase_9_centralizer as phase9
import pumping_simulation as pumping
import well_schematic_adapter as schematic
import well_schematic_engine as schematic_engine
import project_state as ps
import phase_5_cement as phase5
import phase_10_procedure as report
from centralizer_fixtures import execute_centralizer
from test_uat_round10c_integration import save_project, restore_project
from test_uat_round12a5_diesel_word_e2e import projects as diesel_projects, word, sequence_table
from test_uat_round12b_lnr_4_1_2 import word_project as fractional_project


CASE = audit.AuditRegressions()
CSG = 'CSG 9 5/8"'
LNR = 'LNR 4 1/2"'
JOBS = (CSG, LNR, 'LNR 5"', 'LNR 7"', 'TIE BACK LNR 5"', 'TIE BACK LNR 7"', 'CMT PLUG', 'CMT SQUEEZE')


def context(state):
    with patch.object(report.st, "session_state", state):
        return report.build_master_context()


def sources(state):
    return ps.fingerprint({key: state.get(key) for key in (
        "hardware_table", "placement_config", "fluids_config", "well_geometry", "formation_pressure_config",
        "cement_additives_dfs", "lab_qc_params", "executive_summary_config", "centralizer_config",
        "cement_params", "schematic_config")})


def save_restore(state):
    raw = save_project(SimpleNamespace(session_state=state))
    return restore_project(raw)


@pytest.fixture(scope="module")
def projects():
    cache = {}
    diesel = diesel_projects.__wrapped__()
    def get(job=CSG):
        if job not in cache:
            state = fractional_project.__wrapped__() if job == LNR else diesel(job)
            if chart.chart_job_kind(job) is not None:
                # Exercise computed engineering intervals; report-only overrides
                # have their own independence test below.
                for params in state["cement_params"].values():
                    params.update(top_mode="Not entered", top_depth=None, report_bottom_depth=None)
            app = save_restore(state)
            if job == CSG:
                app.session_state["cement_additives_dfs"]["Main"] = phase5._normalize_additive_rows(pd.DataFrame([{
                    "Material Type": "Other (Custom)", "Name": "Sodium Silicate", "Physical State": "Powder",
                    "Mix Method": "Dry Blend", "User Input": .50, "Density": None}]))
                CASE.phase(app, "phase5")
                CASE.phase(app, "phase7")
                next(b for b in app.button if b.key.startswith("_sync_btn_main_")).click().run()
                next(b for b in app.button if b.key.startswith("_confirm_lab_btn_main_")).click().run()
            execute_centralizer(app.session_state, physics="dynamic")
            CASE.export(app)
            state = app.session_state.to_dict()
            assert ps.prepare_calculations(state) == []
            assert state["_compiled_doc_bytes"].startswith(b"PK")
            cache[job] = state
        return deepcopy(cache[job])
    return get


def select_target(state):
    choices = placement.hardware_choices(state["hardware_table"], placement.target_descriptions(state["job_type"]))
    assert choices
    state["placement_config"]["target_row"] = max(choices, key=lambda key: choices[key]["depth"])


def target_description(state):
    return {"casing": "Casing", "liner": "Liner", "tieback": "Tie Back"}[chart.chart_job_kind(state["job_type"])]


def pump_geometry(state):
    return pumping.prepare_pumping_simulation(state)["inputs"]["geometry"]


def assert_shared_geometry(state, ctx):
    geometry = chart.resolve_chart_geometry(state)
    kind = geometry["kind"]
    expected = next(row for row in geometry["active_string"] if row["description"] == target_description(state))
    translated = centralizer.to_engine_geometry(state)
    assert translated["casing"]["od"] == pytest.approx(expected["od_in"] * .0254)
    assert translated["casing"]["id"] == pytest.approx(expected["id_in"] * .0254)
    prepared = pump_geometry(state)
    section = next(row for row in prepared.pipe_sections if row.z_top <= geometry["shoe_md"] - 1 < row.z_bot)
    assert (section.od_in, section.id_in) == (expected["od_in"], expected["id_in"])
    assert prepared.track_depth_m == geometry["shoe_md"]
    drawing = schematic.prepare_schematic(state)
    target = drawing["casings"][drawing["current_job"]["target_index"]]
    assert target["size_in"] == expected["od_in"]
    assert target["shoe_md"] == geometry["shoe_md"]
    source_hardware = state["hardware_table"].copy(deep=True)
    expected_hardware = deepcopy(source_hardware.to_dict("records"))
    for row in expected_hardware:
        for field, value in row.items():
            if pd.isna(value) or (isinstance(value, str) and not value.strip()):
                row[field] = "-"
        depth = row.get("MD (m)")
        try:
            scalar = float(depth)
        except (TypeError, ValueError, OverflowError):
            scalar = None
        if scalar is not None and math.isfinite(scalar):
            row["MD (m)"] = "%.1f" % scalar
        if row.get("Description") == "Open Hole Size":
            row.update({field: "-" for field in
                        ("Joint (m)", "Weight (ppf)", "Grade", "Collapse (psi)", "Burst (psi)")})
    assert ctx["hardware"] == expected_hardware
    assert context(state)["hardware"] == expected_hardware
    pd.testing.assert_frame_equal(state["hardware_table"], source_hardware, check_exact=True)
    assert ctx["cement_chart_report"]["model"] == state["_cement_chart_model"]
    intervals = {row["fluid"]: row for row in chart.resolve_cement_placement(state)}
    for slurry in ctx["slurries"]:
        row = intervals[slurry["name"]]
        assert slurry["top"] == ("surface" if row["top_md"] == 0 else f"{row['top_md']:.1f} m")
        assert slurry["bottom"] == f"{row['bottom_md']:.1f} m MD"
    assert kind == drawing["current_job"]["type"]


@pytest.mark.parametrize("job", JOBS)
def test_executed_job_families_share_canonical_facts_and_real_word(projects, job):
    state = projects(job)
    ctx = context(state)
    drawing = schematic.prepare_schematic(state)
    assert schematic_engine.check_well_design(drawing) == []
    assert [row["name"] for row in ctx["fluids_train"]] == state["fluids_config"]["active"]
    assert len(Document(BytesIO(state["_compiled_doc_bytes"])).tables) > 0
    if job in ("CMT PLUG", "CMT SQUEEZE"):
        assert chart.chart_readiness(state, include_fluids=True)["model"] is None
        assert centralizer.to_engine_geometry(state) is None
        assert not pumping.prepare_pumping_simulation(state)["applicable"]
        assert not ctx["has_centralizer"]
    else:
        assert phase9.result_is_current(state)
        assert_shared_geometry(state, ctx)


@pytest.mark.parametrize("field,value", [("ID (in)", 8.4), ("Weight (ppf)", 104.0), ("Joint (m)", 11.5),
                                         ("MD (m)", "2900")])
def test_valid_hardware_edit_after_execution_reaches_every_consumer(projects, field, value):
    state = projects()
    before_pump = pumping.source_signature(state)
    before_drawing = schematic.source_signature(state)
    frame = state["hardware_table"]
    frame.loc[frame["Description"] == "Casing", field] = value
    select_target(state)  # OD/MD changes retire the old explicit row token.
    assert not phase9.result_is_current(state)
    expected_source = sources(state)
    pump_geometry(state)  # Deliberately prepare Pumping before the common rebuild.
    schematic.prepare_schematic(state)
    assert sources(state) == expected_source
    physical_dimension = field not in ("Weight (ppf)", "Joint (m)")
    assert (pumping.source_signature(state) != before_pump) == physical_dimension
    assert (schematic.source_signature(state) != before_drawing) == physical_dimension
    assert ps.prepare_calculations(state) == []
    assert "_compiled_doc_bytes" not in state
    with pytest.raises(ValueError):
        context(state)  # Stale Centralizer cannot silently reach Word.
    execute_centralizer(state, physics="dynamic")
    assert_shared_geometry(state, context(state))


def test_basic_formation_new_geometry_rebuilds_domain_and_shared_placement(projects):
    state = projects()
    old = chart.current_open_hole_domain(state)
    state["hardware_table"].loc[state["hardware_table"]["Description"] == "Previous Casing", "MD (m)"] = "1600"
    assert not phase9.result_is_current(state)
    prepared = pumping.prepare_pumping_simulation(state)
    assert prepared["inputs"]["geometry"].ann_sections[0].z_bot == 1600.0
    assert ps.prepare_calculations(state) == []
    domain = chart.current_open_hole_domain(state)
    assert domain["top_md"] == 1600.0 != old["top_md"]
    assert state["_cement_chart_model"]["formation_pressure"]["domain"] == domain
    execute_centralizer(state, physics="dynamic")
    assert_shared_geometry(state, context(state))


def test_nominal_od_mismatch_is_blocked_without_fabricating_geometry_and_recovers(projects):
    state = projects()
    frame = state["hardware_table"]
    frame.loc[frame["Description"] == "Casing", "Size (in)"] = "9.5"
    select_target(state)
    assert centralizer.to_engine_geometry(state)["casing"]["od"] == pytest.approx(9.5 * .0254)
    assert pump_geometry(state).pipe_sections[-1].od_in == 9.5
    assert not phase9.result_is_current(state)
    with pytest.raises(ValueError, match="nominal size"):
        schematic.prepare_schematic(state)
    assert any("nominal size" in issue for issue in ps.prepare_calculations(state))
    assert "_compiled_doc_bytes" not in state
    with pytest.raises(ValueError):
        context(state)
    frame.loc[frame["Description"] == "Casing", "Size (in)"] = "9.625"
    select_target(state)
    assert ps.prepare_calculations(state) == []
    assert phase9.result_is_current(state)
    assert_shared_geometry(state, context(state))


@pytest.mark.parametrize("job,id_in", [(LNR, 4.05), ('LNR 5"', 4.4), ('LNR 7"', 6.3)])
def test_liner_actual_id_and_hanger_depth_after_execution_drive_shared_geometry(projects, job, id_in):
    state = projects(job)
    frame = state["hardware_table"]
    frame.loc[frame["Description"].isin(["Liner", "Liner Hanger"]), "ID (in)"] = id_in
    frame.loc[frame["Description"] == "Liner Hanger", "MD (m)"] = "1050"
    frame.loc[frame["Description"] == "Drill Pipes", "MD (m)"] = "1050"
    assert not phase9.result_is_current(state)
    drawing = schematic.prepare_schematic(state)
    assert drawing["casings"][drawing["current_job"]["target_index"]]["id_in"] == id_in
    prepared = pump_geometry(state)
    assert prepared.pipe_sections[-1].id_in == id_in
    assert prepared.pipe_sections[-1].z_top == 1050.0
    assert ps.prepare_calculations(state) == []
    execute_centralizer(state, physics="dynamic")
    assert_shared_geometry(state, context(state))
    assert db.DEFAULT_HOLE_SIZES[LNR] == '6"'
    assert 'TIE BACK LNR 4 1/2"' not in db.JOB_TYPES


def signatures(state):
    return (ps.fingerprint(chart.build_chart_model(state)), phase9.source_signature(state),
            pumping.source_signature(state), schematic.source_signature(state), ps.lab_source_signature(state, "Main"))


MATRIX = [
    ("pipe ID", (True, True, True, True, False), True),
    ("pipe Weight", (False, True, False, False, False), True),
    ("pipe Joint", (False, True, False, False, False), True),
    ("Main volume", (True, True, True, False, False), True),
    ("Diesel density", (True, True, True, False, False), True),
    ("Diesel bhct_down", (False, True, True, False, False), True),
    ("Diesel Surface", (False, False, False, False, False), False),
    ("Diesel gel", (False, False, True, False, False), True),
    ("report placement", (False, False, False, False, False), True),
    ("hardware Grade", (False, False, False, False, False), True),
    ("Basic pore source", (True, False, True, False, False), True),
    ("Advanced formation", (True, False, False, False, False), True),
]


def mutate(state, change):
    frame = state["hardware_table"]
    if change.startswith("pipe "):
        field, value = {"pipe ID": ("ID (in)", 8.4), "pipe Weight": ("Weight (ppf)", 104.),
                        "pipe Joint": ("Joint (m)", 11.5)}[change]
        frame.loc[frame["Description"] == "Casing", field] = value
    elif change == "Main volume":
        state["fluids_config"]["params"]["Main"]["volume"] += 1.
    elif change == "Diesel density":
        state["fluids_config"]["params"]["Diesel"]["density"] = "54"
    elif change in ("Diesel bhct_down", "Diesel Surface", "Diesel gel"):
        dataset = "surface_down" if change == "Diesel Surface" else "bhct_down"
        row = state["lab_qc_params"]["Diesel"]["rheology"][dataset]
        if change == "Diesel gel":
            row["gel_10_min"] = "13"
        else:
            row["readings"]["300"] = "169"
    elif change == "report placement":
        state["cement_params"]["Main"].update(top_mode="Depth (m MD)", top_depth=2100., report_bottom_depth=2900.)
    elif change == "hardware Grade":
        frame["Grade"] = "N-80"
    elif change == "Basic pore source":
        state[chart.PORE_FIELD] = state["well_data"][chart.PORE_FIELD] = 71.
    elif change == "Advanced formation":
        domain = chart.current_open_hole_domain(state)
        state["formation_pressure_config"] = {"mode": "advanced", "profile": [
            {"md_m": domain["top_md"], "pore_psi": 2000., "frac_psi": 6000.},
            {"md_m": domain["bottom_md"], "pore_psi": 4000., "frac_psi": 10000.}]}
    else:
        raise AssertionError(change)


@pytest.mark.parametrize("change,expected,word_changes", MATRIX, ids=[row[0] for row in MATRIX])
def test_source_to_consumer_relevance_matrix_after_execution(projects, change, expected, word_changes):
    state = projects()
    before = signatures(state)
    old_document = report.document_signature(context(state), report.TEMPLATE_PATH)
    mutate(state, change)
    after = signatures(state)
    assert tuple(a != b for a, b in zip(after, before)) == expected
    assert phase9.result_is_current(state) == (not expected[1])
    assert ps.prepare_calculations(state) == []
    if word_changes:
        assert "_compiled_doc_bytes" not in state
    else:
        assert state["_compiled_doc_bytes"].startswith(b"PK")
    if expected[1]:
        with pytest.raises(ValueError):
            context(state)
        execute_centralizer(state, physics="dynamic")
    ctx = context(state)
    assert (report.document_signature(ctx, report.TEMPLATE_PATH) != old_document) == word_changes
    if change == "report placement":
        assert ctx["slurries"][0]["top"] == "2100.0 m"
        assert ctx["slurries"][0]["bottom"] == "2900.0 m MD"


@pytest.mark.parametrize("job", [CSG, LNR, 'TIE BACK LNR 7"'])
def test_capacity_excess_and_report_metadata_remain_distinct_after_calculation(projects, job):
    state = projects(job)
    old = chart.resolve_cement_placement(state)
    original_pump = pump_geometry(state)
    original_schematic = schematic.prepare_schematic(state)
    state["placement_config"].update(chart_excess_mode="apply", excess_csg_oh_pct=10., excess_csg_csg_pct=20.)
    state["cement_params"]["Main"].update(top_mode="Depth (m MD)", top_depth=2100., report_bottom_depth=2900.)
    state["executive_summary_config"] = {"version": 1, "job_type": job, "per_slurry": {"Main": {"wet_stands": 3.}}}
    if "TIE BACK" in job:
        state["executive_summary_config"]["tie_back_origin"] = "Shoe"
    assert pump_geometry(state) == original_pump  # Nominal frozen geometry stays independent of capacity excess.
    assert schematic.prepare_schematic(state) == original_schematic
    assert chart.resolve_cement_placement(state) != old
    assert not phase9.result_is_current(state)
    assert ps.prepare_calculations(state) == []
    execute_centralizer(state, physics="dynamic")
    ctx = context(state)
    assert ctx["slurries"][0]["top"] == "2100.0 m"
    assert ctx["slurries"][0]["bottom"] == "2900.0 m MD"
    assert "3 stands wet" in ctx["exec_summary"]
    assert ctx["cement_chart_report"]["model"]["cement_placement"] == chart.resolve_cement_placement(state)


def test_advanced_formation_domain_changes_preserve_sources_and_block_extrapolation(projects):
    state = projects()
    mutate(state, "Advanced formation")
    state["formation_pressure_config"]["profile"].insert(1, {"md_m": 2000., "pore_psi": 3000., "frac_psi": 8000.})
    assert ps.prepare_calculations(state) == []
    ps.reconcile_formation_profile(state)
    interior = deepcopy(state["formation_pressure_config"]["profile"][1])
    for previous_shoe in (1600., 1700.):
        frame = state["hardware_table"]
        frame.loc[frame["Description"] == "Previous Casing", "MD (m)"] = str(previous_shoe)
        assert any("endpoints" in issue for issue in ps.prepare_calculations(state))
        with pytest.raises(ValueError):
            context(state)
        ps.reconcile_formation_profile(state)
        profile = state["formation_pressure_config"]["profile"]
        assert profile[0] == {"md_m": previous_shoe, "pore_psi": None, "frac_psi": None}
        assert profile[1] == interior
        assert profile[-1] == {"md_m": 3000., "pore_psi": 4000., "frac_psi": 10000.}
        assert chart.chart_readiness(state, include_fluids=True)["model"] is None
        app = save_restore(state)
        state = app.session_state.to_dict()
        assert state["formation_pressure_config"]["profile"] == profile
        assert ps.prepare_calculations(state)
        state["formation_pressure_config"]["profile"][0].update(pore_psi=2200., frac_psi=6500.)
        ps.reconcile_formation_profile(state)
        assert ps.prepare_calculations(state) == []
        execute_centralizer(state, physics="dynamic")
        assert context(state)["cement_chart_report"]["model"]["formation_pressure"]["domain"]["top_md"] == previous_shoe


@pytest.fixture(scope="module")
def annular_diesel(projects):
    state = projects(LNR)
    capacity = centralizer.build_fluid_snapshot(state)["capacity_bbl"]["inside"]
    state["fluids_config"]["params"]["Displacement Fluid"]["volume"] = capacity
    assert ps.prepare_calculations(state) == []
    execute_centralizer(state, physics="dynamic")
    snapshot = centralizer.build_fluid_snapshot(state)
    assert {"Diesel", "Pre Flush"} <= {row["fluid"] for row in snapshot["annulus_intervals"]}
    assert {"Diesel", "Pre Flush"} <= {row["fluid"] for row in state["_cement_chart_model"]["annular_intervals"]}
    assert word(context(state)).tables
    return state


@pytest.mark.parametrize("change", ["Diesel volume", "Diesel 51", "Diesel 54", "Diesel rheology",
                                    "Pre Flush volume", "Pre Flush density", "inactive Diesel", "WBM"])
def test_annular_diesel_preflush_changes_remain_independent_across_consumers(annular_diesel, change):
    state = deepcopy(annular_diesel)
    original_static = chart.build_chart_model(state)
    original_lab = deepcopy(state["lab_qc_params"]["Diesel"])
    original_diesel = deepcopy(state["fluids_config"]["params"]["Diesel"])
    if change in ("Diesel volume", "Pre Flush volume"):
        name = "Diesel" if change.startswith("Diesel") else "Pre Flush"
        state["fluids_config"]["params"][name]["volume"] += 1.
    elif change in ("Diesel 51", "Diesel 54", "Pre Flush density"):
        name = "Pre Flush" if change.startswith("Pre Flush") else "Diesel"
        state["fluids_config"]["params"][name]["density"] = "81" if name == "Pre Flush" else change.split()[-1]
    elif change == "Diesel rheology":
        state["lab_qc_params"]["Diesel"]["rheology"]["bhct_down"]["readings"]["300"] = "169"
    elif change == "inactive Diesel":
        state["fluids_config"]["active"].remove("Diesel")
    else:
        state["mud_type"] = state["well_data"]["mud_type"] = "WBM"
        ps.restore_canonical_fields(state)
    active = "Diesel" in state["fluids_config"]["active"]
    assert not phase9.result_is_current(state)
    prepared = pumping.prepare_pumping_simulation(state)["inputs"]
    assert ("Diesel" in prepared["rheology"]) == active
    assert list(prepared["fluid_data"]) == [name for name in db.HYDRAULIC_EXECUTION_ORDER if name in state["fluids_config"]["active"]]
    assert ps.prepare_calculations(state) == []
    current = chart.build_chart_model(state)
    if change == "Diesel rheology":
        assert current == original_static
        assert state["fluids_config"]["params"]["Diesel"] == original_diesel
    else:
        assert state["lab_qc_params"]["Diesel"] == original_lab
    execute_centralizer(state, physics="dynamic")
    assert ("Diesel" in state["_centralizer_result"]["diagnostics"]["rheology"]) == active
    if active:
        snapshot = centralizer.build_fluid_snapshot(state)
        intervals = snapshot["annulus_intervals"]
        diesel = [row for row in intervals if row["fluid"] == "Diesel"]
        preflush = [row for row in intervals if row["fluid"] == "Pre Flush"]
        assert max(row["bottom_md"] for row in diesel) <= min(row["top_md"] for row in preflush)
    ctx = context(state)
    rows = [[cell.text for cell in row.cells] for row in sequence_table(word(ctx)).rows]
    names = [row[0] for row in rows if row[0] in db.HYDRAULIC_EXECUTION_ORDER]
    assert (names.count("Diesel") == 1) == active
    if active:
        assert names.index("Diesel") + 1 == names.index("Pre Flush")
    assert ctx["has_preflush"]


def test_fully_returned_diesel_has_no_ghost_annulus_but_retains_active_sequence(annular_diesel):
    state = deepcopy(annular_diesel)
    geometry = chart.resolve_chart_geometry(state)
    annulus = chart.interval_volume(geometry["capacity_segments"], 0., geometry["shoe_md"])
    state["fluids_config"]["params"]["Main"]["volume"] = annulus + geometry["shoe_track_bbl"] + 10.
    assert ps.prepare_calculations(state) == []
    assert all(row["fluid"] != "Diesel" for row in state["_cement_chart_model"]["annular_intervals"])
    snapshot = centralizer.build_fluid_snapshot(state)
    assert snapshot["per_fluid"]["Diesel"]["in_path_bbl"] == 0.
    assert snapshot["per_fluid"]["Diesel"]["returned_bbl"] == state["fluids_config"]["params"]["Diesel"]["volume"]
    assert all(row["fluid"] != "Diesel" for row in snapshot["annulus_intervals"])
    assert "Diesel" in pumping.prepare_pumping_simulation(state)["inputs"]["rheology"]
    execute_centralizer(state, physics="dynamic")
    assert [row["name"] for row in context(state)["fluids_train"]].count("Diesel") == 1


@pytest.mark.parametrize("density,expected", [(1.47, 1.47), (None, 1.50)])
def test_sodium_multi_edit_checkpoint_changes_only_formulation_consumers(projects, density, expected):
    state = projects()
    physical = (chart.build_chart_model(state), phase9.source_signature(state), pumping.source_signature(state),
                schematic.prepare_schematic(state), chart.formation_source(state))
    old_lab = ps.lab_source_signature(state, "Main")
    df = state["cement_additives_dfs"]["Main"]
    df.at[0, "User Input"] = .75
    df.at[0, "Density"] = density
    raw_source = df.to_dict("records")
    pumping.prepare_pumping_simulation(state)
    assert df.to_dict("records") == raw_source
    app = save_restore(state)
    CASE.phase(app, "phase5")
    assert app.session_state["cement_additives_dfs"]["Main"].at[0, "Density"] == expected
    CASE.phase(app, "phase7")
    assert any("Formulation Drift" in item.value for item in app.warning)
    assert ps.lab_source_signature(app.session_state, "Main") != old_lab
    next(b for b in app.button if b.key.startswith("_sync_btn_main_")).click().run()
    next(b for b in app.button if b.key.startswith("_confirm_lab_btn_main_")).click().run()
    state = app.session_state.to_dict()
    assert ps.prepare_calculations(state) == []
    assert (chart.build_chart_model(state), phase9.source_signature(state), pumping.source_signature(state),
            schematic.prepare_schematic(state), chart.formation_source(state)) == physical
    execute_centralizer(state, physics="dynamic")
    ctx = context(state)
    assert ctx["slurries"][0]["additives"][0]["Name"] == "Sodium Silicate"
    assert ctx["slurries"][0]["additives"][0]["(lbs or gal)/sk"] == "0.75 gal/sk"
    assert "Sodium Silicate" not in [row["name"] for row in ctx["fluids_train"]]
    assert "Sodium Silicate" not in centralizer.build_fluid_snapshot(state)["per_fluid"]
    assert "Sodium Silicate" not in pumping.prepare_pumping_simulation(state)["inputs"]["fluid_data"]
    assert word(ctx).tables


@pytest.mark.parametrize("change", ["geometry", "placement", "Diesel"])
def test_save_after_source_change_before_rebuild_then_different_consumer_first(projects, change):
    state = projects()
    if change == "geometry":
        state["hardware_table"].loc[state["hardware_table"]["Description"] == "Casing", "MD (m)"] = "2900"
        select_target(state)
        state["fluids_config"]["params"]["Main"]["volume"] = 51.
    elif change == "placement":
        state["placement_config"].update(chart_excess_mode="apply", excess_csg_oh_pct=10., excess_csg_csg_pct=20.)
    else:
        state["fluids_config"]["params"]["Diesel"]["density"] = "54"
        state["lab_qc_params"]["Diesel"]["rheology"]["bhct_down"]["readings"]["300"] = "169"
    expected = sources(state)
    app = save_restore(state)
    restored = app.session_state.to_dict()
    assert "_compiled_doc_bytes" not in restored and "_centralizer_result" not in restored
    assert sources(restored) == expected
    prepared = pumping.prepare_pumping_simulation(restored)
    drawing = schematic.prepare_schematic(restored)
    assert sources(restored) == expected
    assert prepared["inputs"]["geometry"].track_depth_m == drawing["casings"][drawing["current_job"]["target_index"]]["shoe_md"]
    with pytest.raises(ValueError):
        context(restored)
    execute_centralizer(restored, physics="dynamic")
    assert_shared_geometry(restored, context(restored))


def test_real_engines_render_after_multiple_source_edits_in_unusual_order(projects, tmp_path):
    state = projects(LNR)
    frame = state["hardware_table"]
    frame.loc[frame["Description"] == "Liner", "ID (in)"] = 4.05
    frame.loc[frame["Description"] == "Liner Hanger", "ID (in)"] = 4.05
    state["fluids_config"]["params"]["Diesel"].update(density="54", volume=21.)
    state["lab_qc_params"]["Diesel"]["rheology"]["bhct_down"]["readings"]["300"] = "169"
    expected = sources(state)
    drawing = schematic.prepare_schematic(state)
    schematic_engine.render_well_schematic(drawing, str(tmp_path / "updated-liner.png"), dpi=72)
    assert sources(state) == expected
    execute_centralizer(state, physics="dynamic")
    assert phase9.result_is_current(state)
    assert ps.prepare_calculations(state) == []
    ctx = context(state)
    assert_shared_geometry(state, ctx)
    assert word(ctx).tables
    rendered = pumping.run_pumping_simulation(state)
    try:
        assert all(math.isfinite(value) for value in rendered["result"].hydro)
        assert rendered["result"].dyn is not None
    finally:
        pumping.fcm.plt.close(rendered["figure"])
    assert (tmp_path / "updated-liner.png").read_bytes().startswith(b"\x89PNG")
    assert chart_report.render_chart_png(state["_cement_chart_model"]).startswith(b"\x89PNG")


@pytest.mark.parametrize("job", ['TIE BACK LNR 5"', 'TIE BACK LNR 7"'])
def test_explicit_new_tieback_host_is_shared_without_rewriting_source_history(projects, job):
    state = projects(job)
    old_geometry = chart.resolve_chart_geometry(state)
    frame = state["hardware_table"]
    row = frame.loc[frame["Description"] == "Previous Casing"].iloc[0].copy()
    row.update({"Description": "Casing", "Size (in)": "11.75", "ID (in)": 10.75})
    state["hardware_table"] = pd.concat([frame, pd.DataFrame([row])], ignore_index=True)
    state["placement_config"]["host_row"] = next(iter(placement.hardware_choices(
        state["hardware_table"], {"Casing"})))
    expected = sources(state)
    prepared = pumping.prepare_pumping_simulation(state)
    assert prepared["inputs"]["geometry"].ann_sections[0].do_in == 10.75
    assert sources(state) == expected
    geometry = chart.resolve_chart_geometry(state)
    assert geometry["capacity_segments"] != old_geometry["capacity_segments"]
    assert centralizer.to_engine_geometry(state)["wellbore"]["prev_casing_id"] == pytest.approx(10.75 * .0254)
    assert not phase9.result_is_current(state)
    assert ps.prepare_calculations(state) == []
    execute_centralizer(state, physics="dynamic")
    assert_shared_geometry(state, context(state))
    assert state["hardware_table"].iloc[:len(frame)].to_dict("records") == frame.to_dict("records")


def test_retired_target_selection_blocks_each_admission_then_recovers(projects):
    state = projects()
    original = state["placement_config"]["target_row"]
    state["placement_config"]["target_row"] = "retired-hardware-row"
    assert not phase9.result_is_current(state)
    assert chart.chart_readiness(state, include_fluids=True)["model"] is None
    with pytest.raises(ValueError):
        pumping.prepare_pumping_simulation(state)
    with pytest.raises(ValueError):
        schematic.prepare_schematic(state)
    assert ps.prepare_calculations(state)
    with pytest.raises(ValueError):
        context(state)
    state["placement_config"]["target_row"] = original
    assert ps.prepare_calculations(state) == []
    assert phase9.result_is_current(state)
    assert_shared_geometry(state, context(state))


def test_archived_sodium_slurry_does_not_leak_into_active_engineering_or_word(projects):
    state = projects()
    app = save_restore(state)
    CASE.phase(app, "phase4")
    app.checkbox(key="chk_fluid_Tail").check().run()
    app.number_input(key="vol_Tail").set_value(50.).run()
    app.text_input(key="den_Tail").set_value("118").run()
    app.text_input(key="rate_Tail").set_value("4").run()
    app.session_state["cement_params"]["Tail"] = deepcopy(app.session_state["cement_params"]["Main"])
    app.session_state["lab_qc_params"]["Tail"] = deepcopy(app.session_state["lab_qc_params"]["Main"])
    CASE.phase(app, "phase5")
    CASE.phase(app, "phase7")
    next(b for b in app.button if b.key.startswith("_confirm_lab_btn_tail_")).click().run()
    CASE.phase(app, "phase4")
    app.checkbox(key="chk_fluid_Main").uncheck().run()
    state = save_restore(app.session_state.to_dict()).session_state.to_dict()
    assert state["inactive_slurry_drafts"]["Main"]["cement_additives_dfs"].at[0, "Name"] == "Sodium Silicate"
    assert ps.prepare_calculations(state) == []
    assert "Main" not in pumping.prepare_pumping_simulation(state)["inputs"]["fluid_data"]
    assert "Main" not in centralizer.build_fluid_snapshot(state)["per_fluid"]
    assert schematic.prepare_schematic(state) == schematic.prepare_schematic(projects())
    execute_centralizer(state, physics="dynamic")
    ctx = context(state)
    doc = word(ctx)
    assert "Sodium Silicate" not in "\n".join(node.text or "" for node in doc.element.iter()
                                               if node.tag.endswith("}t"))
    assert [row["name"] for row in ctx["slurries"]] == ["Tail"]


def test_conflicting_tieback_lap_blocks_drawing_and_word_without_rewriting_physics(projects):
    state = projects('TIE BACK LNR 7"')
    original = (chart.build_chart_model(state), phase9.source_signature(state),
                pumping.source_signature(state), pump_geometry(state))
    state["placement_config"]["tie_back_liner_lap_m"] = 0.
    assert schematic.prepare_schematic(state)["current_job"]["type"] == "tieback"
    assert ps.prepare_calculations(state) == []
    state["placement_config"]["tie_back_liner_lap_m"] = 10.
    assert (chart.build_chart_model(state), phase9.source_signature(state),
            pumping.source_signature(state), pump_geometry(state)) == original
    assert phase9.result_is_current(state)
    with pytest.raises(ValueError, match="contradicts"):
        schematic.prepare_schematic(state)
    assert ps.prepare_calculations(state)
    with pytest.raises(ValueError):
        context(state)
    restored = save_restore(state).session_state.to_dict()
    assert restored["placement_config"]["tie_back_liner_lap_m"] == 10.
    assert ps.prepare_calculations(restored)
    restored["placement_config"]["tie_back_liner_lap_m"] = 0.
    assert ps.prepare_calculations(restored) == []
    execute_centralizer(restored, physics="dynamic")
    assert_shared_geometry(restored, context(restored))
