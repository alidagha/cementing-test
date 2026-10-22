"""A4: canonical Spacer formulation and reviewed laboratory in actual Word."""
from copy import deepcopy
from io import BytesIO
from unittest.mock import patch
import re
import pandas as pd
import hashlib
from zipfile import ZipFile
from lxml import etree

from docx import Document
from docxtpl import DocxTemplate
import pytest

import phase_10_procedure as report
import project_state as ps
import standalone_spacer as sp
import pumping_simulation as pumping
from well_schematic_adapter import prepare_schematic
import test_audit_regressions as audit
from centralizer_fixtures import ensure_current_centralizer
from test_round13a_a2_spacer_lab import measured


@pytest.fixture(scope="module")
def source():
    app = audit.AuditRegressions().configured_app('CSG 9 5/8"')
    state = audit.round_trip(app.session_state.to_dict())
    state["fluids_config"]["active"].insert(0, "Spacer")
    state["fluids_config"]["params"]["Spacer"] = {
        "volume": 20., "density": "110", "pump_rate": "4", "material_name": "Spacer"}
    state[sp.CONFIG_KEY] = {"mode": "Barite", "dead_vol": 5., "tank_name": "Batch Mixer Tanks"}
    ps.refresh_fluids(state)
    measured(state)
    ensure_current_centralizer(state)
    return state


def context(state, prepared=False):
    with patch.object(report.st, "session_state", state):
        return report.build_master_context(calculations_prepared=prepared)


def word(ctx, state=None):
    doc = DocxTemplate(report.TEMPLATE_PATH)
    doc.render(report._word_quantity_context(ctx))
    stream = BytesIO()
    doc.save(stream)
    data = stream.getvalue()
    if state is not None:
        simulation = pumping.run_pumping_simulation(state)
        try:
            data = report.postprocess_well_report(data, ctx["well_profile_report"], ctx["note_vertical"],
                ctx["cement_chart_report"], ctx["centralizer_report"], pumping_report=simulation,
                schematic_report=prepare_schematic(state))
        finally:
            if simulation.get("figure") is not None:
                pumping.fcm.plt.close(simulation["figure"])
    return Document(BytesIO(data))


def cement_lab_display(sample):
    # The unchanged Phase-VII view follows Cement's concentration/mass rules.
    from phase_7_lab import _spacer_lab_display
    view = _spacer_lab_display(pd.DataFrame(sample["rows"]).astype({"Concentration": str, "Mass": str}))
    view.loc[view["Material"] == "Fresh Water", "Concentration"] = "1.0"
    return view


@pytest.mark.parametrize("mode,density", [("Limestone", 95.), ("Barite", 110.),
    ("Ferro Bar", 140.), ("Limestone + Barite", 110.)])
def test_lab_word_units_are_presentation_only_and_freshwater_is_exact(source, mode, density):
    state = deepcopy(source)
    state[sp.CONFIG_KEY].update(mode=mode, limestone_available_lb=20000., barite_available_lb=20000.)
    state["fluids_config"]["params"]["Spacer"]["density"] = str(density)
    measured(state)
    ensure_current_centralizer(state)
    ctx = context(state)
    raw = deepcopy(ctx)
    signature = ps.fingerprint(state)
    sample = sp.lab_sample(sp.formulation_from_state(state))
    rendered = report._word_quantity_context(ctx)
    assert rendered == report._word_quantity_context(ctx)
    doc = word(ctx)
    lab = next(t for t in doc.tables if t.cell(0, 0).text == "Spacer Lab Results")
    rows = lab.rows[2:-1]
    assert [cell.text for cell in rows[0].cells[:5]] == ["Water", "Fresh Water", "", "1.0", "1.0"]
    units = {"Salt": "FW lb/bbl", "O-Spacer": "FW lb/bbl", "Surfactant": "FW gal/bbl",
             "Defoamer": "FW gal/bbl", "Limestone": "FW lb/bbl", "Barite": "FW lb/bbl",
             "Ferro Bar": "FW lb/bbl"}
    expected = cement_lab_display(sample)
    for row, ingredient in zip(rows[1:], sample["rows"][1:]):
        assert row.cells[1].text == ingredient["Material"]
        assert row.cells[4].text == units[ingredient["Material"]]
    assert [row.cells[3].text for row in rows] == list(expected["Concentration"])
    assert [row.cells[5].text for row in rows] == list(expected["Mass"])
    assert ctx == raw
    assert ps.fingerprint(state) == signature
    assert sp.lab_sample(sp.formulation_from_state(state)) == sample
    assert ctx["standalone_spacer"]["lab_rows"][0]["Concentration"] != 1.0


@pytest.mark.parametrize("value,concentration,mass", [
    (.005, "0.005", "0.01"), (.009, "0.009", "0.01"),
    (1.2345, "1.234", "1.23"), (1.2355, "1.236", "1.24"),
    (9.995, "9.995", "9.99"), (9.999, "9.999", "10.00"),
    (9.9995, "9.999", "10.00"), (10., "10.00", "10.0"),
    (10.005, "10.005", "10.0"), (126., "126.00", "126.0")])
def test_spacer_word_matches_cement_lab_precision_boundaries(source, value, concentration, mass):
    ctx = context(deepcopy(source))
    for row in ctx["standalone_spacer"]["lab_rows"]:
        row.update(Concentration=value, Mass=value)
    # A deliberate display-boundary fixture: the provided unrounded total,
    # not a sum of displayed ingredient masses, is the total's sole authority.
    ctx["standalone_spacer"]["lab_mass_g"] = 10.051
    raw = deepcopy(ctx)
    doc = word(ctx)
    lab = next(t for t in doc.tables if t.cell(0, 0).text == "Spacer Lab Results")
    assert lab.rows[2].cells[3].text == "1.0"
    assert lab.rows[2].cells[4].text == "1.0"
    assert lab.rows[2].cells[5].text == f"{value:.1f}"
    assert [r.cells[3].text for r in lab.rows[3:-1]] == [concentration] * 5
    assert [r.cells[5].text for r in lab.rows[3:-1]] == [mass] * 5
    assert lab.rows[-1].cells[5].text == "10.1"
    assert ctx == raw


def test_barite_lab_word_reference_uses_unrounded_total(source):
    state = deepcopy(source)
    recipe = sp.formulation_from_state(state)
    sample = sp.lab_sample(recipe)
    ctx = context(state)
    signature = ps.fingerprint(state)
    doc = word(ctx)
    lab = next(t for t in doc.tables if t.cell(0, 0).text == "Spacer Lab Results")
    assert [r.cells[3].text for r in lab.rows[2:-1]] == ["1.0", "126.00", "10.00", "2.00", "0.50", "413.461"]
    assert [r.cells[5].text for r in lab.rows[2:-1]] == ["402.4", "144.9", "11.5", "18.0", "4.80", "475.5"]
    assert lab.rows[-1].cells[5].text == "1057.2" == f"{sample['mass_g']:.1f}"
    assert f"{sum(float(r.cells[5].text) for r in lab.rows[2:-1]):.1f}" == "1057.1"
    assert ps.fingerprint(state) == signature
    assert sp.formulation_from_state(state) == recipe
    assert sp.lab_sample(recipe) == sample


@pytest.mark.parametrize("mode,density", [("Limestone", 90.), ("Barite", 110.),
    ("Ferro Bar", 140.), ("Limestone + Barite", 110.)])
@pytest.mark.parametrize("dead", [0., 5.])
def test_canonical_field_table_and_reviewed_lab(source, mode, density, dead):
    state = deepcopy(source)
    state[sp.CONFIG_KEY].update(mode=mode, dead_vol=dead,
        limestone_available_lb=20000., barite_available_lb=20000.)
    state["fluids_config"]["params"]["Spacer"]["density"] = str(density)
    ps.refresh_fluids(state)
    measured(state)
    ensure_current_centralizer(state)
    ctx = context(state)
    recipe = sp.formulation_from_state(state)
    doc = word(ctx)
    tables = [t for t in doc.tables if t.cell(0, 0).text == "Spacer Additives Data"]
    assert len(tables) == 1
    table = tables[0]
    assert len(table.columns) == 6
    assert table._tbl.xpath("./w:tblPr/w:tblLayout/@w:type") == ["fixed"]
    assert table._tbl.xpath("./w:tr[1]/w:tc/w:tcPr/w:gridSpan/@w:val") == ["6"]
    assert table._tbl.xpath("./w:tr[2]/w:tc/w:tcPr/w:gridSpan/@w:val") == ["2", "2"]
    assert [c.text for c in table.rows[2].cells] == [c.replace("\n", " ") for c in sp.FIELD_COLUMNS]
    assert len(table.rows) == 3 + len(recipe["materials"])
    assert [r.cells[1].text for r in table.rows[3:]] == [m["name"] for m in recipe["materials"]]
    assert [r.cells[3].text for r in table.rows[3:]] == [r[sp.FIELD_COLUMNS[3]] for r in sp.field_rows(recipe)]
    assert [r.cells[4].text for r in table.rows[3:]] == [
        f"{material['dose']:.1f}" + (" gal/bbl" if material["unit"] == "gal" else "")
        for row, material in zip(sp.field_rows(recipe), recipe["materials"])]
    assert [r.cells[5].text for r in table.rows[3:]] == [r[sp.FIELD_COLUMNS[5]] for r in sp.field_rows(recipe)]
    assert "20.0 bbl" in [c.text for c in table.rows[1].cells]
    assert "Fresh Water" not in [r.cells[1].text for r in table.rows[3:]]
    body = "\n".join("".join(n.itertext()) for n in doc.element.body)
    assert "Spacer Formulation Data" not in body
    assert "{{" not in body and "{%" not in body
    assert recipe["note"] in body
    assert table._tbl.getnext().tag.endswith("}p")
    assert recipe["note"] in "".join(table._tbl.getnext().itertext())
    lab = next(t for t in doc.tables if t.cell(0, 0).text == "Spacer Lab Results")
    assert len(lab.columns) == 6
    assert [c.text for c in lab.rows[1].cells] == ["Material Type", "Code", "Lot Number",
        "Concentration", "Design Unit", "Mass (g/600 mL)"]
    sample = sp.lab_sample(recipe)
    types = {m["name"]: m["material_type"] for m in recipe["materials"]}
    assert [r.cells[0].text for r in lab.rows[2:-1]] == ["Water"] + [types[r["Material"]] for r in sample["rows"][1:]]
    assert [r.cells[1].text for r in lab.rows[2:-1]] == [r["Material"] for r in sample["rows"]]
    assert [r.cells[4].text for r in lab.rows[2:-1]] == ["1.0"] + [
        "FW " + r["Unit"].split()[0] for r in sample["rows"][1:]]
    assert lab.rows[2].cells[3].text == "1.0"
    expected = cement_lab_display(sample)
    assert [r.cells[3].text for r in lab.rows[2:-1]] == list(expected["Concentration"])
    assert [r.cells[5].text for r in lab.rows[2:-1]] == list(expected["Mass"])
    assert ctx["standalone_spacer"]["lab_rows"] == [dict(row, **{
        "Material Type": "Water" if row["Material"] == "Fresh Water" else types[row["Material"]],
        "Code": row["Material"], "Lot Number": row["Lot No"], "Design Unit": row["Unit"],
    }) for row in state["lab_grid_dfs"]["Spacer"].to_dict("records")]
    assert lab.rows[-1].cells[0].text == "Total Mass"
    assert lab.rows[-1].cells[5].text == f"{sample['mass_g']:.1f}"
    assert ctx["standalone_spacer"]["lab_mass_g"] == sample["mass_g"]
    heading = lab._tbl.getprevious()
    assert "".join(heading.xpath(".//w:t/text()")) == "Spacer Results"
    assert not any(p.text == "Fann 35 / R1B1 / Spring 1.0" for p in doc.paragraphs)
    spacer_region = body.split("Spacer Results", 1)[1].split("VIII. Directional Well Survey Data", 1)[0]
    assert not any(title in spacer_region for title in ("Test Basic Data", "Conventional", "Free Water Test",
        "Thickening Time", "Fluid Loss Test", "Compressive Strength"))
    rheology = next(t for t in doc.tables if t.cell(0, 0).text == "Spacer Rheology")
    assert not "".join(lab._tbl.getnext().itertext())
    assert lab._tbl.getnext().getnext() == rheology._tbl
    assert "IOD" not in [r.cells[0].text for r in rheology.rows]
    assert ctx["standalone_spacer"]["lab"]["rheology"] == ps.spacer_lab_payload(state, rebuild=False)["rheology"]


@pytest.mark.parametrize("prepared", [False, True])
@pytest.mark.parametrize("change", ["review", "reading", "density", "lot", "cache"])
def test_report_rejects_stale_sources_and_retires_old_bytes(source, prepared, change):
    state = deepcopy(source)
    state["_compiled_doc_bytes"] = b"old"
    if change == "review":
        state["lab_qc_params"]["Spacer"]["reviewed"] = False
    elif change == "reading":
        state["lab_qc_params"]["Spacer"]["rheology"]["bhct_down"]["readings"]["300"] = "169"
    elif change == "density":
        state["fluids_config"]["params"]["Spacer"]["density"] = "111"
    elif change == "lot":
        state["lab_grid_dfs"]["Spacer"].loc[0, "Lot No"] = "new lot"
    else:
        state["lab_payload_Spacer"] = {"grid": "forged", "rheology": "forged"}
        # Cached payloads have no admission authority; a genuine current source wins.
        assert context(state, prepared)["standalone_spacer"]["lab"]["rheology"] != "forged"
        return
    with pytest.raises(ValueError, match="Spacer|Phase VII"):
        context(state, prepared)
    assert "_compiled_doc_bytes" not in state


def test_barite_benchmark_and_note_order(source):
    state = deepcopy(source)
    ctx = context(state)
    recipe = sp.formulation_from_state(state)
    assert recipe["water_base_bbl"] == pytest.approx(13.4378344880)
    assert recipe["water_total_bbl"] == pytest.approx(18.4378344880)
    stage = next(row for row in ctx["fluids_train"] if row["name"] == "Spacer")
    assert stage["volume"] == 20. and stage["duration_min"] == 5.
    assert len(ctx["fluids_train"]) == 3  # Ingredients never become pumping stages.
    doc = word(ctx, state)
    table = next(t for t in doc.tables if t.cell(0, 0).text == "Spacer Additives Data")
    assert table.rows[6].cells[3].text == "6.7 gal"
    paragraphs = [p.text for p in doc.paragraphs if re.match(r"NOTE \d+:", p.text.strip())]
    assert [int(re.search(r"NOTE (\d+)", p).group(1)) for p in paragraphs] == list(range(1, len(paragraphs) + 1))
    assert any("18.4 bbl Fresh Water at Batch Mixer Tanks." in p for p in paragraphs)


@pytest.mark.parametrize("standalone", [False, True])
@pytest.mark.parametrize("siblings", [[], ["Spacer Ahead"], ["Spacer Behind"], ["Spacer Ahead", "Spacer Behind"]])
def test_sibling_operations_remain_without_formulation_tables(source, siblings, standalone):
    state = deepcopy(source)
    if not standalone:
        state["fluids_config"]["active"].remove("Spacer")
    for fluid in siblings:
        state["fluids_config"]["active"].insert(0, fluid)
        state["fluids_config"]["params"][fluid] = {
            "volume": 10., "density": "80", "pump_rate": "4", "material_name": fluid}
        state.setdefault("spacer_dfs", {})[fluid] = pd.DataFrame([{
            "Chemical": "Hematite", "Weighting Agent Type": "-", "User Input (% or gal)": 1.25}])
        state["centralizer_config"].setdefault("aux_rheology", {})[fluid] = dict(zip(
            ("300", "200", "100", "60", "30", "6", "3"), (168., 116., 64., 41., 24., 8., 6.)))
    ps.refresh_fluids(state)
    ensure_current_centralizer(state)
    ctx = context(state)
    assert ctx["has_spacer"] is standalone
    assert {r["name"] for r in ctx["fluids_train"]} >= set(siblings)
    text = " ".join(c.text for t in word(ctx).tables for r in t.rows for c in r.cells)
    assert ("Spacer Additives Data" in text) is standalone
    assert "Spacer Formulation Data" not in text
    assert "Hematite" not in text
    assert all(fluid in text for fluid in siblings)
    assert all(state["spacer_dfs"][fluid].iloc[0]["User Input (% or gal)"] == 1.25 for fluid in siblings)


def test_restore_rebuild_and_repeat_are_semantically_current(source):
    state = deepcopy(source)
    state["spacer_dfs"] = {"Spacer": pd.DataFrame([{"Chemical": "FORGED OLD SPACER", "User Input (% or gal)": 123}])}
    state["lab_grid_dfs"]["Spacer"].loc[1, "Lot No"] = "Salt & <lot> 1"
    from engineering_tools import lab_review_signature
    qc = state["lab_qc_params"]["Spacer"]
    qc["review_signature"] = lab_review_signature(qc, state["lab_grid_dfs"]["Spacer"])
    before = context(state)
    restored = audit.round_trip(state)
    ensure_current_centralizer(restored)
    after = context(restored)
    for ctx in (before, after, context(restored)):
        doc = word(ctx)
        lab = next(t for t in doc.tables if t.cell(0, 0).text == "Spacer Lab Results")
        assert lab.rows[3].cells[2].text == "Salt & <lot> 1"
        assert "FORGED OLD SPACER" not in doc.element.xml
        assert ctx["standalone_spacer"] == before["standalone_spacer"]


@pytest.mark.parametrize("change", ["volume", "dead", "tank", "mode", "density", "reading", "lot"])
def test_current_source_changes_document_signature_after_required_review(source, change):
    state = deepcopy(source)
    old = report.document_signature(context(state), report.TEMPLATE_PATH)
    if change in ("volume", "density"):
        state["fluids_config"]["params"]["Spacer"][change] = 21. if change == "volume" else "111"
    elif change in ("dead", "tank", "mode"):
        key, value = {"dead": ("dead_vol", 6.), "tank": ("tank_name", "Mud Reserve Tanks"), "mode": ("mode", "Ferro Bar")}[change]
        state[sp.CONFIG_KEY][key] = value
    elif change == "reading":
        state["lab_qc_params"]["Spacer"]["rheology"]["bhct_down"]["readings"]["300"] = "169"
    else:
        state["lab_grid_dfs"]["Spacer"].loc[0, "Lot No"] = "updated lot"
    ps.refresh_fluids(state)
    if change == "reading":
        from engineering_tools import lab_review_signature
        qc = state["lab_qc_params"]["Spacer"]
        qc["review_signature"] = lab_review_signature(qc, state["lab_grid_dfs"]["Spacer"])
    else:
        measured(state)
    ensure_current_centralizer(state)
    assert report.document_signature(context(state), report.TEMPLATE_PATH) != old


def test_real_build_and_source_edit_retire_download(source):
    from test_uat_round10b_pumping_word import run_build
    app, calls = run_build(source)
    assert len(calls) == 1
    assert "_compiled_doc_bytes" in app.session_state
    doc = Document(BytesIO(app.session_state["_compiled_doc_bytes"]))
    assert sum(t.cell(0, 0).text == "Spacer Additives Data" for t in doc.tables) == 1
    assert sum(t.cell(0, 0).text == "Spacer Rheology" for t in doc.tables) == 1
    audit.AuditRegressions().phase(app, "phase6")
    next(w for w in app.number_input if w.label == "Dead Volume (bbl) - Spacer").set_value(6.).run()
    assert "_compiled_doc_bytes" not in app.session_state


def test_spacer_lab_reuses_cement_heading_and_material_table_styles():
    with ZipFile(report.TEMPLATE_PATH) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": root.nsmap["w"]}
    body = root.find("w:body", ns)
    text = lambda n: "".join(n.xpath(".//w:t/text()", namespaces=ns))
    cement = next(n for n in body if text(n) == "{{ s.name }} Cement Slurry Results")
    spacer = next(n for n in body if text(n) == "Spacer Results")
    canonical = lambda n: etree.tostring(n, method="c14n")
    assert canonical(cement.find("w:pPr", ns)) == canonical(spacer.find("w:pPr", ns))
    assert canonical(cement.find("w:r/w:rPr", ns)) == canonical(spacer.find("w:r/w:rPr", ns))
    cement_table = next(n for n in body if text(n).startswith("ConventionalMaterial Type"))
    spacer_table = spacer.getnext()
    for tag in ("tblPr", "tblGrid"):
        assert canonical(cement_table.find("w:" + tag, ns)) == canonical(spacer_table.find("w:" + tag, ns))
    cement_rows = cement_table.findall("w:tr", ns)
    spacer_rows = spacer_table.findall("w:tr", ns)
    for original, copied in ((cement_rows[0], spacer_rows[0]), (cement_rows[1], spacer_rows[1]),
                             (cement_rows[3], spacer_rows[3]), (cement_rows[3], spacer_rows[-1])):
        for a, b in zip(original.findall("w:tc", ns), copied.findall("w:tc", ns)):
            assert canonical(a.find("w:tcPr", ns)) == canonical(b.find("w:tcPr", ns))
            assert canonical(a.find("w:p/w:pPr", ns)) == canonical(b.find("w:p/w:pPr", ns))


@pytest.mark.parametrize("mode,density", [("Limestone", 95.), ("Barite", 110.),
    ("Ferro Bar", 140.), ("Limestone + Barite", 110.)])
def test_unified_lab_retains_all_lots_rheology_and_unrounded_total(source, mode, density):
    state = deepcopy(source)
    state[sp.CONFIG_KEY].update(mode=mode, limestone_available_lb=20000., barite_available_lb=20000.)
    state["fluids_config"]["params"]["Spacer"]["density"] = str(density)
    measured(state)
    grid = state["lab_grid_dfs"]["Spacer"]
    for index, material in enumerate(grid["Material"]):
        grid.loc[index, "Lot No"] = f"{material} & <lot> {index}"
    from engineering_tools import lab_review_signature
    qc = state["lab_qc_params"]["Spacer"]
    qc["review_signature"] = lab_review_signature(qc, grid)
    ensure_current_centralizer(state)
    ctx = context(state)
    raw = deepcopy(ctx["standalone_spacer"])
    rendered = report._word_quantity_context(ctx)["standalone_spacer"]
    assert ctx["standalone_spacer"] == raw  # Display rounding cannot become source.
    doc = word(ctx)
    lab = next(t for t in doc.tables if t.cell(0, 0).text == "Spacer Lab Results")
    assert [r.cells[2].text for r in lab.rows[2:-1]] == list(grid["Lot No"])
    assert lab.rows[-1].cells[5].text == f"{raw['lab_mass_g']:.1f}"
    assert raw["lab_mass_g"] == sp.lab_sample(sp.formulation_from_state(state))["mass_g"]
    rheology = next(t for t in doc.tables if t.cell(0, 0).text == "Spacer Rheology")
    rows = {r.cells[0].text: r for r in rheology.rows}
    for dataset, column in (("surface_down", 1), ("bhct_up", 2), ("bhct_down", 3)):
        values = rendered["lab"]["rheology"][dataset]
        for rpm in ("300", "200", "100", "60", "30", "6", "3"):
            assert rows[rpm].cells[column].text == values[rpm]
        for label, key in (("Pv (cP)", "pv_cp"), ("Ty (lbf/100ft2)", "ty_lbf_100ft2"),
                           ("10 Sec Gel (deg)", "gel_10_sec"), ("10 Min Gel (deg)", "gel_10_min")):
            assert rows[label].cells[column].text == values[key]
        assert values["iod"] == report._word_rheology(raw["lab"]["rheology"])[dataset]["iod"]
    assert "IOD" not in rows


def test_template_change_is_confined_to_authorized_spacer_regions():
    from test_round13b_post_b4_word_quality import approved_template_xml
    with ZipFile(report.TEMPLATE_PATH) as archive:
        root = etree.fromstring(approved_template_xml(archive.read("word/document.xml")))
        ns = {"w": root.nsmap["w"]}
        body = root.find("w:body", ns)
        text = lambda node: "".join(node.xpath(".//w:t/text()", namespaces=ns))
        for child in list(body):
            if text(child).startswith("Spacer Additives Data") or "standalone_spacer.note_mix" in text(child):
                body.remove(child)
        lab = next(i for i, child in enumerate(body) if text(child) == "Spacer Results")
        for child in list(body)[lab - 1:lab + 5]:
            body.remove(child)
        assert len(body) == 137
        digest = hashlib.sha256()
        for info in archive.infolist():
            data = etree.tostring(root, method="c14n") if info.filename == "word/document.xml" else archive.read(info.filename)
            digest.update(info.filename.encode() + b"\0" + data)
        # Verified A3 template with only its obsolete Spacer table excluded.
        assert digest.hexdigest() == "b3a41cadbbdec15f0229b1aff0eb1164b2ace066f3d18a5120b99f4de43bd64e"


@pytest.mark.parametrize("job", ["CMT PLUG", 'LNR 7"'])
def test_job_valid_cement_unit_tank_and_preflush_coexistence(job):
    app = audit.AuditRegressions().configured_app(job)
    state = app.session_state.to_dict()
    state["fluids_config"]["active"][:0] = ["Pre Flush", "Spacer"]
    state["fluids_config"]["params"].update({
        "Pre Flush": {"volume": 10., "density": "80", "pump_rate": "3", "material_name": "Actual Wash"},
        "Spacer": {"volume": 20., "density": "110", "pump_rate": "4", "material_name": "Weighted Spacer"}})
    state[sp.CONFIG_KEY] = {"mode": "Barite", "dead_vol": 5., "tank_name": "Cement Unit Tanks"}
    state["preflush_config"] = {"type": "Brine (Water + NaCl)", "nacl_multiplier": 126.}
    if report.centralizer_applicable(job):
        state["centralizer_config"].setdefault("aux_rheology", {})["Pre Flush"] = dict(zip(
            ("300", "200", "100", "60", "30", "6", "3"), (168., 116., 64., 41., 24., 8., 6.)))
    ps.refresh_fluids(state)
    measured(state)
    ensure_current_centralizer(state)
    ctx = context(state)
    doc = word(ctx)
    assert sum(t.cell(0, 0).text == "Pre Flush Data" for t in doc.tables) == 1
    assert sum(t.cell(0, 0).text == "Spacer Additives Data" for t in doc.tables) == 1
    assert "18.4 bbl Fresh Water at Cement Unit Tanks." in " ".join(p.text for p in doc.paragraphs)
    assert [r["material_name"] for r in ctx["fluids_train"][:2]] == ["Actual Wash", "Weighted Spacer"]


def test_spacer_note_uses_document_order_before_scavenger_and_lab(source):
    slurry = {"mixing_tank": "Mud Reserve Tanks", "total_water_bbl": 10., "lab": {"surface_hardened_hours": 8.}}
    scavenger = {"mixing_tank": "Mud Reserve Tanks", "total_water_bbl": 2.}
    spacer = {"note_body": sp.formulation_from_state(source)["note"]}
    notes = report.build_ordered_notes(job_type=source["job_type"], well_data=source["well_data"],
        total_pump_time_min=30., slurries_payload=[slurry], has_open_hole_notes=True,
        scavenger_payload=scavenger, spacer_payload=spacer)
    order = notes["all_notes"]
    assert order.index(slurry["note_mix"]) + 1 == order.index(spacer["note_mix"])
    assert order.index(spacer["note_mix"]) + 1 == order.index(scavenger["note_mix"])
    assert order.index(scavenger["note_mix"]) + 1 == order.index(slurry["lab"]["note_freshwater"])


def test_prepared_context_keeps_hydraulic_admission_and_retires_previous_doc(source):
    state = deepcopy(source)
    qc = state["lab_qc_params"]["Spacer"]
    qc["rheology"]["bhct_down"]["selected"] = False
    from engineering_tools import lab_review_signature
    qc["review_signature"] = lab_review_signature(qc, state["lab_grid_dfs"]["Spacer"])
    # Surface data is a valid reviewed A2 lab, but cannot satisfy A3 hydraulics.
    ps.spacer_lab_payload(state, rebuild=False)
    state["_compiled_doc_bytes"] = b"old"
    with pytest.raises(ValueError, match="Spacer.*bhct_down|Spacer.*BHCT"):
        context(state, prepared=True)
    assert "_compiled_doc_bytes" not in state


def test_word_source_signature_retains_precision_beyond_display_rounding(source):
    state = deepcopy(source)
    before = context(state)
    state[sp.CONFIG_KEY]["dead_vol"] += .000001
    after = context(state)
    assert before["standalone_spacer"]["rows"] == after["standalone_spacer"]["rows"]
    assert before["standalone_spacer"]["note_mix"] == after["standalone_spacer"]["note_mix"]
    assert report.document_signature(before, report.TEMPLATE_PATH) != report.document_signature(after, report.TEMPLATE_PATH)


@pytest.mark.parametrize("element", ["IOD", "Fann 35 / R1B1 / Spring 1.0"])
def test_spacer_word_hides_only_unwanted_display_elements(source, element):
    state = deepcopy(source)
    ctx = context(state)
    doc = word(ctx)
    rheology = next(t for t in doc.tables if t.cell(0, 0).text == "Spacer Rheology")
    if element == "IOD":
        assert element not in [r.cells[0].text for r in rheology.rows]
        assert ctx["standalone_spacer"]["lab"]["rheology"]["bhct_down"]["iod"] is not None
    else:
        lab = next(t for t in doc.tables if t.cell(0, 0).text == "Spacer Lab Results")
        assert not "".join(lab._tbl.getnext().itertext())
        assert lab._tbl.getnext().getnext() == rheology._tbl
        assert not any(p.text == element for p in doc.paragraphs)


def test_spacer_field_ui_uses_cement_liquid_and_dosage_precision(source):
    case = audit.AuditRegressions()
    app = case.app(deepcopy(source))
    case.phase(app, "phase6")
    table = next(t.value for t in app.table if "Name" in t.value.columns)
    row = table.loc[table["Name"] == "Defoamer"].iloc[0]
    assert row[sp.FIELD_COLUMNS[3]] == "6.72 gal"
    assert row[sp.FIELD_COLUMNS[4]] == "0.500 gal/bbl"
    assert row[sp.FIELD_COLUMNS[5]] == "9.22 gal"


def test_spacer_lab_ui_formats_without_changing_reviewed_sources(source):
    case = audit.AuditRegressions()
    state = deepcopy(source)
    canonical = state["lab_grid_dfs"]["Spacer"].copy(deep=True)
    qc = deepcopy(state["lab_qc_params"]["Spacer"])
    app = case.app(state)
    case.phase(app, "phase7")
    editor = next(item for item in app.dataframe if "_editor_lab_tbl_spacer_" in item.proto.id)
    grid = editor.value.set_index("Material")
    assert grid.loc["Defoamer", "Mass"] == "4.80"
    assert grid.loc["Fresh Water", "Mass"] == "402.4"
    assert grid.loc["Salt", "Concentration"] == "126.00"
    assert any("Total Mass: 1057.2 g" in item.value for item in app.caption)
    pd.testing.assert_frame_equal(app.session_state["lab_grid_dfs"]["Spacer"], canonical)
    current = app.session_state["lab_qc_params"]["Spacer"]
    assert current["reviewed"] == qc["reviewed"]
    assert current["review_signature"] == qc["review_signature"]
    for dataset, values in qc["rheology"].items():
        assert current["rheology"][dataset] == values
    # The existing shared UI seeds unused dataset widgets on first mount.
    # Once mounted, repeated navigation must preserve the entire record.
    mounted = deepcopy(current)
    case.phase(app, "phase6")
    case.phase(app, "phase7")
    app.run()
    pd.testing.assert_frame_equal(app.session_state["lab_grid_dfs"]["Spacer"], canonical)
    assert app.session_state["lab_qc_params"]["Spacer"] == mounted


@pytest.mark.parametrize("value,solid,liquid,lab_mass,concentration", [
    (.005, "0.0", "0.01", "0.01", "0.005"),
    (.009, "0.0", "0.01", "0.01", "0.009"),
    (9.995, "10.0", "10.0", "9.99", "9.995"),
    (10.005, "10.0", "10.01", "10.0", "10.005"),
    (2., "2.0", "2.0", "2.00", "2.00"),
])
def test_cement_category_display_boundaries_preserve_sources(value, solid, liquid, lab_mass, concentration):
    from phase_6_spacer import _standalone_field_display
    from phase_7_lab import _spacer_lab_display
    recipe = {"materials": [{"name": unit, "material_type": "test", "unit": unit,
                              "base": value, "gross": value, "dose": value} for unit in ("lb", "gal")]}
    original = deepcopy(recipe)
    rows = _standalone_field_display(recipe)
    assert rows[0][sp.FIELD_COLUMNS[3]] == rows[0][sp.FIELD_COLUMNS[5]] == solid + " lb"
    assert rows[1][sp.FIELD_COLUMNS[3]] == rows[1][sp.FIELD_COLUMNS[5]] == liquid + " gal"
    assert rows[0][sp.FIELD_COLUMNS[4]] == f"{value:.3f} lb/bbl"
    assert recipe == original
    grid = pd.DataFrame([{"Material": name, "Concentration": str(value), "Mass": str(value),
                          "Unit": "unchanged", "Lot No": "lot"} for name in ("Fresh Water", "Defoamer")])
    original_grid = grid.copy(deep=True)
    display = _spacer_lab_display(grid)
    assert display.loc[0, "Mass"] == f"{value:.1f}"
    assert display.loc[1, "Mass"] == lab_mass
    assert list(display["Concentration"]) == [concentration, concentration]
    pd.testing.assert_frame_equal(grid, original_grid)


@pytest.mark.parametrize("mode,density", [("Limestone", 95.), ("Barite", 110.),
    ("Ferro Bar", 140.), ("Limestone + Barite", 110.)])
def test_four_mode_display_cannot_become_recipe_or_lab_authority(source, mode, density):
    from phase_6_spacer import _standalone_field_display
    from phase_7_lab import _spacer_lab_display, _commit_spacer_lots
    state = deepcopy(source)
    state[sp.CONFIG_KEY].update(mode=mode, limestone_available_lb=20000., barite_available_lb=20000.)
    state["fluids_config"]["params"]["Spacer"]["density"] = str(density)
    measured(state)
    recipe = sp.formulation_from_state(state)
    original_recipe = deepcopy(recipe)
    grid = state["lab_grid_dfs"]["Spacer"].copy(deep=True)
    signature = ps.lab_source_signature(state, "Spacer")
    _standalone_field_display(recipe)
    display = _spacer_lab_display(grid)
    state["_compiled_doc_bytes"] = b"current"
    with patch("phase_7_lab.st.session_state", state):
        saved = _commit_spacer_lots(display, display)
        pd.testing.assert_frame_equal(saved, grid)
        assert state["lab_qc_params"]["Spacer"]["reviewed"]
        assert state["_compiled_doc_bytes"] == b"current"
        display.loc[0, "Lot No"] = "Water & <lot>"
        display.loc[0, "Mass"] = "999"  # Disabled computed cells have no source authority.
        saved = _commit_spacer_lots(display, display)
        pd.testing.assert_frame_equal(saved.drop(columns="Lot No"), grid.drop(columns="Lot No"))
        assert saved.loc[0, "Lot No"] == "Water & <lot>"
        assert not state["lab_qc_params"]["Spacer"]["reviewed"]
        assert "_compiled_doc_bytes" not in state
    assert ps.lab_source_signature(state, "Spacer") == signature
    assert sp.formulation_from_state(state) == recipe == original_recipe


def test_display_mapping_preserves_explicit_lot_clear(source):
    from phase_7_lab import _spacer_lab_display, _commit_spacer_lots
    state = deepcopy(source)
    grid = state["lab_grid_dfs"]["Spacer"]
    grid.loc[0, "Lot No"] = "previous water lot"
    view = _spacer_lab_display(grid)
    view.loc[0, "Lot No"] = None  # The editor's explicit Delete action.
    with patch("phase_7_lab.st.session_state", state):
        saved = _commit_spacer_lots(view, view)
    assert saved.loc[0, "Lot No"] is None
    pd.testing.assert_frame_equal(saved.drop(columns="Lot No"), grid.drop(columns="Lot No"))
    state["lab_grid_dfs"]["Spacer"] = saved
    ps.refresh_spacer_lab(state)
    assert state["lab_grid_dfs"]["Spacer"].loc[0, "Lot No"] == ""


def test_spacer_lab_and_rheology_remain_separate_word_tables(source):
    doc = word(context(deepcopy(source)))
    lab = next(t for t in doc.tables if t.cell(0, 0).text == "Spacer Lab Results")
    rheology = next(t for t in doc.tables if t.cell(0, 0).text == "Spacer Rheology")
    separator = lab._tbl.getnext()
    assert etree.QName(separator).localname == "p"
    assert not "".join(separator.itertext())
    assert separator.getnext() == rheology._tbl
