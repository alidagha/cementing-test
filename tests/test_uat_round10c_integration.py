"""Round 10C: real Build / project trust-boundary / fresh Build integration."""
import ast
from copy import deepcopy
from io import BytesIO
import json
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
import numpy as np
import pytest
from streamlit.testing.v1 import AppTest

import phase_9_centralizer as phase9
import phase_4_fluids as phase4
import pumping_simulation as pumping
from project_io import content_signature, decode_project, replace_project_state
from project_state import prepare_calculations
import test_audit_regressions as audit
from centralizer_fixtures import execute_centralizer
from test_uat_round9g_centralizer_word import ix_region
from test_uat_round10b_pumping_word import (
    ANCHOR, NA, DESCRIPTION, JOBS, drawing_paragraph, image_bytes,
)


@pytest.fixture(scope="module")
def projects():
    case = audit.AuditRegressions()
    states = {}
    for kind, job in JOBS.items():
        app = case.configured_app(job, optional_fluids=kind == "casing")
        if kind == "casing":
            # Ordinary-size measured program; retain the shared real UI setup.
            case.phase(app, "phase4")
            for fluid in ("Pre Flush", "Spacer", "Spacer Ahead", "Spacer Behind"):
                app.number_input(key="vol_" + fluid).set_value(5.).run()
            app.checkbox(key="chk_fluid_Scavenger").check().run()
            app.number_input(key="vol_Scavenger").set_value(30.).run()
            app.text_input(key="den_Scavenger").set_value("118").run()
            app.text_input(key="rate_Scavenger").set_value("3-5").run()
            case.phase(app, "phase5")
            next(w for w in app.number_input if w.label == "Cement SG - Scavenger").set_value(3.20).run()
            next(w for w in app.number_input if w.label == "Dead Vol (bbl) - Scavenger").set_value(2.).run()
            app.session_state["lab_qc_params"]["Scavenger"] = {"rheology": audit.valid_export_rheology_fixture()}
            case.phase(app, "phase7")
            next(b for b in app.button if b.label == "Confirm measured lab results").click().run()
            assert "Scavenger" in app.session_state["fluids_config"]["active"]
        execute_centralizer(app.session_state)
        state = app.session_state.to_dict()
        assert prepare_calculations(state) == []
        assert phase9.result_is_current(state)
        states[kind] = state
    return states


def save_project(app):
    """Execute the application's save function, then its real JSON encoder."""
    namespace = {**audit.namespace, "st": SimpleNamespace(session_state=app.session_state),
                 "is_project_key": audit.is_project_key}
    functions = [node for node in audit.tree.body if isinstance(node, ast.FunctionDef)
                 and node.name == "get_serializable_state"]
    exec(compile(ast.Module(body=functions, type_ignores=[]), str(audit.ROOT / "main.py"), "exec"), namespace)
    saved = namespace["get_serializable_state"]()
    assert not any(key.startswith("_") for key in saved)
    raw = json.dumps(saved, ensure_ascii=False, sort_keys=True, allow_nan=False).encode()
    assert all(token not in raw for token in (b"SimResult", b"Figure(", b"_compiled_doc_", b"_pumping", b"_centralizer_"))
    return raw


def restore_project(raw):
    app = AppTest.from_file(str(audit.ROOT / "main.py"), default_timeout=30)
    replace_project_state(app.session_state, decode_project(raw, audit.namespace["deserialize_item"]),
                          content_signature(raw))
    assert not any(key.startswith(("_centralizer_", "_compiled_doc_", "_pumping")) for key in app.session_state.to_dict())
    app.run()
    audit.AuditRegressions().healthy(app)
    return app


def build(app):
    case = audit.AuditRegressions()
    case.phase(app, "phase10")
    assert not app.error, [error.value for error in app.error]
    button = next(b for b in app.button if b.label == "Build Word Document")
    assert not button.disabled
    button.click().run()
    case.healthy(app)
    assert not app.error, [error.value for error in app.error]
    return app.session_state["_compiled_doc_bytes"]


def observe_solver(records):
    original = pumping.fcm.build_three_charts

    def observed(**kwargs):
        inputs = deepcopy(kwargs)
        figure, result, warnings = original(**kwargs)
        arrays = {key: value.copy() for key, value in vars(result).items() if isinstance(value, np.ndarray)}
        assert arrays and all(np.isfinite(value).all() for value in arrays.values())
        records.append(dict(inputs=inputs, arrays=arrays, figure=figure))
        return figure, result, warnings

    return patch.object(pumping.fcm, "build_three_charts", side_effect=observed)


def inspect_word(data, applicable=True):
    doc = Document(BytesIO(data))
    paragraphs = [p.text.strip() for p in doc.paragraphs]
    assert ANCHOR not in paragraphs
    assert paragraphs.count("IV. Fluids Sequence") == paragraphs.count("V. Cementing Program") == 1
    assert paragraphs.count("IX. Centralizer Placement Recommendation") == 1
    body = list(doc._element.body)
    v = next(p for p in doc.paragraphs if p.text.strip() == "V. Cementing Program")
    region = ix_region(doc)
    # Keep the existing Section-X page transition and sequential note content.
    assert len(region[-1].xpath('.//w:br[@w:type="page"]')) == 1
    assert "{{" not in "\n".join(paragraphs) and "{%" not in "\n".join(paragraphs)
    if applicable:
        pump = drawing_paragraph(doc, DESCRIPTION)
        diagnostic = drawing_paragraph(doc, "Static Well Security and Fluid Sequence")
        iv = next(p for p in doc.paragraphs if p.text.strip() == "IV. Fluids Sequence")
        assert body.index(iv._p) < body.index(diagnostic._p) < body.index(pump._p) < body.index(v._p)
        assert len(pump._p.xpath('.//w:drawing')) == 1
        assert image_bytes(doc, pump).startswith(b"\x89PNG\r\n\x1a\n")
        drawing_paragraph(doc, "Pipe Standoff")
        return image_bytes(doc, pump), image_bytes(doc, diagnostic), [e.xml for e in region]
    assert paragraphs.count(NA) == 1
    assert not any(p._p.xpath(f'.//wp:docPr[@descr="{DESCRIPTION}"]') for p in doc.paragraphs)
    assert not any(e.xpath('.//w:drawing') or e.tag.endswith('}tbl') for e in region)
    return None


@pytest.fixture(scope="module")
def built_projects(projects):
    built = {}
    for kind, state in projects.items():
        app = audit.AuditRegressions().app(deepcopy(state))
        records = []
        before = pumping.fcm.plt.get_fignums()
        with observe_solver(records):
            data = build(app)
        assert len(records) == 1 and pumping.fcm.plt.get_fignums() == before
        assert records[0]["figure"].number not in pumping.fcm.plt.get_fignums()
        inspect_word(data)
        built[kind] = (app.session_state.to_dict(), data, records[0])
    return built


@pytest.mark.parametrize("kind", JOBS)
def test_save_fresh_restore_real_execute_and_build_are_deterministic(built_projects, kind):
    state, first_doc, first = built_projects[kind]
    case = audit.AuditRegressions()
    app = case.app(deepcopy(state))
    raw = save_project(app)
    saved = json.loads(raw)
    restored = restore_project(raw)
    assert restored.session_state["centralizer_config"] == saved["centralizer_config"]
    assert restored.session_state["placement_config"] == state["placement_config"]
    assert pumping.source_signature(restored.session_state) == pumping.source_signature(state)
    assert not phase9.result_is_current(restored.session_state)
    records = []
    before = pumping.fcm.plt.get_fignums()
    with observe_solver(records):
        case.phase(restored, "phase10")
        assert next(b for b in restored.button if b.label == "Build Word Document").disabled
        assert not records
        case.phase(restored, "phase9")
        restored.button(key="_centralizer_execute").click().run()
        assert not restored.error and phase9.result_is_current(restored.session_state)
        assert not records
        second_doc = build(restored)
    assert len(records) == 1 and pumping.fcm.plt.get_fignums() == before
    assert inspect_word(second_doc) == inspect_word(first_doc)
    assert records[0]["inputs"]["geometry"] == first["inputs"]["geometry"]
    for field, values in first["arrays"].items():
        np.testing.assert_array_equal(records[0]["arrays"][field], values)
    assert save_project(restored) == raw
    geometry = records[0]["inputs"]["geometry"]
    if kind == "casing":
        assert "Scavenger" in records[0]["inputs"]["fluid_data"]
        assert "Scavenger" in records[0]["inputs"]["rheology"]
        assert all(name in records[0]["inputs"]["fluid_data"] for name in pumping.fcm.SPACER_TYPES)
    elif kind == "liner":
        assert len(geometry.pipe_sections) == 2
        assert geometry.pipe_sections[0].z_bot == geometry.pipe_sections[1].z_top == 1000.
        assert geometry.pipe_sections[0].id_in != geometry.pipe_sections[1].id_in
    else:
        assert geometry.track_depth_m == 3000. and geometry.td_m == 3300.
        assert geometry.pipe_sections[-1].z_top == geometry.track_depth_m
        assert geometry.pipe_sections[-1].z_bot == geometry.td_m
        assert geometry.ann_sections[0].do_in == 12.25  # Explicit Previous Casing host.


@pytest.mark.parametrize("missing, phase, fluid", [
    ("mud", "Phase II", "YP > 0"), ("slurry", "Phase VII", "Main"),
    ("aux", "Phase IX", "Spacer"), ("scavenger", "Phase VII", "Scavenger"),
])
def test_missing_required_sources_block_before_solver_and_retire_real_doc(built_projects, missing, phase, fluid):
    state = deepcopy(built_projects["casing"][0])
    assert "_compiled_doc_bytes" in state
    if missing == "mud":
        state["yield_point"] = state["well_data"]["yield_point"] = "0"
    elif missing == "aux":
        del state["centralizer_config"]["aux_rheology"]["Spacer"]
    else:
        del state["lab_qc_params"]["Main" if missing == "slurry" else "Scavenger"]["rheology"]["bhct_down"]
    records = []
    with observe_solver(records):
        app = audit.AuditRegressions().app(state)
        audit.AuditRegressions().phase(app, "phase10")
        assert next(b for b in app.button if b.label == "Build Word Document").disabled
        issues = prepare_calculations(app.session_state.to_dict())
    assert any("Pumping Simulation" in issue and phase in issue and fluid in issue for issue in issues), issues
    assert not records
    assert not any(key.startswith("_compiled_doc_") for key in app.session_state.to_dict())


def test_repeated_build_then_real_source_edit_executes_only_on_build(built_projects):
    case = audit.AuditRegressions()
    app = case.app(deepcopy(built_projects["casing"][0]))
    records = []
    before = pumping.fcm.plt.get_fignums()
    with observe_solver(records):
        first = build(app)
        second = build(app)
        assert len(records) == 2
        assert inspect_word(first) == inspect_word(second)
        for field in records[0]["arrays"]:
            np.testing.assert_array_equal(records[0]["arrays"][field], records[1]["arrays"][field])
        old_signature = app.session_state["_compiled_doc_signature"]
        old_source = pumping.source_signature(app.session_state)
        case.phase(app, "phase4")
        app.text_input(key="rate_Displacement Fluid").set_value("5.0").run()
        assert app.session_state["fluids_config"]["params"]["Displacement Fluid"]["pump_rate"] == "5.0"
        assert pumping.source_signature(app.session_state) != old_source
        assert phase9.result_is_current(app.session_state)
        assert "_compiled_doc_bytes" not in app.session_state
        app.run()
        assert len(records) == 2
        changed = build(app)
    assert len(records) == 3
    assert records[-1]["inputs"]["fluid_data"]["Displacement Fluid"]["min_rate"] == 5.
    assert app.session_state["_compiled_doc_signature"] != old_signature
    assert inspect_word(changed)[0] != inspect_word(first)[0]
    assert pumping.fcm.plt.get_fignums() == before


@pytest.mark.parametrize("job", ["CMT PLUG", "CMT SQUEEZE"])
def test_special_job_save_restore_and_build_remain_na_without_simulation(job):
    case = audit.AuditRegressions()
    app = case.configured_app(job)
    records = []
    with observe_solver(records):
        first = build(app)
        raw = save_project(app)
        restored = restore_project(raw)
        second = build(restored)
    assert not records
    inspect_word(first, applicable=False)
    inspect_word(second, applicable=False)
    assert save_project(restored) == raw


@pytest.mark.parametrize("fluid,field,old,new", [
    ("Main", "volume", 50., 51.),
    ("Spacer", "density", "80", "81"),
    ("Displacement Fluid", "pump_rate", "4", "5"),
    ("Main", "material_name", "Cement Slurry", "Measured Cement"),
])
@pytest.mark.parametrize("changed", [False, True])
def test_fluid_callback_invalidates_only_changed_canonical_source(fluid, field, old, new, changed):
    state = {
        "fluids_config": {"active": [fluid], "params": {fluid: {field: old}}},
        "fluid_data": {fluid: {field: old}}, "_edit": new if changed else old,
        "_compiled_doc_bytes": b"existing document", "_compiled_doc_filename": "existing.docx",
        "_compiled_doc_signature": "existing signature",
    }
    before = deepcopy(state)
    with patch.object(phase4.st, "session_state", state):
        phase4._commit_fluid_widget(fluid, field, "_edit")
    if not changed:
        if field == "density":
            before["fluid_data"][fluid]["effective_density"] = float(old)
        assert state == before
    else:
        assert not any(key.startswith("_compiled_doc_") for key in state)
        before["fluids_config"]["params"][fluid][field] = new
        before["fluid_data"][fluid][field] = new
        if field == "density":
            before["fluid_data"][fluid]["effective_density"] = float(new)
        for key in ("_compiled_doc_bytes", "_compiled_doc_filename", "_compiled_doc_signature"):
            del before[key]
        assert state == before


def test_inactive_fluid_callback_leaves_canonical_and_document_unchanged():
    state = {"fluids_config": {"active": [], "params": {"Main": {"volume": 50.}}},
             "_edit": 51., "_compiled_doc_bytes": b"existing document"}
    before = deepcopy(state)
    with patch.object(phase4.st, "session_state", state):
        phase4._commit_fluid_widget("Main", "volume", "_edit")
    assert state == before
