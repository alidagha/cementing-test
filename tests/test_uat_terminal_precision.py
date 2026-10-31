"""Full-precision arithmetic; independently formatted terminal quantities."""
from copy import deepcopy
from io import BytesIO
import inspect
import unittest
from unittest.mock import patch

import pandas as pd
from docx import Document
import engineering_tools as eng
import phase_5_cement as cement
import phase_10_procedure as report
import test_audit_regressions as audit
from project_state import prepare_calculations, lab_source_signature, fingerprint
from test_uat_round4_mass_balance import formulation


def state_for(auto=True):
    return {'fluids_config': {'active': ['Main']},
            'fluid_data': {'Main': {'volume': 731., 'density': '152'}},
            'cement_params': {'Main': {'auto_calc': auto, 'cmt_sg': 3.2,
                                      'yield': 1.321, 'mix_water': 123.4567,
                                      'dead_vol': 31.1234, 'tank_name': 'Test'}},
            'cement_additives_dfs': {'Main': formulation()}}


def additive(state, mix, dosage):
    return pd.DataFrame([{'Material Type': 'Local Additive', 'Name': 'FIELD-X',
                         'Physical State': state, 'Mix Method': mix,
                         'User Input': dosage, 'Density': 2.2 if state == 'Powder' else 1.1}])


class TerminalPrecision(unittest.TestCase):
    app = audit.AuditRegressions.app
    phase = audit.AuditRegressions.phase
    healthy = audit.AuditRegressions.healthy
    configured_app = audit.AuditRegressions.configured_app
    export = audit.AuditRegressions.export

    def refreshed(self, auto=True):
        state = state_for(auto)
        self.assertEqual(cement.refresh_cement_calculations(state), [])
        result = eng.calculate_slurry_from_components(152., 731., cmt_sg=3.2,
                                                      powders=cement.build_components(formulation())[0],
                                                      liquids=cement.build_components(formulation())[1],
                                                      salt_pct=18.)
        return state['cement_params']['Main'], result

    def test_auto_fresh_water_is_raw(self):
        p, result = self.refreshed()
        self.assertEqual(p['mix_water'], result['field_water_bbl'])

    def test_auto_sacks_are_raw(self):
        p, result = self.refreshed()
        self.assertEqual(p['total_sacks'], result['field_sacks'])

    def test_solution_is_raw(self):
        p, result = self.refreshed()
        self.assertEqual(p['solution'], result['field_solution_bbl'])

    def test_manual_sacks_are_raw(self):
        p, _ = self.refreshed(False)
        self.assertEqual(p['total_sacks'], 731. * eng.BBL_TO_CUFT / 1.321)
        self.assertEqual((p['yield'], p['mix_water']), (1.321, 123.4567))

    def tables(self, state, mix, dosage, auto=False):
        df = additive(state, mix, dosage)
        p = {'auto_calc': auto, 'mix_water': 123.4567, 'total_sacks': 2345.6789,
             'dead_vol': 31.1234, 'tank_name': 'Test'}
        raw = {'field_water_bbl': 123.4567, 'field_sacks': 2345.6789}
        before = deepcopy(p)
        blend, adds, note = cement.build_cement_tables(p, df, raw)
        self.assertEqual(p, before)
        return blend, adds, note, p

    def test_dry_blend_total_uses_raw_concentration(self):
        blend, _, _, p = self.tables('Powder', 'Dry Blend', 2.123)
        self.assertEqual(blend.iloc[1]['sks or lbs'],
                         f"{eng.round_half_up(2.123 * 1.1 * p['total_sacks'], 1):.1f} lb")

    def assert_wet(self, physical, dosage, decimals):
        _, adds, _, p = self.tables(physical, 'In Mix Water', dosage)
        concentration = dosage * 1.1 if physical == 'Powder' else dosage
        per_bbl = concentration * p['total_sacks'] / p['mix_water']
        with_dead = per_bbl * (p['mix_water'] + p['dead_vol'])
        self.assertEqual(float(adds.iloc[0]['(lbs or gal)/bbl'].split()[0]),
                         eng.round_half_up(per_bbl, 3))
        self.assertEqual(float(adds.iloc[0]['lbs or gal (with dead Vol.)'].split()[0]),
                         eng.round_half_up(with_dead, decimals))

    def test_wet_powder_has_no_cascading_rounding(self):
        self.assert_wet('Powder', 2.123, 1)

    def test_liquid_has_no_cascading_rounding(self):
        self.assert_wet('Liquid', .042, 2)

    def test_master_context_fallback_uses_exact_conversion(self):
        self.assertNotIn('5.6146', inspect.getsource(report.build_master_context))

    def test_logistics_fallback_uses_exact_conversion(self):
        self.assertNotIn('5.6146', inspect.getsource(report.render))

    def test_runtime_sacks_fallback_and_logistics_sum_raw(self):
        app = self.configured_app(audit.BATCH1_JOBS[0])
        state = deepcopy(app.session_state.to_dict())
        state['cement_params']['Main'].pop('total_sacks')
        p = state['cement_params']['Main']
        expected = 50. * eng.BBL_TO_CUFT / p['yield']
        with patch.object(report.st, 'session_state', state):
            payload = report.build_master_context(calculations_prepared=True)['slurries'][0]
        self.assertEqual(payload['total_sacks'], expected)
        app.session_state['cement_params']['Main'].pop('total_sacks')
        with patch.object(report, 'prepare_calculations', return_value=[]):
            self.phase(app, 'phase10')
        info = next(i.value for i in app.info if 'Total Cement Sacks' in i.value)
        self.assertIn(f'{expected:.1f} sacks', info)
        # Two raw counts round to 10.0 individually but must total 20.1.
        state['cement_params']['Main']['total_sacks'] = 10.04
        state['cement_params']['Lead'] = dict(p, total_sacks=10.04, top_mode='Surface')
        state['cement_params']['Main'].update(top_mode='Depth (m MD)', top_depth=1000.)
        state['fluid_data']['Lead'] = dict(state['fluid_data']['Main'])
        state['fluids_config']['active'] = ['Lead', 'Main']
        state['fluids_config']['params']['Lead'] = dict(state['fluids_config']['params']['Main'])
        state['cement_additives_dfs']['Lead'] = state['cement_additives_dfs']['Main'].copy()
        app = self.app(state)
        with patch.object(report, 'prepare_calculations', return_value=[]):
            self.phase(app, 'phase10')
        info = next(i.value for i in app.info if 'Total Cement Sacks' in i.value)
        self.assertIn('20.1 sacks', info)

    def test_word_formatting_is_a_terminal_copy(self):
        context = {'slurries': [{'yield_cuft_sk': 1.811171284123,
                                'solution': 123.45678,
                                'lab': {'base_fluid': 6.25412345, 'mix_fluid': 7.15945678}}]}
        before = deepcopy(context)
        terminal = report._word_quantity_context(context)
        self.assertEqual(context, before)
        self.assertEqual(terminal['slurries'][0]['yield_cuft_sk'], '1.811')
        self.assertEqual(terminal['slurries'][0]['solution'], '123.5')
        self.assertEqual(terminal['slurries'][0]['lab']['base_fluid'], '6.254')

    def test_auto_tables_use_engine_sources_and_salt_is_unchanged(self):
        df = pd.concat([additive('Powder', 'Dry Blend', 2.123),
                        additive('Powder', 'In Mix Water', 2.123),
                        additive('Liquid', 'In Mix Water', .413), formulation().query("Name == 'SALT'")],
                       ignore_index=True)
        raw = {'field_water_bbl': 123.4567, 'field_sacks': 2345.6789}
        p = {'auto_calc': True, 'mix_water': 123.5, 'total_sacks': 2345.7,
             'dead_vol': 31.1234, 'tank_name': 'Test'}
        blend, adds, note = cement.build_cement_tables(p, df, raw)
        self.assertEqual(blend.iloc[1]['sks or lbs'],
                         f"{eng.round_half_up(2.123 * 1.1 * raw['field_sacks'], 1):.1f} lb")
        salt = eng.calculate_salt_field_amounts(18., raw['field_water_bbl'], raw['field_sacks'], p['dead_vol'])
        self.assertEqual(adds.iloc[-1].to_dict(), cement.build_salt_additive_row(
            3, 'SALT', df.iloc[-1]['Material Type'], 18., raw['field_water_bbl'], raw['field_sacks'], p['dead_vol']))
        self.assertEqual(float(adds.iloc[-1]['lbs or gal']), eng.round_half_up(salt['base_lb'], 1))
        self.assertIn(f"{eng.round_half_up(raw['field_water_bbl'] + p['dead_vol'], 1):.1f} bbl", note)

    def test_current_review_navigation_restore_and_terminal_word(self):
        app = self.configured_app(audit.BATCH1_JOBS[0], additives=True)
        source = lab_source_signature(app.session_state, 'Main')
        reviewed = deepcopy(app.session_state['lab_qc_params']['Main'])
        self.phase(app, 'phase5')
        p = deepcopy(app.session_state['cement_params']['Main'])
        result = cement.calculate_base_results(deepcopy(p), 50., 118.,
                    *cement.build_components(app.session_state['cement_additives_dfs']['Main']))
        for field, key in (('yield', 'yield_ft3_per_sk'), ('mix_water', 'field_water_bbl'),
                           ('total_sacks', 'field_sacks'), ('solution', 'field_solution_bbl')):
            self.assertEqual(p[field], result[key])
        self.assertEqual(next(m.value for m in app.metric if m.label == 'Fresh Water'),
                         f"{result['field_water_bbl']:.2f} bbl")
        self.assertEqual(source, lab_source_signature(app.session_state, 'Main'))
        self.assertEqual(reviewed, app.session_state['lab_qc_params']['Main'])
        self.phase(app, 'phase2_3'); self.phase(app, 'phase5')
        self.assertEqual(app.session_state['cement_params']['Main'], p)
        state = audit.round_trip(app.session_state.to_dict())
        legacy = deepcopy(state)
        legacy['cement_params']['Main'].update(mix_water=round(p['mix_water'], 1),
                                              total_sacks=round(p['total_sacks'], 1), solution='Historical text')
        self.assertEqual(prepare_calculations(legacy), [])
        self.assertEqual(legacy['cement_params']['Main'], p)
        restored = self.app(audit.round_trip(legacy))
        self.export(restored)
        for _ in range(2):
            snapshot = audit.round_trip(restored.session_state.to_dict())
            signature = fingerprint(snapshot)
            self.phase(restored, 'phase5'); self.phase(restored, 'phase10')
            self.assertEqual(fingerprint(audit.round_trip(restored.session_state.to_dict())), signature)
        doc = Document(BytesIO(restored.session_state['_compiled_doc_bytes']))
        text = '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        self.assertIn(f"{p['yield']:.3f} cuft/sk", text)
        self.assertNotIn(str(p['yield']), text)
        self.assertIn(f"{p['solution']:.1f} bbl", text)
        self.assertIn(f"{p['base_fluid_gal_sk']:.3f} gal/sk", text)


if __name__ == '__main__':
    unittest.main()


# WNF-01: actual Word projection; all synthetic input edits use existing admission.
import ast
from pathlib import Path
import pickle
import subprocess
import pytest
from docxtpl import DocxTemplate
import project_state as ps


@pytest.fixture(scope='module')
def wnf_baseline():
    source = subprocess.check_output(['git', 'show',
        '60dffefb8c941ad90cb7caba0992126a35803717:phase_10_procedure.py'])
    names = {'generate_official_procedure', 'synchronize_report_texts',
             'document_signature', '_word_quantity_context'}
    functions = [n for n in ast.parse(source).body if isinstance(n, ast.FunctionDef) and n.name in names]
    namespace = dict(vars(report))
    exec(compile(ast.Module(body=functions, type_ignores=[]), '<WNF approved baseline>', 'exec'), namespace)
    return namespace


@pytest.fixture(scope='module')
def wnf_v2():
    # Reconstruct only the approved v2 presentation differences, not engineering code.
    source = subprocess.check_output(['git','show','24c578bb0d2c4465b94f4515c67911be93b434b3:phase_10_procedure.py']).decode()
    reversals = [
        ("        if not density:\n            fraction = fraction.rstrip('0')\n", ""),
        ('(("pump_rate", 1), ("density", 1))', '(("pump_rate", 2), ("density", 1))'),
        ('_word_fluid_number(min(configured_rates), 1)', '_word_fluid_number(min(configured_rates), 2)'),
        ('"plug_numeric_display_revision": 2', '"plug_numeric_display_revision": 1'),
        ('"word_output_revision": "WNF-01-r2"', '"word_output_revision": "WNF-01"'),
    ]
    for current, previous in reversals:
        assert source.count(current) == 1
        source = source.replace(current, previous)
    names = {'_word_fluid_number', '_word_quantity_context', 'generate_official_procedure',
             'synchronize_report_texts', 'document_signature'}
    functions = [n for n in ast.parse(source).body if isinstance(n, ast.FunctionDef) and n.name in names]
    namespace = dict(vars(report))
    exec(compile(ast.Module(body=functions, type_ignores=[]), '<approved WNF v2 presentation>', 'exec'), namespace)
    return namespace


@pytest.fixture(scope='module')
def wnf_plug_source():
    app = audit.AuditRegressions().configured_app('CMT PLUG', optional_fluids=True,
                                                 include_standalone_spacer=False)
    state = app.session_state.to_dict()
    state['fluids_config']['params']['Spacer Ahead']['pump_rate'] = '0.7'
    assert ps.prepare_calculations(state) == []
    state.pop('procedure_text', None)
    state.get('report_text_state', {}).pop('procedure_text', None)
    return state


def wnf_sync(state, function=None):
    with patch.object(report.st, 'session_state', state):
        return (function or report.synchronize_report_texts)()


def wnf_word(context, formatter=None):
    doc = DocxTemplate(report.TEMPLATE_PATH)
    doc.render((formatter or report._word_quantity_context)(context))
    stream = BytesIO(); doc.save(stream)
    return stream.getvalue()


@pytest.mark.parametrize('rate,density,w_rate,w_density,minimum', [
    ('5', '118', '5.0', '118.0', 5.),
    ('5.0', '118', '5.0', '118.0', 5.),
    ('5.00', '118', '5.0', '118.0', 5.),
    ('0.70', '118.250', '0.7', '118.250', .7),
    ('0.050', '118', '0.05', '118.0', .05),
    ('0.0050', '118', '0.005', '118.0', .005),
    ('7.0', '118.0', '7.0', '118.0', 7.),
    ('8.00', '118.25', '8.0', '118.25', 8.),
    ('5.000', '118.250', '5.0', '118.250', 5.),
    ('0.7', '80-82', '0.7', '80.0–82.0', .7),
    ('0.05', '80/84', '0.05', '80.0/84.0', .05),
    ('0.005', '80/82-84', '0.005', '80.0/82.0–84.0', .005),
    ('3-5', '118', '3.0–5.0', '118.0', 3.),
    ('0.7-1.2', '118.25', '0.7–1.2', '118.25', .7),
    ('5-3', '84/80', '5.0–3.0', '84.0/80.0', 3.),
    ('+.5-+1.250', '80.00', '+0.5–+1.25', '80.00', .5),
    ('۰٫۷', '۱۱۸٫۲۵', '0.7', '118.25', .7),
    ('١٫٢', '٨٠/٨٤', '1.2', '80.0/84.0', 1.2),
    ('0.000005', '118', '0.000005', '118.0', .000005),
])
def test_wnf_actual_docx_numeric_cells_preserve_raw_context_and_engine_values(
        wnf_plug_source, wnf_baseline, rate, density, w_rate, w_density, minimum, tmp_path):
    state = deepcopy(wnf_plug_source)
    state['fluids_config']['params']['Spacer Ahead'].update(pump_rate=rate, density=density)
    with patch.object(report.st, 'session_state', state):
        ctx = report.build_master_context()
    assert ctx['fluids_train'] == [state['fluid_data'][f] for f in state['fluids_config']['active']]
    raw = state['fluid_data']['Spacer Ahead']
    assert (raw['pump_rate'], raw['density'], raw['min_rate']) == (rate, density, minimum)
    assert raw['duration_min'] == raw['volume'] / minimum
    before, source = deepcopy(ctx), pickle.dumps(state)
    data = wnf_word(ctx)
    assert ctx == before and pickle.dumps(state) == source
    doc = Document(BytesIO(data))
    table = next(t for t in doc.tables if t.rows[0].cells[0].text.strip() == 'Fluids')
    row = next(r for r in table.rows if r.cells[0].text == 'Spacer Ahead')
    assert (row.cells[2].text, row.cells[4].text) == (w_density, w_rate)
    # Compare every other existing cell/step/NOTE against the original terminal renderer.
    comparison = deepcopy(ctx)
    # Independent approved temperature literals for this unchanged 150/200 fixture.
    assert comparison['well_data']['bhst'] == 200
    comparison['well_data']['bhst'] = '200.0'
    for slurry in comparison['slurries']:
        for key, raw, expected in [('bhct',150,'150.0'),('bhst',200,'200.0'),('uca_temp',200,'200.0'),('crush_temp',200,'200.0')]:
            value=slurry['lab'].get(key,'-')
            if value != '-':
                assert value == raw
                slurry['lab'][key] = expected
    original = Document(BytesIO(wnf_word(comparison, wnf_baseline['_word_quantity_context'])))
    old_table = next(t for t in original.tables if t.rows[0].cells[0].text.strip() == 'Fluids')
    for old, new in zip(old_table.rows, table.rows, strict=True):
        assert [c.text for i,c in enumerate(old.cells) if i not in (2,4,5)] == [c.text for i,c in enumerate(new.cells) if i not in (2,4,5)]
    previous_minutes = 0
    for source_row in ctx['fluids_train']:
        h,m=map(int,eng.format_to_hr_mm(source_row['cumul_time_min']).split(':'))
        boundary=h*60+m
        actual=next(r for r in table.rows if r.cells[0].text==source_row['name'])
        assert actual.cells[5].text==eng.format_to_hr_mm(boundary-previous_minutes)
        previous_minutes=boundary
    assert [p.text for p in original.paragraphs] == [p.text for p in doc.paragraphs]
    for old, new in zip(original.tables, doc.tables, strict=True):
        if new is table or new._tbl is table._tbl: continue
        assert [[c.text for c in r.cells] for r in old.rows] == [[c.text for c in r.cells] for r in new.rows]
    (tmp_path / 'word-numeric-cells.docx').write_bytes(data)


@pytest.mark.parametrize('rate,display', [('5','5.0'),('0.7','0.7'),('0.05','0.05'),('0.005','0.005'),('0.000005','0.000005')])
def test_wnf_plug_minimum_docx_changes_only_numeric_token(wnf_plug_source, wnf_baseline, rate, display, tmp_path):
    state = deepcopy(wnf_plug_source)
    state['fluids_config']['params']['Spacer Ahead']['pump_rate'] = rate
    if rate == '5':
        for name in state['fluids_config']['active']:
            state['fluids_config']['params'][name]['pump_rate'] = rate
    ps.refresh_fluids(state)
    old_state = deepcopy(state)
    old = wnf_sync(old_state, wnf_baseline['synchronize_report_texts'])['procedure_text']
    with patch.object(report.st, 'session_state', old_state), \
         patch.object(report, 'synchronize_report_texts', wnf_baseline['synchronize_report_texts']):
        old_ctx = report.build_master_context()
    old_bytes = wnf_word(old_ctx, wnf_baseline['_word_quantity_context'])
    old_line = next(l for l in old['generated'].splitlines() if 'Min rate for pump slurry / displacement,' in l)
    assert old_line in [p.text for p in Document(BytesIO(old_bytes)).paragraphs]
    (tmp_path / 'baseline-plug-minimum.docx').write_bytes(old_bytes)
    fresh = wnf_sync(state)['procedure_text']
    line = next(l for l in fresh['generated'].splitlines() if 'Min rate for pump slurry / displacement,' in l)
    assert line.endswith(f'Min rate for pump slurry / displacement, {display} bbl/ min.')
    old_line = next(l for l in old['generated'].splitlines() if 'Min rate for pump slurry / displacement,' in l)
    expected = old['generated'].replace(old_line, line)
    if rate == '5':
        expected = expected.replace('at 5 bpm;', 'at 5.0 bpm;')
    assert expected == fresh['generated']
    with patch.object(report.st, 'session_state', state): ctx=report.build_master_context()
    data=wnf_word(ctx);doc=Document(BytesIO(data))
    assert line in [p.text for p in doc.paragraphs]
    (tmp_path/'plug-minimum.docx').write_bytes(data)


def test_wnf_signed_auto_migrates_but_generator_only_update_does_not(wnf_plug_source, wnf_baseline):
    original = deepcopy(wnf_plug_source)
    old = wnf_sync(original, wnf_baseline['synchronize_report_texts'])
    saved = deepcopy(original)
    # A real generator-only code change leaves the original signature and saved auto text stale.
    namespace = dict(wnf_baseline)
    namespace['generate_official_procedure'] = report.generate_official_procedure
    node = next(n for n in ast.parse(subprocess.check_output(['git','show',
        '60dffefb8c941ad90cb7caba0992126a35803717:phase_10_procedure.py'])).body
        if isinstance(n,ast.FunctionDef) and n.name=='synchronize_report_texts')
    exec(compile(ast.Module(body=[node],type_ignores=[]),'<generator-only control>','exec'),namespace)
    control = deepcopy(saved)
    wnf_sync(control, namespace['synchronize_report_texts'])
    assert control['procedure_text'] == saved['procedure_text']
    fresh = wnf_sync(saved)
    assert fresh['procedure_text']['signature'] != old['procedure_text']['signature']
    assert fresh['exec_summary_text']['signature'] == old['exec_summary_text']['signature']
    assert not fresh['procedure_text']['pending']
    assert '0.7 bbl/ min.' in saved['procedure_text']
    stable = pickle.dumps(saved)
    wnf_sync(saved)
    assert pickle.dumps(saved) == stable
    assert saved['fluids_config'] == original['fluids_config']
    assert saved['fluid_data'] == original['fluid_data']


@pytest.mark.parametrize('prior', ['main','v2'])
@pytest.mark.parametrize('kind', ['manual','reviewed_manual','legacy','legacy_equal','partial_provenance'])
def test_wnf_plug_manual_legacy_keep_regenerate_restore(wnf_plug_source, wnf_baseline, wnf_v2, kind, prior):
    wnf_baseline = wnf_baseline if prior == 'main' else wnf_v2
    state = deepcopy(wnf_plug_source)
    old = wnf_sync(state, wnf_baseline['synchronize_report_texts'])['procedure_text']
    if kind in ('manual','reviewed_manual'):
        state['procedure_text'] += '\nOwner-reviewed operational instruction: retain 1.25 bpm.'
        if kind=='reviewed_manual':ps.accept_report_text(state,'procedure_text',old['generated'],old['signature'])
    elif kind=='legacy_equal':
        with patch.object(report.st,'session_state',state):
            state['procedure_text']=report.generate_official_procedure(state['job_type'],state['fluid_data'],state['fluids_config']['active'],state['cement_params'],state['cement_additives_dfs'],state.get('preflush_calc'),state.get('geo_md'),state.get('total_pump_time_min'),state.get('placement_config'))
        state['report_text_state'].pop('procedure_text')
    elif kind=='legacy':state['report_text_state'].pop('procedure_text')
    else:state['report_text_state']['procedure_text']['source_signature']=None
    manual=state['procedure_text'];fresh=wnf_sync(state)['procedure_text']
    assert fresh['pending'] and state['procedure_text']==manual
    with patch.object(report.st,'session_state',state):
        with pytest.raises(ValueError,match='Review|review'):report.build_master_context()
    ps.accept_report_text(state,'procedure_text',fresh['generated'],fresh['signature'])
    assert not wnf_sync(state)['procedure_text']['pending'] and state['procedure_text']==manual
    # A subsequent genuine source change still requires the established review.
    state['fluids_config']['params']['Spacer Ahead']['pump_rate']='0.8'
    ps.refresh_fluids(state)
    changed=wnf_sync(state)['procedure_text']
    if manual != fresh['generated']:assert changed['pending']
    assert state['procedure_text']==manual or manual==fresh['generated']
    ps.accept_report_text(state,'procedure_text',changed['generated'],changed['signature'],replace=True)
    assert state['procedure_text']==changed['generated']
    if manual != fresh['generated']:
        assert state['report_text_state']['procedure_text']['history'][-1]==manual
        ps.restore_previous_text(state,'procedure_text')
        assert state['procedure_text']==manual and wnf_sync(state)['procedure_text']['pending']


@pytest.mark.parametrize('prior', ['main','v2'])
def test_wnf_renderer_revision_retires_actual_old_docx_without_churn(wnf_baseline, wnf_v2, tmp_path, prior):
    wnf_baseline = wnf_baseline if prior == 'main' else wnf_v2
    from test_uat_round10c_integration import save_project,restore_project
    case=audit.AuditRegressions();app=case.configured_app('CSG 9 5/8"')
    with patch.object(report,'synchronize_report_texts',wnf_baseline['synchronize_report_texts']), \
         patch.object(report,'_word_quantity_context',wnf_baseline['_word_quantity_context']), \
         patch.object(report,'document_signature',wnf_baseline['document_signature']):
        case.export(app)
        old_bytes=app.session_state['_compiled_doc_bytes'];old_sig=app.session_state['_compiled_doc_signature']
    inputs=deepcopy(app.session_state['fluids_config'])
    with patch.object(report.st,'session_state',app.session_state.to_dict()):ctx=report.build_master_context()
    assert wnf_baseline['document_signature'](ctx,str(report.TEMPLATE_PATH))==old_sig
    assert report.document_signature(ctx,str(report.TEMPLATE_PATH))!=old_sig
    case.phase(app,'phase10')
    assert '_compiled_doc_bytes' not in app.session_state
    assert '_compiled_doc_signature' not in app.session_state
    case.export(app)
    new_bytes=app.session_state['_compiled_doc_bytes'];new_sig=app.session_state['_compiled_doc_signature']
    assert new_sig!=old_sig and new_bytes!=old_bytes
    assert app.session_state['fluids_config']==inputs
    case.phase(app,'phase1');case.phase(app,'phase10')
    assert app.session_state['_compiled_doc_signature']==new_sig
    assert app.session_state['_compiled_doc_bytes']==new_bytes
    raw=save_project(app);restored=restore_project(raw)
    assert restored.session_state['fluids_config']==inputs
    case.export(restored)
    assert restored.session_state['_compiled_doc_signature']==new_sig
    assert [p.text for p in Document(BytesIO(new_bytes)).paragraphs]==[p.text for p in Document(BytesIO(restored.session_state['_compiled_doc_bytes'])).paragraphs]
    import json
    (tmp_path/'cache-signatures.json').write_text(json.dumps({'old_signature':old_sig,'new_signature':new_sig,
        'stable_rerun':True,'restore_signature_equal':True,'same_raw_context_under_old_signer':True},indent=2))
    (tmp_path/'old-renderer.docx').write_bytes(old_bytes);(tmp_path/'new-renderer.docx').write_bytes(new_bytes)


@pytest.mark.parametrize('value', [None, '', '   ', '-', '.... [PUMP RATE NOT SET IN PHASE IV]',
                                    'bad', '1e-3', '3--5', '0', '-0.7', False])
def test_wnf_terminal_view_preserves_missing_and_invalid_markers(value):
    context={'fluids_train':[{'pump_rate':value,'density':value,'volume':1.234567,'duration_min':.123456789,'nested':{'raw':'kept'}}]}
    before=deepcopy(context);view=report._word_quantity_context(context)
    assert context==before
    assert view['fluids_train']==before['fluids_train']
    view['fluids_train'][0]['nested']['raw']='copy-only'
    assert context==before


@pytest.mark.parametrize('kind',['manual','legacy_equal'])
def test_wnf_actual_plug_review_buttons_preserve_unknown_text(wnf_plug_source,wnf_baseline,kind,tmp_path):
    state=deepcopy(wnf_plug_source)
    old=wnf_sync(state,wnf_baseline['synchronize_report_texts'])['procedure_text']
    if kind=='manual':state['procedure_text']+='\nOwner-reviewed manual instruction, never replace silently.'
    else:
        fresh=deepcopy(state);fresh.pop('procedure_text');fresh['report_text_state'].pop('procedure_text')
        state['procedure_text']=wnf_sync(fresh)['procedure_text']['generated']
        state['report_text_state'].pop('procedure_text')
    manual=state['procedure_text'];case=audit.AuditRegressions();app=case.app(state)
    case.phase(app,'phase10')
    assert app.session_state['procedure_text']==manual
    assert next(b for b in app.button if b.label=='Build Word Document').disabled
    assert '_compiled_doc_bytes' not in app.session_state
    app.button(key='_keep_procedure_text').click().run()
    case.healthy(app)
    assert app.session_state['procedure_text']==manual
    assert not next(b for b in app.button if b.label=='Build Word Document').disabled
    if kind=='manual':
        app.button(key='_regenerate_procedure_text').click().run()
        assert '0.7 bbl/ min.' in app.session_state['procedure_text']
        assert app.session_state['report_text_state']['procedure_text']['history'][-1]==manual
        app.button(key='_restore_procedure_text').click().run()
        assert app.session_state['procedure_text']==manual
        assert next(b for b in app.button if b.label=='Build Word Document').disabled
        app.button(key='_keep_procedure_text').click().run()
    case.export(app)
    doc=Document(BytesIO(app.session_state['_compiled_doc_bytes']))
    for line in manual.splitlines():
        if line and 'Cementing Procedure' not in line:assert line in [p.text for p in doc.paragraphs]
    (tmp_path/'reviewed-plug.docx').write_bytes(app.session_state['_compiled_doc_bytes'])


def test_wnf_v2_signed_auto_migrates_once_with_summary_and_sources_unchanged(wnf_plug_source, wnf_v2, tmp_path):
    state = deepcopy(wnf_plug_source)
    old = wnf_sync(state, wnf_v2['synchronize_report_texts'])
    assert '0.70 bbl/ min.' in state['procedure_text']
    with patch.object(report.st, 'session_state', state), patch.object(report, 'synchronize_report_texts', wnf_v2['synchronize_report_texts']):
        old_ctx = report.build_master_context()
    old_docx = wnf_word(old_ctx, wnf_v2['_word_quantity_context'])
    source = pickle.dumps((state['fluids_config'], state['fluid_data']))
    old_text = state['procedure_text']
    fresh = wnf_sync(state)
    assert state['procedure_text'] == old_text.replace('0.70 bbl/ min.', '0.7 bbl/ min.')
    assert fresh['procedure_text']['signature'] != old['procedure_text']['signature']
    assert fresh['exec_summary_text'] == old['exec_summary_text']
    assert not fresh['procedure_text']['pending']
    assert pickle.dumps((state['fluids_config'], state['fluid_data'])) == source
    stable = pickle.dumps(state)
    wnf_sync(state)
    assert pickle.dumps(state) == stable
    with patch.object(report.st, 'session_state', state): new_ctx = report.build_master_context()
    new_docx = wnf_word(new_ctx)
    for data, token in ((old_docx, '0.70'), (new_docx, '0.7')):
        assert any(p.text.endswith(f'Min rate for pump slurry / displacement, {token} bbl/ min.')
                   for p in Document(BytesIO(data)).paragraphs)
    (tmp_path/'v2-auto-plug.docx').write_bytes(old_docx)
    (tmp_path/'v3-auto-plug.docx').write_bytes(new_docx)
    import json
    (tmp_path/'auto-signatures.json').write_text(json.dumps({
        'v2_procedure_signature': old['procedure_text']['signature'],
        'v3_procedure_signature': fresh['procedure_text']['signature'],
        'summary_unchanged': fresh['exec_summary_text'] == old['exec_summary_text'],
        'stable_rerun': True, 'sources_unchanged': True}, indent=2))


# WNF-02: report-only rates, additive minute boundaries and source temperatures.
@pytest.fixture(scope='module')
def wnf02_previous():
    source = subprocess.check_output(['git','show','24c578bb0d2c4465b94f4515c67911be93b434b3:phase_10_procedure.py'])
    names={'_word_fluid_number','_word_quantity_context','generate_official_procedure',
           '_ct_procedure','synchronize_report_texts','document_signature'}
    namespace=dict(vars(report))
    nodes=[n for n in ast.parse(source).body if isinstance(n,ast.FunctionDef) and n.name in names]
    exec(compile(ast.Module(body=nodes,type_ignores=[]),'<WNF-01 v3 exact prior>','exec'),namespace)
    return namespace


def wnf02_rows(durations):
    cumulative=0.;rows=[]
    for i,d in enumerate(durations):
        cumulative+=d
        rows.append({'name':f'Stage {i+1}','material_name':'Synthetic admitted fluid','volume':d,
            'pump_rate':'1.00','density':'80.00','min_rate':1.,'duration_min':d,
            'cumul_time_min':cumulative,'duration_str':eng.format_to_hr_mm(d),
            'cumul_time_str':eng.format_to_hr_mm(cumulative)})
    return rows


@pytest.mark.parametrize('durations,stage,cumulative',[
    ([12.5,12.5],['00:12','00:13'],['00:12','00:25']),
    ([59.5,1.5,60.25],['00:59','00:02','01:00'],['00:59','01:01','02:01']),
    ([1.,60.,120.],['00:01','01:00','02:00'],['00:01','01:01','03:01']),
])
def test_wnf02_actual_word_additive_minutes(wnf_plug_source,durations,stage,cumulative,tmp_path):
    s=deepcopy(wnf_plug_source)
    with patch.object(report.st,'session_state',s):ctx=report.build_master_context()
    ctx['fluids_train']=wnf02_rows(durations);ctx['total_pump_time_min']=sum(durations)
    ctx['total_pump_time_hhmm']=eng.format_to_hr_mm(sum(durations))
    before=pickle.dumps(ctx)
    view=report._word_quantity_context(ctx)
    assert [r['duration_str'] for r in view['fluids_train']]==stage
    assert [r['cumul_time_str'] for r in view['fluids_train']]==cumulative
    assert view['total_pump_time_hhmm']==cumulative[-1]
    data=wnf_word(ctx);doc=Document(BytesIO(data))
    table=next(t for t in doc.tables if t.rows[0].cells[0].text.strip()=='Fluids')
    actual=[[c.text for c in r.cells] for r in table.rows]
    for i,(d,c) in enumerate(zip(stage,cumulative)):
        row=next(r for r in actual if r[0]==f'Stage {i+1}')
        assert row[5:7]==[d,c]
    assert pickle.dumps(ctx)==before
    assert not any(__import__('re').search(r'\b\d{2}:\d{2}:\d{2}\b',c.text) for r in table.rows for c in r.cells)
    (tmp_path/'additive-word.docx').write_bytes(data)


@pytest.mark.parametrize('durations,blocked', [([.999,.992],'Stage 2'),([.5],'Stage 1')])
@pytest.mark.parametrize('entry',['terminal','signature','context'])
def test_wnf02_zero_display_word_only_block_and_cache_retirement(durations,blocked,entry):
    rows=wnf02_rows(durations);rows[0]['name']='Main'
    if len(rows)>1:rows[1]['name']='Displacement Fluid';blocked='Displacement Fluid'
    else:blocked='Main'
    state={'job_type':'CSG 20"','mud_density':80.,'fluids_config':{
        'active':[r['name'] for r in rows], 'params':{r['name']:{'volume':r['volume'],'density':r['density'],'pump_rate':r['pump_rate']} for r in rows}},
        '_compiled_doc_bytes':b'previous compiled cache','_compiled_doc_signature':'old',
        'procedure_text':'Owner instruction must stay verbatim'}
    ps.refresh_fluids(state)  # Actual Phase IV still admits the unchanged plan.
    expected=pickle.dumps((state['fluids_config'],state['fluid_data'],state['total_pump_time_min'],state['procedure_text']))
    context={'fluids_train':list(state['fluid_data'].values())}
    with patch.object(report.st,'session_state',state):
        with pytest.raises(ValueError,match=blocked+'.*HH:MM|HH:MM.*'+blocked) as error:
            if entry=='terminal':report._word_quantity_context(context)
            elif entry=='signature':report.document_signature(context,str(report.TEMPLATE_PATH))
            else:report.build_master_context(calculations_prepared=True)
    assert 'adjust' in str(error.value).lower() and '00:00' in str(error.value)
    assert '_compiled_doc_bytes' not in state and '_compiled_doc_signature' not in state
    assert pickle.dumps((state['fluids_config'],state['fluid_data'],state['total_pump_time_min'],state['procedure_text']))==expected


@pytest.mark.parametrize('bhct,bhst,ct_display,st_display',[(150,200,'150.0','200.0'),(150.25,200.25,'150.3','200.3')])
def test_wnf02_actual_word_temperatures_not_rheology(wnf_plug_source,wnf02_previous,bhct,bhst,ct_display,st_display,tmp_path):
    s=deepcopy(wnf_plug_source)
    with patch.object(report.st,'session_state',s):ctx=report.build_master_context()
    ctx['well_data']=dict(ctx['well_data'],bhst=bhst)
    for slurry in ctx['slurries']:slurry['lab'].update(bhct=bhct,bhst=bhst,uca_temp=bhst,crush_temp=bhst)
    before=pickle.dumps(ctx);view=report._word_quantity_context(ctx)
    assert view['well_data']['bhst']==st_display
    old=wnf02_previous['_word_quantity_context'](ctx)
    for a,z in zip(old['slurries'],view['slurries']):
        assert z['lab']['bhct']==ct_display and z['lab']['bhst']==st_display
        assert z['lab']['uca_temp']==st_display and z['lab']['crush_temp']==st_display
        assert z['lab']['rheology']==a['lab']['rheology']
        for key in ('yield_cuft_sk','total_sacks','mix_water_bbl','lab_conventional','lab_additives'):assert a[key]==z[key]
    data=wnf_word(ctx);doc=Document(BytesIO(data))
    text='\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert ct_display+' degF' in text and st_display+' degF' in text
    assert f'Tstart: {ct_display} degF' in text
    assert pickle.dumps(ctx)==before
    (tmp_path/'temperatures-word.docx').write_bytes(data)


@pytest.mark.parametrize('job',['CSG 9 5/8"','LNR 7"','TIE BACK LNR 7"','CMT PLUG','CMT SQUEEZE'])
@pytest.mark.parametrize('rate,display',[('5.00','5.0'),('0.70','0.7'),('0.050','0.05'),('0.0050','0.005'),('0.7-1.2','0.7–1.2')])
def test_wnf02_operational_procedure_tokens_preserve_sources(job,rate,display,wnf02_previous):
    from test_round13b_b4_operational_spacers import operational
    s=operational(job);cfg=s['fluids_config'];cfg['active']=['Main','Spacer Behind','Displacement Fluid']
    minimum=eng.require_positive_pump_rate(rate)
    for name in cfg['active']:cfg['params'][name].update(volume=12.5*minimum,pump_rate=rate)
    ps.refresh_fluids(s);source=pickle.dumps((s['fluid_data'],s['cement_params'],cfg))
    with patch.object(report.st,'session_state',s):
        args=(job,s['fluid_data'],cfg['active'],s['cement_params'],s.get('cement_additives_dfs',{}),s.get('preflush_calc'),s.get('geo_md'),s.get('total_pump_time_min'),s.get('placement_config'))
        old=wnf02_previous['generate_official_procedure'](*args);new=report.generate_official_procedure(*args)
    old_line=next(l for l in old.splitlines() if 'as Spacer Behind' in l)
    line=next(l for l in new.splitlines() if 'as Spacer Behind' in l)
    assert f'at {display} bpm; duration 00:13 (HH:MM).' in line
    assert old.replace(old_line,line)==new
    assert pickle.dumps((s['fluid_data'],s['cement_params'],cfg))==source


from test_round13b_b4_ct_report import valid_ct as wnf02_ct
from test_uat_round12a5_diesel_word_e2e import projects as wnf02_projects


@pytest.mark.parametrize('rate,display',[('5.00','5.0'),('0.0050','0.005'),('+.5-+1.250','+0.5–+1.25')])
def test_wnf02_actual_ct_word_rate_only_procedure_change(wnf02_ct,wnf02_previous,rate,display,tmp_path):
    s=deepcopy(wnf02_ct)
    for name in s['fluids_config']['active']:s['fluids_config']['params'][name]['pump_rate']=rate
    assert ps.prepare_calculations(s)==[]
    before=pickle.dumps((s['fluids_config'],s['fluid_data'],s['cement_params'],s['lab_qc_params']))
    old=wnf02_previous['_ct_procedure'](s);new=report._ct_procedure(s)
    assert old.replace(f'at {rate} bpm',f'at {display} bpm')==new
    for r in s['fluid_data'].values():assert f"calculated duration {r['duration_min']:.2f} min." in new
    with patch.object(report.st,'session_state',s):ctx=report.build_master_context()
    data=wnf_word(ctx);doc=Document(BytesIO(data));text='\n'.join(p.text for p in doc.paragraphs)
    assert f'at {display} bpm' in text and old!=new
    assert pickle.dumps((s['fluids_config'],s['fluid_data'],s['cement_params'],s['lab_qc_params']))==before
    (tmp_path/'ct-rate.docx').write_bytes(data)


def test_wnf02_actual_standalone_temperature_not_dose_or_rheology(wnf02_projects,wnf02_previous,tmp_path):
    s=wnf02_projects('CMT PLUG',include_standalone_spacer=True)
    with patch.object(report.st,'session_state',s):ctx=report.build_master_context()
    ctx['standalone_spacer']['lab']['bhct']=150.25
    before=pickle.dumps(ctx);view=report._word_quantity_context(ctx);old=wnf02_previous['_word_quantity_context'](ctx)
    assert view['standalone_spacer']['lab']['bhct']=='150.3'
    assert view['standalone_spacer']['rows']==old['standalone_spacer']['rows']
    assert view['standalone_spacer']['lab_rows']==old['standalone_spacer']['lab_rows']
    assert view['standalone_spacer']['lab']['rheology']==old['standalone_spacer']['lab']['rheology']
    data=wnf_word(ctx);doc=Document(BytesIO(data))
    assert 'Tstart: 150.3 degF' in '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert pickle.dumps(ctx)==before
    (tmp_path/'standalone-temperature.docx').write_bytes(data)


@pytest.mark.parametrize('value',[None,'-','unfinished',False])
def test_wnf02_temperature_missing_invalid_markers_and_other_fields(value):
    ctx={'well_data':{'bhst':value,'surface_temp':20.123,'geo_gradient':1.234},
         'slurries':[{'lab':dict(bhct=value,bhst=value,uca_temp=value,crush_temp=value)}]}
    before=deepcopy(ctx);v=report._word_quantity_context(ctx)
    assert v['well_data']==before['well_data']
    assert v['slurries'][0]['lab']==before['slurries'][0]['lab']
    assert ctx==before


@pytest.mark.parametrize('kind',['auto','manual','reviewed','legacy','legacy_equal','partial'])
@pytest.mark.parametrize('job',['CSG 9 5/8"','CMT PLUG WITH CT'])
def test_wnf02_prior_auto_and_unknown_ownership(wnf02_projects,wnf02_ct,wnf02_previous,kind,job):
    s=deepcopy(wnf02_ct) if job=='CMT PLUG WITH CT' else wnf02_projects(job)
    for name in s['fluids_config']['active']:s['fluids_config']['params'][name]['pump_rate']='5.00'
    ps.refresh_fluids(s);s.pop('procedure_text',None);s.get('report_text_state',{}).pop('procedure_text',None)
    old=wnf_sync(s,wnf02_previous['synchronize_report_texts'])
    if kind in ('manual','reviewed'):
        s['procedure_text']+='\nOwner reviewed operational instruction, preserve 0.70 literally.'
        if kind=='reviewed':ps.accept_report_text(s,'procedure_text',old['procedure_text']['generated'],old['procedure_text']['signature'])
    if kind=='legacy_equal':
        fresh=deepcopy(s);fresh.pop('procedure_text');fresh['report_text_state'].pop('procedure_text')
        s['procedure_text']=wnf_sync(fresh)['procedure_text']['generated']
    if kind in ('legacy','legacy_equal'):s['report_text_state'].pop('procedure_text')
    if kind=='partial':s['report_text_state']['procedure_text']['source_signature']=None
    text=s['procedure_text'];sources=pickle.dumps((s['fluids_config'],s['fluid_data'],s['lab_qc_params']))
    new=wnf_sync(s)
    assert new['exec_summary_text']==old['exec_summary_text']
    if kind=='auto':
        assert not new['procedure_text']['pending']
        assert new['procedure_text']['signature']!=old['procedure_text']['signature']
        assert s['procedure_text']==new['procedure_text']['generated'] and '5.00 bpm' not in s['procedure_text']
        stable=pickle.dumps(s);wnf_sync(s);assert pickle.dumps(s)==stable
        restored=audit.round_trip(s);assert not wnf_sync(restored)['procedure_text']['pending']
    else:
        assert new['procedure_text']['pending'] and s['procedure_text']==text
        ps.accept_report_text(s,'procedure_text',new['procedure_text']['generated'],new['procedure_text']['signature'])
        assert not wnf_sync(s)['procedure_text']['pending'] and s['procedure_text']==text
        if text!=new['procedure_text']['generated']:
            ps.accept_report_text(s,'procedure_text',new['procedure_text']['generated'],new['procedure_text']['signature'],replace=True)
            assert s['report_text_state']['procedure_text']['history'][-1]==text
            ps.restore_previous_text(s,'procedure_text');assert s['procedure_text']==text
            assert wnf_sync(s)['procedure_text']['pending']
    assert pickle.dumps((s['fluids_config'],s['fluid_data'],s['lab_qc_params']))==sources


@pytest.mark.parametrize('job',['CSG 9 5/8"','LNR 7"','TIE BACK LNR 7"','CMT PLUG','CMT SQUEEZE'])
def test_wnf02_no_operational_spacer_sources_do_not_churn(wnf02_projects,wnf02_previous,job):
    s=wnf02_projects(job);s['fluids_config']['active']=[n for n in s['fluids_config']['active'] if n not in ('Spacer Ahead','Spacer Behind')]
    ps.refresh_fluids(s)
    previous=deepcopy(s);current=deepcopy(s)
    old=wnf_sync(previous,wnf02_previous['synchronize_report_texts']);new=wnf_sync(current)
    for key in ('exec_summary_text','procedure_text'):
        assert new[key]['signature']==old[key]['signature']
        assert new[key]['generated']==old[key]['generated']


def test_wnf02_real_old_cache_rebuild_and_zero_plan_ui_block(wnf02_projects,wnf02_previous,tmp_path):
    from test_uat_round10c_integration import save_project,restore_project
    case=audit.AuditRegressions();app=case.app(wnf02_projects('CSG 9 5/8"'))
    with patch.object(report,'synchronize_report_texts',wnf02_previous['synchronize_report_texts']), \
         patch.object(report,'_word_quantity_context',wnf02_previous['_word_quantity_context']), \
         patch.object(report,'document_signature',wnf02_previous['document_signature']):
        case.export(app)
        old_bytes=app.session_state['_compiled_doc_bytes'];old_sig=app.session_state['_compiled_doc_signature']
    source=pickle.dumps((app.session_state['fluids_config'],app.session_state['fluid_data'],app.session_state['lab_qc_params']))
    case.phase(app,'phase10');assert '_compiled_doc_bytes' not in app.session_state
    case.export(app);new_bytes=app.session_state['_compiled_doc_bytes'];new_sig=app.session_state['_compiled_doc_signature']
    assert old_sig!=new_sig and old_bytes!=new_bytes
    assert pickle.dumps((app.session_state['fluids_config'],app.session_state['fluid_data'],app.session_state['lab_qc_params']))==source
    case.phase(app,'phase1');case.phase(app,'phase10')
    assert app.session_state['_compiled_doc_signature']==new_sig and app.session_state['_compiled_doc_bytes']==new_bytes
    raw=save_project(app);restored=restore_project(raw);case.phase(restored,'phase10')
    assert '_compiled_doc_bytes' not in restored.session_state
    # Re-execute the saved Dynamic design, not the generic helper's Static default.
    import phase_9_centralizer as phase9
    from centralizer_fixtures import _copy_execution
    replay=restored.session_state.to_dict();config=deepcopy(replay['centralizer_config'])
    phase9.execute(replay)
    assert replay['centralizer_config']==config
    _copy_execution(replay,restored.session_state)
    with patch.object(report.st,'session_state',restored.session_state.to_dict()):ctx=report.build_master_context()
    assert report.document_signature(ctx,str(report.TEMPLATE_PATH))==new_sig
    (tmp_path/'previous-wnf01.docx').write_bytes(old_bytes);(tmp_path/'wnf02.docx').write_bytes(new_bytes)
    import json
    (tmp_path/'cache-signatures.json').write_text(json.dumps({'old':old_sig,'new':new_sig,'stable':True,'json_restore_equal':True},indent=2))
    # Keep Main above its real shoe-track requirement; only the operational stage is short.
    app.session_state['fluids_config']['params']['Spacer Behind'].update(volume=.1,pump_rate='1')
    case.phase(app,'phase10');case.healthy(app)
    diagnostic=deepcopy(app.session_state.to_dict())
    (tmp_path/'blocked-ui-diagnostic.json').write_text(json.dumps({
        'errors':[e.value for e in app.error], 'warnings':[e.value for e in app.warning],
        'active':diagnostic['fluids_config']['active'], 'fluid_data':diagnostic['fluid_data'],
        'calculation_issues':ps.prepare_calculations(diagnostic),
        'procedure_pending':wnf_sync(diagnostic)['procedure_text']['pending'],
        'stale_doc_present':'_compiled_doc_bytes' in app.session_state,
        'buttons':[(button.label,button.disabled) for button in app.button]},indent=2))
    assert any('HH:MM' in e.value and 'Spacer Behind' in e.value for e in app.error)
    assert '_compiled_doc_bytes' not in app.session_state and '_compiled_doc_signature' not in app.session_state
    assert next(b for b in app.button if b.label=='Build Word Document').disabled
    assert not any('Download Final Document' in e.proto.label for e in app.get('download_button'))
    assert app.session_state['fluid_data']['Main']['volume']==50.
    assert app.session_state['fluid_data']['Spacer Behind']['duration_min']==.1
