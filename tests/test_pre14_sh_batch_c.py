"""Pre Flush: one measured Phase-VII authority, legacy-only auxiliary fallback."""
from copy import deepcopy
import pickle

import pytest

import centralizer_adapter as ca
import engineering_tools as et
import phase_9_centralizer as p9
import pumping_simulation as pumping
import rheology
import test_audit_regressions as audit
from test_round13a_a3_spacer_dynamic import project, centralizer
from test_uat_round10c_integration import save_project, restore_project

CASE = audit.AuditRegressions()


def current(state):
    readings = deepcopy(state['centralizer_config']['aux_rheology']['Pre Flush'])
    state['lab_qc_params']['Pre Flush'] = {'rheology': {'bhct_down': {
        'selected': True, 'readings': readings, 'gel_10_sec': '-', 'gel_10_min': '-'}}}
    return state


def test_real_phase7_preflush_tab_preserves_legacy_measurements():
    state = project()
    original = deepcopy(state['centralizer_config']['aux_rheology']['Pre Flush'])
    app = CASE.app(state)
    CASE.phase(app, 'phase7')
    assert 'Pre Flush' in [tab.label for tab in app.tabs]
    row = app.session_state['lab_qc_params']['Pre Flush']['rheology']['bhct_down']
    assert row['selected'] is True and row['readings'] == original
    assert row['gel_10_sec'] == row['gel_10_min'] == '-'
    assert 'Pre Flush' not in app.session_state['lab_grid_dfs']
    assert app.session_state['centralizer_config']['aux_rheology']['Pre Flush'] == original


@pytest.mark.parametrize('consumer', [centralizer, pumping.prepare_pumping_simulation])
def test_current_invalid_data_cannot_fall_back_to_legacy(consumer):
    state = current(project())
    state['lab_qc_params']['Pre Flush']['rheology']['bhct_down']['readings']['300'] = ''
    before = pickle.dumps(state)
    with pytest.raises(ValueError, match='Pre Flush'):
        consumer(state)
    assert pickle.dumps(state) == before


@pytest.mark.parametrize('kind', ['casing', 'liner', 'tieback'])
def test_both_consumers_use_current_seven_readings_and_preserve_legacy(kind):
    state = current(project(kind))
    row = state['lab_qc_params']['Pre Flush']['rheology']['bhct_down']
    row['readings']['300'] = '169'
    original = deepcopy(state['centralizer_config']['aux_rheology']['Pre Flush'])
    expected = ca._dynamic_fann(row['readings'], 'Pre Flush')
    dynamic = centralizer(state)['diagnostics']['rheology']['Pre Flush']
    inputs = pumping.prepare_pumping_simulation(state)['inputs']
    assert dynamic == {'source': 'lab_bhct_down', **expected}
    assert inputs['fluid_data']['Pre Flush']['pv_cp'] == expected['pv_cp']
    assert inputs['fluid_data']['Pre Flush']['yp_lbf100ft2'] == expected['ty_lbf_100ft2']
    assert state['centralizer_config']['aux_rheology']['Pre Flush'] == original


@pytest.mark.parametrize('consumer', [centralizer, pumping.prepare_pumping_simulation])
@pytest.mark.parametrize('defect', ['unselected', 'empty', 'missing', 'extra', 'zero', 'negative', 'bool', 'nan', 'infinity', 'text'])
def test_invalid_current_source_is_never_replaced_by_legacy(consumer, defect):
    state = current(project())
    row = state['lab_qc_params']['Pre Flush']['rheology']['bhct_down']
    if defect == 'unselected':
        row['selected'] = False
    elif defect == 'missing':
        del row['readings']['3']
    elif defect == 'extra':
        row['readings']['600'] = 300
    else:
        row['readings']['300'] = {'empty': '', 'zero': 0, 'negative': -1, 'bool': True,
            'nan': float('nan'), 'infinity': float('inf'), 'text': 'invalid'}[defect]
    with pytest.raises(ValueError, match='Pre Flush'):
        consumer(state)
    assert et.compute_phase_status(state)['phase7']['level'] == 'warning'


@pytest.mark.parametrize('signature', [p9.source_signature, pumping.source_signature])
def test_signatures_track_only_consumed_current_data_and_preserve_legacy_migration(signature):
    state = project()
    old = signature(state)
    current(state)
    assert signature(state) == old
    state['centralizer_config']['aux_rheology']['Pre Flush']['300'] = 999
    assert signature(state) == old
    qc = state['lab_qc_params']['Pre Flush']
    qc['rheology']['surface_down'] = deepcopy(qc['rheology']['bhct_down'])
    qc['rheology']['surface_down']['readings']['300'] = 170
    assert signature(state) == old
    qc['rheology']['bhct_down']['readings']['300'] = '169'
    assert signature(state) != old


def select_and_fill(app, dataset, readings):
    label = dict(rheology.RHEOLOGY_DATASETS)[dataset]
    next(w for w in app.checkbox if w.label == f'{label} - Pre Flush').check().run()
    for rpm, value in readings.items():
        next(w for w in app.text_input if w.label == f'{rpm} RPM - {label} - Pre Flush').set_value(str(value))
    # Real widget callbacks run before immediate cross-phase navigation.
    CASE.phase(app, 'phase1')


def test_real_fann_datasets_navigation_restore_and_document_retirement():
    state = project()
    readings = state['centralizer_config']['aux_rheology'].pop('Pre Flush')
    state['fluids_config']['active'] = ['Pre Flush', 'Displacement Fluid']
    app = CASE.app(state)
    CASE.phase(app, 'phase7')
    assert 'Pre Flush' not in app.session_state['lab_qc_params']
    assert [tab.label for tab in app.tabs] == ['Pre Flush']
    assert not app.dataframe
    assert not any(b.label == 'Confirm measured lab results' for b in app.button)
    assert et.compute_phase_status(app.session_state)['phase7']['level'] == 'warning'
    with pytest.raises(ValueError, match='Pre Flush'):
        pumping.prepare_pumping_simulation(app.session_state)
    for dataset, label in rheology.RHEOLOGY_DATASETS:
        CASE.phase(app, 'phase7')
        app.session_state['_compiled_doc_bytes'] = b'stale'
        app.session_state['_compiled_doc_filename'] = 'stale.docx'
        app.session_state['_compiled_doc_signature'] = 'stale'
        select_and_fill(app, dataset, readings)
        qc = app.session_state['lab_qc_params']['Pre Flush']
        assert qc['rheology'][dataset]['readings'] == {k: str(v) for k, v in readings.items()}
        if dataset == 'bhct_down':
            assert not any(k.startswith('_compiled_doc_') for k in app.session_state.to_dict())
    assert set(qc) == {'rheology'}  # No cement tests, grids or operator-review obligation.
    assert 'Pre Flush' not in app.session_state['lab_grid_dfs']
    assert et.compute_phase_status(app.session_state)['phase7']['level'] == 'ok'
    snapshot = deepcopy(qc)
    restored = restore_project(save_project(app))
    CASE.phase(restored, 'phase7')
    assert restored.session_state['lab_qc_params']['Pre Flush'] == snapshot
    for _, label in rheology.RHEOLOGY_DATASETS:
        assert len([w for w in restored.text_input if f'RPM - {label} - Pre Flush' in w.label]) == 7
    assert pumping.prepare_pumping_simulation(restored.session_state)['applicable']


@pytest.mark.parametrize('job', ['CMT PLUG', 'CMT SQUEEZE', 'CMT PLUG WITH CT'])
def test_nonapplicable_jobs_acquire_no_preflush_hydraulic_blocker(job):
    state = project()
    state['job_type'] = job
    state['fluids_config']['active'] = ['Pre Flush', 'Displacement Fluid']
    state['lab_qc_params']['Pre Flush'] = {'rheology': {}}
    assert pumping.prepare_pumping_simulation(state)['applicable'] is False
    assert centralizer(state) is None
    assert et.compute_phase_status(state)['phase7']['level'] == 'ok'


@pytest.mark.parametrize('physics', ['static', 'dynamic'])
def test_current_and_legacy_control_same_solver_without_mutating_measurements(physics):
    state = project()
    state['centralizer_config']['physics_mode'] = physics
    legacy = p9.execute(state)
    current(state)
    assert p9.result_is_current(state)
    measured = p9.execute(state)
    assert measured['profile'] == legacy['profile']
    state['lab_qc_params']['Pre Flush']['rheology']['bhct_down']['readings']['300'] = '169'
    assert p9.result_is_current(state) is (physics == 'static')


def test_valid_preflush_lab_enables_real_word_and_invalid_edit_retires_it(tmp_path):
    from io import BytesIO
    from docx import Document
    import project_state as ps
    app = CASE.configured_app('CSG 9 5/8"', optional_fluids=True, include_standalone_spacer=False)
    CASE.phase(app, 'phase7')
    assert 'Pre Flush' in [tab.label for tab in app.tabs]
    assert et.compute_phase_status(app.session_state)['phase7']['level'] == 'ok'
    assert ps.prepare_calculations(app.session_state.to_dict()) == []
    CASE.export(app)
    raw = app.session_state['_compiled_doc_bytes']
    doc = Document(BytesIO(raw))
    assert any(p._p.xpath('.//wp:docPr[@descr="Pumping Simulation Charts"]') for p in doc.paragraphs)
    assert any('Pre Flush' in cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    assert not any(t.rows[0].cells[0].text == 'Pre Flush Lab Data' for t in doc.tables)
    (tmp_path / 'preflush-lab-current.docx').write_bytes(raw)
    CASE.phase(app, 'phase7')
    w = next(w for w in app.text_input if w.label == '300 RPM - BHCT — Ramp-Down - Pre Flush')
    w.set_value('')
    CASE.phase(app, 'phase1')
    assert not any(k.startswith('_compiled_doc_') for k in app.session_state.to_dict())
    assert any('Pre Flush' in issue for issue in ps.prepare_calculations(app.session_state.to_dict()))
    CASE.phase(app, 'phase10')
    assert next(b for b in app.button if b.label == 'Build Word Document').disabled
    CASE.phase(app, 'phase7')
    next(w for w in app.text_input if w.label == '300 RPM - BHCT — Ramp-Down - Pre Flush').set_value('168.0')
    CASE.phase(app, 'phase1')
    assert ps.prepare_calculations(app.session_state.to_dict()) == []
    CASE.phase(app, 'phase10')
    assert not next(b for b in app.button if b.label == 'Build Word Document').disabled
    assert not any(k.startswith('_compiled_doc_') for k in app.session_state.to_dict())


def test_current_blank_gel_draft_is_not_replaced_by_an_rpm_callback():
    state = current(project())
    row = state['lab_qc_params']['Pre Flush']['rheology']['bhct_down']
    row['gel_10_sec'] = ''
    state['fluids_config']['active'] = ['Pre Flush', 'Displacement Fluid']
    app = CASE.app(state)
    CASE.phase(app, 'phase7')
    next(w for w in app.text_input if w.label == '300 RPM - BHCT — Ramp-Down - Pre Flush').set_value('169')
    CASE.phase(app, 'phase1')
    row = app.session_state['lab_qc_params']['Pre Flush']['rheology']['bhct_down']
    assert row['readings']['300'] == '169'
    assert row['gel_10_sec'] == ''
    assert pumping.prepare_pumping_simulation(app.session_state)['applicable']


def test_inactive_current_preflush_drafts_do_not_change_either_hydraulic_consumer():
    state = current(project())
    state['fluids_config']['active'].remove('Pre Flush')
    before = p9.source_signature(state)
    signature = pumping.source_signature(state)
    dynamic = centralizer(state)
    prepared = pumping.prepare_pumping_simulation(state)
    state['lab_qc_params']['Pre Flush']['rheology'] = {}
    state['centralizer_config']['aux_rheology']['Pre Flush'] = {}
    assert p9.source_signature(state) == before
    assert pumping.source_signature(state) == signature
    assert centralizer(state) == dynamic
    assert pumping.prepare_pumping_simulation(state)['inputs']['fluid_data'] == prepared['inputs']['fluid_data']


def test_authorized_phase7_integrity_pin_rejects_one_changed_byte(monkeypatch):
    from pathlib import Path
    from test_round13b_b4_ct_report import test_unrelated_guardrails_and_all_frozen_authorities_are_unchanged as identity
    read_bytes = Path.read_bytes
    checked = []
    def changed(path):
        data = read_bytes(path)
        if path.name == 'phase_7_lab.py':
            checked.append(path)
            return bytes([data[0] ^ 1]) + data[1:]
        return data
    identity()
    monkeypatch.setattr(Path, 'read_bytes', changed)
    with pytest.raises(AssertionError):
        identity()
    assert len(checked) == 1
